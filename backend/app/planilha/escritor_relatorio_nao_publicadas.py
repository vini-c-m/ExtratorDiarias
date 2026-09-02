"""Sprint 4 (plano "XLSX Operacional de Diárias Não Publicadas") — escritor
do SEGUNDO artefato do módulo, `Diarias_Nao_Publicadas_<competencia>.xlsx`.

**OPERACIONAL/INTERNO — nunca o produto oficial do Portal da Transparência.**
Existe só para facilitar a investigação/encaminhamento administrativo das
ocorrências que não entraram no XLSX oficial após a aplicação das decisões
humanas. Nenhuma informação aqui é publicada externamente.

Deliberadamente um workbook construído do zero (`openpyxl.Workbook()`), não
uma cópia do template oficial — usar o template oficial aqui criaria
acoplamento artificial entre dois produtos conceitualmente separados (este
arquivo não precisa da aba/layout/fórmulas do modelo do Portal, que existem
só para reproduzir o padrão visual já aprovado pelo órgão externo).

Consome diretamente `list[OcorrenciaNaoPublicavel]` (`decisao.py`,
`ResultadoRevisado.ocorrencias_nao_publicaveis`) — a classificação em
Categoria A/B nasce em `aplicar_decisoes` (`decisao.CategoriaNaoPublicada`),
nunca reconstruída aqui nem no frontend. Todo campo com resolução canônica
já homologada (Nome, Cargo, Transporte, Quantidade de Diárias, Valor da
Passagem) reaproveita EXATAMENTE as mesmas funções puras do XLSX oficial
(`linha_publicavel.py`, `quantidade_diarias.py`) — zero lógica paralela.

Sprint 5 (fora do escopo deste módulo): geração conjunta com o XLSX
oficial, promoção atômica coordenada entre os dois arquivos, UX pós-
geração. Este módulo só sabe gerar o SEU arquivo, isoladamente — reaproveita
os MESMOS primitivos de segurança já homologados no escritor oficial (lock
por destino, escrita em temporário + rename atômico, no-clobber), nunca uma
segunda implementação concorrente."""

import os
import tempfile
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.entities.linha_passagem import EstadoPassagem
from app.domain.relatorio_transparencia.decisao import CategoriaNaoPublicada, OcorrenciaNaoPublicavel
from app.domain.relatorio_transparencia.linha_publicavel import (
    nome_publicado,
    valor_do_controle_com_fallback,
    valor_passagem_publicavel,
)
from app.domain.relatorio_transparencia.quantidade_diarias import resolver_quantidade_diarias
from app.planilha.escritor_relatorio_transparencia import (  # noqa: PLC2701 — mesmo reuso já praticado em outros módulos (ex. `_cpf_normalizado`): primitivos de segurança já homologados, nunca uma segunda implementação concorrente
    DestinoJaExiste,
    GeracaoConcorrenteEmAndamento,
    TemplateEstruturalmenteInvalido,
    _adquirir_trava_destino,
)

__all__ = [
    "DestinoJaExiste",
    "GeracaoConcorrenteEmAndamento",
    "RelatorioNaoPublicadasGeradoResultado",
    "gerar_xlsx_diarias_nao_publicadas",
]

_NOME_ABA = "Não Publicadas"

_CABECALHOS = (
    ("Categoria", 34),
    ("Regra", 8),
    ("SEI", 24),
    ("DH", 14),
    ("OB", 16),
    ("Nome", 30),
    ("Cargo", 26),
    ("Unidade de Origem", 42),
    ("Período Inicial", 14),
    ("Período Final", 14),
    ("Trecho", 32),
    ("Transporte", 12),
    ("Nº Diárias Integral", 10),
    ("Nº Diárias Parcial", 10),
    ("Valor das Diárias", 14),
    ("Valor da Passagem", 14),
    ("Situação da Passagem", 18),
    ("Responsável pela Decisão", 26),
    ("Data/Hora da Decisão", 18),
    ("Observação/Justificativa", 42),
)

_ROTULOS_CATEGORIA = {
    CategoriaNaoPublicada.NAO_PUBLICADA_POR_DECISAO_HUMANA: "Não publicada por decisão humana",
    CategoriaNaoPublicada.DUPLICATA_FISICA_IDENTICA_SUPRIMIDA: "Duplicata física idêntica suprimida",
}

_ROTULOS_ESTADO_PASSAGEM = {
    EstadoPassagem.NUMERICO: "NUMÉRICO",
    EstadoPassagem.AUSENTE: "AUSENTE",
    EstadoPassagem.NAO_SE_APLICA: "NÃO SE APLICA",
    EstadoPassagem.INVALIDO: "INVÁLIDO",
}
_ROTULO_PASSAGEM_SEM_CORRESPONDENCIA = "SEM CORRESPONDÊNCIA"


@dataclass(frozen=True)
class RelatorioNaoPublicadasGeradoResultado:
    caminho_arquivo: str
    total_ocorrencias: int
    total_categoria_a: int
    total_categoria_b: int


def _situacao_passagem(oc) -> str:
    """Sprint 4 — rótulo humano do estado de passagem (A-E), reaproveitando
    o MESMO `EstadoPassagem` já resolvido pelo cruzamento (Sprint 1/2) —
    nunca uma nova classificação. `oc.passagem is None` é, literalmente, o
    estado E (SEM CORRESPONDÊNCIA) — mesma convenção de
    `linha_publicavel.valor_passagem_publicavel`."""
    if oc.passagem is None:
        return _ROTULO_PASSAGEM_SEM_CORRESPONDENCIA
    return _ROTULOS_ESTADO_PASSAGEM[oc.passagem.estado]


def _escrever_linha(ws: Worksheet, linha: int, item: OcorrenciaNaoPublicavel) -> None:
    oc = item.ocorrencia
    ordem = oc.ordem_bancaria
    quantidade = resolver_quantidade_diarias(oc.registros_controle)
    cargo = valor_do_controle_com_fallback(oc.registros_controle, "cargo", ordem.cargo)
    transporte = valor_do_controle_com_fallback(oc.registros_controle, "transporte", ordem.transporte)
    # Unidade de origem (item 7 do plano): `RegistroControle.doc_origem` via
    # o MESMO padrão de fallback já usado para Cargo/Transporte — a Base OB
    # não tem um campo equivalente (achado da investigação de origem deste
    # campo), então o "fallback" aqui é sempre `None` quando o Controle
    # DIMCAC não resolve um valor único — nunca inferido de outra fonte.
    unidade_origem = valor_do_controle_com_fallback(oc.registros_controle, "doc_origem", None)
    valor_passagem = valor_passagem_publicavel(oc.passagem)

    ws[f"A{linha}"] = _ROTULOS_CATEGORIA[item.categoria]
    ws[f"B{linha}"] = item.decisao.regra
    ws[f"C{linha}"] = ordem.sei
    ws[f"D{linha}"] = oc.dh
    ws[f"E{linha}"] = ordem.ob
    ws[f"F{linha}"] = nome_publicado(ordem)
    ws[f"G{linha}"] = cargo
    ws[f"H{linha}"] = unidade_origem
    ws[f"I{linha}"] = ordem.periodo_inicial
    ws[f"J{linha}"] = ordem.periodo_final
    ws[f"K{linha}"] = ordem.trecho
    ws[f"L{linha}"] = transporte
    ws[f"M{linha}"] = quantidade.integral if quantidade else None
    ws[f"N{linha}"] = quantidade.parcial if quantidade else None
    ws[f"O{linha}"] = ordem.valor
    ws[f"P{linha}"] = valor_passagem.valor
    ws[f"Q{linha}"] = _situacao_passagem(oc)
    ws[f"R{linha}"] = item.decisao.autor
    # `DecisaoHumana.decidido_em` é sempre timezone-aware (UTC —
    # `registrar_decisao` grava `datetime.now(timezone.utc)`), mas o
    # formato XLSX (openpyxl) não aceita `tzinfo` em células de data —
    # escreve o mesmo instante, sem reinterpretar fuso (nenhuma conversão
    # para horário local inventada; mesmo valor gravado no JSON de
    # decisões, só sem o marcador de fuso que o Excel não representa).
    decidido_em = item.decisao.decidido_em
    ws[f"S{linha}"] = decidido_em.replace(tzinfo=None) if decidido_em.tzinfo else decidido_em
    ws[f"T{linha}"] = item.decisao.observacao

    ws[f"I{linha}"].number_format = "DD/MM/YYYY"
    ws[f"J{linha}"].number_format = "DD/MM/YYYY"
    ws[f"O{linha}"].number_format = "#,##0.00"
    ws[f"P{linha}"].number_format = "#,##0.00"
    ws[f"S{linha}"].number_format = "DD/MM/YYYY HH:MM:SS"


def _construir_workbook(ocorrencias_classificadas: list[OcorrenciaNaoPublicavel]) -> Workbook:
    workbook = Workbook()
    ws = workbook.active
    ws.title = _NOME_ABA

    for indice, (cabecalho, largura) in enumerate(_CABECALHOS, start=1):
        celula = ws.cell(row=1, column=indice, value=cabecalho)
        celula.font = Font(bold=True)
        celula.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[celula.column_letter].width = largura
    ws.row_dimensions[1].height = 30

    for indice, item in enumerate(ocorrencias_classificadas, start=2):
        _escrever_linha(ws, indice, item)

    ultima_linha = max(len(ocorrencias_classificadas) + 1, 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(_CABECALHOS)).column_letter}{ultima_linha}"

    return workbook


def gerar_xlsx_diarias_nao_publicadas(
    ocorrencias_classificadas: list[OcorrenciaNaoPublicavel],
    *,
    caminho_destino: str,
) -> RelatorioNaoPublicadasGeradoResultado:
    """Gera o XLSX operacional a partir das ocorrências JÁ classificadas
    (`decisao.CategoriaNaoPublicada`) — nunca reclassifica, nunca lê fontes.
    Reaproveita os MESMOS primitivos de segurança do escritor oficial:
    trava cooperativa por destino (`_adquirir_trava_destino`), escrita em
    arquivo temporário na mesma pasta do destino, validação pós-escrita,
    promoção via `os.rename` (nunca sobrescreve silenciosamente — mesma
    proteção contra TOCTOU já homologada). Sprint 5 cuida da coordenação
    ENTRE este arquivo e o oficial (promoção conjunta); esta função só
    garante a segurança do SEU PRÓPRIO destino, isoladamente."""
    destino_dir = os.path.dirname(os.path.abspath(caminho_destino)) or "."
    os.makedirs(destino_dir, exist_ok=True)

    caminho_lock = _adquirir_trava_destino(caminho_destino)
    try:
        if os.path.exists(caminho_destino):
            raise DestinoJaExiste(f"Já existe um arquivo no destino: {caminho_destino}")

        descritor_tmp, caminho_tmp = tempfile.mkstemp(suffix=".xlsx", dir=destino_dir)
        os.close(descritor_tmp)
        try:
            workbook = _construir_workbook(ocorrencias_classificadas)
            try:
                workbook.save(caminho_tmp)
            finally:
                # Sempre fecha, mesmo se `save` falhar no meio — caso
                # contrário o handle interno do openpyxl mantém o
                # temporário travado (Windows) e o `finally` externo
                # (`os.remove(caminho_tmp)`) falharia com `PermissionError`
                # em vez de propagar o erro real da escrita.
                workbook.close()

            validacao = load_workbook(caminho_tmp)
            if _NOME_ABA not in validacao.sheetnames:
                raise TemplateEstruturalmenteInvalido(
                    "Falha ao validar o arquivo operacional gerado (aba ausente)."
                )
            validacao.close()

            try:
                os.rename(caminho_tmp, caminho_destino)
            except FileExistsError as erro:
                raise DestinoJaExiste(f"Já existe um arquivo no destino: {caminho_destino}") from erro
        finally:
            if os.path.exists(caminho_tmp):
                os.remove(caminho_tmp)
    finally:
        os.remove(caminho_lock)

    total_a = sum(
        1
        for item in ocorrencias_classificadas
        if item.categoria is CategoriaNaoPublicada.NAO_PUBLICADA_POR_DECISAO_HUMANA
    )
    total_b = len(ocorrencias_classificadas) - total_a
    return RelatorioNaoPublicadasGeradoResultado(
        caminho_arquivo=caminho_destino,
        total_ocorrencias=len(ocorrencias_classificadas),
        total_categoria_a=total_a,
        total_categoria_b=total_b,
    )
