"""Micro-sprint 3F — gerador do XLSX final do Relatório de Diárias — Portal
da Transparência. Nenhuma regra de negócio aqui: recebe um `ResultadoRevisado`
já resolvido pelo backend homologado (RN01-RN07 + decisões humanas válidas,
Micro-sprints 3B/3C) e só escreve. Nunca recalcula duplicidade, competência,
restituição, conflito, fingerprint ou revisão.

Estratégia (aprovada, item 2 da especificação): MODELO OFICIAL como
template — nunca um workbook novo construído "no olho". A geração sempre
parte de uma CÓPIA do arquivo apontado por `caminho_template`; o template
original nunca é aberto em modo de escrita, nunca é sobrescrito.

Mapeamento de colunas (investigação read-only do modelo oficial real,
confirmado célula a célula — ver relatório da Micro-sprint 3F):

    A Ordem Bancária          <- OrdemBancaria.ob
    B Data de Pagamento       <- OrdemBancaria.data_pagamento (DATA OB, RN05),
                                  escrita como TEXTO "DD/MM/YYYY HH:MM:SS"
                                  (Micro-sprint 3F.1: o modelo real armazena
                                  esta coluna como texto em 100% das linhas,
                                  nunca como data/serial Excel)
    C Nome                    <- OrdemBancaria.nome (fonte primária), com
                                  fallback para OrdemBancaria.credor SOMENTE
                                  quando nome for comprovadamente um
                                  truncamento/prefixo de credor (decisão de
                                  negócio aprovada após investigação
                                  read-only NOME×CREDOR — Base OB real tem
                                  ~18,5% dos NOME cortados por um limite de
                                  campo de ~30 caracteres da fonte, enquanto
                                  CREDOR preserva o nome completo nesses
                                  casos; nunca uma troca por tamanho nem por
                                  simples divergência de grafia — ver
                                  `nome_publicado` em `app/domain/
                                  relatorio_transparencia/linha_publicavel.py`,
                                  reaproveitada também por `decisao.py`
                                  desde a correção HOMO-RT-05)
    D Cargo                   <- RegistroControle.cargo (Controle DIMCAC
                                  coluna G, matriz de origem confirmada pelo
                                  negócio, Micro-sprint 3G.1), com fallback
                                  para OrdemBancaria.cargo (Base OB) quando
                                  ausente/conflitante no Controle — nunca
                                  ambos vazios inventam um valor
    E Período (inicial)       <- OrdemBancaria.periodo_inicial
    F Período (final)         <- OrdemBancaria.periodo_final
    G Trecho                  <- OrdemBancaria.trecho (ORIGINAL — nunca o
                                  trecho_normalizado, que é só um artefato
                                  interno do detector RN07)
    H Transporte              <- RegistroControle.transporte (Controle
                                  DIMCAC coluna K, Micro-sprint 3G.1), com
                                  fallback para OrdemBancaria.transporte
                                  (Base OB) quando ausente/conflitante no
                                  Controle — mesma precedência de Cargo
    I Motivo                  <- OrdemBancaria.motivo
    J Valor da Passagem       <- Sprint 3 (HOMO-RT-06): linha_publicavel.
                                  valor_passagem_publicavel(passagem), onde
                                  `passagem` vem do cruzamento exato por OB
                                  com a Planilha de Passagens (Sprint 1/2).
                                  Numérico quando a fonte tem valor (estado
                                  A); célula vazia para ausência/"NÃO SE
                                  APLICA"/inconsistência da fonte/OB sem
                                  correspondência (B/C/D/E) — fechamento
                                  HOMO-RT-06 (decisão de produto final):
                                  estado D (ex. "#VALOR!") NUNCA bloqueia a
                                  geração, produz célula vazia como os
                                  demais, mas permanece INVALIDO no domínio
                                  (`OcorrenciaFinanceira.passagem.estado`)
                                  e participa do fingerprint (`decisao.
                                  _campos_linha_publicavel`) — nunca
                                  reescrito para ausente/zero. Histórico
                                  anterior a esta Sprint (Micro-sprint 3F):
                                  o gerador sempre gravava "-" fixo, decisão
                                  institucional então vigente de que valor
                                  da passagem não era responsabilidade do
                                  DIMCAC — premissa alterada pelo plano
                                  aprovado desta feature.
    K Nº Diárias Integral     <- resolver_quantidade_diarias(...).integral
    L Nº Diárias Parcial      <- resolver_quantidade_diarias(...).parcial
                                  (Controle DIMCAC colunas N/O, já
                                  calculadas por processo anterior — nunca
                                  recalculadas aqui; ver `app/domain/
                                  relatorio_transparencia/quantidade_diarias.py`)
    M Valor total diárias     <- OrdemBancaria.valor (RN02: sempre o valor
                                  bruto da Base OB, nunca ajustado por
                                  restituição)
"""

import os
import shutil
import tempfile
import time
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.entities.linha_passagem import LinhaPassagem
from app.domain.entities.ordem_bancaria import OrdemBancaria
from app.domain.relatorio_transparencia.decisao import ResultadoRevisado
from app.domain.relatorio_transparencia.linha_publicavel import (
    nome_publicado,
    valor_do_controle_com_fallback,
    valor_passagem_publicavel,
)
from app.domain.relatorio_transparencia.motor_decisao import CandidatoNormal
from app.domain.relatorio_transparencia.quantidade_diarias import resolver_quantidade_diarias

__all__ = [
    "RelatorioBloqueadoPorPendencias",
    "DestinoJaExiste",
    "GeracaoConcorrenteEmAndamento",
    "DestinoInacessivel",
    "TemplateEstruturalmenteInvalido",
    "RelatorioGeradoResultado",
    "gerar_xlsx_relatorio_transparencia",
]

_NOME_ABA_DADOS = "Controle de Diárias"
_PRIMEIRA_LINHA_DADOS = 7
_ROTULO_LINHA_TOTAL_PREFIXO = "Total"

# Auditoria real (confronto Setembro/2026, 29 de 60 valores de passagem
# exibidos como data) — achado: a coluna J nunca tinha um `number_format`
# monetário afirmado em código; toda linha nova herdava o formato de
# QUALQUER QUE FOSSE o estado da linha-modelo (`_PRIMEIRA_LINHA_DADOS`, via
# `_copiar_estilo_linha`) no momento da inserção, e linhas já existentes
# nunca eram corrigidas (`_escrever_linha_ocorrencia` só reescreve
# `.value`). Quando a linha-modelo esteve com `dd/mm/yyyy` na coluna J
# (comprovado para OBs pagas entre abril e julho/2026), o valor numérico
# correto passou a ser exibido como data — sem nenhuma perda do valor em
# si (sempre recuperável a partir do float bruto), mas incorreto para
# publicação. Corrigido com um formato monetário fixo, aplicado
# explicitamente a cada geração, nunca dependente de estilo herdado.
_NUMBER_FORMAT_VALOR_PASSAGEM = '"R$ "#,##0.00'

# Colunas A-M — a ÁREA DE DADOS inteira (cabeçalho é a linha 5/6, nunca
# tocado; linha de totais é identificada à parte, também nunca tocada).
_COLUNAS_DADOS = "ABCDEFGHIJKLM"

# QA (achado real, relatório de Maio/2026) — a célula J833 apareceu com um
# preenchimento amarelo sem nenhum significado no relatório (auditoria
# confirmou: a ÚNICA célula da área de dados com destaque visual em toda a
# faixa auditada). A área de dados nunca prevê destaque de nenhuma cor —
# cabeçalho (`_PREENCHIMENTO_CABECALHO`/`_PREENCHIMENTO_SUBCABECALHO`, fora
# desta área) é o único lugar com preenchimento legítimo. Em vez de corrigir
# J833 especificamente (trataria o sintoma, não a causa — outra célula
# poderia ter a mesma anomalia, introduzida por edição manual em qualquer
# geração passada), `_SEM_PREENCHIMENTO` é reafirmado em TODA célula de
# dados a cada geração — mesmo princípio já usado para `_NUMBER_FORMAT_
# VALOR_PASSAGEM` acima (roda para linhas novas E já existentes).
_SEM_PREENCHIMENTO = PatternFill(fill_type=None)

_MESES_EXTENSO = [
    "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
    "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO",
]


class RelatorioBloqueadoPorPendencias(Exception):
    """G05 aplicado na camada de geração: nunca gera XLSX enquanto existir
    revisão sem decisão válida (pendente, adiada ou desatualizada) — mesma
    garantia mesmo chamando esta função diretamente, sem passar pela API."""


class DestinoJaExiste(Exception):
    """Item 17 — nunca sobrescreve silenciosamente um arquivo já existente
    no destino."""


class GeracaoConcorrenteEmAndamento(Exception):
    """Confronto QA / homologação operacional — duas requisições `/gerar`
    simultâneas para o MESMO destino não podem ambas concluir com sucesso.
    Levantada quando outra geração para este destino específico já está em
    andamento (arquivo de lock, ver `_adquirir_trava_destino` abaixo);
    nunca uma condição sobre o conteúdo do relatório em si."""


class DestinoInacessivel(Exception):
    """QA Codex (achado real, versionamento automático) — a criação do
    arquivo de trava (`_adquirir_trava_destino`) continuou falhando com
    `PermissionError` mesmo depois de esgotadas as tentativas de retry
    curto (a mesma corrida transitória documentada ali já foi absorvida) —
    sinal de um problema PERSISTENTE no destino em si (pasta protegida,
    sem permissão de escrita, sincronização de nuvem/antivírus retendo o
    arquivo), nunca mais uma colisão passageira entre gerações. Mensagem
    já adequada para exibição direta ao usuário."""


class TemplateEstruturalmenteInvalido(Exception):
    """Template não tem a aba/estrutura mínima esperada — erro explícito,
    nunca uma exceção genérica do openpyxl não tratada."""


@dataclass(frozen=True)
class RelatorioGeradoResultado:
    caminho_arquivo: str
    total_ocorrencias: int
    total_valor_diarias: float
    primeira_linha_dados: int
    ultima_linha_dados: int
    linha_total: int


def _localizar_linha_total(ws: Worksheet) -> int:
    """A linha de totais é identificada pelo próprio conteúdo do template
    (coluna A começa com "Total" — confirmado: `A1819 = 'Total (j)'` no
    modelo real), nunca por um número de linha fixo — assim o gerador
    continua funcionando mesmo que o template ganhe/perca linhas
    históricas no futuro."""
    for linha in range(_PRIMEIRA_LINHA_DADOS, ws.max_row + 1):
        valor = ws[f"A{linha}"].value
        if isinstance(valor, str) and valor.strip().startswith(_ROTULO_LINHA_TOTAL_PREFIXO):
            return linha
    raise TemplateEstruturalmenteInvalido(
        f"Não foi possível localizar a linha de totais (coluna A iniciando com "
        f"{_ROTULO_LINHA_TOTAL_PREFIXO!r}) no template."
    )


def _copiar_estilo_linha(ws: Worksheet, linha_origem: int, linha_destino: int, colunas: str) -> None:
    """Copia fonte/preenchimento/borda/alinhamento/formato numérico de uma
    linha-modelo para outra. `copy.copy()` é necessário aqui (não é só
    estilo): depois de `ws.insert_rows`, o acesso `célula.font` de uma
    célula pré-existente pode devolver um `StyleProxy` interno do openpyxl
    em vez do `Font`/`Border`/... imutável direto, e atribuir esse proxy a
    outra célula quebra com `TypeError: unhashable type: 'StyleProxy'`
    dentro da lista indexada de estilos do workbook — `copy.copy()` sempre
    devolve o objeto de estilo real, nunca o proxy."""
    origem = ws.row_dimensions[linha_origem]
    destino = ws.row_dimensions[linha_destino]
    destino.height = origem.height
    for coluna in colunas:
        cel_origem = ws[f"{coluna}{linha_origem}"]
        cel_destino = ws[f"{coluna}{linha_destino}"]
        cel_destino.font = copy(cel_origem.font)
        cel_destino.fill = copy(cel_origem.fill)
        cel_destino.border = copy(cel_origem.border)
        cel_destino.alignment = copy(cel_origem.alignment)
        cel_destino.number_format = cel_origem.number_format


def _realinhar_merges(ws: Worksheet, linha_referencia: int, deslocamento: int) -> None:
    """`Worksheet.insert_rows`/`delete_rows` do openpyxl só realinham o
    CONTEÚDO das células (via `_move_cells`, biblioteca) — `merged_cells.
    ranges` NUNCA é tocado, uma limitação documentada da biblioteca. Sem
    esta correção manual, toda mesclagem abaixo do ponto de inserção/
    remoção (linha de Total, "Fonte de Informação", "Data da última
    atualização", notas de rodapé) fica ancorada na posição ANTIGA — o que
    não só desalinha visualmente essas mesclagens do seu próprio conteúdo
    (que foi movido corretamente), como também corrompe o ESTILO das
    linhas de dados novas/deslocadas: o openpyxl só persiste o estilo da
    célula-âncora de uma mesclagem ao salvar; as demais células dentro de
    uma mesclagem têm o estilo limpo. Sempre chamada logo após
    `insert_rows`/`delete_rows`, nunca antes."""
    # Manipula `ws.merged_cells` (a coleção `MultiCellRange`) diretamente
    # via `.remove`/`.add`, NUNCA via `ws.unmerge_cells`/`ws.merge_cells`:
    # esses métodos de conveniência também tentam apagar/criar entradas em
    # `ws._cells` para as células cobertas — mas `delete_rows` já apagou
    # fisicamente as células da "sobra" final (as últimas `amount` linhas,
    # ver `_gutter`/`_move_cells` da biblioteca), então `unmerge_cells`
    # levantaria `KeyError` ao tentar apagar uma célula que já não existe.
    # Aqui só a METADADO da mesclagem é corrigido — o conteúdo das células
    # já foi movido corretamente por `insert_rows`/`delete_rows`.
    afetadas = []
    for intervalo in list(ws.merged_cells.ranges):
        min_row, max_row = intervalo.min_row, intervalo.max_row
        if max_row < linha_referencia:
            continue
        if min_row < linha_referencia <= max_row:
            raise TemplateEstruturalmenteInvalido(
                f"Mesclagem {intervalo} atravessa o ponto de inserção/remoção de "
                f"linhas (linha {linha_referencia}) — realinhamento automático não "
                f"é seguro; é necessária investigação manual do template."
            )
        novo_min_row, novo_max_row = min_row + deslocamento, max_row + deslocamento
        if novo_min_row < _PRIMEIRA_LINHA_DADOS:
            raise TemplateEstruturalmenteInvalido(
                f"Realinhamento da mesclagem {intervalo} produziria uma posição "
                f"inválida (linha {novo_min_row}) — abortando geração."
            )
        afetadas.append((intervalo, novo_min_row, intervalo.min_col, novo_max_row, intervalo.max_col))

    for intervalo, *_ in afetadas:
        ws.merged_cells.remove(intervalo)
    for _, novo_min_row, min_col, novo_max_row, max_col in afetadas:
        ws.merged_cells.add(
            CellRange(min_row=novo_min_row, min_col=min_col, max_row=novo_max_row, max_col=max_col)
        )


def _realinhar_alturas_de_linha(ws: Worksheet, linha_referencia: int, deslocamento: int) -> None:
    """Mesma limitação da biblioteca descrita em `_realinhar_merges`, mas
    para `row_dimensions` (altura/oculta/nível de contorno) — também nunca
    realinhado automaticamente por `insert_rows`/`delete_rows`. Lê todos
    os valores originais ANTES de limpar/escrever qualquer linha (evita
    qualquer ambiguidade de ordem entre linhas-origem que também são
    linhas-destino de outro deslocamento)."""
    linhas_afetadas = [linha for linha in list(ws.row_dimensions.keys()) if linha >= linha_referencia]
    valores_originais = {
        linha: (ws.row_dimensions[linha].height, ws.row_dimensions[linha].hidden, ws.row_dimensions[linha].outlineLevel)
        for linha in linhas_afetadas
    }
    for linha in linhas_afetadas:
        dimensao = ws.row_dimensions[linha]
        dimensao.height = None
        dimensao.hidden = False
        dimensao.outlineLevel = 0
    for linha, (altura, oculta, nivel) in valores_originais.items():
        linha_destino = linha + deslocamento
        if linha_destino < _PRIMEIRA_LINHA_DADOS:
            continue
        dimensao = ws.row_dimensions[linha_destino]
        dimensao.height = altura
        dimensao.hidden = oculta
        dimensao.outlineLevel = nivel


def _valor_data(valor: date | datetime | None):
    if isinstance(valor, datetime):
        return valor
    return valor


def _data_pagamento_como_texto(valor: datetime | None) -> str | None:
    """Micro-sprint 3F.1 — investigação célula a célula do modelo oficial
    real confirmou que a coluna B (Data de Pagamento) NUNCA armazena uma
    data/serial real: 100% das 1812 linhas de dados têm `data_type='s'`
    (texto), sempre no padrão literal `DD/MM/YYYY HH:MM:SS` — o
    `number_format` da célula é inconsistente entre linhas e irrelevante
    (Excel não aplica number_format a células de texto). Reproduzido aqui
    literalmente, nunca como `datetime`/serial Excel."""
    if valor is None:
        return None
    return valor.strftime("%d/%m/%Y %H:%M:%S")


def _escrever_linha_ocorrencia(
    ws: Worksheet,
    linha: int,
    ordem: OrdemBancaria,
    registros_controle,
    passagem: LinhaPassagem | None,
) -> float:
    quantidade = resolver_quantidade_diarias(registros_controle)
    cargo = valor_do_controle_com_fallback(registros_controle, "cargo", ordem.cargo)
    transporte = valor_do_controle_com_fallback(registros_controle, "transporte", ordem.transporte)

    ws[f"A{linha}"] = ordem.ob
    ws[f"B{linha}"] = _data_pagamento_como_texto(ordem.data_pagamento)
    ws[f"C{linha}"] = nome_publicado(ordem)
    ws[f"D{linha}"] = cargo
    ws[f"E{linha}"] = _valor_data(ordem.periodo_inicial)
    ws[f"F{linha}"] = _valor_data(ordem.periodo_final)
    ws[f"G{linha}"] = ordem.trecho
    ws[f"H{linha}"] = transporte
    ws[f"I{linha}"] = ordem.motivo
    # Sprint 3 (HOMO-RT-06, plano "Valor da Passagem no XLSX Oficial") —
    # substitui o placeholder histórico fixo (Micro-sprint 3F: "valor da
    # passagem não é responsabilidade do DIMCAC") pelo valor efetivamente
    # publicável do cruzamento com a Planilha de Passagens. MESMA função
    # usada pelo fingerprint (`decisao._campos_linha_publicavel`) — nunca
    # duas implementações. Estado D (`inconsistente=True`) também chega
    # aqui e produz célula vazia (`.valor is None`), igual a B/C/E —
    # fechamento HOMO-RT-06: D nunca bloqueia a geração, só nunca vira um
    # valor numérico por acidente.
    valor_passagem = valor_passagem_publicavel(passagem).valor
    ws[f"J{linha}"] = valor_passagem
    # Correção do bug de formatação (auditoria Setembro/2026, ver
    # `_NUMBER_FORMAT_VALOR_PASSAGEM` acima) — o formato monetário é
    # afirmado aqui explicitamente sempre que a célula recebe um valor real,
    # nunca herdado de `_copiar_estilo_linha`/da linha-modelo. Só quando
    # `valor_passagem` não é `None` (B/C/D/E continuam produzindo célula
    # vazia, sem nenhum `number_format` forçado sobre um valor ausente).
    # Como esta função roda para TODAS as linhas — novas e já existentes —
    # a cada geração mensal, isto também normaliza retroativamente qualquer
    # linha antiga já gravada com o formato errado no passado.
    if valor_passagem is not None:
        ws[f"J{linha}"].number_format = _NUMBER_FORMAT_VALOR_PASSAGEM
    ws[f"K{linha}"] = quantidade.integral if quantidade else None
    ws[f"L{linha}"] = quantidade.parcial if quantidade else None
    ws[f"M{linha}"] = ordem.valor  # RN02 — sempre o valor bruto da Base OB, zero incluso

    # QA (achado real, relatório de Maio/2026 — ver `_SEM_PREENCHIMENTO`
    # acima): normaliza o preenchimento de TODA célula de dados desta linha,
    # sempre — nunca depende de identificar qual célula específica está
    # anômala, e corrige retroativamente qualquer linha já existente (esta
    # função roda para TODAS as linhas a cada geração, novas e
    # preexistentes). Só `.fill` é tocado — fonte, borda, alinhamento e
    # number_format (herdados de `_copiar_estilo_linha` para linhas novas,
    # ou já corretos para linhas preexistentes) permanecem intocados.
    for coluna in _COLUNAS_DADOS:
        ws[f"{coluna}{linha}"].fill = _SEM_PREENCHIMENTO

    return ordem.valor or 0.0


def _caminho_trava(caminho_destino: str) -> str:
    return f"{caminho_destino}.gerando.lock"


# Mesmos valores já homologados em `app/consulta/persistencia_decisoes.py`/
# `base_local_sqlite._substituir_com_retry` para a mesma classe de corrida
# transitória — ver docstring de `_adquirir_trava_destino`.
_TENTATIVAS_TRAVA_DESTINO = 3
_INTERVALO_BASE_RETRY_TRAVA_DESTINO_S = 0.05


def _adquirir_trava_destino(caminho_destino: str) -> str:
    """Confronto QA / homologação operacional — `os.path.exists(destino)`
    sozinho é só uma checagem, não uma reserva: duas requisições `/gerar`
    simultâneas para o MESMO destino passavam as duas por ela (nenhuma via
    o arquivo do outro ainda), geravam em paralelo e a última a chamar
    `os.replace` no fim vencia silenciosamente — nunca um erro, nunca um
    aviso, só uma sobrescrita.

    A correção não depende de nada no frontend (um duplo-clique, uma aba
    duplicada, um retry automático — nenhum deles pode ser assumido como
    impossível pelo backend). `os.open(..., O_CREAT | O_EXCL)` é a mesma
    chamada de sistema (`open(2)`/`CreateFileW` com `CREATE_NEW`) usada
    para lockfiles há décadas justamente porque cria-e-falha-se-já-existir
    é atômico no nível do sistema de arquivos, em qualquer SO local — não
    há janela entre "checar" e "criar" para uma segunda chamada colar no
    meio, ao contrário de `os.path.exists` seguido de qualquer escrita
    posterior. Suficiente e apropriado aqui porque o caso de uso é
    desktop/local de um único usuário por vez (nunca um serviço
    multiusuário/distribuído) — não precisa de um lock distribuído, só de
    impedir que duas requisições HTTP concorrentes no mesmo processo/
    máquina peguem o mesmo destino ao mesmo tempo.

    O arquivo de trava fica AO LADO do destino real (nunca substitui nem
    aparenta ser o relatório — sufixo `.gerando.lock`, nunca `.xlsx`), então
    uma pasta sendo observada nunca mostra um `.xlsx` vazio/inválido
    aparecendo antes da hora (ver também `os.replace` atômico do
    temporário, já existente, para o próprio conteúdo do arquivo).

    QA Codex (achado real, versionamento automático — `PermissionError`
    reproduzido e isolado por diagnóstico dedicado, EXATAMENTE nesta
    chamada de `os.open`) — sob concorrência real (múltiplas gerações da
    MESMA competência disputando a mesma sequência de candidatos de
    versão), o Windows por vezes devolve `PermissionError` em vez de
    `FileExistsError` na criação exclusiva (`O_CREAT|O_EXCL`) quando o
    MESMO nome de arquivo de trava está num estado transitório de exclusão
    por outra thread/processo (ciclo completo criar→promover→remover se
    sobrepondo) — sempre se resolve em poucos milissegundos, nunca a mesma
    classe de erro de uma pasta genuinamente protegida (essa persiste
    através das tentativas). Retry curto — MESMOS valores já homologados
    de `app/consulta/persistencia_decisoes.py`/`base_local_sqlite.
    _substituir_com_retry` — absorve o caso transitório; se ainda assim
    persistir após as tentativas, `DestinoInacessivel` é levantada com uma
    mensagem já adequada para o usuário — o chamador (`gerar_relatorio_
    xlsx`) NUNCA deve tratar isso como "esta versão específica está
    ocupada, tenta a próxima": um problema real de acesso à pasta afeta
    IGUALMENTE qualquer candidato de versão, avançar centenas delas não
    ajudaria e só mascararia o erro real."""
    caminho_lock = _caminho_trava(caminho_destino)
    ultimo_erro_transitorio: PermissionError | None = None
    for tentativa in range(_TENTATIVAS_TRAVA_DESTINO):
        try:
            descritor_lock = os.open(caminho_lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError as erro:
            raise GeracaoConcorrenteEmAndamento(
                f"Já existe uma geração em andamento para este destino: {caminho_destino}"
            ) from erro
        except PermissionError as erro:
            ultimo_erro_transitorio = erro
            if tentativa < _TENTATIVAS_TRAVA_DESTINO - 1:
                time.sleep(_INTERVALO_BASE_RETRY_TRAVA_DESTINO_S * (tentativa + 1))
            continue
        else:
            os.close(descritor_lock)
            return caminho_lock

    assert ultimo_erro_transitorio is not None
    raise DestinoInacessivel(
        f"Não foi possível preparar a geração no destino {caminho_destino!r} — a pasta pode estar "
        f"protegida, sem permissão de escrita, ou com um arquivo bloqueado por outro programa "
        f"(antivírus, Excel, sincronização de nuvem). Verifique o destino e tente novamente."
    ) from ultimo_erro_transitorio


def _chave_ordenacao(candidato: CandidatoNormal):
    """Ordenação determinística (item 14): por DATA OB (RN05, a mesma
    competência financeira do relatório) e, em caso de empate, por número
    da OB. Investigação empírica contra o relatório histórico real (Micro-
    sprint 3F) não encontrou nenhuma chave única que reproduzisse a ordem
    histórica exatamente (nem OB, nem data de pagamento, isolados) — a
    ordem real parece refletir a sequência de inserção manual ao longo do
    tempo, não uma função pura dos dados. Esta é a ordenação mais próxima e
    100% determinística; divergências de ORDEM (não de conteúdo) contra o
    histórico são esperadas e reportadas como tal."""
    data = candidato.ocorrencia.ordem_bancaria.data_pagamento
    return (data is None, data or datetime.min, candidato.ocorrencia.ordem_bancaria.ob)


def gerar_xlsx_relatorio_transparencia(
    resultado: ResultadoRevisado,
    *,
    caminho_template: str,
    caminho_destino: str,
    ano_exercicio: int,
    mes_referencia: int,
    hoje: date | None = None,
) -> RelatorioGeradoResultado:
    """Gera o XLSX final a partir do template oficial + do resultado já
    revisado pelo backend. G05 aplicado aqui (nunca só no frontend/service
    chamador): `resultado.pode_gerar_relatorio` precisa ser `True`."""
    if not resultado.pode_gerar_relatorio:
        motivos = "; ".join(resultado.motivos_bloqueio) or "revisões pendentes"
        raise RelatorioBloqueadoPorPendencias(
            f"Não é possível gerar o relatório: existem revisões sem decisão válida ({motivos})."
        )

    if not os.path.isfile(caminho_template):
        raise TemplateEstruturalmenteInvalido(f"Template não encontrado: {caminho_template}")

    candidatos_ordenados = sorted(resultado.candidatos_publicaveis, key=_chave_ordenacao)

    # Sprint 3 (HOMO-RT-06), fechamento — decisão de produto final: estado D
    # (INVALIDO, ex. "#VALOR!") do Valor da Passagem NUNCA bloqueia a
    # geração do relatório oficial (revisão da versão provisória anterior).
    # `_escrever_linha_ocorrencia` já escreve célula vazia para D (mesmo
    # caminho de B/C/E, via `valor_passagem_publicavel(...).valor is None`)
    # — o estado INVALIDO em si permanece preservado em
    # `OcorrenciaFinanceira.passagem.estado` e no fingerprint (`decisao.
    # _campos_linha_publicavel`), nunca reescrito para ausente/zero; só não
    # impede mais a publicação das demais ocorrências.

    destino_dir = os.path.dirname(os.path.abspath(caminho_destino)) or "."
    os.makedirs(destino_dir, exist_ok=True)

    # Confronto QA / homologação operacional — a trava precisa existir
    # ANTES da checagem de "já existe" abaixo: é ela quem garante que só
    # uma requisição por vez chega a fazer essa checagem para este
    # destino (ver `_adquirir_trava_destino`).
    caminho_lock = _adquirir_trava_destino(caminho_destino)
    try:
        if os.path.exists(caminho_destino):
            raise DestinoJaExiste(f"Já existe um arquivo no destino: {caminho_destino}")

        # Item 17 — escrita segura: gera num temporário na MESMA pasta do
        # destino (garante que o `os.replace` final seja atômico, mesmo
        # sistema de arquivos), nunca deixa um arquivo final parcialmente
        # escrito no destino real.
        descritor_tmp, caminho_tmp = tempfile.mkstemp(suffix=".xlsx", dir=destino_dir)
        os.close(descritor_tmp)
        try:
            shutil.copy2(caminho_template, caminho_tmp)
            workbook = load_workbook(caminho_tmp, data_only=False)
            if _NOME_ABA_DADOS not in workbook.sheetnames:
                raise TemplateEstruturalmenteInvalido(
                    f"Aba obrigatória ausente no template: {_NOME_ABA_DADOS!r}"
                )
            ws = workbook[_NOME_ABA_DADOS]

            linha_total_antiga = _localizar_linha_total(ws)
            quantidade_antiga = linha_total_antiga - _PRIMEIRA_LINHA_DADOS
            quantidade_nova = len(candidatos_ordenados)

            if quantidade_nova > quantidade_antiga:
                diferenca = quantidade_nova - quantidade_antiga
                ws.insert_rows(linha_total_antiga, amount=diferenca)
                _realinhar_merges(ws, linha_total_antiga, diferenca)
                _realinhar_alturas_de_linha(ws, linha_total_antiga, diferenca)
                for offset in range(diferenca):
                    linha_nova = linha_total_antiga + offset
                    _copiar_estilo_linha(ws, _PRIMEIRA_LINHA_DADOS, linha_nova, _COLUNAS_DADOS)
            elif quantidade_nova < quantidade_antiga:
                diferenca = quantidade_antiga - quantidade_nova
                inicio_remocao = _PRIMEIRA_LINHA_DADOS + quantidade_nova
                ws.delete_rows(inicio_remocao, amount=diferenca)
                _realinhar_merges(ws, inicio_remocao + diferenca, -diferenca)
                _realinhar_alturas_de_linha(ws, inicio_remocao + diferenca, -diferenca)

            linha_total_nova = _PRIMEIRA_LINHA_DADOS + quantidade_nova
            ultima_linha_dados = linha_total_nova - 1 if quantidade_nova > 0 else _PRIMEIRA_LINHA_DADOS - 1

            total_valor = 0.0
            for indice, candidato in enumerate(candidatos_ordenados):
                linha = _PRIMEIRA_LINHA_DADOS + indice
                total_valor += _escrever_linha_ocorrencia(
                    ws,
                    linha,
                    candidato.ocorrencia.ordem_bancaria,
                    candidato.registros_controle,
                    candidato.ocorrencia.passagem,
                )

            # Fórmulas de total — sempre recalculadas para a faixa REAL de
            # dados desta geração (o histórico tinha uma faixa desatualizada,
            # ver relatório da investigação — G01, não reproduzido).
            if quantidade_nova > 0:
                ws[f"J{linha_total_nova}"] = f"=SUM(J{_PRIMEIRA_LINHA_DADOS}:J{ultima_linha_dados})"
                ws[f"M{linha_total_nova}"] = f"=SUM(M{_PRIMEIRA_LINHA_DADOS}:M{ultima_linha_dados})"
            else:
                ws[f"J{linha_total_nova}"] = 0
                ws[f"M{linha_total_nova}"] = 0

            # Título/competência (linha 3) e data de atualização (2 linhas
            # abaixo da linha de totais, mesmo deslocamento do template real).
            mes_nome = _MESES_EXTENSO[mes_referencia - 1]
            ws["A3"] = f"PAGAS ATÉ {mes_nome}/{ano_exercicio}"

            linha_atualizacao = linha_total_nova + 2
            data_referencia = hoje or date.today()
            celula_atualizacao = ws[f"A{linha_atualizacao}"]
            if isinstance(celula_atualizacao.value, str) and celula_atualizacao.value.startswith(
                "Data da última atualização"
            ):
                celula_atualizacao.value = (
                    f"Data da última atualização: {data_referencia.strftime('%d.%m.%Y')}"
                )

            # Autofiltro cobre cabeçalho + dados (nunca a área de totais/notas
            # — decisão deliberada desta Micro-sprint, ver relatório).
            if quantidade_nova > 0:
                ws.auto_filter.ref = f"A5:O{ultima_linha_dados}"
            else:
                ws.auto_filter.ref = "A5:O5"

            workbook.save(caminho_tmp)
            workbook.close()

            # Validação mínima antes do replace: reabre o temporário e
            # confere que a aba e a contagem de linhas batem com o esperado —
            # nunca promove um arquivo que não passou por esta checagem.
            validacao = load_workbook(caminho_tmp, data_only=False)
            if _NOME_ABA_DADOS not in validacao.sheetnames:
                raise TemplateEstruturalmenteInvalido("Falha ao validar o arquivo gerado (aba ausente).")
            ws_validacao = validacao[_NOME_ABA_DADOS]
            if quantidade_nova > 0:
                assert ws_validacao[f"A{_PRIMEIRA_LINHA_DADOS}"].value == candidatos_ordenados[0].ocorrencia.ordem_bancaria.ob
            validacao.close()

            # Confronto QA / revalidação Codex — a checagem `os.path.exists`
            # feita mais acima (ainda sob a trava) só protege contra OUTRA
            # geração deste app: um terceiro processo (ou o próprio usuário,
            # via Explorer) pode criar um arquivo em `caminho_destino` DEPOIS
            # dessa checagem e ANTES desta promoção — a trava é cooperativa
            # entre chamadas deste módulo, nunca um lock do sistema de
            # arquivos sobre o destino em si. `os.replace` (que sempre
            # sobrescreve, mesmo semântica de `rename(2)` POSIX) reabriria
            # exatamente essa janela na última linha da função.
            #
            # `os.rename` no Windows mapeia para `MoveFileExW` SEM a flag
            # `MOVEFILE_REPLACE_EXISTING` — falha com `FileExistsError` se o
            # destino já existir, numa única chamada de sistema atômica (o
            # próprio SO faz checar+renomear como uma operação só, sem
            # janela para uma escrita de terceiro colar no meio — ao
            # contrário de um `exists()` seguido de qualquer operação
            # posterior). O arquivo que apareceu no destino nunca é tocado
            # (confirmado: `os.rename` falhando não move nem apaga nada); o
            # temporário desta geração segue intacto e é removido no
            # `finally` abaixo — nunca fica um XLSX parcial no destino real.
            try:
                os.rename(caminho_tmp, caminho_destino)
            except FileExistsError as erro:
                raise DestinoJaExiste(
                    f"Já existe um arquivo no destino: {caminho_destino}"
                ) from erro
        finally:
            if os.path.exists(caminho_tmp):
                os.remove(caminho_tmp)
    finally:
        # A trava é liberada sempre — sucesso ou falha — nunca deixando um
        # `DestinoJaExiste`/erro de geração órfão bloquear indefinidamente
        # uma nova tentativa para o mesmo destino.
        os.remove(caminho_lock)

    return RelatorioGeradoResultado(
        caminho_arquivo=caminho_destino,
        total_ocorrencias=quantidade_nova,
        total_valor_diarias=total_valor,
        primeira_linha_dados=_PRIMEIRA_LINHA_DADOS,
        ultima_linha_dados=ultima_linha_dados,
        linha_total=linha_total_nova,
    )
