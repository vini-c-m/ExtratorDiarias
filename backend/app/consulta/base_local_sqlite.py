"""Base Local de Consulta em SQLite (Sprint 17.3) — a segunda
materialização da Base Local, ao lado da cópia `.xlsx` já existente
(`app/consulta/base_local.py`, Sprint 13). Mesma pasta, mesmo diretório de
dados por instalação: as duas são representações de *uma* Base Local, não
dados de módulos diferentes (decisão 9, `docs/central-analitica/README.md
§ 0.1` — o único ponto de compartilhamento entre a Central Analítica e a
Central de Conferência é `app/consulta/`).

Reconstrução sempre total (decisão 4): este módulo nunca faz `UPDATE`
incremental, sempre recria o banco do zero num arquivo temporário e troca
atomicamente com `os.replace` — mesma técnica já homologada em
`base_local.py`/`configuracao_app.py`. Uma falha em qualquer ponto antes do
`commit()`+`close()`+`os.replace` deixa o `.tmp` órfão e nunca toca o
destino — a última Base Local válida nunca é apagada/sobrescrita por uma
reconstrução malsucedida (decisão 5)."""

import os
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.consulta import base_local
from app.consulta.repositorio_diarias import ResultadoDiariasCompletas
from app.domain.analitica.categoria_funcional import classificar_categoria_funcional
from app.domain.entities.aviso_linha_ignorada import AvisoLinhaIgnorada
from app.domain.entities.servidor import Servidor
from app.infra.db.esquema import VERSAO_ESQUEMA, criar_esquema
from app.planilha.escritor_planilha import NOME_ABA

NOME_ARQUIVO_BASE_LOCAL_SQLITE = "base_local.sqlite"


def caminho_base_local_sqlite(*, diretorio_base: Path | str | None = None) -> Path:
    """Mesma pasta de `base_local.caminho_base_local()` — reaproveita a
    resolução de diretório já existente em vez de duplicá-la uma terceira
    vez (já duplicada entre `base_local.py` e `configuracao_app.py`)."""
    return base_local.caminho_base_local(diretorio_base=diretorio_base).parent / NOME_ARQUIVO_BASE_LOCAL_SQLITE


class SinalFrescorIndeterminado(Exception):
    """Micro-sprint 18.7.1 — nem o caminho rápido (`read_only=True`) nem o
    fallback completo (`read_only=False`) conseguiram determinar um
    inteiro válido para o número de linhas de `Controle de Diárias`.
    Levantada por `capturar_sinal_frescor` antes de devolver qualquer
    coisa — nunca silenciosamente substituída por `0` ou `None`, para não
    corromper `sincronizacoes.linhas_controle_diarias_fonte` (`NOT NULL`)
    com um valor que não representa a quantidade real de linhas. Cenário
    extremamente raro em produção (mesmo planilhas sem a tag `<dimension>`
    do XML — ex. exportadas pelo Google Sheets — resolvem no fallback,
    que escaneia as células de verdade em vez de confiar em metadado);
    existe como rede de segurança explícita, não como caminho esperado."""


@dataclass(frozen=True)
class SinalFrescor:
    """Sinal de frescor de **conteúdo** — deliberadamente separado da
    validação **estrutural** já existente (`FingerprintEstrutural`, Sprint
    15, `app/infra/config/configuracao_app.py`, hash de cabeçalho/abas —
    decisão 3). Barato de calcular: tamanho + data de modificação do
    arquivo, mais a contagem de linhas de `Controle de Diárias` (lida com
    `read_only=True`, sem parse completo)."""

    tamanho_bytes: int
    mtime: float
    linhas_controle_diarias: int


def capturar_sinal_frescor(caminho_planilha: str) -> SinalFrescor:
    """Micro-sprint 18.7.1 — `read_only=True` (caminho rápido, sem parse
    completo) depende da tag `<dimension>` do XML da aba para calcular
    `max_row` sem escanear célula por célula; planilhas que chegaram a
    passar pelo Google Sheets (confirmado numa planilha real desta
    instalação, via a relação `.../customschemas.google.com/...` em
    `xl/_rels/workbook.xml.rels`) podem não gravar essa tag, e `max_row`
    volta `None`. Sem este fallback, `None` atravessava `SinalFrescor`
    (dataclass sem validação de tipo em runtime) até o `INSERT INTO
    sincronizacoes` de `reconstruir()`, violando a coluna `NOT NULL`
    `linhas_controle_diarias_fonte` — bug confirmado (Sprint 18.7,
    auditoria do Gatilho A) que impedia `base_local.sqlite` de ser
    publicado para qualquer planilha nessas condições, sempre da mesma
    forma (determinístico, não uma condição de corrida).

    Fallback só quando `max_row` não é um `int` (nunca corre no caminho
    normal): reabre a mesma planilha com `read_only=False` — nesse modo o
    openpyxl calcula `max_row` varrendo as células de fato, sem depender
    do metadado `<dimension>` — e usa esse valor. Os dois `Workbook`s são
    sempre fechados (um por `try/finally` próprio) antes de seguir. Se
    mesmo assim `max_row` não vier um `int` (planilha genuinamente
    indeterminável), levanta `SinalFrescorIndeterminado` — nunca devolve
    `None` nem inventa `0` (§4 do pedido)."""
    estatisticas = os.stat(caminho_planilha)

    workbook = load_workbook(caminho_planilha, read_only=True, data_only=False)
    try:
        linhas = workbook[NOME_ABA].max_row
    finally:
        workbook.close()

    if not isinstance(linhas, int):
        workbook_completo = load_workbook(caminho_planilha, read_only=False, data_only=False)
        try:
            linhas = workbook_completo[NOME_ABA].max_row
        finally:
            workbook_completo.close()

    if not isinstance(linhas, int):
        raise SinalFrescorIndeterminado(
            f"Não foi possível determinar o número de linhas da aba '{NOME_ABA}' em "
            f"{caminho_planilha!r} — max_row indeterminado mesmo fora do modo read_only."
        )

    return SinalFrescor(
        tamanho_bytes=estatisticas.st_size,
        mtime=estatisticas.st_mtime,
        linhas_controle_diarias=linhas,
    )


def sinal_mudou(atual: SinalFrescor, registrado: SinalFrescor | None) -> bool:
    """Função pura, sem I/O. `None` (nunca sincronizado) sempre `True`.
    Fundação de dados para a verificação automática ao entrar no módulo
    (decisão 1) e o botão manual (decisão 2) — nenhum dos dois é disparado
    por código nesta Sprint; ambos ficam para a Sprint 17.4 (API/UI), que
    chamará esta mesma função pura."""
    if registrado is None:
        return True
    return atual != registrado


def ultimo_sinal_frescor(caminho_sqlite: Path) -> SinalFrescor | None:
    """Lê o sinal de frescor registrado na sincronização mais recente —
    `None` se a Base Local em SQLite ainda não existe (nunca sincronizado).
    Não chamada por nenhum código nesta Sprint (sem UI/API ainda); existe
    para a Sprint 17.4 comparar contra `capturar_sinal_frescor` do estado
    atual da Planilha Oficial via `sinal_mudou`."""
    if not caminho_sqlite.is_file():
        return None
    conexao = sqlite3.connect(caminho_sqlite)
    try:
        linha = conexao.execute(
            """
            SELECT tamanho_bytes_fonte, mtime_fonte, linhas_controle_diarias_fonte
            FROM sincronizacoes
            ORDER BY id DESC
            LIMIT 1
            """
        ).fetchone()
    finally:
        conexao.close()
    if linha is None:
        return None
    tamanho_bytes, mtime, linhas_controle_diarias = linha
    return SinalFrescor(tamanho_bytes=tamanho_bytes, mtime=mtime, linhas_controle_diarias=linhas_controle_diarias)


def esquema_desatualizado(caminho_sqlite: Path) -> bool:
    """Sprint 19.9 (achado da auditoria — Electron real + Planilha Oficial
    real): `PRAGMA user_version` do `.sqlite` existente comparado a
    `esquema.VERSAO_ESQUEMA` — eixo de frescor **independente** do sinal de
    conteúdo (`sinal_mudou`/`SinalFrescor`, que só compara a Planilha
    Oficial) e da validação estrutural (`FingerprintEstrutural`, decisão 3,
    que só valida a planilha em si). Sem este terceiro eixo, um `.sqlite`
    de uma instalação anterior a uma mudança de schema (ex.: coluna `dh`,
    Sprint 19.5.1) nunca era detectado como desatualizado quando a
    Planilha Oficial em si não mudava — `GET /sincronizacao` respondia
    "atualizado" enquanto `GET /diarias`/`/indicadores`/etc. quebravam com
    `sqlite3.OperationalError` (coluna ausente). `False` quando o arquivo
    ainda não existe — "nunca sincronizado" é um estado próprio
    (`ultimo_sinal_frescor` devolvendo `None`), não uma variante de
    schema desatualizado."""
    if not caminho_sqlite.is_file():
        return False
    conexao = sqlite3.connect(caminho_sqlite)
    try:
        (versao_gravada,) = conexao.execute("PRAGMA user_version").fetchone()
    finally:
        conexao.close()
    return versao_gravada != VERSAO_ESQUEMA


@dataclass(frozen=True)
class ResumoSincronizacao:
    sincronizado_em: str
    origem: str
    linhas_diarias: int
    linhas_servidores_lidos: int
    linhas_ignoradas: int
    status: str
    avisos: list[AvisoLinhaIgnorada] = field(default_factory=list)


def reconstruir(
    resultado_diarias: ResultadoDiariasCompletas,
    servidores: dict[str, Servidor],
    sinal_frescor: SinalFrescor,
    *,
    origem: str,
    caminho_destino: Path,
    versao_app: str = "",
) -> ResumoSincronizacao:
    """Reconstrução total da Base Local em SQLite. `versao_app` (Sprint
    2026-08-06, achado real — build empacotado desatualizado) — a versão do
    Electron que disparou esta sincronização (`EXTRATOR_DIARIAS_APP_VERSION`,
    resolvida em `base_local.sincronizar`); `""` quando ausente (ex.: backend
    rodando standalone em desenvolvimento/testes, sem Electron) — nunca
    inventada, só o que a variável de ambiente realmente trouxe. `diarias` é
    denormalizada — os campos de servidor vêm inline em cada linha (via
    `servidores.get(cpf)`), sem uma tabela `servidores` separada: a
    reconstrução é sempre total, então não existe o cenário que
    justificaria normalizar (nenhuma "propagação" futura de mudança), e
    nenhum caso de uso aprovado precisa de um servidor sem diária. Um CPF
    de diária ausente em `servidores` (situação real já documentada, RN-01)
    grava os campos `servidor_*` como `NULL`, sem gerar aviso — não é dado
    malformado, é uma condição conhecida.

    Journal mode padrão do driver (`DELETE`), nunca `WAL` — WAL deixaria
    sidecars `-wal`/`-shm` que quebrariam a garantia "um arquivo, uma troca
    atômica" que sustenta a decisão 5. Sem `PRAGMA foreign_keys` — nunca
    habilitado (ver docstring do módulo `esquema.py`)."""
    caminho_destino.parent.mkdir(parents=True, exist_ok=True)
    caminho_tmp = caminho_destino.with_suffix(caminho_destino.suffix + ".tmp")
    if caminho_tmp.exists():
        caminho_tmp.unlink()

    conexao = sqlite3.connect(caminho_tmp)
    try:
        criar_esquema(conexao)

        for registro in resultado_diarias.diarias.values():
            servidor = servidores.get(registro.cpf)
            conexao.execute(
                """
                INSERT INTO diarias (
                    dh, sei, doc_origem, cpf, periodo_inicial, periodo_final,
                    cidade_origem, cidade_destino, motivo, transporte,
                    valor_total_diarias, linha_planilha,
                    servidor_nome, servidor_cargo, servidor_previdencia,
                    servidor_matricula, servidor_situacao, categoria_funcional
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    registro.dh,
                    registro.sei,
                    registro.doc_origem,
                    registro.cpf,
                    registro.periodo_inicial.isoformat(),
                    registro.periodo_final.isoformat(),
                    registro.cidade_origem,
                    registro.cidade_destino,
                    registro.motivo,
                    registro.transporte,
                    registro.valor_total_diarias,
                    registro.linha,
                    servidor.nome if servidor else None,
                    servidor.cargo if servidor else None,
                    servidor.previdencia if servidor else None,
                    servidor.matricula if servidor else None,
                    servidor.situacao if servidor else None,
                    classificar_categoria_funcional(servidor.cargo if servidor else None),
                ),
            )

        sincronizado_em = datetime.now().isoformat(timespec="seconds")
        cursor = conexao.execute(
            """
            INSERT INTO sincronizacoes (
                sincronizado_em, origem, linhas_diarias, linhas_servidores_lidos,
                linhas_ignoradas, status, tamanho_bytes_fonte, mtime_fonte,
                linhas_controle_diarias_fonte, versao_app
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                sincronizado_em,
                origem,
                len(resultado_diarias.diarias),
                len(servidores),
                resultado_diarias.linhas_ignoradas,
                "sucesso",
                sinal_frescor.tamanho_bytes,
                sinal_frescor.mtime,
                sinal_frescor.linhas_controle_diarias,
                versao_app,
            ),
        )
        sincronizacao_id = cursor.lastrowid

        for aviso in resultado_diarias.avisos:
            conexao.execute(
                """
                INSERT INTO avisos_sincronizacao (sincronizacao_id, aba, linha, motivo)
                VALUES (?, ?, ?, ?)
                """,
                (sincronizacao_id, aviso.aba, aviso.linha, aviso.motivo),
            )

        conexao.commit()
    finally:
        conexao.close()

    os.replace(caminho_tmp, caminho_destino)

    return ResumoSincronizacao(
        sincronizado_em=sincronizado_em,
        origem=origem,
        linhas_diarias=len(resultado_diarias.diarias),
        linhas_servidores_lidos=len(servidores),
        linhas_ignoradas=resultado_diarias.linhas_ignoradas,
        status="sucesso",
        avisos=resultado_diarias.avisos,
    )
