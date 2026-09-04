"""Plano de Migração E-Fisco, Sprint A — leitura isolada da Base E-Fisco
(export financeiro bruto do MPPE: contratos, diárias, suprimentos etc.,
TODAS as despesas — não só diárias). Só recebe/valida/lê/normaliza; nunca
decide se uma linha é diária (isso é `app/domain/relatorio_transparencia/
filtro_diaria_efisco.py`) nem produz `OrdemBancaria` (isso é a Sprint B,
ainda não implementada). Esta Sprint não conecta nada ao Relatório de
Diárias — Portal da Transparência; a Base OB (ClickView) continua sendo a
única fonte usada em produção.

Mesma FILOSOFIA arquitetural de `repositorio_ordens_bancarias.py` (fonte
externa, posição de coluna e nome de aba sem garantia de estabilidade),
mas SEM NENHUMA dependência de código dele — Micro-sprint de fechamento da
Sprint A (auditoria de desacoplamento): a Base OB/ClickView está a caminho
da aposentadoria (Plano de Migração E-Fisco) e este leitor não pode
depender de código que será removido.

Auditoria dos utilitários que antes eram importados de
`repositorio_ordens_bancarias.py` (`_chave_cabecalho`, `_indice_colunas`,
`_campo`, `_texto_ou_none`, `_data_pagamento_como_datetime`): todos são,
na prática, normalização GENÉRICA de planilha (resolução de coluna por
nome de cabeçalho, limpeza de texto, parsing de data/hora em texto) — sem
NENHUMA regra de negócio específica do ClickView embutida. Ainda assim,
optei por reimplementá-los aqui, localmente, em vez de extraí-los para um
módulo neutro compartilhado: é exatamente o padrão JÁ em uso neste projeto
entre leitores-irmãos de fonte externa — `repositorio_planilha_passagens.py`
já tem sua PRÓPRIA cópia de `_chave_cabecalho` (docstring de lá: "Mesma
normalização de `repositorio_ordens_bancarias._chave_cabecalho`") em vez de
importar da Base OB, pela mesma razão de isolamento arquitetural. Criar um
módulo neutro só para isso seria refatoração ampla por estética (visando um
único consumidor extra, `repositorio_efisco.py`) sobre um código
(`repositorio_ordens_bancarias.py`) que o próprio plano de migração já
prevê remover — investir em compartilhamento de infraestrutura com código
que será apagado não compensa o risco de tocar num leitor já homologado.

Achado real da auditoria (arquivo real de referência, `OB2026.xlsx`, 17.461
linhas): `Número do DH` vem vazio em despesas administrativas sem Documento
Hábil (ex.: transferências entre contas da própria UG) — ao contrário da
Base OB (ClickView), que já chega pré-filtrada só para diárias, aqui isso é
esperado e não descarta a linha; só `Ordem Bancária` ausente descarta."""

import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.consulta.repositorio_diarias import _valor_como_numero  # noqa: PLC2701 — repositorio_diarias.py não é ClickView (lê a Planilha Oficial, controlada por este projeto) e já é reaproveitado por múltiplos leitores-irmãos (repositorio_ordens_bancarias.py, repositorio_controle_dimcac.py) — não está sujeito à aposentadoria da Base OB
from app.domain.entities.aviso_linha_ignorada import AvisoLinhaIgnorada
from app.domain.entities.linha_efisco import LinhaEfisco

__all__ = [
    "RUBRICAS_OBRIGATORIAS_EFISCO",
    "EfiscoEstruturalmenteInvalido",
    "ResultadoLeituraEfisco",
    "abrir_base_efisco",
    "linhas_efisco",
    "validar_estrutura",
]

# Rótulos exigidos no cabeçalho (linha 1) da Base E-Fisco, na forma EXATA
# como aparecem na fonte real (`OB2026.xlsx`) — comparação sempre via
# `_chave_cabecalho` (normalizada), nunca por igualdade literal. Só os
# campos necessários à Sprint A (leitura + filtro de diária); campos usados
# só pela Sprint B (ex. Empenho, Nome Fantasia) não entram aqui — exigir
# coluna que ninguém lê ainda rejeitaria arquivos válidos à toa.
RUBRICAS_OBRIGATORIAS_EFISCO = (
    "Ordem Bancária",
    "Número do DH",
    "Situação",
    "Código da Natureza de Despesa",
    "Número do documento do Credor",
    "Razão Social do Credor",
    "Data da OB",
    "Data Pagto./Devol.",
    "Valor Bruto",
    "Observação",
)


def _chave_cabecalho(texto: object) -> str:
    """Normaliza um rótulo de cabeçalho para comparação: remove todo espaço
    em branco e converte para maiúsculas. Deliberadamente não remove
    acentos — os rótulos reais e os exigidos usam a mesma acentuação, só a
    presença/ausência de espaço poderia divergir entre eles.

    Implementação local e independente (ver docstring do módulo) — mesma
    normalização já usada por `repositorio_ordens_bancarias._chave_
    cabecalho`/`repositorio_planilha_passagens._chave_cabecalho`, cada
    leitor com sua própria cópia por isolamento arquitetural, nunca
    compartilhada entre fontes distintas."""
    return re.sub(r"\s+", "", str(texto or "")).upper()


def _indice_colunas(planilha: Worksheet) -> dict[str, int]:
    """Cabeçalho (linha 1) → posição (0-based) na tupla de valores que
    `iter_rows(values_only=True)` devolve por linha. Resolvido uma única
    vez por leitura, nunca por linha de dado. Implementação local (ver
    docstring do módulo)."""
    cabecalho = next(planilha.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {_chave_cabecalho(valor): posicao for posicao, valor in enumerate(cabecalho) if valor}


def _campo(valores: tuple, indice: dict[str, int], rotulo: str) -> object:
    posicao = indice.get(_chave_cabecalho(rotulo))
    if posicao is None or posicao >= len(valores):
        return None
    return valores[posicao]


def _texto_ou_none(valor: object) -> str | None:
    texto = str(valor or "").strip()
    return texto or None


class EfiscoEstruturalmenteInvalido(Exception):
    """Mesma filosofia de `OrdensBancariasEstruturalmenteInvalida`
    (`repositorio_ordens_bancarias.py`): erro estrutural que impede leitura
    confiável — arquivo ausente, extensão errada, workbook corrompido/
    bloqueado, ou colunas obrigatórias ausentes na aba lida. Nunca é uma
    decisão de regra de negócio (filtro de diária é responsabilidade de
    `filtro_diaria_efisco.py`) — só sinaliza que a Base E-Fisco fornecida
    não tem o mínimo necessário para ser lida com confiança. Mensagem já
    adequada para exibição direta ao usuário."""


@dataclass(frozen=True)
class ResultadoLeituraEfisco:
    """Mesma filosofia de `ResultadoLeituraOrdensBancarias`: `avisos` é um
    superconjunto de `linhas_ignoradas` — hoje (Sprint A) a única condição
    que descarta uma linha por completo é `Ordem Bancária` ausente; nenhum
    outro campo malformado descarta a linha, só fica `None` (representado
    tal como está, sem inventar aviso adicional nesta Sprint — a Sprint B,
    ao cruzar com o Controle DIMCAC e produzir `OrdemBancaria`, é quem
    decide o que vira aviso de negócio)."""

    linhas: list[LinhaEfisco] = field(default_factory=list)
    avisos: list[AvisoLinhaIgnorada] = field(default_factory=list)
    linhas_ignoradas: int = 0


def _selecionar_aba(workbook: Workbook) -> Worksheet:
    """A Base E-Fisco real de referência tem uma única aba ('OB') — mesmo
    critério de `repositorio_ordens_bancarias._selecionar_aba`: não é seguro
    exigir esse nome literal como estrutura obrigatória. Se houver só uma
    aba, usa-a; se houver mais de uma, usa a primeira cujo cabeçalho
    satisfaça `RUBRICAS_OBRIGATORIAS_EFISCO`."""
    abas = workbook.worksheets
    if len(abas) == 1:
        return abas[0]
    for aba in abas:
        try:
            validar_estrutura(aba)
        except EfiscoEstruturalmenteInvalido:
            continue
        return aba
    raise EfiscoEstruturalmenteInvalido(
        "Nenhuma aba do arquivo tem as colunas mínimas exigidas da Base E-Fisco."
    )


def validar_estrutura(planilha: Worksheet) -> None:
    """Confere só que as colunas de `RUBRICAS_OBRIGATORIAS_EFISCO` existem
    no cabeçalho (linha 1) — comparação por nome normalizado, nunca por
    posição (mesmo raciocínio da Base OB: fonte de sistema externo, ordem
    de coluna sem garantia de estabilidade)."""
    cabecalho = next(planilha.iter_rows(min_row=1, max_row=1, values_only=True), ())
    presentes = {_chave_cabecalho(valor) for valor in cabecalho if valor}
    faltantes = [
        rotulo for rotulo in RUBRICAS_OBRIGATORIAS_EFISCO if _chave_cabecalho(rotulo) not in presentes
    ]
    if faltantes:
        raise EfiscoEstruturalmenteInvalido(
            "Base E-Fisco: colunas obrigatórias ausentes no cabeçalho — " + ", ".join(faltantes)
        )


def abrir_base_efisco(caminho: str) -> Worksheet:
    """Abre e valida um arquivo `.xlsx` de Base E-Fisco, devolvendo a
    `Worksheet` pronta para `linhas_efisco`. Mesma ordem de checagem de
    `abrir_base_ordens_bancarias` (existência → extensão → abertura →
    estrutura).

    `read_only=True` (leitura sequencial — a fonte real tem 17.461 linhas,
    todas as despesas do MPPE, não só diárias) e `data_only=True` (valor já
    calculado da célula, nunca o texto de uma fórmula) — mesmos parâmetros
    da Base OB, mesma justificativa (export de sistema externo)."""
    if not os.path.isfile(caminho):
        raise EfiscoEstruturalmenteInvalido(f"Arquivo não encontrado: {caminho}")
    if not caminho.lower().endswith(".xlsx"):
        raise EfiscoEstruturalmenteInvalido("Formato não suportado — selecione um arquivo .xlsx")
    try:
        workbook = load_workbook(caminho, read_only=True, data_only=True)
    except Exception as erro:  # abrangente de propósito — qualquer falha de
        # leitura (arquivo corrompido, senha, permissão) vira o mesmo erro
        # estrutural explícito, nunca uma exceção genérica não tratada.
        raise EfiscoEstruturalmenteInvalido(
            f"Não foi possível abrir a Base E-Fisco — arquivo corrompido ou bloqueado ({erro})"
        ) from erro

    planilha = _selecionar_aba(workbook)
    validar_estrutura(planilha)
    return planilha


_PADRAO_DATA = re.compile(r"^(\d{2})/(\d{2})/(\d{4})$")
_PADRAO_DATA_HORA = re.compile(r"^(\d{2})/(\d{2})/(\d{4})[ T](\d{2}):(\d{2}):(\d{2})$")


def _data_ob_como_date(valor: object) -> date | None:
    """Coluna DATA DA OB, na fonte real, vem como texto `'dd/mm/aaaa'`
    (sem horário — ao contrário de DATA PAGTO./DEVOL., abaixo). Nunca
    propaga exceção: texto fora desse formato (ou vazio) vira `None`."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    if not texto:
        return None
    correspondencia = _PADRAO_DATA.match(texto)
    if not correspondencia:
        return None
    dia, mes, ano = (int(grupo) for grupo in correspondencia.groups())
    try:
        return date(ano, mes, dia)
    except ValueError:
        return None


def _data_hora_como_datetime(valor: object) -> datetime | None:
    """Texto no formato `'dd/mm/aaaa HH:MM:SS'` (espaço ou `'T'` como
    separador) vira `datetime`; qualquer outro formato, ou célula vazia
    (achado real: uma OB `Gerada`/ainda não liquidada não tem DATA PAGTO./
    DEVOL. preenchida — condição esperada, nunca um erro), vira `None`.

    Implementação local e independente (ver docstring do módulo) — o
    mesmo formato de texto é usado tanto pela coluna DATA OB da Base OB
    (ClickView) quanto por DATA PAGTO./DEVOL. do E-Fisco, mas isso é
    coincidência de dado real entre duas fontes independentes, nunca uma
    dependência entre os dois leitores: se uma das duas fontes mudar de
    formato amanhã, a outra não pode quebrar."""
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)
    texto = str(valor or "").strip()
    if not texto:
        return None
    correspondencia = _PADRAO_DATA_HORA.match(texto)
    if not correspondencia:
        return None
    dia, mes, ano, hora, minuto, segundo = (int(grupo) for grupo in correspondencia.groups())
    try:
        return datetime(ano, mes, dia, hora, minuto, segundo)
    except ValueError:
        return None


def linhas_efisco(planilha: Worksheet) -> ResultadoLeituraEfisco:
    """Lê toda a Base E-Fisco já aberta (recebe a `Worksheet`, nunca o
    caminho — mesmo padrão de `ordens_bancarias_por_dh`), devolvendo uma
    `LinhaEfisco` por linha física, na ordem em que aparecem na planilha.

    A única condição que descarta uma linha por completo é `Ordem Bancária`
    ausente (a chave sem a qual a linha não pode ser posicionada em lugar
    nenhum, mesmo critério de `ordens_bancarias_por_dh` para DH/OB) —
    `Número do DH` ausente é condição esperada nesta fonte (despesas sem
    Documento Hábil, ex. transferência entre contas da própria UG) e NUNCA
    descarta a linha nesta Sprint (a decisão sobre o que fazer com DH
    ausente é da Sprint B, ao cruzar com o Controle DIMCAC)."""
    indice = _indice_colunas(planilha)
    aba = planilha.title or ""

    linhas: list[LinhaEfisco] = []
    avisos: list[AvisoLinhaIgnorada] = []
    linhas_ignoradas = 0

    for numero_linha, valores in enumerate(planilha.iter_rows(min_row=2, values_only=True), start=2):
        ob = _texto_ou_none(_campo(valores, indice, "Ordem Bancária"))
        if ob is None:
            avisos.append(AvisoLinhaIgnorada(aba=aba, linha=numero_linha, motivo="Ordem Bancária ausente"))
            linhas_ignoradas += 1
            continue

        linhas.append(
            LinhaEfisco(
                linha=numero_linha,
                ob=ob,
                dh=_texto_ou_none(_campo(valores, indice, "Número do DH")),
                situacao=_texto_ou_none(_campo(valores, indice, "Situação")),
                natureza=_texto_ou_none(_campo(valores, indice, "Código da Natureza de Despesa")),
                documento_credor=_texto_ou_none(_campo(valores, indice, "Número do documento do Credor")),
                credor=_texto_ou_none(_campo(valores, indice, "Razão Social do Credor")),
                data_ob=_data_ob_como_date(_campo(valores, indice, "Data da OB")),
                data_pagamento_devolucao=_data_hora_como_datetime(
                    _campo(valores, indice, "Data Pagto./Devol.")
                ),
                valor_bruto=_valor_como_numero(_campo(valores, indice, "Valor Bruto")),
                observacao=_texto_ou_none(_campo(valores, indice, "Observação")),
            )
        )

    return ResultadoLeituraEfisco(linhas=linhas, avisos=avisos, linhas_ignoradas=linhas_ignoradas)
