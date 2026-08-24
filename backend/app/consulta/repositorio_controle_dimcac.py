"""Micro-sprint 3 (Relatório de Diárias — Portal da Transparência) —
leitura própria e não-destrutiva do Controle DIMCAC, para alimentar o lado
"controle" do motor de reconciliação (`domain/relatorio_transparencia/
motor_reconciliacao.py`) em vez de `repositorio_diarias.diarias_completas`.

Achado real da auditoria que motivou este módulo: `diarias_completas`
indexa por SEI e descarta silenciosamente qualquer linha cujo SEI já
apareceu antes (`repositorio_diarias.py`, aviso "SEI duplicado, mantida a
primeira ocorrência") — 13 DHs reais do Controle desapareciam do universo
de reconciliação só por causa disso, mesmo tendo DH preenchido. Aqui SEI
NUNCA é chave de deduplicação: a leitura agrupa por DH, preservando
**todas** as linhas com DH preenchido, não importa quantas vezes um SEI se
repita dentro do mesmo DH, entre DHs diferentes, ou apareça malformado —
mesma filosofia de `repositorio_ordens_bancarias.ordens_bancarias_por_dh`
(1 DH → N OBs preservado), aplicada agora ao lado do Controle.

`repositorio_diarias.py`/`diarias_completas` continuam 100% intocados —
usados pela Central de Conferência/Base Local, com seu próprio critério já
homologado (mais tolerante, descarte silencioso por design). Este módulo é
uma leitura paralela e isolada, só para o Relatório de Diárias — nenhum
fluxo homologado (PDF → extração → Planilha Oficial, Central de
Conferência, Central Analítica, Central de Regularização, Base Local) é
alterado ou compartilha este código."""

import logging
import os
import re
from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.consulta.repositorio_diarias import (
    _como_data,  # noqa: PLC2701 — mesmo reuso já praticado em repositorio_ordens_bancarias.py
    _dividir_trecho,  # noqa: PLC2701
    _valor_como_numero,  # noqa: PLC2701
)
from app.domain.entities.aviso_linha_ignorada import AvisoLinhaIgnorada
from app.domain.entities.registro_controle import ClassificacaoSei, RegistroControle
from app.planilha.escritor_planilha import (
    COLUNA_CPF,
    COLUNA_DH,
    COLUNA_DOC_ORIGEM,
    COLUNA_MOTIVO,
    COLUNA_PERIODO_FINAL,
    COLUNA_PERIODO_INICIAL,
    COLUNA_SEI,
    COLUNA_TRANSPORTE,
    COLUNA_TRECHO,
    COLUNA_VALOR_DIARIAS,
    NOME_ABA,
    PRIMEIRA_LINHA_DADOS,
    _cpf_normalizado,  # noqa: PLC2701 — nunca duplicar a normalização de CPF já homologada
)

__all__ = [
    "ControleDimcacEstruturalmenteInvalido",
    "ResultadoLeituraControleDimcac",
    "abrir_controle_dimcac",
    "controle_dimcac_por_dh",
]

_logger = logging.getLogger(__name__)

# NUP do SEI, formato estrutural gravado pelo próprio SEI em toda a base
# real de referência: NN.NN.NNNN(...N).NNNNNNN/NNNN-NN — o grupo da
# unidade (terceiro) varia de tamanho entre achados reais (`sei_parser_v1.
# py`), por isso `\d{4,}`, nunca `\d{4}` fixo. Usado só para CLASSIFICAR
# (`_classificar_sei`), nunca para corrigir ou rejeitar uma linha.
_PADRAO_SEI_NUP = re.compile(r"^\d{2}\.\d{2}\.\d{4,}\.\d{7}/\d{4}-\d{2}$")


class ControleDimcacEstruturalmenteInvalido(Exception):
    """Mesma filosofia de `OrdensBancariasEstruturalmenteInvalida`
    (`repositorio_ordens_bancarias.py`): erro estrutural que impede leitura
    confiável — arquivo ausente, extensão errada, workbook corrompido/
    bloqueado, ou aba obrigatória ausente. Nunca uma decisão de regra de
    negócio (isso é o motor de reconciliação)."""


def _classificar_sei(sei_bruto: object) -> tuple[str | None, ClassificacaoSei]:
    """Nunca corrige nem infere — só rotula o texto tal como está na
    célula (normalização mínima de `strip()`, mesma da coluna SEI em
    `repositorio_diarias.diarias_completas`)."""
    sei = str(sei_bruto or "").strip()
    if not sei:
        return None, ClassificacaoSei.AUSENTE
    if _PADRAO_SEI_NUP.match(sei):
        return sei, ClassificacaoSei.VALIDO
    return sei, ClassificacaoSei.MALFORMADO


def abrir_controle_dimcac(caminho: str) -> Worksheet:
    """Abre o Controle DIMCAC em modo NORMAL do openpyxl (`read_only`
    padrão, `False`) — mesmo modo que `base_local.py` já usa para abrir a
    Planilha Oficial antes de chamar `repositorio_diarias.diarias_completas`
    (`load_workbook(caminho, data_only=True)`, sem `read_only=True`). Ao
    contrário de `repositorio_ordens_bancarias.abrir_base_ordens_bancarias`
    (Base OB, arquivo externo lido sequencialmente), o Controle DIMCAC é o
    arquivo controlado por este próprio projeto (linhas escritas por
    `escritor_planilha`, algumas digitadas manualmente ao longo de anos) —
    modo normal preserva o mesmo comportamento de leitura já homologado
    para essa mesma aba em outros pontos do sistema. `data_only=True`:
    valor já calculado da célula, nunca o texto de uma fórmula."""
    if not os.path.isfile(caminho):
        raise ControleDimcacEstruturalmenteInvalido(f"Arquivo não encontrado: {caminho}")
    if not caminho.lower().endswith(".xlsx"):
        raise ControleDimcacEstruturalmenteInvalido(
            "Formato não suportado — selecione um arquivo .xlsx"
        )
    try:
        workbook = load_workbook(caminho, data_only=True)
    except Exception as erro:  # abrangente de propósito — qualquer falha de
        # leitura (arquivo corrompido, senha, permissão) vira o mesmo erro
        # estrutural explícito, nunca uma exceção genérica não tratada.
        raise ControleDimcacEstruturalmenteInvalido(
            "Não foi possível abrir o Controle DIMCAC — arquivo corrompido "
            f"ou bloqueado ({erro})"
        ) from erro

    if NOME_ABA not in workbook.sheetnames:
        raise ControleDimcacEstruturalmenteInvalido(
            f"Aba obrigatória ausente no Controle DIMCAC: '{NOME_ABA}'"
        )
    return workbook[NOME_ABA]


@dataclass(frozen=True)
class ResultadoLeituraControleDimcac:
    """`por_dh` agrupa **todas** as linhas com DH preenchido, na ordem em
    que aparecem na planilha — nunca deduplicado por SEI, nem colapsado
    quando o próprio DH se repete (achado real: o Controle às vezes tem o
    mesmo DH duplicado em duas linhas). `avisos` é só diagnóstico (SEI
    malformado, período/trecho/valor não reconhecidos) — nunca remove uma
    linha de `por_dh`.

    As únicas duas contagens que representam linhas fora de `por_dh`:
    `linhas_sem_dh` (SEI e/ou CPF preenchido, mas sem DH — condição
    esperada, mesma semântica de `RegistroDiariaIndexada.dh is None`:
    diária ainda não paga, fora do escopo desta projeção DH-a-DH, nunca um
    erro) e `linhas_vazias_ignoradas` (sem DH, sem SEI e sem CPF —
    placeholder/rascunho da planilha, a única condição que descarta uma
    linha por completo aqui). `total_linhas` é sempre igual à soma das
    três: linhas em `por_dh` + `linhas_sem_dh` + `linhas_vazias_
    ignoradas` — nenhuma linha lida desaparece sem estar contabilizada em
    algum dos três."""

    por_dh: dict[str, list[RegistroControle]] = field(default_factory=dict)
    avisos: list[AvisoLinhaIgnorada] = field(default_factory=list)
    total_linhas: int = 0
    linhas_sem_dh: int = 0
    linhas_vazias_ignoradas: int = 0


def controle_dimcac_por_dh(planilha: Worksheet) -> ResultadoLeituraControleDimcac:
    """Leitura completa e não-destrutiva do Controle DIMCAC, agrupada por
    DH, para o motor de reconciliação do Relatório de Diárias (Micro-
    sprint 3). Ao contrário de `diarias_completas` (indexado por SEI, uma
    linha por SEI — descarta duplicatas), SEI nunca é chave aqui: nenhuma
    linha com DH preenchido é descartada, sobrescrita ou colapsada, não
    importa quantas vezes o SEI ou o próprio DH se repitam, nem se o SEI
    vier malformado ou ausente."""
    por_dh: dict[str, list[RegistroControle]] = {}
    avisos: list[AvisoLinhaIgnorada] = []
    total_linhas = 0
    linhas_sem_dh = 0
    linhas_vazias_ignoradas = 0

    for linha in range(PRIMEIRA_LINHA_DADOS, planilha.max_row + 1):
        total_linhas += 1

        dh_bruto = planilha[f"{COLUNA_DH}{linha}"].value
        sei_bruto = planilha[f"{COLUNA_SEI}{linha}"].value
        cpf_bruto = planilha[f"{COLUNA_CPF}{linha}"].value

        dh = str(dh_bruto or "").strip() or None
        sei, sei_classificacao = _classificar_sei(sei_bruto)
        cpf = _cpf_normalizado(cpf_bruto) or None if cpf_bruto not in (None, "") else None

        # Linha realmente vazia — placeholder/rascunho da planilha, sem
        # nenhum dos três campos que poderiam torná-la relevante. Única
        # condição que descarta uma linha por completo nesta leitura;
        # nunca aplicada quando DH está preenchido (ver docstring do
        # módulo e da entidade).
        if dh is None and sei is None and cpf is None:
            linhas_vazias_ignoradas += 1
            continue

        # Sem DH, mas com SEI e/ou CPF: diária ainda não chegou à fase de
        # liquidação — condição esperada (mesma semântica de
        # `RegistroDiariaIndexada.dh is None`), fora do escopo desta
        # projeção DH-a-DH. Não é erro, não gera aviso, só não entra em
        # `por_dh` (não há DH para agrupar).
        if dh is None:
            linhas_sem_dh += 1
            continue

        periodo_inicial_bruto = planilha[f"{COLUNA_PERIODO_INICIAL}{linha}"].value
        periodo_final_bruto = planilha[f"{COLUNA_PERIODO_FINAL}{linha}"].value
        periodo_inicial = _como_data(periodo_inicial_bruto)
        periodo_final = _como_data(periodo_final_bruto)
        if periodo_inicial_bruto not in (None, "") and periodo_inicial is None:
            avisos.append(
                AvisoLinhaIgnorada(
                    aba=NOME_ABA,
                    linha=linha,
                    motivo=f"Período inicial não reconhecido como data: {periodo_inicial_bruto!r}",
                )
            )
        if periodo_final_bruto not in (None, "") and periodo_final is None:
            avisos.append(
                AvisoLinhaIgnorada(
                    aba=NOME_ABA,
                    linha=linha,
                    motivo=f"Período final não reconhecido como data: {periodo_final_bruto!r}",
                )
            )

        trecho = str(planilha[f"{COLUNA_TRECHO}{linha}"].value or "").strip() or None
        cidade_origem: str | None = None
        cidade_destino: str | None = None
        if trecho:
            cidade_origem, cidade_destino = _dividir_trecho(trecho)
            if cidade_origem is None:
                avisos.append(
                    AvisoLinhaIgnorada(
                        aba=NOME_ABA,
                        linha=linha,
                        motivo=f"Trecho não separável em origem/destino: {trecho!r}",
                    )
                )

        if sei_classificacao is ClassificacaoSei.MALFORMADO:
            avisos.append(
                AvisoLinhaIgnorada(aba=NOME_ABA, linha=linha, motivo=f"SEI malformado: {sei!r}")
            )

        doc_origem = str(planilha[f"{COLUNA_DOC_ORIGEM}{linha}"].value or "").strip() or None
        motivo = str(planilha[f"{COLUNA_MOTIVO}{linha}"].value or "").strip() or None
        transporte = str(planilha[f"{COLUNA_TRANSPORTE}{linha}"].value or "").strip() or None

        valor_bruto = planilha[f"{COLUNA_VALOR_DIARIAS}{linha}"].value
        valor_total_diarias = _valor_como_numero(valor_bruto)
        if valor_total_diarias is None and valor_bruto not in (None, ""):
            avisos.append(
                AvisoLinhaIgnorada(
                    aba=NOME_ABA,
                    linha=linha,
                    motivo=f"Valor total das diárias não reconhecido como número: {valor_bruto!r}",
                )
            )

        por_dh.setdefault(dh, []).append(
            RegistroControle(
                linha=linha,
                dh=dh,
                sei=sei,
                sei_classificacao=sei_classificacao,
                doc_origem=doc_origem,
                cpf=cpf,
                periodo_inicial=periodo_inicial,
                periodo_final=periodo_final,
                cidade_origem=cidade_origem,
                cidade_destino=cidade_destino,
                trecho=trecho,
                motivo=motivo,
                transporte=transporte,
                valor_total_diarias=valor_total_diarias,
            )
        )

    if linhas_vazias_ignoradas:
        _logger.info(
            "%s linha(s) realmente vazias ignoradas na leitura do Controle DIMCAC "
            "(sem DH, sem SEI e sem CPF).",
            linhas_vazias_ignoradas,
        )

    return ResultadoLeituraControleDimcac(
        por_dh=por_dh,
        avisos=avisos,
        total_linhas=total_linhas,
        linhas_sem_dh=linhas_sem_dh,
        linhas_vazias_ignoradas=linhas_vazias_ignoradas,
    )
