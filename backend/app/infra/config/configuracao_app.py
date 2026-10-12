"""Configuração local da instalação (Sprint 15) — o caminho da planilha
oficial que o usuário selecionou. Ver `docs/configuracao-inicial/README.md`.

Persistência simples em JSON, sem SQLite: um único registro por
instalação, sempre lido/escrito por inteiro (§ 9 do doc arquitetural). A
planilha escolhida nunca é copiada nem movida — só o caminho é guardado.
"""

import hashlib
import json
import os
import time
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path

import filelock
from openpyxl import load_workbook

from app.planilha.escritor_planilha import NOME_ABA

# Auditoria de concorrência cross-process (QA Codex, mesma sprint que
# corrigiu `persistencia_decisoes.py`/`configuracao_fontes_relatorio_
# transparencia.py`/`configuracao_template_relatorio_transparencia.py`) —
# achado: `persistir` e `persistir_pastas_operacionais` (abaixo) fazem cada
# uma um ciclo read-modify-write + `.tmp` fixo + `os.replace` sobre o MESMO
# arquivo (`configuracao.json`) sem NENHUMA trava — nem em memória sequer
# (pior que os 3 módulos já citados, que ao menos tinham `threading.Lock`
# antes desta correção). É literalmente o módulo citado como "técnica já
# homologada" pelos outros — a origem do padrão nunca tinha a própria
# proteção. Mesmo remédio: `filelock.FileLock` (trava real do SO, protege
# threads E processos; ver `app/consulta/persistencia_decisoes.py` para a
# análise completa da escolha e da garantia de não deixar lock órfão após
# um crash). Uma trava só, compartilhada pelas duas funções — protegem o
# MESMO arquivo de destino, então uma escrita de `persistir` nunca pode
# correr ao mesmo tempo que uma de `persistir_pastas_operacionais`.
_TIMEOUT_TRAVA_CONFIGURACAO_S = 10


class ConfiguracaoAppEmUso(Exception):
    """A trava do `configuracao.json` não foi obtida dentro do timeout de
    10s — outro processo/thread mantém a seção crítica aberta por tempo
    anormal. Mensagem já adequada para exibição direta ao usuário, nunca a
    exceção técnica (`filelock.Timeout`) crua."""


def _trava_configuracao(destino: Path) -> filelock.FileLock:
    return filelock.FileLock(f"{destino}.lock", timeout=_TIMEOUT_TRAVA_CONFIGURACAO_S)


# Achado real durante a homologação (volume alto de requisições HTTP
# concorrentes, servidor Uvicorn multi-worker de verdade) — mesma classe de
# `PermissionError: [WinError 5]` documentada em
# `app/consulta/persistencia_decisoes.py` (leitor×escritor no `os.replace`
# do Windows). Mesmo remédio já homologado em `base_local_sqlite.
# _substituir_com_retry`, reaproveitado aqui — compartilhado pelas duas
# funções deste módulo (`persistir`/`persistir_pastas_operacionais`), que
# já compartilham a mesma trava por escreverem o mesmo arquivo de destino.
_TENTATIVAS_REPLACE = 3
_INTERVALO_BASE_RETRY_REPLACE_S = 0.05


def _substituir_com_retry(caminho_tmp: Path, caminho_destino: Path) -> None:
    ultimo_erro: PermissionError | None = None
    for tentativa in range(_TENTATIVAS_REPLACE):
        try:
            os.replace(caminho_tmp, caminho_destino)
            return
        except PermissionError as erro:
            ultimo_erro = erro
            if tentativa < _TENTATIVAS_REPLACE - 1:
                time.sleep(_INTERVALO_BASE_RETRY_REPLACE_S * (tentativa + 1))
    assert ultimo_erro is not None
    raise ultimo_erro


# Mesmo achado, lado LEITOR (ver `app/consulta/persistencia_decisoes.py`
# para a análise completa): `carregar` também pode colidir com um
# `os.replace` em andamento e receber `PermissionError`. Este módulo, ao
# contrário dos 3 irmãos do Relatório de Transparência, não tem uma
# exceção de domínio própria para "configuração corrompida" (`carregar`
# nunca teve tratamento de erro de leitura — pré-existente, fora do escopo
# desta correção de concorrência) — o retry aqui só absorve a corrida
# TRANSIENTE; um `PermissionError` genuinamente persistente continua
# propagando cru, comportamento inalterado.
def _ler_texto_com_retry(caminho: Path) -> str:
    ultimo_erro: PermissionError | None = None
    for tentativa in range(_TENTATIVAS_REPLACE):
        try:
            return caminho.read_text(encoding="utf-8")
        except PermissionError as erro:
            ultimo_erro = erro
            if tentativa < _TENTATIVAS_REPLACE - 1:
                time.sleep(_INTERVALO_BASE_RETRY_REPLACE_S * (tentativa + 1))
    assert ultimo_erro is not None
    raise ultimo_erro


NOME_PASTA_DADOS_FALLBACK = "ExtratorDiarias"
NOME_ARQUIVO_CONFIGURACAO = "configuracao.json"
VAR_AMBIENTE_USER_DATA = "EXTRATOR_DIARIAS_USER_DATA"

# Mesmas colunas que `escritor_planilha.preencher` de fato escreve, com os
# rótulos reais documentados em `docs/central-conferencia/README.md` — ver
# § 5 do doc arquitetural desta Sprint. Colunas de fórmula/manuais
# (E/F/G/M/N/O/R/S/...) nunca são tocadas pelo escritor, então não entram
# aqui.
CABECALHO_EXIGIDO = {
    "A": "DH",
    "B": "SEI",
    "C": "DOC ORIGEM",
    "D": "CPF",
    "H": "PERÍODO INICIAL",
    "I": "PERÍODO FINAL",
    "J": "TRECHO",
    "K": "TRANSPORTE",
    "L": "MOTIVO",
    "P": "VALOR TOTAL DAS DIÁRIAS",
    "Q": "LOCAL",
}


class PlanilhaInvalida(Exception):
    """A mensagem já é adequada para exibição direta ao usuário (§5)."""


@dataclass(frozen=True)
class FingerprintEstrutural:
    """Só elementos estruturais que não mudam no uso normal — nunca
    tamanho ou data de modificação do arquivo (§6 do doc arquitetural:
    ambos mudam a cada gravação legítima, então não servem para detectar
    substituição)."""

    abas: list[str]
    hash_cabecalho: str

    def to_dict(self) -> dict:
        return {"abas": self.abas, "hash_cabecalho": self.hash_cabecalho}

    @staticmethod
    def from_dict(dados: dict) -> "FingerprintEstrutural":
        return FingerprintEstrutural(abas=list(dados["abas"]), hash_cabecalho=dados["hash_cabecalho"])


# Organização Automática de PDFs — quatro pastas operacionais configuráveis
# independentemente da planilha oficial (podem ser definidas depois, uma de
# cada vez, via `POST /api/v1/configuracao/pastas`). `None` = "ainda não
# configurada" — nunca um valor obrigatório em `Configuracao.from_dict`,
# para uma `configuracao.json` gravada antes desta feature continuar
# carregando normalmente (compatibilidade com configuração antiga).
CAMPOS_PASTAS_OPERACIONAIS = (
    "caminho_pdfs_a_processar",
    "caminho_pdfs_complemento_a_processar",
    "caminho_pdfs_processados",
    "caminho_pdfs_diarias_a_confirmar",
)

ROTULOS_PASTAS_OPERACIONAIS = {
    "caminho_pdfs_a_processar": "PDFs a Processar",
    "caminho_pdfs_complemento_a_processar": "PDFs de Complemento a Processar",
    "caminho_pdfs_processados": "PDFs Processados",
    "caminho_pdfs_diarias_a_confirmar": "PDFs de Diárias a Confirmar",
}


@dataclass(frozen=True)
class Configuracao:
    caminho_planilha_oficial: str
    configurado_em: str
    fingerprint_estrutural: FingerprintEstrutural
    caminho_pdfs_a_processar: str | None = None
    caminho_pdfs_complemento_a_processar: str | None = None
    caminho_pdfs_processados: str | None = None
    caminho_pdfs_diarias_a_confirmar: str | None = None

    def to_dict(self) -> dict:
        return {
            "caminho_planilha_oficial": self.caminho_planilha_oficial,
            "configurado_em": self.configurado_em,
            "fingerprint_estrutural": self.fingerprint_estrutural.to_dict(),
            "caminho_pdfs_a_processar": self.caminho_pdfs_a_processar,
            "caminho_pdfs_complemento_a_processar": self.caminho_pdfs_complemento_a_processar,
            "caminho_pdfs_processados": self.caminho_pdfs_processados,
            "caminho_pdfs_diarias_a_confirmar": self.caminho_pdfs_diarias_a_confirmar,
        }

    @staticmethod
    def from_dict(dados: dict) -> "Configuracao":
        return Configuracao(
            caminho_planilha_oficial=dados["caminho_planilha_oficial"],
            configurado_em=dados["configurado_em"],
            fingerprint_estrutural=FingerprintEstrutural.from_dict(dados["fingerprint_estrutural"]),
            caminho_pdfs_a_processar=dados.get("caminho_pdfs_a_processar"),
            caminho_pdfs_complemento_a_processar=dados.get("caminho_pdfs_complemento_a_processar"),
            caminho_pdfs_processados=dados.get("caminho_pdfs_processados"),
            caminho_pdfs_diarias_a_confirmar=dados.get("caminho_pdfs_diarias_a_confirmar"),
        )


@dataclass(frozen=True)
class ProntidaoPastas:
    liberado: bool
    pastas_faltando: list[str]


@dataclass(frozen=True)
class ProntidaoFuncionalidade:
    processamento_padrao: ProntidaoPastas
    complemento: ProntidaoPastas
    finalizacao: ProntidaoPastas


def _pasta_valida(caminho: str | None) -> bool:
    return bool(caminho) and os.path.isdir(caminho)


def calcular_prontidao_pastas(configuracao: Configuracao | None) -> ProntidaoFuncionalidade:
    """Backend é a única autoridade sobre a prontidão de cada funcionalidade
    — o frontend só consome o resultado. Revalida o filesystem a cada
    chamada (nunca confia num booleano persistido): se um drive/pasta
    desaparecer depois de configurado, a próxima leitura já reflete o
    bloqueio automaticamente."""

    def _avaliar(campos: tuple[str, ...]) -> ProntidaoPastas:
        faltando = [
            ROTULOS_PASTAS_OPERACIONAIS[campo]
            for campo in campos
            if not _pasta_valida(getattr(configuracao, campo, None) if configuracao else None)
        ]
        return ProntidaoPastas(liberado=not faltando, pastas_faltando=faltando)

    return ProntidaoFuncionalidade(
        processamento_padrao=_avaliar(("caminho_pdfs_a_processar", "caminho_pdfs_processados")),
        complemento=_avaliar(("caminho_pdfs_complemento_a_processar", "caminho_pdfs_processados")),
        finalizacao=_avaliar(("caminho_pdfs_diarias_a_confirmar",)),
    )


class PastaInvalida(Exception):
    """A mensagem já é adequada para exibição direta ao usuário."""


def _normalizar_pasta(caminho: str) -> str:
    return str(Path(caminho).resolve())


def validar_pasta_operacional(caminho: str) -> str:
    """Critérios do § "Configuração Inválida": caminho existente, diretório
    acessível. Devolve o caminho resolvido/normalizado (nunca o texto cru
    do usuário) — mesmo espírito do fingerprint estrutural da planilha:
    persistir sempre o valor já validado."""
    if not os.path.isdir(caminho):
        raise PastaInvalida(f"Pasta não encontrada ou inacessível: {caminho}")
    return _normalizar_pasta(caminho)


def persistir_pastas_operacionais(
    valores: dict[str, str], *, diretorio_base: Path | str | None = None
) -> Configuracao:
    """Atualiza uma ou mais pastas operacionais, preservando as demais já
    configuradas e todo o resto de `Configuracao` (planilha oficial,
    fingerprint) — mesma técnica de escrita atômica (`.tmp` + `os.replace`)
    já usada por `persistir`. `valores` só contém os campos que o usuário
    de fato selecionou nesta chamada (ver `app/schemas/configuracao.
    RequisicaoPastasOperacionais`).

    Auditoria de concorrência (QA Codex) — todo o ciclo leitura+validação+
    mesclagem+escrita+promoção corre sob `_trava_configuracao(destino)`
    (`filelock.FileLock`, cross-process): duas atualizações parciais
    concorrentes (campos diferentes), mesmo vindas de processos distintos,
    agora são serializadas e ficam mescladas no arquivo final, em vez de
    uma descartar a outra."""
    destino = caminho_configuracao(diretorio_base=diretorio_base)
    try:
        with _trava_configuracao(destino):
            atual = carregar(diretorio_base=diretorio_base)
            if atual is None:
                raise PastaInvalida(
                    "Configure a planilha oficial antes de configurar as pastas operacionais."
                )

            for campo in valores:
                if campo not in CAMPOS_PASTAS_OPERACIONAIS:
                    raise PastaInvalida(f"Campo de pasta operacional desconhecido: {campo}")

            normalizados = {campo: validar_pasta_operacional(caminho) for campo, caminho in valores.items()}

            pastas_finais = {campo: getattr(atual, campo) for campo in CAMPOS_PASTAS_OPERACIONAIS}
            pastas_finais.update(normalizados)

            # `normcase` (case-insensitive no Windows, idêntico a `os.path`
            # em outras plataformas) — mesma noção de "mesmo caminho" já
            # usada pela guarda de pertencimento
            # (`organizador_arquivos.pertence_a_pasta`).
            valores_definidos = [caminho for caminho in pastas_finais.values() if caminho]
            valores_normalizados = [os.path.normcase(caminho) for caminho in valores_definidos]
            if len(valores_normalizados) != len(set(valores_normalizados)):
                raise PastaInvalida("As pastas operacionais não podem apontar para o mesmo caminho.")

            nova_configuracao = replace(atual, **pastas_finais)

            destino.parent.mkdir(parents=True, exist_ok=True)
            temporario = destino.with_suffix(destino.suffix + ".tmp")
            temporario.write_text(
                json.dumps(nova_configuracao.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8"
            )
            _substituir_com_retry(temporario, destino)

            return nova_configuracao
    except filelock.Timeout as erro:
        raise ConfiguracaoAppEmUso(
            "A configuração está sendo usada por outro processo. Tente novamente em instantes."
        ) from erro


@dataclass(frozen=True)
class EstadoConfiguracao:
    """Espelha o contrato do `GET /api/v1/configuracao` (§4) — booleanos
    independentes, sem estados de texto."""

    configurado: bool
    planilha_acessivel: bool = False
    requer_revalidacao: bool = False
    configuracao: Configuracao | None = None
    # Organização Automática de PDFs — independente de `planilha_acessivel`
    # (as pastas operacionais não têm relação com a planilha em si), sempre
    # populado quando `configuracao` existe (mesmo com a planilha
    # inacessível) para o Hub conseguir mostrar o aviso de pendência em
    # qualquer estado.
    prontidao_pastas: ProntidaoFuncionalidade | None = None


def _diretorio_dados_usuario() -> Path:
    """Mesma resolução já usada por `app/consulta/base_local.py` — mesma
    variável de ambiente, mesmo fallback — para as duas partes do app
    concordarem sobre o mesmo diretório sem duplicar a lógica."""
    caminho_env = os.environ.get(VAR_AMBIENTE_USER_DATA)
    if caminho_env:
        return Path(caminho_env)
    appdata = os.environ.get("APPDATA")
    base = Path(appdata) if appdata else Path.home()
    return base / NOME_PASTA_DADOS_FALLBACK


def caminho_configuracao(*, diretorio_base: Path | str | None = None) -> Path:
    diretorio = Path(diretorio_base) if diretorio_base is not None else _diretorio_dados_usuario()
    return diretorio / NOME_ARQUIVO_CONFIGURACAO


def _calcular_fingerprint(workbook) -> FingerprintEstrutural:
    planilha = workbook[NOME_ABA]
    valores = [str(planilha[f"{coluna}1"].value or "").strip() for coluna in sorted(CABECALHO_EXIGIDO)]
    digest = hashlib.sha256("|".join(valores).encode("utf-8")).hexdigest()
    return FingerprintEstrutural(abas=list(workbook.sheetnames), hash_cabecalho=f"sha256:{digest}")


def validar_planilha(caminho_planilha: str) -> FingerprintEstrutural:
    """Valida os 5 critérios do §5, nesta ordem, parando no primeiro que
    falhar. Levanta `PlanilhaInvalida` com mensagem específica; devolve o
    fingerprint estrutural quando a planilha passa em todos."""
    if not os.path.isfile(caminho_planilha):
        raise PlanilhaInvalida(f"Arquivo não encontrado: {caminho_planilha}")

    if not caminho_planilha.lower().endswith(".xlsx"):
        raise PlanilhaInvalida("Formato não suportado — selecione um arquivo .xlsx")

    try:
        workbook = load_workbook(caminho_planilha, read_only=True, data_only=False)
    except Exception as erro:
        raise PlanilhaInvalida(
            f"Não foi possível abrir a planilha — arquivo corrompido ou bloqueado ({erro})"
        ) from erro

    try:
        if NOME_ABA not in workbook.sheetnames:
            raise PlanilhaInvalida(f"Aba '{NOME_ABA}' não encontrada na planilha")

        planilha = workbook[NOME_ABA]
        for coluna, rotulo_esperado in CABECALHO_EXIGIDO.items():
            valor = str(planilha[f"{coluna}1"].value or "").strip()
            if valor != rotulo_esperado:
                raise PlanilhaInvalida(
                    f"Cabeçalho inesperado na coluna {coluna}1 "
                    f"(esperado '{rotulo_esperado}', encontrado '{valor}')"
                )

        return _calcular_fingerprint(workbook)
    finally:
        workbook.close()


def persistir(
    caminho_planilha: str,
    fingerprint: FingerprintEstrutural,
    *,
    diretorio_base: Path | str | None = None,
) -> Configuracao:
    """Escrita segura contra corrupção (§9): arquivo temporário no mesmo
    diretório + `os.replace` — mesma técnica já homologada em
    `app/consulta/base_local.py::sincronizar`.

    Organização Automática de PDFs — preserva as pastas operacionais já
    configuradas (independentes da planilha oficial): trocar a planilha via
    "Alterar Planilha Oficial" nunca pode apagar silenciosamente as pastas
    já selecionadas pelo usuário.

    Auditoria de concorrência (QA Codex) — todo o ciclo leitura+mesclagem+
    escrita+promoção corre sob `_trava_configuracao(destino)` (`filelock.
    FileLock`, cross-process, a MESMA trava de `persistir_pastas_
    operacionais` — protegem o mesmo arquivo de destino, nunca podem
    correr ao mesmo tempo uma com a outra)."""
    destino = caminho_configuracao(diretorio_base=diretorio_base)
    try:
        with _trava_configuracao(destino):
            anterior = carregar(diretorio_base=diretorio_base)
            pastas_preservadas = (
                {campo: getattr(anterior, campo) for campo in CAMPOS_PASTAS_OPERACIONAIS}
                if anterior is not None
                else {}
            )
            configuracao = Configuracao(
                caminho_planilha_oficial=caminho_planilha,
                configurado_em=datetime.now().isoformat(timespec="seconds"),
                fingerprint_estrutural=fingerprint,
                **pastas_preservadas,
            )

            destino.parent.mkdir(parents=True, exist_ok=True)

            temporario = destino.with_suffix(destino.suffix + ".tmp")
            temporario.write_text(
                json.dumps(configuracao.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8"
            )
            _substituir_com_retry(temporario, destino)

            return configuracao
    except filelock.Timeout as erro:
        raise ConfiguracaoAppEmUso(
            "A configuração está sendo usada por outro processo. Tente novamente em instantes."
        ) from erro


def carregar(*, diretorio_base: Path | str | None = None) -> Configuracao | None:
    origem = caminho_configuracao(diretorio_base=diretorio_base)
    if not origem.is_file():
        return None
    return Configuracao.from_dict(json.loads(_ler_texto_com_retry(origem)))


def avaliar_estado_atual(*, diretorio_base: Path | str | None = None) -> EstadoConfiguracao:
    """Única leitura de disco/planilha que alimenta o `GET
    /api/v1/configuracao` inteiro — evita reabrir o arquivo duas vezes
    para `planilha_acessivel` e `requer_revalidacao` separadamente."""
    configuracao = carregar(diretorio_base=diretorio_base)
    if configuracao is None:
        return EstadoConfiguracao(configurado=False)

    # Organização Automática de PDFs — revalida o filesystem a cada leitura
    # (nunca um booleano persistido): independente de `planilha_acessivel`,
    # já que as pastas operacionais não têm relação com a planilha em si.
    prontidao_pastas = calcular_prontidao_pastas(configuracao)

    if not os.path.isfile(configuracao.caminho_planilha_oficial):
        return EstadoConfiguracao(
            configurado=True,
            planilha_acessivel=False,
            configuracao=configuracao,
            prontidao_pastas=prontidao_pastas,
        )

    try:
        workbook = load_workbook(configuracao.caminho_planilha_oficial, read_only=True, data_only=False)
    except Exception:
        return EstadoConfiguracao(
            configurado=True,
            planilha_acessivel=False,
            configuracao=configuracao,
            prontidao_pastas=prontidao_pastas,
        )

    try:
        if NOME_ABA not in workbook.sheetnames:
            fingerprint_atual = None
        else:
            fingerprint_atual = _calcular_fingerprint(workbook)
    finally:
        workbook.close()

    exige_revalidacao = fingerprint_atual is None or fingerprint_atual != configuracao.fingerprint_estrutural

    return EstadoConfiguracao(
        configurado=True,
        planilha_acessivel=True,
        requer_revalidacao=exige_revalidacao,
        configuracao=configuracao,
        prontidao_pastas=prontidao_pastas,
    )


class PlanilhaOficialIndisponivel(Exception):
    """A mensagem já é adequada para exibição direta ao usuário — quem
    chama `resolver_caminho_para_preenchimento` não precisa saber qual das
    três checagens (configurado/acessível/revalidação) falhou."""


def resolver_caminho_para_preenchimento(*, diretorio_base: Path | str | None = None) -> str:
    """Único ponto de entrada que a camada de API de preenchimento deveria
    conhecer (§7.1 do doc arquitetural — Sprint 15): resolve e valida a
    configuração atual, devolvendo um caminho pronto para uso, ou levanta
    `PlanilhaOficialIndisponivel`. Quem chama esta função nunca precisa
    saber que existe um `configuracao.json`, qual sua estrutura, ou como o
    fingerprint é calculado — isso permanece inteiramente encapsulado
    aqui."""
    estado = avaliar_estado_atual(diretorio_base=diretorio_base)

    if not estado.configurado:
        raise PlanilhaOficialIndisponivel(
            "Nenhuma planilha oficial configurada. Configure em 'Alterar Planilha Oficial'."
        )
    if not estado.planilha_acessivel:
        raise PlanilhaOficialIndisponivel(
            "A planilha oficial configurada está inacessível — foi movida, renomeada ou "
            "excluída? Selecione-a novamente em 'Alterar Planilha Oficial'."
        )
    if estado.requer_revalidacao:
        raise PlanilhaOficialIndisponivel(
            "A estrutura da planilha oficial mudou desde a última configuração. "
            "Revalide em 'Alterar Planilha Oficial' antes de continuar."
        )

    assert estado.configuracao is not None
    return estado.configuracao.caminho_planilha_oficial
