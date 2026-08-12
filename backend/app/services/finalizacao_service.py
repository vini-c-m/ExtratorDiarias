import logging
import os
from dataclasses import dataclass
from enum import Enum

from openpyxl import load_workbook

from app.consulta import base_local
from app.domain.rules.finalizacao_diarias import AvaliacaoElegibilidade, avaliar
from app.extraction import pdf_reader
from app.extraction.parsers.parser_factory import obter_parser
from app.planilha.escritor_planilha import (
    NOME_ABA,
    _capturar_cache_formulas,  # noqa: PLC2701 — mesmo reuso já praticado em repositorio_analitico.py
    _reinjetar_cache_formulas,  # noqa: PLC2701 — evita duplicar a correção da ETAPA 1 (Auditoria P0)
)
from app.planilha.finalizador_planilha import (
    CandidatoLinhaFinalizacao,
    EstadoLinhaFinalizacao,
    gravar_finalizacao,
    localizar_candidatos,
)

_logger = logging.getLogger(__name__)


class ErroFinalizacaoDiarias(Exception):
    pass


# Regra de negócio definitiva (Sprint pré-build, 2026-08-11) — a decisão de
# finalizar é do servidor do DIMCAC, não do sistema. A avaliação documental
# (`avaliar`, em `finalizacao_diarias.py`) nunca mais decide o `status`
# abaixo — ela só popula `ItemAnaliseFinalizacao.avaliacao.alertas`,
# sempre informativos, nunca bloqueantes. O ÚNICO critério que decide se um
# processo pode ser finalizado é a localização do SEI na Planilha Oficial:
#   - SEI não extraído do PDF com segurança → SEI_NAO_EXTRAIDO (bloqueia)
#   - SEI não encontrado na planilha → SEI_NAO_ENCONTRADO (bloqueia)
#   - falha técnica de leitura do PDF → ERRO_LEITURA (bloqueia)
#   - linha com valor fora do padrão em X/Y (ex. "SM") → ESTADO_PARCIAL
#     (bloqueia — nunca sobrescreve um valor que não reconhece)
#   - SEI localizado, alguma linha ainda precisa de X/Y → PRONTO (pode gravar)
#   - SEI localizado, todas as linhas já X=SIM/Y=SIM → JA_FINALIZADO (NO-OP)
class StatusFinalizacao(str, Enum):
    PRONTO = "pronto"
    JA_FINALIZADO = "ja_finalizado"
    SEI_NAO_EXTRAIDO = "sei_nao_extraido"
    SEI_NAO_ENCONTRADO = "sei_nao_encontrado"
    ESTADO_PARCIAL = "estado_parcial"
    ERRO_LEITURA = "erro_leitura"


_MENSAGENS_LINHA = {
    StatusFinalizacao.SEI_NAO_ENCONTRADO: "SEI não localizado na Planilha Oficial.",
    StatusFinalizacao.JA_FINALIZADO: "Já confirmado na Planilha Oficial.",
    StatusFinalizacao.ESTADO_PARCIAL: "Estado parcial — revisão manual necessária.",
}


@dataclass(frozen=True)
class ItemAnaliseFinalizacao:
    arquivo_origem: str
    sei: str | None
    status: StatusFinalizacao
    motivo: str | None
    avaliacao: AvaliacaoElegibilidade | None
    candidatos: list[CandidatoLinhaFinalizacao]


def _status_linha(
    candidatos: list[CandidatoLinhaFinalizacao],
) -> StatusFinalizacao:
    """A regra é do PROCESSO (SEI), não da linha isolada (ajuste de
    2026-08-07): quando o mesmo SEI corresponde a mais de uma linha (ex.:
    diária original + Complemento de Diárias), todas são tratadas como um
    conjunto atômico — nunca uma escolha manual de qual linha vale.

    Qualquer linha em estado PARCIAL (valor fora do padrão, ex. "SM")
    bloqueia o SEI inteiro (revisão manual — nunca sobrescrito). Sem
    nenhuma PARCIAL: se TODAS já estiverem SIM/SIM, o SEI conta como já
    confirmado (NO-OP); caso contrário (alguma PRONTA ou INCOMPLETA) o SEI
    está pronto — `confirmar` grava só o que falta em cada linha, nunca
    reescreve o que já está certo."""
    if not candidatos:
        return StatusFinalizacao.SEI_NAO_ENCONTRADO
    estados = {candidato.estado for candidato in candidatos}
    if EstadoLinhaFinalizacao.PARCIAL in estados:
        return StatusFinalizacao.ESTADO_PARCIAL
    if estados == {EstadoLinhaFinalizacao.JA_FINALIZADA}:
        return StatusFinalizacao.JA_FINALIZADO
    return StatusFinalizacao.PRONTO


def analisar_lote(caminhos: list[str], caminho_planilha: str) -> list[ItemAnaliseFinalizacao]:
    """Para cada PDF: extrai o SEI (mesmo parser já homologado), localiza a
    linha na Planilha Oficial e SEMPRE roda a avaliação documental — mas
    esta última nunca decide `status` (regra definitiva, ver comentário da
    classe `StatusFinalizacao` acima): populam só `avaliacao.alertas`,
    exibidos ao servidor sem nunca bloquear. Nunca grava nada — só
    leitura."""
    if not os.path.isfile(caminho_planilha):
        raise ErroFinalizacaoDiarias(f"Planilha não encontrada: {caminho_planilha}")

    parser = obter_parser()
    try:
        workbook = load_workbook(caminho_planilha)
    except Exception as erro:
        raise ErroFinalizacaoDiarias(f"Não foi possível abrir a planilha oficial: {erro}") from erro

    try:
        planilha = workbook[NOME_ABA]
    except KeyError:
        workbook.close()
        raise ErroFinalizacaoDiarias(f"Aba '{NOME_ABA}' não encontrada na planilha") from None

    itens: list[ItemAnaliseFinalizacao] = []
    for caminho in caminhos:
        nome_arquivo = os.path.basename(caminho)
        try:
            texto = pdf_reader.ler_texto(caminho)
        except Exception as erro:
            itens.append(
                ItemAnaliseFinalizacao(
                    arquivo_origem=nome_arquivo,
                    sei=None,
                    status=StatusFinalizacao.ERRO_LEITURA,
                    motivo=f"Falha ao ler o PDF: {erro}",
                    avaliacao=None,
                    candidatos=[],
                )
            )
            continue

        sei = parser.extrair(texto).sei
        if not sei:
            itens.append(
                ItemAnaliseFinalizacao(
                    arquivo_origem=nome_arquivo,
                    sei=None,
                    status=StatusFinalizacao.SEI_NAO_EXTRAIDO,
                    motivo="Número SEI não identificado no processo.",
                    avaliacao=None,
                    candidatos=[],
                )
            )
            continue

        # Sempre roda — nunca decide `status`, só popula alertas informativos.
        avaliacao = avaliar(texto)

        candidatos = localizar_candidatos(planilha, sei)
        status = _status_linha(candidatos)
        itens.append(
            ItemAnaliseFinalizacao(
                arquivo_origem=nome_arquivo,
                sei=sei,
                status=status,
                motivo=_MENSAGENS_LINHA.get(status),
                avaliacao=avaliacao,
                candidatos=candidatos,
            )
        )

    workbook.close()
    return itens


@dataclass(frozen=True)
class ItemConfirmacao:
    """Um SEI confirmado pelo usuário — nunca uma linha isolada (ajuste de
    2026-08-07): `confirmar` sempre resolve, de novo, TODAS as linhas
    daquele SEI na planilha, no momento da gravação."""

    arquivo_origem: str
    sei: str


class CategoriaConfirmacao(str, Enum):
    FINALIZADO = "finalizado"
    # Regra definitiva (Sprint pré-build, 2026-08-11) — distinto de
    # IGNORADO: aqui o SEI foi localizado e já estava inteiramente SIM/SIM
    # (NO-OP, nenhuma célula tocada) — nunca um bloqueio, nunca um erro.
    JA_CONFIRMADO = "ja_confirmado"
    IGNORADO = "ignorado"


@dataclass(frozen=True)
class ItemResultadoConfirmacao:
    arquivo_origem: str
    sei: str
    categoria: CategoriaConfirmacao
    # Linhas que receberam alguma escrita nesta chamada — o SIM/SIM
    # completo (linha PRONTA) ou só a célula que faltava (linha INCOMPLETA).
    linhas_atualizadas: list[int]
    # Linhas que já estavam SIM/SIM antes desta chamada (mantidas como
    # estão, nunca regravadas).
    linhas_mantidas: list[int]
    motivo: str | None = None


@dataclass(frozen=True)
class ResultadoConfirmacao:
    total_finalizado: int
    # Sprint pré-build, 2026-08-11 — contagem separada de "já estava
    # SIM/SIM" (NO-OP) — nunca somado a `total_ignorado` (que continua só
    # os bloqueios reais: SEI não localizado/estado parcial).
    total_ja_confirmado: int
    total_ignorado: int
    itens: list[ItemResultadoConfirmacao]
    # ETAPA 3 (Auditoria P0, 2026-08-09) — mesmo princípio da ETAPA 2
    # (`escritor_planilha.ResultadoPreenchimento.sincronizacao_analitica_
    # falhou`): a gravação da Confirmação e a sincronização da Central
    # Analítica são resultados distintos. Default `False` preserva 100%
    # os construtores/testes existentes.
    sincronizacao_analitica_falhou: bool = False


def confirmar(itens: list[ItemConfirmacao], caminho_planilha: str) -> ResultadoConfirmacao:
    """Grava X/Y em TODAS as linhas de cada SEI confirmado pelo usuário —
    nunca uma linha isolada. Resolve as linhas do SEI de novo, no momento
    da gravação (nunca confia só na análise anterior, que pode estar
    desatualizada).

    Regra definitiva (Sprint pré-build, 2026-08-11): alertas documentais
    NUNCA impedem a gravação — quem chega até aqui já foi decidido pelo
    servidor do DIMCAC (a análise só bloqueia por SEI não localizado/não
    extraído/estado parcial, nunca por pendência documental). Por linha:
    PRONTA grava X=SIM/Y=SIM; INCOMPLETA grava só a célula vazia, sem
    reescrever a que já é "SIM" (`gravar_finalizacao`); JA_FINALIZADA nunca
    é tocada. Atômico por SEI (ajuste de 2026-08-07): se qualquer linha do
    SEI estiver em estado PARCIAL (valor fora do padrão), nenhuma linha
    daquele SEI é gravada. Um SEI bloqueado nunca impede os demais itens do
    lote de serem gravados."""
    if not os.path.isfile(caminho_planilha):
        raise ErroFinalizacaoDiarias(f"Planilha não encontrada: {caminho_planilha}")

    try:
        workbook = load_workbook(caminho_planilha)
    except Exception as erro:
        raise ErroFinalizacaoDiarias(f"Não foi possível abrir a planilha oficial: {erro}") from erro

    try:
        planilha = workbook[NOME_ABA]
    except KeyError:
        workbook.close()
        raise ErroFinalizacaoDiarias(f"Aba '{NOME_ABA}' não encontrada na planilha") from None

    resultados: list[ItemResultadoConfirmacao] = []
    algum_gravado = False

    for item in itens:
        candidatos = localizar_candidatos(planilha, item.sei)
        if not candidatos:
            resultados.append(
                ItemResultadoConfirmacao(
                    arquivo_origem=item.arquivo_origem,
                    sei=item.sei,
                    categoria=CategoriaConfirmacao.IGNORADO,
                    linhas_atualizadas=[],
                    linhas_mantidas=[],
                    motivo="SEI não localizado na Planilha Oficial.",
                )
            )
            continue

        estados = {candidato.estado for candidato in candidatos}
        if EstadoLinhaFinalizacao.PARCIAL in estados:
            resultados.append(
                ItemResultadoConfirmacao(
                    arquivo_origem=item.arquivo_origem,
                    sei=item.sei,
                    categoria=CategoriaConfirmacao.IGNORADO,
                    linhas_atualizadas=[],
                    linhas_mantidas=[],
                    motivo="Estado parcial — revisão manual necessária.",
                )
            )
            continue

        # PRONTA (as duas vazias) e INCOMPLETA (uma já "SIM") precisam de
        # alguma escrita — `gravar_finalizacao` decide, célula a célula,
        # o que já está certo e nunca reescreve.
        linhas_a_gravar = [
            c.linha
            for c in candidatos
            if c.estado in (EstadoLinhaFinalizacao.PRONTA, EstadoLinhaFinalizacao.INCOMPLETA)
        ]
        linhas_ja_finalizadas = [
            c.linha for c in candidatos if c.estado is EstadoLinhaFinalizacao.JA_FINALIZADA
        ]

        if not linhas_a_gravar:
            resultados.append(
                ItemResultadoConfirmacao(
                    arquivo_origem=item.arquivo_origem,
                    sei=item.sei,
                    categoria=CategoriaConfirmacao.JA_CONFIRMADO,
                    linhas_atualizadas=[],
                    linhas_mantidas=linhas_ja_finalizadas,
                    motivo="Já confirmado na Planilha Oficial.",
                )
            )
            continue

        for linha in linhas_a_gravar:
            gravar_finalizacao(planilha, linha)
        algum_gravado = True
        resultados.append(
            ItemResultadoConfirmacao(
                arquivo_origem=item.arquivo_origem,
                sei=item.sei,
                categoria=CategoriaConfirmacao.FINALIZADO,
                linhas_atualizadas=linhas_a_gravar,
                linhas_mantidas=linhas_ja_finalizadas,
            )
        )

    sincronizacao_analitica_falhou = False
    if algum_gravado:
        # ETAPA 1 (Auditoria P0) reaproveitada aqui: este `workbook.save()`
        # tem a MESMA vulnerabilidade que `escritor_planilha.preencher()`
        # tinha antes da correção — `openpyxl` carregado sem
        # `data_only=True` (necessário para não converter as fórmulas que
        # este módulo não toca) apaga o cache de TODA fórmula pré-
        # existente ao salvar o workbook inteiro, não só das células que
        # `gravar_finalizacao` de fato grava (X/Y de confirmação, nunca
        # fórmula). Sem isto, ligar a sincronização abaixo reintroduziria
        # o mesmo bug de perda de valor financeiro (941k→482k) por um
        # caminho novo. Mesmas duas funções já usadas por `preencher()` —
        # nenhuma lógica duplicada.
        cache_formulas = _capturar_cache_formulas(workbook, caminho_planilha)
        workbook.save(caminho_planilha)
        _reinjetar_cache_formulas(caminho_planilha, cache_formulas)
        workbook.close()

        # ETAPA 3 (Auditoria P0) — lacuna confirmada pela auditoria:
        # Finalizar/Confirmar alterava a Planilha Oficial mas nunca
        # disparava sincronização da Base Local, deixando a Central
        # Analítica desatualizada até o próximo processamento/complemento
        # ou clique manual. Reaproveita literalmente `base_local.
        # sincronizar()` — mesmo `LOCK_SINCRONIZACAO`, mesmo retry de
        # `PermissionError`, mesma validação estrutural, mesmo Last Known
        # Good da ETAPA 2 — nenhum mecanismo novo. Uma falha aqui nunca
        # desfaz a finalização já concluída (mesmo contrato de
        # `planilha_service.preencher_planilha`): só sinaliza.
        def _marcar_falha_sincronizacao(erro: Exception) -> None:
            nonlocal sincronizacao_analitica_falhou
            sincronizacao_analitica_falhou = True

        try:
            base_local.sincronizar(
                caminho_planilha, origem="finalizacao", ao_falhar_sqlite=_marcar_falha_sincronizacao
            )
        except Exception:
            _logger.exception("Falha ao sincronizar a Base Local de Consulta após Finalizar/Confirmar")
            sincronizacao_analitica_falhou = True
    else:
        workbook.close()

    total_finalizado = sum(1 for r in resultados if r.categoria is CategoriaConfirmacao.FINALIZADO)
    total_ja_confirmado = sum(1 for r in resultados if r.categoria is CategoriaConfirmacao.JA_CONFIRMADO)
    return ResultadoConfirmacao(
        total_finalizado=total_finalizado,
        total_ja_confirmado=total_ja_confirmado,
        total_ignorado=len(resultados) - total_finalizado - total_ja_confirmado,
        itens=resultados,
        sincronizacao_analitica_falhou=sincronizacao_analitica_falhou,
    )
