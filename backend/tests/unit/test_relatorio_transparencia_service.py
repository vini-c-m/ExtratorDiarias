from datetime import date

import pytest
from openpyxl import Workbook

from app.domain.relatorio_transparencia.decisao import AcaoDecisao
from app.services import relatorio_transparencia_service as service

_CABECALHO_BASE_OB = [
    "DH",
    "OB",
    "SITUAÇÃO",
    "DATA OB",
    "SEI",
    "CPF",
    "NOME",
    "CARGO",
    "PERÍODO INICIAL",
    "PERÍODO FINAL",
    "TRECHO",
    "TRANSPORTE",
    "MOTIVO",
    "VALOR DA DIÁRIA",
    "RESTITUIÇÃO",
    "DUPLICIDADE",
    "ITEM DE GASTO",
]


def _linha_ob(
    *,
    dh,
    ob,
    cpf=11111111111,
    periodo_inicial=date(2026, 5, 10),
    periodo_final=date(2026, 5, 10),
    trecho="RECIFE - GARANHUNS",
    valor=154.09,
    data_ob="15/05/2026 09:00:00",
    situacao="Paga",
    restituicao=None,
    duplicidade="Único",
    motivo="Motivo de teste",
):
    return [
        dh,
        ob,
        situacao,
        data_ob,
        "19.20.0001.0000001/2026-01",
        cpf,
        "FULANO DE TAL",
        "Cargo de Teste",
        periodo_inicial,
        periodo_final,
        trecho,
        "CARRO",
        motivo,
        valor,
        restituicao,
        duplicidade,
        "33901401",
    ]


def _criar_base_ob(caminho, linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(_CABECALHO_BASE_OB)
    for linha in linhas:
        ws.append(linha)
    wb.save(caminho)


def _criar_controle_vazio(caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = "Controle de Diárias"
    wb.save(caminho)


# HOMO-RT-05 — layout de colunas do Controle DIMCAC real (fixo por posição,
# nunca por busca de cabeçalho — `app/consulta/repositorio_controle_dimcac.py`
# lê por letra de coluna: A=DH, B=SEI, C=DOC ORIGEM, D=CPF, G=CARGO,
# H=PERÍODO INICIAL, I=PERÍODO FINAL, J=TRECHO, K=TRANSPORTE, L=MOTIVO,
# N=QUANT. DIÁRIA INTEGRAL, O=QUANT. DIÁRIA PARCIAL, P=VALOR DA DIÁRIA).
def _criar_controle_com_linha(
    caminho,
    *,
    dh,
    cargo=None,
    transporte="CARRO",
    quantidade_integral=None,
    quantidade_parcial=None,
    cpf=11111111111,
    periodo_inicial=date(2026, 5, 10),
    periodo_final=date(2026, 5, 10),
    trecho="RECIFE - GARANHUNS",
    motivo="Motivo de teste",
    valor=154.09,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Controle de Diárias"
    ws.append(
        [
            "DH", "SEI", "DOC ORIGEM", "CPF", "", "", "CARGO", "PERÍODO INICIAL",
            "PERÍODO FINAL", "TRECHO", "TRANSPORTE", "MOTIVO", "",
            "QUANT. DIÁRIA INTEGRAL", "QUANT. DIÁRIA PARCIAL", "VALOR DA DIÁRIA",
        ]
    )
    ws.append(
        [
            dh, "19.20.0001.0000001/2026-01", "PJTESTE", cpf, None, None, cargo,
            periodo_inicial, periodo_final, trecho, transporte, motivo, None,
            quantidade_integral, quantidade_parcial, valor,
        ]
    )
    wb.save(caminho)


def _preparar(tmp_path, linhas_ob, *, mes_referencia=5, caminho_decisoes=None):
    caminho_base_ob = tmp_path / "base_ob.xlsx"
    caminho_controle = tmp_path / "controle.xlsx"
    _criar_base_ob(caminho_base_ob, linhas_ob)
    _criar_controle_vazio(caminho_controle)
    caminho_decisoes = caminho_decisoes or (tmp_path / "decisoes.json")
    return service.preparar_relatorio(
        caminho_base_ordens_bancarias=str(caminho_base_ob),
        caminho_controle_dimcac=str(caminho_controle),
        ano_exercicio=2026,
        mes_referencia=mes_referencia,
        caminho_decisoes=str(caminho_decisoes),
    ), str(caminho_base_ob), str(caminho_controle), str(caminho_decisoes)


# --- 1. Preparar relatório sem revisão ------------------------------------


def test_preparar_relatorio_sem_revisao(tmp_path):
    resumo, *_ = _preparar(tmp_path, [_linha_ob(dh="2026DH000001", ob="2026OB000001")])

    assert resumo.total_candidatos_normais == 1
    assert resumo.total_revisoes == 0
    assert resumo.pode_gerar_relatorio is True


# --- 2-3. Preparar com revisão / resumo correto ---------------------------


def test_preparar_relatorio_com_revisao_e_resumo_correto(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000001", ob="2026OB000001"),  # normal
        _linha_ob(
            dh="2026DH000002", ob="2026OB000002", cpf=22222222222,
            periodo_inicial=date(2026, 5, 11), periodo_final=date(2026, 5, 11),
            trecho="RECIFE - ARCOVERDE", valor=308.19, data_ob="16/05/2026 09:00:00",
        ),
        _linha_ob(
            dh="2026DH000003", ob="2026OB000003", cpf=22222222222,
            periodo_inicial=date(2026, 5, 11), periodo_final=date(2026, 5, 11),
            trecho="RECIFE - ARCOVERDE", valor=308.19, data_ob="17/05/2026 09:00:00",
        ),
    ]
    resumo, *_ = _preparar(tmp_path, linhas)

    assert resumo.total_candidatos_normais == 1
    assert resumo.total_revisoes == 1
    assert resumo.total_pendentes == 1
    assert resumo.pode_gerar_relatorio is False
    revisao = resumo.revisoes[0]
    assert revisao.regra == "RN07"
    assert {oc.ob for oc in revisao.ocorrencias} == {"2026OB000002", "2026OB000003"}


# --- 4. pode_gerar_relatorio=False com pendência --------------------------


def test_pode_gerar_relatorio_false_com_pendencia(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, *_ = _preparar(tmp_path, linhas)
    assert resumo.pode_gerar_relatorio is False


# --- 5. Listar revisões -----------------------------------------------------


def test_listar_revisoes(tmp_path):
    caminho_base_ob = tmp_path / "base_ob.xlsx"
    caminho_controle = tmp_path / "controle.xlsx"
    _criar_base_ob(
        caminho_base_ob,
        [
            _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
            _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
        ],
    )
    _criar_controle_vazio(caminho_controle)
    caminho_decisoes = tmp_path / "decisoes.json"

    lista = service.listar_revisoes(
        caminho_base_ordens_bancarias=str(caminho_base_ob),
        caminho_controle_dimcac=str(caminho_controle),
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=str(caminho_decisoes),
    )

    assert len(lista) == 1
    assert lista[0].regra == "RN07"


# --- 6. Contrato da revisão contém campos obrigatórios --------------------


def test_contrato_da_revisao_contem_campos_obrigatorios(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, *_ = _preparar(tmp_path, linhas)
    revisao = resumo.revisoes[0]

    assert revisao.identidade_grupo
    assert revisao.regra
    assert revisao.estado
    assert revisao.explicacao
    assert isinstance(revisao.campos_iguais, list)
    assert isinstance(revisao.campos_diferentes, list)
    ocorrencia = revisao.ocorrencias[0]
    for campo in ("ob", "dh", "beneficiario", "sei", "data_ob", "periodo_inicial", "periodo_final", "trecho", "valor", "motivo", "restituicao"):
        assert hasattr(ocorrencia, campo)


# --- 7. Fingerprint não é exposto -------------------------------------------


def test_fingerprint_nao_e_exposto(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, *_ = _preparar(tmp_path, linhas)
    revisao = resumo.revisoes[0]

    assert not hasattr(revisao, "fingerprint")
    assert not hasattr(revisao, "fingerprint_evidencias")
    assert "fingerprint" not in repr(revisao).lower()


# --- 8-9. PUBLICAR TODAS / subconjunto --------------------------------------


def test_publicar_todas(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002", "2026OB000003"],
        autor="servidor.teste",
    )

    assert atualizado.total_candidatos_normais == 2
    assert atualizado.pode_gerar_relatorio is True


def test_publicar_subconjunto(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002"],
        autor="servidor.teste",
    )

    assert atualizado.total_candidatos_normais == 1
    assert atualizado.total_ocorrencias_nao_publicaveis == 1
    assert atualizado.pode_gerar_relatorio is True


# --- 10. ADIAR ---------------------------------------------------------------


def test_adiar(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.ADIAR,
        autor="servidor.teste",
    )

    assert atualizado.revisoes[0].estado == "adiada"
    assert atualizado.pode_gerar_relatorio is False


# --- 11-12. Decisão atualiza estado / todas decididas -----------------------


def test_decisao_atualiza_estado_e_todas_decididas_libera(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002", "2026OB000003"],
        autor="servidor.teste",
    )

    assert atualizado.revisoes[0].estado == "decidida"
    assert atualizado.pode_gerar_relatorio is True


# --- 13. Decisão reaproveitada em nova preparação ---------------------------


def test_decisao_reaproveitada_em_nova_preparacao(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo
    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002", "2026OB000003"],
        autor="servidor.teste",
    )

    # "mês seguinte": mesma fonte, mesma decisão persistida
    resumo_de_novo = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=8,
        caminho_decisoes=caminho_decisoes,
    )

    assert resumo_de_novo.revisoes[0].estado == "decidida"
    assert resumo_de_novo.pode_gerar_relatorio is True


# --- 14. Evidência alterada -> DESATUALIZADA --------------------------------


def test_evidencia_alterada_vira_desatualizada(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo
    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002", "2026OB000003"],
        autor="servidor.teste",
    )

    # fonte muda materialmente: valor de uma das OBs é alterado
    linhas_alteradas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=999.99),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    _criar_base_ob(caminho_ob, linhas_alteradas)

    resumo_depois = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )

    # o grupo em si deixou de existir (valor não bate mais para RN07) —
    # portanto zero revisões RN07 agora, e nenhuma decisão foi aplicada
    # silenciosamente a um grupo diferente.
    assert resumo_depois.pode_gerar_relatorio is True
    assert resumo_depois.total_revisoes == 0


def test_evidencia_alterada_em_campo_nao_estrutural_vira_desatualizada(tmp_path):
    """Mudar o MOTIVO (não afeta a chave RN07) preserva o grupo, mas muda
    o fingerprint — decisão antiga deve virar desatualizada, não ser
    reaplicada cegamente."""
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19, motivo="Motivo A"),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19, motivo="Motivo A"),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo
    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002", "2026OB000003"],
        autor="servidor.teste",
    )

    linhas_com_motivo_diferente = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19, motivo="Motivo B, revisado"),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19, motivo="Motivo A"),
    ]
    _criar_base_ob(caminho_ob, linhas_com_motivo_diferente)

    resumo_depois = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )

    assert resumo_depois.total_revisoes == 1
    assert resumo_depois.revisoes[0].estado == "desatualizada"
    assert resumo_depois.pode_gerar_relatorio is False


# --- 15. Grupo inexistente ---------------------------------------------------


def test_grupo_inexistente_levanta_erro(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    )

    with pytest.raises(service.RevisaoNaoEncontrada):
        service.registrar_decisao_revisao(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            identidade_grupo="grupo-que-nao-existe",
            acao=AcaoDecisao.PUBLICAR,
            obs_selecionadas=["2026OB000001"],
            autor="servidor.teste",
        )


# --- 16. OB fora do grupo ----------------------------------------------------


def test_ob_fora_do_grupo_levanta_erro(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000001", ob="2026OB000001"),
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    with pytest.raises(service.DecisaoInvalida):
        service.registrar_decisao_revisao(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            identidade_grupo=identidade,
            acao=AcaoDecisao.PUBLICAR,
            obs_selecionadas=["2026OB000001"],  # não pertence a este grupo
            autor="servidor.teste",
        )


# --- 17. Ação inválida (adiar + OBs selecionadas = payload ambíguo) --------


def test_adiar_com_obs_selecionadas_e_payload_ambiguo(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    with pytest.raises(service.DecisaoInvalida):
        service.registrar_decisao_revisao(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            identidade_grupo=identidade,
            acao=AcaoDecisao.ADIAR,
            obs_selecionadas=["2026OB000002"],
            autor="servidor.teste",
        )


# --- 17b. Micro-sprint 3E.2 — PUBLICAR com zero OBs exige justificativa ----
# Garantia REAL: chamada direta a `registrar_decisao_revisao` (sem passar
# pelo schema Pydantic da API) não consegue contornar a regra — é o único
# ponto de entrada desta operação.


def _preparar_com_uma_revisao(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    return resumo.revisoes[0].identidade_grupo, caminho_ob, caminho_controle, caminho_decisoes


def test_publicar_zero_obs_sem_observacao_levanta_erro(tmp_path):
    identidade, caminho_ob, caminho_controle, caminho_decisoes = _preparar_com_uma_revisao(tmp_path)

    with pytest.raises(service.DecisaoInvalida):
        service.registrar_decisao_revisao(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            identidade_grupo=identidade,
            acao=AcaoDecisao.PUBLICAR,
            obs_selecionadas=[],
            autor="servidor.teste",
            observacao=None,
        )


def test_publicar_zero_obs_com_observacao_vazia_levanta_erro(tmp_path):
    identidade, caminho_ob, caminho_controle, caminho_decisoes = _preparar_com_uma_revisao(tmp_path)

    with pytest.raises(service.DecisaoInvalida):
        service.registrar_decisao_revisao(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            identidade_grupo=identidade,
            acao=AcaoDecisao.PUBLICAR,
            obs_selecionadas=[],
            autor="servidor.teste",
            observacao="",
        )


def test_publicar_zero_obs_com_observacao_so_espacos_levanta_erro(tmp_path):
    identidade, caminho_ob, caminho_controle, caminho_decisoes = _preparar_com_uma_revisao(tmp_path)

    with pytest.raises(service.DecisaoInvalida):
        service.registrar_decisao_revisao(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            identidade_grupo=identidade,
            acao=AcaoDecisao.PUBLICAR,
            obs_selecionadas=[],
            autor="servidor.teste",
            observacao="   \n\t  ",
        )


def test_publicar_zero_obs_com_justificativa_valida_e_aceito(tmp_path):
    identidade, caminho_ob, caminho_controle, caminho_decisoes = _preparar_com_uma_revisao(tmp_path)

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[],
        autor="servidor.teste",
        observacao="Ambas as OBs são suspeitas de duplicidade; excluídas até apuração interna.",
    )

    assert atualizado.revisoes[0].estado == "decidida"
    assert atualizado.total_ocorrencias_nao_publicaveis == 2
    assert atualizado.pode_gerar_relatorio is True


def test_publicar_uma_ob_sem_observacao_e_aceito(tmp_path):
    identidade, caminho_ob, caminho_controle, caminho_decisoes = _preparar_com_uma_revisao(tmp_path)

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002"],
        autor="servidor.teste",
    )

    assert atualizado.revisoes[0].estado == "decidida"


def test_publicar_subconjunto_sem_observacao_e_aceito(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000004", ob="2026OB000004", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002", "2026OB000003"],
        autor="servidor.teste",
    )

    assert atualizado.revisoes[0].estado == "decidida"
    assert atualizado.total_ocorrencias_nao_publicaveis == 1


def test_publicar_todas_sem_observacao_e_aceito(tmp_path):
    identidade, caminho_ob, caminho_controle, caminho_decisoes = _preparar_com_uma_revisao(tmp_path)

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002", "2026OB000003"],
        autor="servidor.teste",
    )

    assert atualizado.revisoes[0].estado == "decidida"
    assert atualizado.total_ocorrencias_nao_publicaveis == 0


def test_adiar_sem_observacao_continua_aceito_nao_vira_obrigatorio(tmp_path):
    identidade, caminho_ob, caminho_controle, caminho_decisoes = _preparar_com_uma_revisao(tmp_path)

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.ADIAR,
        obs_selecionadas=[],
        autor="servidor.teste",
        observacao=None,
    )

    assert atualizado.revisoes[0].estado == "adiada"


# --- 18. Persistência inexistente (primeira preparação) --------------------


def test_persistencia_inexistente_funciona_normalmente(tmp_path):
    resumo, _, _, caminho_decisoes = _preparar(tmp_path, [_linha_ob(dh="2026DH000001", ob="2026OB000001")])
    assert resumo.pode_gerar_relatorio is True
    import os

    assert not os.path.exists(caminho_decisoes)  # preparar nunca cria o arquivo


# --- 19. Persistência corrompida --------------------------------------------


def test_persistencia_corrompida_levanta_erro(tmp_path):
    caminho_decisoes = tmp_path / "decisoes.json"
    caminho_decisoes.write_text("{isso nao e json valido", encoding="utf-8")

    with pytest.raises(service.PersistenciaDecisoesCorrompida):
        _preparar(tmp_path, [_linha_ob(dh="2026DH000001", ob="2026OB000001")], caminho_decisoes=caminho_decisoes)


# --- 20-21. Base OB / Controle inexistentes ---------------------------------


def test_base_ob_inexistente_levanta_erro(tmp_path):
    caminho_controle = tmp_path / "controle.xlsx"
    _criar_controle_vazio(caminho_controle)

    with pytest.raises(service.BaseOBNaoEncontrada):
        service.preparar_relatorio(
            caminho_base_ordens_bancarias=str(tmp_path / "nao_existe.xlsx"),
            caminho_controle_dimcac=str(caminho_controle),
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=str(tmp_path / "decisoes.json"),
        )


def test_controle_inexistente_levanta_erro(tmp_path):
    caminho_base_ob = tmp_path / "base_ob.xlsx"
    _criar_base_ob(caminho_base_ob, [_linha_ob(dh="2026DH000001", ob="2026OB000001")])

    with pytest.raises(service.ControleDimcacNaoEncontrado):
        service.preparar_relatorio(
            caminho_base_ordens_bancarias=str(caminho_base_ob),
            caminho_controle_dimcac=str(tmp_path / "nao_existe.xlsx"),
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=str(tmp_path / "decisoes.json"),
        )


# --- 22-23. Estrutura inválida -----------------------------------------------


def test_estrutura_base_ob_invalida_levanta_erro(tmp_path):
    caminho_base_ob = tmp_path / "base_ob.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["DH", "OB"])  # faltam colunas obrigatórias
    wb.save(caminho_base_ob)
    caminho_controle = tmp_path / "controle.xlsx"
    _criar_controle_vazio(caminho_controle)

    with pytest.raises(service.EstruturaBaseOBInvalida):
        service.preparar_relatorio(
            caminho_base_ordens_bancarias=str(caminho_base_ob),
            caminho_controle_dimcac=str(caminho_controle),
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=str(tmp_path / "decisoes.json"),
        )


def test_estrutura_controle_invalida_levanta_erro(tmp_path):
    caminho_base_ob = tmp_path / "base_ob.xlsx"
    _criar_base_ob(caminho_base_ob, [_linha_ob(dh="2026DH000001", ob="2026OB000001")])
    caminho_controle = tmp_path / "controle.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Aba Errada"
    wb.save(caminho_controle)

    with pytest.raises(service.EstruturaControleInvalida):
        service.preparar_relatorio(
            caminho_base_ordens_bancarias=str(caminho_base_ob),
            caminho_controle_dimcac=str(caminho_controle),
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=str(tmp_path / "decisoes.json"),
        )


# --- 24. Nenhum candidato normal alterado por decisão de outro grupo -------


def test_decisao_de_um_grupo_nao_afeta_candidatos_normais_de_outro(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000001", ob="2026OB000001"),  # normal, alheio à revisão
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002"],
        autor="servidor.teste",
    )

    assert atualizado.total_candidatos_normais == 2  # o original + o publicado da revisão


# --- 25. Nenhuma OB desaparece sem rastreabilidade --------------------------


def test_nenhuma_ob_desaparece_sem_rastreabilidade(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000001", ob="2026OB000001"),
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, *_ = _preparar(tmp_path, linhas)

    total_em_revisoes = sum(len(r.ocorrencias) for r in resumo.revisoes)
    assert resumo.total_candidatos_normais + total_em_revisoes == 3


# --- 26. Recriar/reiniciar "service" (funções puras, sem estado) -----------


def test_service_e_stateless_preparar_de_novo_funciona_igual(tmp_path):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    resumo_1, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    resumo_2 = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )
    assert resumo_1 == resumo_2


# --- 27. Nenhum estado derivado é persistido --------------------------------


def test_preparar_nunca_persiste_estado_derivado(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    import os

    assert not os.path.exists(caminho_decisoes)

    identidade = resumo.revisoes[0].identidade_grupo
    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002", "2026OB000003"],
        autor="servidor.teste",
    )
    # depois de registrar, o arquivo contém só a DECISÃO — nunca candidatos/
    # ocorrências/reconciliação inteira.
    conteudo = (tmp_path / "decisoes.json").read_text(encoding="utf-8")
    assert "candidatos" not in conteudo
    assert "ocorrencias_nao_publicaveis" not in conteudo


# --- 29-30. ADIAR e DESATUALIZADA continuam bloqueando ----------------------


def test_adiar_continua_bloqueando_apos_nova_preparacao(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo
    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.ADIAR,
        autor="servidor.teste",
    )

    resumo_2 = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=8,
        caminho_decisoes=caminho_decisoes,
    )
    assert resumo_2.revisoes[0].estado == "adiada"
    assert resumo_2.pode_gerar_relatorio is False


# --- Minimização de dados (3D.1): autor nunca aparece no log ---------------


def test_registrar_decisao_nao_loga_autor(tmp_path, caplog):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo

    autor_sentinela = "servidor.nome.super.especifico.para.o.teste"
    with caplog.at_level("INFO", logger="app.services.relatorio_transparencia_service"):
        service.registrar_decisao_revisao(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            identidade_grupo=identidade,
            acao=AcaoDecisao.PUBLICAR,
            obs_selecionadas=["2026OB000002", "2026OB000003"],
            autor=autor_sentinela,
        )

    texto_dos_logs = "\n".join(registro.getMessage() for registro in caplog.records)
    assert autor_sentinela not in texto_dos_logs
    # o autor CONTINUA persistido na decisão em si, só não no log técnico.
    decisoes_persistidas = (tmp_path / "decisoes.json").read_text(encoding="utf-8")
    assert autor_sentinela in decisoes_persistidas


# --- Micro-sprint 3F: gerar_relatorio_xlsx — orquestração fontes+decisões+geração


def _configurar_template(tmp_path, nome_arquivo="template.xlsx"):
    """Micro-sprint 3H — o template deixou de ser um parâmetro direto de
    `gerar_relatorio_xlsx`: agora é resolvido via configuração isolada do
    módulo (`configuracao_template_relatorio_transparencia`), configurada
    uma vez. Devolve `caminho_configuracao_template`, pronto para reuso em
    toda chamada de `gerar_relatorio_xlsx` no teste."""
    from app.consulta.configuracao_template_relatorio_transparencia import persistir_configuracao_template
    from tests.fixtures.template_relatorio_transparencia_sintetico import construir_template_sintetico

    caminho_template = str(tmp_path / nome_arquivo)
    construir_template_sintetico(caminho_template)
    caminho_configuracao_template = str(tmp_path / "template_configuracao.json")
    persistir_configuracao_template(caminho_configuracao_template, caminho_template)
    return caminho_configuracao_template


def test_gerar_relatorio_xlsx_reconstroi_fontes_e_gera_arquivo(tmp_path):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    _, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = str(tmp_path / "saida.xlsx")

    resultado = service.gerar_relatorio_xlsx(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=caminho_destino,
    )

    assert resultado.total_ocorrencias == 1
    import os

    assert os.path.exists(caminho_destino)


def test_gerar_relatorio_xlsx_bloqueia_com_revisao_pendente(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    _, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = str(tmp_path / "saida.xlsx")

    with pytest.raises(service.RelatorioBloqueadoPorPendencias):
        service.gerar_relatorio_xlsx(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            caminho_configuracao_template=caminho_configuracao_template,
            caminho_destino=caminho_destino,
        )


def test_gerar_relatorio_xlsx_sem_template_configurado_levanta_erro_dedicado(tmp_path):
    """Micro-sprint 3H — nunca gera com um template inexistente/ausente:
    `TemplateOficialIndisponivel` é distinto de `RelatorioBloqueadoPorPendencias`
    (G05) para a UI conseguir mostrar a orientação certa em cada caso."""
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    _, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    caminho_configuracao_template = str(tmp_path / "template_configuracao_nunca_criado.json")
    caminho_destino = str(tmp_path / "saida.xlsx")

    with pytest.raises(service.TemplateOficialIndisponivel):
        service.gerar_relatorio_xlsx(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            caminho_configuracao_template=caminho_configuracao_template,
            caminho_destino=caminho_destino,
        )


def test_gerar_relatorio_xlsx_nunca_confia_em_lista_de_obs_do_chamador(tmp_path):
    """`gerar_relatorio_xlsx` não aceita nenhum parâmetro de conteúdo (lista
    de OBs, valores) — o único jeito de mudar o que é publicado é registrar
    uma decisão humana real e deixá-la ser reconstruída a partir das
    fontes."""
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, linhas)
    identidade = resumo.revisoes[0].identidade_grupo
    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=["2026OB000002"],
        autor="servidor.teste",
    )
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = str(tmp_path / "saida.xlsx")

    resultado = service.gerar_relatorio_xlsx(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=caminho_destino,
    )

    assert resultado.total_ocorrencias == 1
    import openpyxl

    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb["Controle de Diárias"]
    assert ws["A7"].value == "2026OB000002"


# --- HOMO-RT-02: identidade por ocorrência (repetição física de OB) --------
# Duas linhas FÍSICAS com a MESMA OB na Base OB (I11), divergindo em
# `motivo` — a forma real como `_resolver_repeticao_fisica` produz um RN04
# (linhas idênticas colapsam antes de virar revisão; só divergência real
# gera RN04). Testa ponta a ponta, via fontes reais em disco, a correção do
# acoplamento de checkboxes: cada ocorrência física precisa ser
# selecionável (publicável) de forma independente.


def _linhas_rn04_repeticao_fisica(motivo_a="Motivo da ocorrência A", motivo_b="Motivo da ocorrência B"):
    return [
        _linha_ob(dh="2026DH000001", ob="2026OB000001", motivo=motivo_a),
        _linha_ob(dh="2026DH000001", ob="2026OB000001", motivo=motivo_b),
    ]


def test_rn04_repeticao_fisica_expoe_identidade_distinta_por_ocorrencia(tmp_path):
    resumo, *_ = _preparar(tmp_path, _linhas_rn04_repeticao_fisica())

    assert resumo.total_revisoes == 1
    revisao = resumo.revisoes[0]
    assert revisao.regra == "RN04"
    identidades = [oc.identidade for oc in revisao.ocorrencias]
    assert len(identidades) == 2
    assert len(set(identidades)) == 2  # nunca duas ocorrências com a mesma identidade
    assert all(identidade.startswith("2026OB000001#") for identidade in identidades)
    # `ob` continua igual para as duas — só `identidade` as distingue.
    assert {oc.ob for oc in revisao.ocorrencias} == {"2026OB000001"}


def test_publicar_somente_ocorrencia_a_de_grupo_rn04(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, _linhas_rn04_repeticao_fisica())
    revisao = resumo.revisoes[0]
    identidade_grupo = revisao.identidade_grupo
    identidade_a = revisao.ocorrencias[0].identidade

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_a],
        autor="servidor.teste",
    )

    assert atualizado.total_candidatos_normais == 1
    assert atualizado.total_ocorrencias_nao_publicaveis == 1
    assert atualizado.pode_gerar_relatorio is True
    assert atualizado.revisoes[0].estado == "decidida"


def test_publicar_somente_ocorrencia_b_de_grupo_rn04(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, _linhas_rn04_repeticao_fisica())
    revisao = resumo.revisoes[0]
    identidade_grupo = revisao.identidade_grupo
    identidade_b = revisao.ocorrencias[1].identidade

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_b],
        autor="servidor.teste",
    )

    assert atualizado.total_candidatos_normais == 1
    assert atualizado.total_ocorrencias_nao_publicaveis == 1
    assert atualizado.pode_gerar_relatorio is True


def test_publicar_ambas_ocorrencias_fisicas_de_grupo_rn04(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, _linhas_rn04_repeticao_fisica())
    revisao = resumo.revisoes[0]
    identidades = [oc.identidade for oc in revisao.ocorrencias]

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=identidades,
        autor="servidor.teste",
    )

    assert atualizado.total_candidatos_normais == 2
    assert atualizado.total_ocorrencias_nao_publicaveis == 0
    assert atualizado.pode_gerar_relatorio is True


def test_ob_pura_e_rejeitada_para_grupo_rn04_com_repeticao_fisica(tmp_path):
    """Uma decisão NOVA sobre um grupo com repetição física precisa
    identificar a ocorrência (`ob#N`) — `ob` pura não pertence mais a este
    grupo (não distingue qual das duas linhas físicas), então é rejeitada
    exatamente como qualquer outro valor fora do grupo."""
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, _linhas_rn04_repeticao_fisica())
    revisao = resumo.revisoes[0]

    with pytest.raises(service.DecisaoInvalida):
        service.registrar_decisao_revisao(
            caminho_base_ordens_bancarias=caminho_ob,
            caminho_controle_dimcac=caminho_controle,
            ano_exercicio=2026,
            mes_referencia=5,
            caminho_decisoes=caminho_decisoes,
            identidade_grupo=revisao.identidade_grupo,
            acao=AcaoDecisao.PUBLICAR,
            obs_selecionadas=["2026OB000001"],
            autor="servidor.teste",
        )


def test_resumo_de_revisao_expoe_historico_da_decisao(tmp_path):
    """Histórico de decisões (micro-sprint corretiva HOMO-RT-02): depois de
    decidir, `RevisaoResumo` expõe autor/decidido_em/observacao/
    obs_publicaveis exatamente como persistido — sem endpoint novo, sem
    sistema de auditoria novo."""
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, _linhas_rn04_repeticao_fisica())
    revisao = resumo.revisoes[0]
    identidade_a = revisao.ocorrencias[0].identidade
    assert resumo.revisoes[0].autor is None  # PENDENTE — nenhuma decisão ainda
    assert resumo.revisoes[0].decidido_em is None

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_a],
        autor="servidor.teste",
        observacao="Só a primeira ocorrência corresponde ao pagamento real.",
    )

    revisao_atualizada = atualizado.revisoes[0]
    assert revisao_atualizada.autor == "servidor.teste"
    assert revisao_atualizada.decidido_em is not None
    assert revisao_atualizada.observacao == "Só a primeira ocorrência corresponde ao pagamento real."
    assert revisao_atualizada.obs_publicaveis == [identidade_a]


def test_decisao_antiga_ambigua_de_grupo_rn04_migra_para_requer_nova_decisao(tmp_path):
    """Migração (HOMO-RT-02): simula uma decisão gravada ANTES desta
    correção (formato "ob pura", sem distinguir ocorrência) — registra uma
    decisão válida hoje (garante fingerprint/identidade_grupo corretos) e
    depois REESCREVE só `obs_publicaveis` no arquivo persistido para o
    formato antigo, reproduzindo exatamente o que uma instalação existente
    teria em disco. `preparar_relatorio` de novo deve reconhecer a
    ambiguidade, bloquear a geração, e preservar autor/data/observação da
    decisão antiga no histórico (nunca apagar silenciosamente)."""
    import json

    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(tmp_path, _linhas_rn04_repeticao_fisica())
    revisao = resumo.revisoes[0]
    identidade_a = revisao.ocorrencias[0].identidade

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_a],
        autor="servidor.antigo",
    )

    # Simula o formato pré-HOMO-RT-02: `obs_publicaveis` gravado por OB
    # pura (sem "#N"), a própria ambiguidade que motivou esta correção.
    with open(caminho_decisoes, encoding="utf-8") as arquivo:
        dados = json.load(arquivo)
    (entrada,) = dados["decisoes"].values()
    assert entrada["autor"] == "servidor.antigo"
    entrada["obs_publicaveis"] = ["2026OB000001"]
    with open(caminho_decisoes, "w", encoding="utf-8") as arquivo:
        json.dump(dados, arquivo)

    resumo_remontado = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )

    revisao_remontada = resumo_remontado.revisoes[0]
    assert revisao_remontada.estado == "requer_nova_decisao"
    assert resumo_remontado.pode_gerar_relatorio is False
    # histórico da decisão antiga preservado, nunca descartado silenciosamente
    assert revisao_remontada.autor == "servidor.antigo"
    assert revisao_remontada.decidido_em is not None
    assert revisao_remontada.obs_publicaveis == ["2026OB000001"]


def test_gerar_relatorio_xlsx_publica_exatamente_a_ocorrencia_fisica_escolhida(tmp_path):
    """Fim a fim: publicar só a ocorrência B (pelo `motivo`, já que a OB é
    igual nas duas linhas) — o XLSX final deve conter exatamente UMA linha
    para esta OB, com o `motivo` da ocorrência escolhida, nunca da outra."""
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_rn04_repeticao_fisica(motivo_a="Motivo da ocorrência A", motivo_b="Motivo da ocorrência B")
    )
    revisao = resumo.revisoes[0]
    identidade_b = revisao.ocorrencias[1].identidade
    assert revisao.ocorrencias[1].motivo == "Motivo da ocorrência B"

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_b],
        autor="servidor.teste",
    )
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = str(tmp_path / "saida.xlsx")

    resultado = service.gerar_relatorio_xlsx(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=caminho_destino,
    )

    assert resultado.total_ocorrencias == 1
    import openpyxl

    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb["Controle de Diárias"]
    assert ws["A7"].value == "2026OB000001"
    assert ws["I7"].value == "Motivo da ocorrência B"


# --- HOMO-RT-03: identidade estável ponta a ponta, à prova de reordenação -
# Reprodução end-to-end (fontes reais em disco, não só objetos de domínio)
# do achado do Codex: decisão publica só a ocorrência A; a Base OB é
# REESCRITA com as duas linhas físicas trocadas de posição (mesmo conteúdo,
# ordem física invertida — o equivalente real a "reexportação com linhas
# deslocadas"); a decisão reaplicada nunca pode publicar B no lugar de A.


def test_decisao_para_ocorrencia_a_nunca_publica_b_apos_reexportar_linhas_trocadas(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path,
        _linhas_rn04_repeticao_fisica(motivo_a="Motivo da ocorrência A", motivo_b="Motivo da ocorrência B"),
    )
    revisao = resumo.revisoes[0]
    # localiza a ocorrência pelo MOTIVO (conteúdo), nunca pela posição na lista
    ocorrencia_a = next(oc for oc in revisao.ocorrencias if oc.motivo == "Motivo da ocorrência A")
    identidade_a = ocorrencia_a.identidade

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_a],
        autor="servidor.teste",
    )

    # "Reexportação com linhas físicas trocadas": mesmo conteúdo lógico,
    # ordem física das duas linhas invertida no arquivo.
    _criar_base_ob(
        caminho_ob,
        [
            _linha_ob(dh="2026DH000001", ob="2026OB000001", motivo="Motivo da ocorrência B"),
            _linha_ob(dh="2026DH000001", ob="2026OB000001", motivo="Motivo da ocorrência A"),
        ],
    )

    resumo_reordenado = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )

    revisao_reordenada = resumo_reordenado.revisoes[0]
    assert revisao_reordenada.estado == "decidida"
    assert resumo_reordenado.pode_gerar_relatorio is True
    assert resumo_reordenado.total_candidatos_normais == 1

    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = str(tmp_path / "saida_reordenada.xlsx")
    resultado_xlsx = service.gerar_relatorio_xlsx(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=caminho_destino,
    )
    assert resultado_xlsx.total_ocorrencias == 1
    import openpyxl

    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb["Controle de Diárias"]
    assert ws["I7"].value == "Motivo da ocorrência A"  # NUNCA "Motivo da ocorrência B"


# --- HOMO-RT-04: canonicalização estrutural, fim a fim via arquivo real ---
# `\x1f`/`\x1e` (o contraexemplo exato do Codex) são bytes de controle
# ILEGAIS em XML — openpyxl rejeita ao gravar (`IllegalCharacterError`),
# então uma planilha real nunca poderia carregá-los. Mas outros
# caracteres "perigosos" para um formato ingênuo por separador — vírgula,
# dois-pontos, aspas, barra invertida, quebra de linha — são texto válido
# em qualquer célula (MOTIVO/SITUAÇÃO são campos livres). Este teste prova
# a correção com um desses casos realistas, através do pipeline real de
# leitura/gravação, não só em objetos de domínio construídos à mão.


def test_canonicalizacao_nunca_colide_com_quebra_de_linha_em_motivo_via_planilha_real(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path,
        [
            _linha_ob(dh="2026DH000001", ob="2026OB000001", motivo="Primeira linha\nSegunda parte", situacao="Paga"),
            _linha_ob(dh="2026DH000001", ob="2026OB000001", motivo="Primeira linha", situacao="Segunda parte\nPaga"),
        ],
    )

    revisao = resumo.revisoes[0]
    assert revisao.regra == "RN04"
    identidades = {oc.identidade for oc in revisao.ocorrencias}
    assert len(identidades) == 2  # nunca colidem, mesmo com o texto "redistribuído" entre os campos

    identidade_primeira = next(
        oc.identidade for oc in revisao.ocorrencias if oc.motivo == "Primeira linha\nSegunda parte"
    )
    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_primeira],
        autor="servidor.teste",
    )

    resumo_apos = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )
    assert resumo_apos.revisoes[0].estado == "decidida"
    assert resumo_apos.total_candidatos_normais == 1
    assert resumo_apos.total_ocorrencias_nao_publicaveis == 1


# --- Micro-sprint "Duplicatas idênticas" — fim a fim via fontes reais -----
# Achado ao desenhar estes testes: duas linhas físicas TOTALMENTE idênticas
# (mesma OB, nenhum campo divergente) nunca chegam a virar uma revisão
# RN04 de verdade — `motor_decisao._resolver_repeticao_fisica` já as
# colapsa automaticamente em UM representante (I11b), sem pedir decisão
# humana, ANTES desta micro-sprint sequer existir. O cenário real que
# precisa de decisão humana é um grupo de 3+ linhas onde PARTE delas é
# idêntica entre si e outra parte diverge — só a divergência real (não a
# duplicata em si) é o que dispara RN04 para o grupo inteiro. Por isso os
# testes fim-a-fim abaixo usam esse formato (A, A idêntica, B divergente);
# os testes de par PURAMENTE idêntico ([A,A] sem nenhum terceiro
# divergente) já estão cobertos no nível de domínio (`test_decisao.py`,
# construindo a `RevisaoNecessaria` diretamente — o mesmo padrão já usado
# por todo o resto desta suíte para casos que não nascem do motor real).


def _linhas_duplicata_identica_mais_divergente(motivo_divergente="Motivo B — divergente"):
    return [
        _linha_ob(dh="2026DH000001", ob="2026OB000001"),  # A
        _linha_ob(dh="2026DH000001", ob="2026OB000001"),  # A — idêntica à primeira em todo campo material
        _linha_ob(dh="2026DH000001", ob="2026OB000001", motivo=motivo_divergente),  # B — diverge só em motivo
    ]


def test_grupo_misto_a_a_b_expoe_identidade_igual_para_as_duas_copias_de_a(tmp_path):
    resumo, *_ = _preparar(tmp_path, _linhas_duplicata_identica_mais_divergente())

    assert resumo.total_revisoes == 1
    revisao = resumo.revisoes[0]
    assert revisao.regra == "RN04"
    assert len(revisao.ocorrencias) == 3
    identidades = [oc.identidade for oc in revisao.ocorrencias]
    assert len(set(identidades)) == 2  # 2 classes: [A, A] idênticas + B distinta
    contagem = {identidade: identidades.count(identidade) for identidade in set(identidades)}
    assert sorted(contagem.values()) == [1, 2]  # uma classe com 1 (B), outra com 2 (A, A)


def test_publicar_classe_identica_e_a_divergente_no_grupo_misto(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    revisao = resumo.revisoes[0]
    identidades = [oc.identidade for oc in revisao.ocorrencias]
    identidade_classe_a = next(i for i in identidades if identidades.count(i) == 2)
    identidade_b = next(i for i in identidades if identidades.count(i) == 1)

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_classe_a, identidade_b],
        autor="servidor.teste",
    )

    assert atualizado.revisoes[0].estado == "decidida"
    assert atualizado.total_candidatos_normais == 2  # 1 representante de A + B — NUNCA 3
    assert atualizado.total_ocorrencias_nao_publicaveis == 1  # a segunda cópia de A


def test_publicar_somente_a_classe_identica_no_grupo_misto(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    revisao = resumo.revisoes[0]
    identidades = [oc.identidade for oc in revisao.ocorrencias]
    identidade_classe_a = next(i for i in identidades if identidades.count(i) == 2)

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_classe_a],
        autor="servidor.teste",
    )

    assert atualizado.total_candidatos_normais == 1
    assert atualizado.total_ocorrencias_nao_publicaveis == 2  # a outra cópia de A + B


def test_nao_publicar_nenhuma_no_grupo_misto(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    identidade_grupo_alvo = resumo.revisoes[0].identidade_grupo

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=identidade_grupo_alvo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[],
        autor="servidor.teste",
        observacao="Grupo todo suspeito — nenhuma ocorrência deve ser publicada.",
    )

    assert atualizado.revisoes[0].estado == "decidida"
    assert atualizado.pode_gerar_relatorio is True
    assert atualizado.total_candidatos_normais == 0
    assert atualizado.total_ocorrencias_nao_publicaveis == 3


def test_adiar_grupo_misto_bloqueia(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    revisao = resumo.revisoes[0]

    atualizado = service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.ADIAR,
        autor="servidor.teste",
    )

    assert atualizado.revisoes[0].estado == "adiada"
    assert atualizado.pode_gerar_relatorio is False


def test_publicar_1_de_n_identicas_sobrevive_a_reload(tmp_path):
    """Persistência: salvar/recarregar "publicar 1 de N idênticas"."""
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    revisao = resumo.revisoes[0]
    identidades = [oc.identidade for oc in revisao.ocorrencias]
    identidade_classe_a = next(i for i in identidades if identidades.count(i) == 2)

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_classe_a],
        autor="servidor.teste",
    )

    resumo_recarregado = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )

    assert resumo_recarregado.revisoes[0].estado == "decidida"
    assert resumo_recarregado.total_candidatos_normais == 1
    assert resumo_recarregado.total_ocorrencias_nao_publicaveis == 2


def test_publicar_zero_sobrevive_a_reload(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    revisao = resumo.revisoes[0]

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[],
        autor="servidor.teste",
        observacao="Nenhuma ocorrência deve ser publicada.",
    )

    resumo_recarregado = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )

    assert resumo_recarregado.revisoes[0].estado == "decidida"
    assert resumo_recarregado.total_candidatos_normais == 0
    assert resumo_recarregado.pode_gerar_relatorio is True


def test_decisao_antiga_com_chave_de_grupo_legada_e_selecao_especifica_exige_nova_decisao(tmp_path):
    """Migração de `identidade_grupo` (corrigida nesta micro-sprint):
    simula uma decisão gravada por uma instalação existente com a chave de
    grupo ANTIGA (formato pré-correção) — registra uma decisão válida hoje
    e depois reescreve só a CHAVE do dicionário `decisoes` no arquivo
    persistido para o formato legado, preservando os demais campos
    intactos, reproduzindo exatamente o que uma instalação existente teria
    em disco."""
    import json

    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    revisao = resumo.revisoes[0]
    identidades = [oc.identidade for oc in revisao.ocorrencias]
    identidade_classe_a = next(i for i in identidades if identidades.count(i) == 2)

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_classe_a],
        autor="servidor.antigo",
    )

    with open(caminho_decisoes, encoding="utf-8") as arquivo:
        dados = json.load(arquivo)
    (_chave_antiga, entrada), = list(dados["decisoes"].items())
    from app.domain.relatorio_transparencia.decisao import _identidade_grupo_legado

    obs_do_grupo = sorted({oc.ob for oc in revisao.ocorrencias})
    nova_chave = _identidade_grupo_legado(revisao.regra, obs_do_grupo)
    entrada["identidade_grupo"] = nova_chave
    dados["decisoes"] = {nova_chave: entrada}
    with open(caminho_decisoes, "w", encoding="utf-8") as arquivo:
        json.dump(dados, arquivo)

    resumo_remontado = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )

    revisao_remontada = resumo_remontado.revisoes[0]
    assert revisao_remontada.estado == "requer_nova_decisao"
    assert resumo_remontado.pode_gerar_relatorio is False
    assert revisao_remontada.autor == "servidor.antigo"  # histórico preservado, nunca descartado


def test_decisao_antiga_com_chave_de_grupo_legada_e_zero_continua_valida(tmp_path):
    """Contraparte segura do teste acima: uma decisão antiga "publicar
    zero" (inequívoca, sem risco material mesmo sob colisão hipotética da
    chave legada) permanece válida sem exigir nova decisão."""
    import json

    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    revisao = resumo.revisoes[0]

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[],
        autor="servidor.antigo",
        observacao="Grupo todo suspeito.",
    )

    with open(caminho_decisoes, encoding="utf-8") as arquivo:
        dados = json.load(arquivo)
    (_chave_antiga, entrada), = list(dados["decisoes"].items())
    from app.domain.relatorio_transparencia.decisao import _identidade_grupo_legado

    obs_do_grupo = sorted({oc.ob for oc in revisao.ocorrencias})
    nova_chave = _identidade_grupo_legado(revisao.regra, obs_do_grupo)
    entrada["identidade_grupo"] = nova_chave
    dados["decisoes"] = {nova_chave: entrada}
    with open(caminho_decisoes, "w", encoding="utf-8") as arquivo:
        json.dump(dados, arquivo)

    resumo_remontado = service.preparar_relatorio(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
    )

    revisao_remontada = resumo_remontado.revisoes[0]
    assert revisao_remontada.estado == "decidida"
    assert resumo_remontado.pode_gerar_relatorio is True
    assert revisao_remontada.autor == "servidor.antigo"


def test_xlsx_publica_exatamente_uma_linha_para_classe_identica(tmp_path):
    """XLSX fim a fim: publicar a classe idêntica + B gera exatamente 2
    linhas no relatório final — nunca 3, mesmo havendo 3 ocorrências
    físicas no grupo."""
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente(motivo_divergente="Motivo B — divergente")
    )
    revisao = resumo.revisoes[0]
    identidades = [oc.identidade for oc in revisao.ocorrencias]
    identidade_classe_a = next(i for i in identidades if identidades.count(i) == 2)
    identidade_b = next(i for i in identidades if identidades.count(i) == 1)

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_classe_a, identidade_b],
        autor="servidor.teste",
    )
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = str(tmp_path / "saida.xlsx")

    resultado = service.gerar_relatorio_xlsx(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=caminho_destino,
    )

    assert resultado.total_ocorrencias == 2  # 1 representante de A + B — NUNCA 3
    import openpyxl

    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb["Controle de Diárias"]
    valores_publicados = {ws["I7"].value, ws["I8"].value}
    assert valores_publicados == {"Motivo de teste", "Motivo B — divergente"}
    assert ws["I9"].value is None  # nunca uma 3ª linha


def test_xlsx_nenhuma_selecionada_gera_zero_linhas(tmp_path):
    resumo, caminho_ob, caminho_controle, caminho_decisoes = _preparar(
        tmp_path, _linhas_duplicata_identica_mais_divergente()
    )
    revisao = resumo.revisoes[0]

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[],
        autor="servidor.teste",
        observacao="Grupo todo suspeito.",
    )
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = str(tmp_path / "saida.xlsx")

    resultado = service.gerar_relatorio_xlsx(
        caminho_base_ordens_bancarias=caminho_ob,
        caminho_controle_dimcac=caminho_controle,
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=caminho_decisoes,
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=caminho_destino,
    )

    assert resultado.total_ocorrencias == 0
    import openpyxl

    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb["Controle de Diárias"]
    assert ws["A7"].value.startswith("Total")  # linha de total logo após o cabeçalho — zero linhas de dados


# --- Micro-sprint corretiva HOMO-RT-05: identidade da LINHA PUBLICÁVEL ----
# Fim a fim, via arquivos reais (Base OB + Controle DIMCAC sintéticos):
# publica uma decisão específica; o Controle DIMCAC é alterado (só o Cargo
# efetivamente publicável muda, a Base OB permanece IDÊNTICA); a decisão
# precisa exigir nova decisão humana (`DESATUALIZADA`) — nunca continuar
# `DECIDIDA` publicando um Cargo diferente do que foi visto. O XLSX final
# só reflete o Cargo novo depois de uma decisão nova.


def test_decisao_especifica_vira_desatualizada_quando_cargo_do_controle_muda(tmp_path):
    linhas_ob = _linhas_rn04_repeticao_fisica(motivo_a="Motivo A", motivo_b="Motivo B")
    caminho_ob = tmp_path / "base_ob.xlsx"
    caminho_controle = tmp_path / "controle.xlsx"
    caminho_decisoes = tmp_path / "decisoes.json"
    _criar_base_ob(caminho_ob, linhas_ob)
    _criar_controle_com_linha(caminho_controle, dh="2026DH000001", cargo="ANALISTA JUDICIÁRIO")

    resumo = service.preparar_relatorio(
        caminho_base_ordens_bancarias=str(caminho_ob),
        caminho_controle_dimcac=str(caminho_controle),
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=str(caminho_decisoes),
    )
    revisao = resumo.revisoes[0]
    assert revisao.regra == "RN04"
    identidade_a = next(oc.identidade for oc in revisao.ocorrencias if oc.motivo == "Motivo A")

    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=str(caminho_ob),
        caminho_controle_dimcac=str(caminho_controle),
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=str(caminho_decisoes),
        identidade_grupo=revisao.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[identidade_a],
        autor="servidor.teste",
    )

    # Controle DIMCAC "corrigido" — MESMO DH, só o Cargo muda; Base OB
    # permanece intocada.
    _criar_controle_com_linha(caminho_controle, dh="2026DH000001", cargo="TÉCNICO JUDICIÁRIO")

    resumo_apos_mudanca_no_controle = service.preparar_relatorio(
        caminho_base_ordens_bancarias=str(caminho_ob),
        caminho_controle_dimcac=str(caminho_controle),
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=str(caminho_decisoes),
    )

    assert resumo_apos_mudanca_no_controle.revisoes[0].estado == "desatualizada"
    assert resumo_apos_mudanca_no_controle.pode_gerar_relatorio is False
    assert resumo_apos_mudanca_no_controle.total_candidatos_normais == 0

    # Nova decisão humana (identidades recalculadas — o conteúdo mudou,
    # então a identidade por ocorrência também muda) libera a geração, e o
    # XLSX final só então reflete o Cargo NOVO.
    revisao_reatualizada = resumo_apos_mudanca_no_controle.revisoes[0]
    nova_identidade_a = next(oc.identidade for oc in revisao_reatualizada.ocorrencias if oc.motivo == "Motivo A")
    service.registrar_decisao_revisao(
        caminho_base_ordens_bancarias=str(caminho_ob),
        caminho_controle_dimcac=str(caminho_controle),
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=str(caminho_decisoes),
        identidade_grupo=revisao_reatualizada.identidade_grupo,
        acao=AcaoDecisao.PUBLICAR,
        obs_selecionadas=[nova_identidade_a],
        autor="servidor.teste",
    )

    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = str(tmp_path / "saida.xlsx")
    resultado = service.gerar_relatorio_xlsx(
        caminho_base_ordens_bancarias=str(caminho_ob),
        caminho_controle_dimcac=str(caminho_controle),
        ano_exercicio=2026,
        mes_referencia=5,
        caminho_decisoes=str(caminho_decisoes),
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=caminho_destino,
    )

    assert resultado.total_ocorrencias == 1
    import openpyxl

    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb["Controle de Diárias"]
    assert ws["D7"].value == "TÉCNICO JUDICIÁRIO"  # NUNCA "ANALISTA JUDICIÁRIO" (o que foi visto na decisão antiga)
