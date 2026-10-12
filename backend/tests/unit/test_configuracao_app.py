import json
import multiprocessing
import os
import time

import pytest
from openpyxl import Workbook, load_workbook

from app.infra.config import configuracao_app as config
from tests._multiprocesso_util import executar_e_reportar, segurar_trava_ate_ser_morto


# --- validar_planilha — critérios do §5, na ordem ---------------------------


def test_validar_planilha_arquivo_inexistente(tmp_path):
    with pytest.raises(config.PlanilhaInvalida, match="não encontrado"):
        config.validar_planilha(str(tmp_path / "nao_existe.xlsx"))


def test_validar_planilha_extensao_nao_suportada(tmp_path):
    caminho = tmp_path / "planilha.xls"
    caminho.write_bytes(b"conteudo qualquer")

    with pytest.raises(config.PlanilhaInvalida, match="Formato não suportado"):
        config.validar_planilha(str(caminho))


def test_validar_planilha_arquivo_corrompido(tmp_path):
    caminho = tmp_path / "corrompida.xlsx"
    caminho.write_bytes(b"isto nao e um zip valido")

    with pytest.raises(config.PlanilhaInvalida, match="Não foi possível abrir"):
        config.validar_planilha(str(caminho))


def test_validar_planilha_aba_esperada_ausente(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Outra Aba Qualquer"
    caminho = tmp_path / "sem_aba.xlsx"
    workbook.save(caminho)

    with pytest.raises(config.PlanilhaInvalida, match="Controle de Diárias"):
        config.validar_planilha(str(caminho))


def test_validar_planilha_cabecalho_incompleto(caminho_planilha_configuracao_teste):
    caminho = caminho_planilha_configuracao_teste(cabecalho={"A": "DH", "B": "SEI"})

    with pytest.raises(config.PlanilhaInvalida, match="Cabeçalho inesperado"):
        config.validar_planilha(caminho)


def test_validar_planilha_cabecalho_com_rotulo_errado(caminho_planilha_configuracao_teste):
    cabecalho_errado = {**config.CABECALHO_EXIGIDO, "Q": "ESTADO"}
    caminho = caminho_planilha_configuracao_teste(cabecalho=cabecalho_errado)

    with pytest.raises(config.PlanilhaInvalida, match=r"coluna Q1"):
        config.validar_planilha(caminho)


def test_validar_planilha_valida_devolve_fingerprint(caminho_planilha_configuracao_teste):
    caminho = caminho_planilha_configuracao_teste()

    fingerprint = config.validar_planilha(caminho)

    assert fingerprint.abas == ["Controle de Diárias", "Dados", "Tabela Valores"]
    assert fingerprint.hash_cabecalho.startswith("sha256:")


# --- fingerprint estrutural — §6 --------------------------------------------


def test_fingerprint_nao_muda_ao_gravar_linhas_de_dados(caminho_planilha_configuracao_teste):
    """Uso normal (o escritor grava a partir da linha 2) não pode disparar
    `requer_revalidacao` — o fingerprint olha só a linha de cabeçalho."""
    caminho = caminho_planilha_configuracao_teste()
    fingerprint_antes = config.validar_planilha(caminho)

    # Reabre e grava uma linha de dados, como o escritor faria.
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    planilha["A2"] = "2026DH000100"
    planilha["B2"] = "19.20.0000.0000001/2026-00"
    workbook.save(caminho)

    fingerprint_depois = config.validar_planilha(caminho)
    assert fingerprint_depois == fingerprint_antes


def test_fingerprint_muda_quando_uma_aba_e_removida(caminho_planilha_configuracao_teste):
    caminho_completo = caminho_planilha_configuracao_teste()
    fingerprint_completo = config.validar_planilha(caminho_completo)

    caminho_sem_dados = caminho_planilha_configuracao_teste(
        nome_arquivo="sem_dados.xlsx", abas_extra=("Tabela Valores",)
    )
    fingerprint_sem_dados = config.validar_planilha(caminho_sem_dados)

    assert fingerprint_sem_dados != fingerprint_completo


# --- persistência (temp + os.replace) e leitura -----------------------------


def test_persistir_e_carregar_configuracao(tmp_path, caminho_planilha_configuracao_teste):
    caminho_planilha = caminho_planilha_configuracao_teste()
    fingerprint = config.validar_planilha(caminho_planilha)

    diretorio = tmp_path / "dados_usuario"
    configuracao = config.persistir(caminho_planilha, fingerprint, diretorio_base=diretorio)

    assert configuracao.caminho_planilha_oficial == caminho_planilha
    assert (diretorio / config.NOME_ARQUIVO_CONFIGURACAO).is_file()
    # nenhum arquivo temporário deve sobrar depois da troca atômica
    assert not any(nome.endswith(".tmp") for nome in os.listdir(diretorio))

    recarregada = config.carregar(diretorio_base=diretorio)
    assert recarregada == configuracao


def test_persistir_sobrescreve_configuracao_anterior(tmp_path, caminho_planilha_configuracao_teste):
    diretorio = tmp_path / "dados_usuario"

    caminho_v1 = caminho_planilha_configuracao_teste(nome_arquivo="v1.xlsx")
    config.persistir(caminho_v1, config.validar_planilha(caminho_v1), diretorio_base=diretorio)

    caminho_v2 = caminho_planilha_configuracao_teste(nome_arquivo="v2.xlsx")
    config.persistir(caminho_v2, config.validar_planilha(caminho_v2), diretorio_base=diretorio)

    recarregada = config.carregar(diretorio_base=diretorio)
    assert recarregada.caminho_planilha_oficial == caminho_v2


def test_carregar_sem_configuracao_devolve_none(tmp_path):
    assert config.carregar(diretorio_base=tmp_path / "vazio") is None


def test_configuracao_json_e_legivel_e_serializado_com_as_chaves_documentadas(
    tmp_path, caminho_planilha_configuracao_teste
):
    caminho_planilha = caminho_planilha_configuracao_teste()
    fingerprint = config.validar_planilha(caminho_planilha)
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, fingerprint, diretorio_base=diretorio)

    conteudo = json.loads((diretorio / config.NOME_ARQUIVO_CONFIGURACAO).read_text(encoding="utf-8"))
    # Organização Automática de PDFs — quatro chaves aditivas, sempre
    # presentes (mesmo `None` quando não configuradas ainda).
    assert set(conteudo.keys()) == {
        "caminho_planilha_oficial",
        "configurado_em",
        "fingerprint_estrutural",
        "caminho_pdfs_a_processar",
        "caminho_pdfs_complemento_a_processar",
        "caminho_pdfs_processados",
        "caminho_pdfs_diarias_a_confirmar",
    }
    assert set(conteudo["fingerprint_estrutural"].keys()) == {"abas", "hash_cabecalho"}


def test_configuracao_legado_sem_campos_de_pasta_carrega_normalmente(
    tmp_path, caminho_planilha_configuracao_teste
):
    """Regressão P3 (auditoria pré-release v2.3.1): uma `configuracao.json`
    gravada antes da Organização Automática de PDFs (v2.3) nunca tinha os
    4 campos de pasta — nem sequer como chave presente com valor `null`,
    já que `Configuracao`/`persistir_pastas_operacionais` simplesmente não
    existiam. `Configuracao.from_dict` usa `.get(...)` com default `None`
    para esses 4 campos exatamente para esse caso; este teste prova isso
    fim a fim: planilha oficial permanece configurada, campos novos vêm
    `None`, `avaliar_estado_atual` não quebra, a prontidão de pastas orienta
    o usuário a configurá-las (todas listadas em `pastas_faltando`), e
    configurar uma pasta em cima dessa configuração legada não apaga nem
    corrompe a planilha oficial já configurada."""
    caminho_planilha = caminho_planilha_configuracao_teste()
    fingerprint = config.validar_planilha(caminho_planilha)
    diretorio = tmp_path / "dados_usuario"
    diretorio.mkdir(parents=True)

    # Formato pré-v2.3: só as 3 chaves que `Configuracao`/`persistir` já
    # gravavam antes desta feature existir — sem as 4 chaves de pasta.
    conteudo_legado = {
        "caminho_planilha_oficial": caminho_planilha,
        "configurado_em": "2026-06-01T09:00:00",
        "fingerprint_estrutural": fingerprint.to_dict(),
    }
    (diretorio / config.NOME_ARQUIVO_CONFIGURACAO).write_text(
        json.dumps(conteudo_legado, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    carregada = config.carregar(diretorio_base=diretorio)
    assert carregada is not None
    assert carregada.caminho_planilha_oficial == caminho_planilha
    assert carregada.caminho_pdfs_a_processar is None
    assert carregada.caminho_pdfs_complemento_a_processar is None
    assert carregada.caminho_pdfs_processados is None
    assert carregada.caminho_pdfs_diarias_a_confirmar is None

    estado = config.avaliar_estado_atual(diretorio_base=diretorio)
    assert estado.configurado is True
    assert estado.planilha_acessivel is True
    assert estado.prontidao_pastas is not None
    assert estado.prontidao_pastas.processamento_padrao.liberado is False
    assert set(estado.prontidao_pastas.processamento_padrao.pastas_faltando) == {
        "PDFs a Processar",
        "PDFs Processados",
    }

    # Configurar uma pasta em cima da configuração legada não apaga nem
    # corrompe a planilha oficial já configurada.
    pasta_a_processar = tmp_path / "a_processar"
    pasta_a_processar.mkdir()
    atualizada = config.persistir_pastas_operacionais(
        {"caminho_pdfs_a_processar": str(pasta_a_processar)}, diretorio_base=diretorio
    )
    assert atualizada.caminho_planilha_oficial == caminho_planilha
    assert atualizada.caminho_pdfs_a_processar == str(pasta_a_processar.resolve())


# --- avaliar_estado_atual — combina os três sinais do contrato -------------


def test_avaliar_estado_nao_configurado(tmp_path):
    estado = config.avaliar_estado_atual(diretorio_base=tmp_path / "vazio")
    assert estado.configurado is False
    assert estado.planilha_acessivel is False
    assert estado.requer_revalidacao is False


def test_avaliar_estado_configurado_e_acessivel(tmp_path, caminho_planilha_configuracao_teste):
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)

    estado = config.avaliar_estado_atual(diretorio_base=diretorio)

    assert estado.configurado is True
    assert estado.planilha_acessivel is True
    assert estado.requer_revalidacao is False
    assert estado.configuracao.caminho_planilha_oficial == caminho_planilha


def test_avaliar_estado_planilha_movida_ou_excluida(tmp_path, caminho_planilha_configuracao_teste):
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)

    os.remove(caminho_planilha)

    estado = config.avaliar_estado_atual(diretorio_base=diretorio)
    assert estado.configurado is True
    assert estado.planilha_acessivel is False


def test_avaliar_estado_arquivo_corrompido_no_caminho_configurado(
    tmp_path, caminho_planilha_configuracao_teste
):
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)

    with open(caminho_planilha, "wb") as arquivo:
        arquivo.write(b"nao e mais um zip valido")

    estado = config.avaliar_estado_atual(diretorio_base=diretorio)
    assert estado.configurado is True
    assert estado.planilha_acessivel is False


def test_avaliar_estado_requer_revalidacao_quando_estrutura_diverge(
    tmp_path, caminho_planilha_configuracao_teste
):
    """Caminho continua o mesmo, mas o arquivo foi substituído por um com
    estrutura diferente — sem bloquear, só sinalizando revalidação (§6)."""
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)

    workbook = Workbook()
    workbook.active.title = "Controle de Diárias"
    workbook.active["A1"] = "outro cabeçalho qualquer"
    workbook.save(caminho_planilha)

    estado = config.avaliar_estado_atual(diretorio_base=diretorio)
    assert estado.configurado is True
    assert estado.planilha_acessivel is True
    assert estado.requer_revalidacao is True


def test_avaliar_estado_nao_requer_revalidacao_apos_gravacao_normal(
    tmp_path, caminho_planilha_configuracao_teste
):
    """Uso normal (gravar novas linhas de dados) não deve exigir
    revalidação — só mudanças na linha de cabeçalho/abas."""
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)

    workbook = load_workbook(caminho_planilha)
    planilha = workbook["Controle de Diárias"]
    planilha["A2"] = "2026DH000100"
    planilha["B2"] = "19.20.0000.0000001/2026-00"
    workbook.save(caminho_planilha)

    estado = config.avaliar_estado_atual(diretorio_base=diretorio)
    assert estado.requer_revalidacao is False


# --- resolver_caminho_para_preenchimento — §7.1 (baixo acoplamento) ---------
#
# Único ponto de entrada que a camada de API de preenchimento deveria
# conhecer — os testes abaixo confirmam que as três checagens (configurado
# / acessível / revalidação) viram uma única exceção com mensagem pronta,
# sem expor a quem chama nada sobre `configuracao.json` ou o fingerprint.


def test_resolver_caminho_sem_configuracao(tmp_path):
    with pytest.raises(config.PlanilhaOficialIndisponivel, match="Nenhuma planilha oficial configurada"):
        config.resolver_caminho_para_preenchimento(diretorio_base=tmp_path / "vazio")


def test_resolver_caminho_planilha_inacessivel(tmp_path, caminho_planilha_configuracao_teste):
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)
    os.remove(caminho_planilha)

    with pytest.raises(config.PlanilhaOficialIndisponivel, match="inacessível"):
        config.resolver_caminho_para_preenchimento(diretorio_base=diretorio)


def test_resolver_caminho_requer_revalidacao(tmp_path, caminho_planilha_configuracao_teste):
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)

    workbook = Workbook()
    workbook.active.title = "Controle de Diárias"
    workbook.save(caminho_planilha)

    with pytest.raises(config.PlanilhaOficialIndisponivel, match="estrutura da planilha oficial mudou"):
        config.resolver_caminho_para_preenchimento(diretorio_base=diretorio)


def test_resolver_caminho_configuracao_valida_devolve_o_caminho(
    tmp_path, caminho_planilha_configuracao_teste
):
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)

    assert config.resolver_caminho_para_preenchimento(diretorio_base=diretorio) == caminho_planilha


# ============================================================================
# QA Codex — achado colateral: `persistir`/`persistir_pastas_operacionais`
# não tinham NENHUMA trava (nem em memória) antes desta correção — pior que
# os outros módulos irmãos já corrigidos. `filelock.FileLock` cross-process,
# mesma técnica; testes com `multiprocessing.Process` de verdade.
# ============================================================================


def test_persistir_pastas_operacionais_multiprocesso_campos_diferentes_mesclados(
    tmp_path, caminho_planilha_configuracao_teste
):
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)

    pasta_a_processar = tmp_path / "a_processar"
    pasta_a_processar.mkdir()
    pasta_processados = tmp_path / "processados"
    pasta_processados.mkdir()

    fila = multiprocessing.Queue()
    p1 = multiprocessing.Process(
        target=executar_e_reportar,
        args=(
            fila,
            config.persistir_pastas_operacionais,
            ({"caminho_pdfs_a_processar": str(pasta_a_processar)},),
            {"diretorio_base": diretorio},
        ),
    )
    p2 = multiprocessing.Process(
        target=executar_e_reportar,
        args=(
            fila,
            config.persistir_pastas_operacionais,
            ({"caminho_pdfs_processados": str(pasta_processados)},),
            {"diretorio_base": diretorio},
        ),
    )
    p1.start()
    p2.start()
    resultados = [fila.get(timeout=20) for _ in range(2)]
    p1.join(timeout=20)
    p2.join(timeout=20)

    assert all(r[0] == "ok" for r in resultados), resultados
    conteudo_bruto = config.caminho_configuracao(diretorio_base=diretorio).read_text(encoding="utf-8")
    json.loads(conteudo_bruto)  # nunca corrompe

    final = config.carregar(diretorio_base=diretorio)
    assert final.caminho_pdfs_a_processar == str(pasta_a_processar.resolve())  # nenhum dos dois foi perdido
    assert final.caminho_pdfs_processados == str(pasta_processados.resolve())
    assert final.caminho_planilha_oficial == caminho_planilha  # preservado, nenhum dos dois apagou


def test_persistir_processo_morto_segurando_trava_recupera_depois(
    tmp_path, caminho_planilha_configuracao_teste
):
    caminho_planilha = caminho_planilha_configuracao_teste()
    diretorio = tmp_path / "dados_usuario"
    destino = config.caminho_configuracao(diretorio_base=diretorio)

    fila = multiprocessing.Queue()
    processo = multiprocessing.Process(
        target=segurar_trava_ate_ser_morto,
        args=(fila, config._trava_configuracao, (destino,)),
    )
    processo.start()
    sinal = fila.get(timeout=15)
    assert sinal == "trava_obtida"

    processo.kill()
    processo.join(timeout=15)
    assert not processo.is_alive()

    inicio = time.monotonic()
    configuracao = config.persistir(caminho_planilha, config.validar_planilha(caminho_planilha), diretorio_base=diretorio)
    duracao = time.monotonic() - inicio
    assert duracao < 8  # recupera bem antes do timeout de 10s — nenhum lock órfão
    assert configuracao.caminho_planilha_oficial == caminho_planilha
