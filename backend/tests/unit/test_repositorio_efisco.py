"""Plano de Migração E-Fisco, Sprint A — testes de `repositorio_efisco.py`
(leitura + validação estrutural da Base E-Fisco). Mesmo padrão de
`test_repositorio_ordens_bancarias.py`: `Workbook`/`Worksheet` em memória,
nunca um arquivo real em disco (isso é `test_abrir_base_efisco.py`,
integração). Dados sintéticos."""

from datetime import date, datetime

import pytest
from openpyxl import Workbook

from app.consulta.repositorio_efisco import (
    EfiscoEstruturalmenteInvalido,
    linhas_efisco,
    validar_estrutura,
)

_CABECALHO = [
    "Unidade Gestora",
    "Ordem Bancária",
    "Número do DH",
    "Situação",
    "Data da OB",
    "Data Pagto./Devol.",
    "Valor Bruto",
    "Valor Líquido",
    "Observação",
    "Código da Natureza de Despesa",
    "Razão Social do Credor",
    "Número do documento do Credor",
]

_OB_PADRAO = "2026OB000001"
_DH_PADRAO = "2026DH000001"


def _nova_planilha(cabecalho: list[str] | None = None):
    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "OB"
    planilha.append(cabecalho if cabecalho is not None else _CABECALHO)
    return workbook, planilha


def _escrever_linha_efisco(
    planilha,
    linha,
    *,
    unidade_gestora="320101",
    ob=_OB_PADRAO,
    dh=_DH_PADRAO,
    situacao="Paga",
    data_ob="10/02/2026",
    data_pagto_devol="10/02/2026 09:01:40",
    valor_bruto=154.09,
    valor_liquido=154.09,
    observacao="PAGTO 2026LE000001 SEI 19.20.0001.0000001/2026-01 FULANO DE TAL UMA DIARIA PARCIAL, DEST RECIFE",
    natureza="33901400",
    credor="FULANO DE TAL",
    documento_credor="111.222.333-44",
):
    linha_valores = {
        "Unidade Gestora": unidade_gestora,
        "Ordem Bancária": ob,
        "Número do DH": dh,
        "Situação": situacao,
        "Data da OB": data_ob,
        "Data Pagto./Devol.": data_pagto_devol,
        "Valor Bruto": valor_bruto,
        "Valor Líquido": valor_liquido,
        "Observação": observacao,
        "Código da Natureza de Despesa": natureza,
        "Razão Social do Credor": credor,
        "Número do documento do Credor": documento_credor,
    }
    for coluna_idx, rotulo in enumerate(_CABECALHO, start=1):
        planilha.cell(row=linha, column=coluna_idx, value=linha_valores.get(rotulo))


# --- 1. Validação estrutural ----------------------------------------------


def test_cabecalho_valido_nao_levanta_erro():
    _, planilha = _nova_planilha()
    validar_estrutura(planilha)  # não deve levantar


def test_coluna_obrigatoria_ausente_levanta_erro_explicito():
    cabecalho_incompleto = [c for c in _CABECALHO if c != "Ordem Bancária"]
    _, planilha = _nova_planilha(cabecalho_incompleto)

    with pytest.raises(EfiscoEstruturalmenteInvalido, match="Ordem Bancária"):
        validar_estrutura(planilha)


def test_multiplas_colunas_obrigatorias_ausentes_sao_todas_listadas():
    cabecalho_incompleto = [c for c in _CABECALHO if c not in ("Situação", "Valor Bruto")]
    _, planilha = _nova_planilha(cabecalho_incompleto)

    with pytest.raises(EfiscoEstruturalmenteInvalido) as excinfo:
        validar_estrutura(planilha)
    assert "Situação" in str(excinfo.value)
    assert "Valor Bruto" in str(excinfo.value)


def test_planilha_vazia_sem_cabecalho_levanta_erro():
    workbook = Workbook()
    planilha = workbook.active

    with pytest.raises(EfiscoEstruturalmenteInvalido):
        validar_estrutura(planilha)


# --- 2. Leitura de linha bem formada ---------------------------------------


def test_linha_bem_formada_e_lida_corretamente():
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2)

    resultado = linhas_efisco(planilha)

    assert len(resultado.linhas) == 1
    assert resultado.linhas_ignoradas == 0
    linha = resultado.linhas[0]
    assert linha.linha == 2
    assert linha.ob == _OB_PADRAO
    assert linha.dh == _DH_PADRAO
    assert linha.situacao == "Paga"
    assert linha.natureza == "33901400"
    assert linha.credor == "FULANO DE TAL"
    assert linha.documento_credor == "111.222.333-44"
    assert linha.data_ob == date(2026, 2, 10)
    assert linha.data_pagamento_devolucao == datetime(2026, 2, 10, 9, 1, 40)
    assert linha.valor_bruto == 154.09
    assert "DIARIA" in linha.observacao.upper()


def test_ordem_bancaria_ausente_e_ignorada_com_aviso():
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, ob=None)

    resultado = linhas_efisco(planilha)

    assert resultado.linhas == []
    assert resultado.linhas_ignoradas == 1
    assert len(resultado.avisos) == 1
    assert "Ordem Bancária" in resultado.avisos[0].motivo


def test_numero_do_dh_ausente_nao_e_ignorado():
    """Achado real da auditoria: despesas administrativas sem Documento
    Hábil (ex. transferência entre contas da própria UG) têm DH vazio — a
    linha continua sendo lida normalmente nesta Sprint."""
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, dh=None)

    resultado = linhas_efisco(planilha)

    assert len(resultado.linhas) == 1
    assert resultado.linhas[0].dh is None
    assert resultado.linhas_ignoradas == 0


def test_multiplas_linhas_preservam_ordem_e_posicao():
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, ob="2026OB000001")
    _escrever_linha_efisco(planilha, 3, ob="2026OB000002")
    _escrever_linha_efisco(planilha, 4, ob="2026OB000003")

    resultado = linhas_efisco(planilha)

    assert [linha.ob for linha in resultado.linhas] == ["2026OB000001", "2026OB000002", "2026OB000003"]
    assert [linha.linha for linha in resultado.linhas] == [2, 3, 4]


# --- 3. Tipos reais vindos do Excel -----------------------------------------


def test_data_da_ob_como_objeto_datetime_do_openpyxl_e_aceita():
    """Defesa por tipo, não só por texto: se a célula vier formatada como
    data no Excel, o openpyxl entrega um `datetime` nativo, não texto."""
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, data_ob=datetime(2026, 2, 10))

    resultado = linhas_efisco(planilha)

    assert resultado.linhas[0].data_ob == date(2026, 2, 10)


def test_data_da_ob_texto_fora_do_formato_vira_none_sem_excecao():
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, data_ob="10-02-2026")

    resultado = linhas_efisco(planilha)

    assert resultado.linhas[0].data_ob is None


def test_data_pagto_devol_vazia_vira_none_sem_excecao():
    """Achado real: uma OB Gerada (ainda não liquidada) não tem Data
    Pagto./Devol. preenchida — condição esperada, nunca um erro."""
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, situacao="Gerada", data_pagto_devol="")

    resultado = linhas_efisco(planilha)

    assert resultado.linhas[0].data_pagamento_devolucao is None


def test_valor_bruto_como_float_nativo_do_openpyxl_e_aceito():
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, valor_bruto=1234.56)

    resultado = linhas_efisco(planilha)

    assert resultado.linhas[0].valor_bruto == 1234.56


def test_valor_bruto_zero_e_distinto_de_ausente():
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, valor_bruto=0)

    resultado = linhas_efisco(planilha)

    assert resultado.linhas[0].valor_bruto == 0.0


def test_natureza_vazia_e_lida_como_none_nao_como_string_vazia():
    _, planilha = _nova_planilha()
    _escrever_linha_efisco(planilha, 2, natureza="")

    resultado = linhas_efisco(planilha)

    assert resultado.linhas[0].natureza is None
