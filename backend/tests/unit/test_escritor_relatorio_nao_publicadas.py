"""Sprint 4 (plano "XLSX Operacional de Diárias Não Publicadas") — testes do
escritor do segundo artefato, `Diarias_Nao_Publicadas_<competencia>.xlsx`.
Fixtures 100% sintéticas (nenhum CPF/nome/OB/SEI real)."""

import glob
import os
from datetime import date, datetime, timezone

import openpyxl
import pytest

from app.domain.entities.linha_passagem import EstadoPassagem, LinhaPassagem
from app.domain.entities.ordem_bancaria import OrdemBancaria
from app.domain.entities.registro_controle import ClassificacaoSei, RegistroControle
from app.domain.relatorio_transparencia.decisao import (
    AcaoDecisao,
    CategoriaNaoPublicada,
    DecisaoHumana,
    OcorrenciaNaoPublicavel,
)
from app.domain.relatorio_transparencia.motor_decisao import OcorrenciaFinanceira
from app.planilha.escritor_relatorio_nao_publicadas import (
    DestinoJaExiste,
    gerar_xlsx_diarias_nao_publicadas,
)

_NOME_ABA = "Não Publicadas"


def _ordem(
    *,
    ob="2026OB000001",
    dh="2026DH000001",
    linha=2,
    sei="19.20.0001.0000001/2026-01",
    nome="FULANO DE TAL",
    credor=None,
    cargo="Cargo de Teste",
    periodo_inicial=date(2026, 5, 10),
    periodo_final=date(2026, 5, 12),
    trecho="RECIFE-GARANHUNS",
    transporte="CARRO",
    motivo="Motivo de teste",
    valor=154.09,
) -> OrdemBancaria:
    return OrdemBancaria(
        linha=linha,
        dh=dh,
        ob=ob,
        credor=credor if credor is not None else nome,
        situacao="Paga",
        data_pagamento=datetime(2026, 5, 15, 9, 0, 0),
        sei=sei,
        cpf="11111111111",
        nome=nome,
        cargo=cargo,
        periodo_inicial=periodo_inicial,
        periodo_final=periodo_final,
        cidade_origem=None,
        cidade_destino=None,
        transporte=transporte,
        motivo=motivo,
        valor=valor,
        valor_bruto_sem_formula=valor,
        restituicao=None,
        duplicidade="Único",
        item_de_gasto="33901401",
        trecho=trecho,
    )


def _registro(
    *,
    dh="2026DH000001",
    linha=2,
    doc_origem="DIMOT - Divisão Ministerial de Operações e Transporte",
    cargo=None,
    transporte="CARRO",
    integral=1,
    parcial=0,
) -> RegistroControle:
    return RegistroControle(
        linha=linha,
        dh=dh,
        sei="19.20.0001.0000001/2026-01",
        sei_classificacao=ClassificacaoSei.VALIDO,
        doc_origem=doc_origem,
        cpf="11111111111",
        periodo_inicial=date(2026, 5, 10),
        periodo_final=date(2026, 5, 12),
        cidade_origem="RECIFE",
        cidade_destino="GARANHUNS",
        trecho="RECIFE-GARANHUNS",
        motivo="Motivo de teste",
        transporte=transporte,
        valor_total_diarias=154.09,
        quantidade_diaria_integral=integral,
        quantidade_diaria_parcial=parcial,
        cargo=cargo,
    )


def _linha_passagem(*, ob="2026OB000001", estado, valor=None, bruto=None, transporte="AVIÃO", linha=10):
    return LinhaPassagem(linha=linha, ob=ob, transporte=transporte, estado=estado, valor=valor, bruto=bruto)


def _decisao(
    *,
    regra="RN07",
    autor="servidor.teste",
    decidido_em=datetime(2026, 5, 20, 14, 30, 0, tzinfo=timezone.utc),
    observacao=None,
    obs_publicaveis=frozenset(),
) -> DecisaoHumana:
    return DecisaoHumana(
        identidade_grupo="grupo-teste",
        regra=regra,
        obs_envolvidas=frozenset({"2026OB000001"}),
        fingerprint_evidencias="fp-teste",
        acao=AcaoDecisao.PUBLICAR,
        autor=autor,
        decidido_em=decidido_em,
        obs_publicaveis=obs_publicaveis,
        observacao=observacao,
    )


def _item(
    *,
    categoria=CategoriaNaoPublicada.NAO_PUBLICADA_POR_DECISAO_HUMANA,
    registros_controle=None,
    passagem=None,
    decisao=None,
    **kwargs_ordem,
) -> OcorrenciaNaoPublicavel:
    ordem = _ordem(**kwargs_ordem)
    registros = registros_controle if registros_controle is not None else [_registro(dh=ordem.dh)]
    ocorrencia = OcorrenciaFinanceira(
        dh=ordem.dh, ordem_bancaria=ordem, registros_controle=registros, passagem=passagem
    )
    return OcorrenciaNaoPublicavel(ocorrencia=ocorrencia, decisao=decisao or _decisao(), categoria=categoria)


def _valores_linha(ws, linha: int) -> dict:
    cabecalhos = [c.value for c in ws[1]]
    return {cab: ws.cell(row=linha, column=i + 1).value for i, cab in enumerate(cabecalhos)}


# --- Estrutura básica -------------------------------------------------------


def test_gera_arquivo_com_cabecalho_e_uma_linha(tmp_path):
    destino = str(tmp_path / "nao_publicadas.xlsx")
    resultado = gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino)

    assert os.path.exists(destino)
    assert resultado.total_ocorrencias == 1
    assert resultado.total_categoria_a == 1
    assert resultado.total_categoria_b == 0

    wb = openpyxl.load_workbook(destino)
    assert _NOME_ABA in wb.sheetnames
    ws = wb[_NOME_ABA]
    assert ws["A1"].value == "Categoria"
    assert ws.cell(row=1, column=1).font.bold is True


def test_lista_vazia_gera_arquivo_so_com_cabecalho(tmp_path):
    destino = str(tmp_path / "vazio.xlsx")
    resultado = gerar_xlsx_diarias_nao_publicadas([], caminho_destino=destino)

    assert resultado.total_ocorrencias == 0
    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA]
    assert ws.max_row == 1  # só o cabeçalho


# --- Categoria/regra ---------------------------------------------------------


def test_categoria_a_e_b_rotuladas_corretamente(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    itens = [
        _item(ob="2026OB000001", categoria=CategoriaNaoPublicada.NAO_PUBLICADA_POR_DECISAO_HUMANA, decisao=_decisao(regra="RN07")),
        _item(ob="2026OB000002", dh="2026DH000002", categoria=CategoriaNaoPublicada.DUPLICATA_FISICA_IDENTICA_SUPRIMIDA, decisao=_decisao(regra="RN04")),
    ]
    gerar_xlsx_diarias_nao_publicadas(itens, caminho_destino=destino)

    ws = openpyxl.load_workbook(destino)[_NOME_ABA]
    linha1 = _valores_linha(ws, 2)
    linha2 = _valores_linha(ws, 3)
    assert linha1["Categoria"] == "Não publicada por decisão humana"
    assert linha1["Regra"] == "RN07"
    assert linha2["Categoria"] == "Duplicata física idêntica suprimida"
    assert linha2["Regra"] == "RN04"


# --- Campos com fonte canônica direta (SEI/DH/OB) ---------------------------


def test_sei_dh_ob_vem_das_fontes_originais(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(ob="2026OB700004", dh="2026DH700007", sei="19.20.0009.0009999/2026-09")
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["SEI"] == "19.20.0009.0009999/2026-09"
    assert linha["DH"] == "2026DH700007"
    assert linha["OB"] == "2026OB700004"


def test_sei_ausente_nunca_e_inventado(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(sei=None)
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["SEI"] is None


# --- Nome/CREDOR (reaproveita linha_publicavel.nome_publicado) --------------


def test_nome_usa_credor_quando_e_continuacao_comprovada(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(nome="FULANO DE TAL COM NOME MUITO LON", credor="FULANO DE TAL COM NOME MUITO LONGO TRUNCADO")
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Nome"] == "FULANO DE TAL COM NOME MUITO LONGO TRUNCADO"


# --- Cargo/Transporte (reaproveita valor_do_controle_com_fallback) ----------


def test_cargo_e_transporte_vem_do_controle_com_precedencia(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(
        cargo="Cargo da Base OB",
        transporte="CARRO",
        registros_controle=[_registro(cargo="Cargo do Controle", transporte="AVIÃO")],
    )
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Cargo"] == "Cargo do Controle"
    assert linha["Transporte"] == "AVIÃO"


def test_cargo_cai_para_base_ob_quando_controle_nao_resolve(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(cargo="Cargo da Base OB", registros_controle=[_registro(cargo=None)])
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Cargo"] == "Cargo da Base OB"


# --- Unidade de origem (DOC ORIGEM via RegistroControle, item 7 do plano) --


def test_unidade_de_origem_vem_do_doc_origem_do_controle(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(registros_controle=[_registro(doc_origem="GABPGJ - Gabinete do Procurador-Geral de Justiça")])
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Unidade de Origem"] == "GABPGJ - Gabinete do Procurador-Geral de Justiça"


def test_unidade_de_origem_nunca_inferida_do_nome_quando_ausente(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(registros_controle=[])  # nenhum Controle associado
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Unidade de Origem"] is None


def test_unidade_de_origem_ausente_quando_controle_diverge(tmp_path):
    """Mesmo padrão de Cargo/Transporte: dois registros de Controle com
    DOC ORIGEM divergente nunca escolhem um à toa."""
    destino = str(tmp_path / "saida.xlsx")
    item = _item(
        registros_controle=[
            _registro(linha=2, doc_origem="GABPGJ - Gabinete do Procurador-Geral de Justiça"),
            _registro(linha=3, doc_origem="GAECO"),
        ]
    )
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Unidade de Origem"] is None


# --- Quantidade de diárias / Valor das diárias ------------------------------


def test_quantidade_de_diarias_vem_do_controle(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(registros_controle=[_registro(integral=2, parcial=1)])
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Nº Diárias Integral"] == 2
    assert linha["Nº Diárias Parcial"] == 1


def test_valor_das_diarias_e_o_valor_bruto_da_base_ob(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(valor=999.99)
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Valor das Diárias"] == 999.99
    assert isinstance(linha["Valor das Diárias"], float)


# --- Valor/Situação da Passagem (reaproveita valor_passagem_publicavel) ----


@pytest.mark.parametrize(
    "estado,valor,bruto,valor_esperado,situacao_esperada",
    [
        (EstadoPassagem.NUMERICO, 1657.76, "R$ 1.657,76", 1657.76, "NUMÉRICO"),
        (EstadoPassagem.AUSENTE, None, "-", None, "AUSENTE"),
        (EstadoPassagem.NAO_SE_APLICA, None, "NÃO SE APLICA", None, "NÃO SE APLICA"),
        (EstadoPassagem.INVALIDO, None, "#VALOR!", None, "INVÁLIDO"),
    ],
)
def test_passagem_a_b_c_d(tmp_path, estado, valor, bruto, valor_esperado, situacao_esperada):
    destino = str(tmp_path / "saida.xlsx")
    passagem = _linha_passagem(estado=estado, valor=valor, bruto=bruto)
    item = _item(passagem=passagem)
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Valor da Passagem"] == valor_esperado
    assert linha["Situação da Passagem"] == situacao_esperada


def test_passagem_estado_e_sem_correspondencia(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item(passagem=None)
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Valor da Passagem"] is None
    assert linha["Situação da Passagem"] == "SEM CORRESPONDÊNCIA"


def test_valor_da_passagem_nunca_zero_para_estados_sem_valor(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    itens = [
        _item(ob=f"2026OB{i:06d}", dh=f"2026DH{i:06d}", passagem=p)
        for i, p in enumerate(
            [
                _linha_passagem(estado=EstadoPassagem.AUSENTE),
                _linha_passagem(estado=EstadoPassagem.NAO_SE_APLICA),
                _linha_passagem(estado=EstadoPassagem.INVALIDO, bruto="#VALOR!"),
                None,
            ]
        )
    ]
    gerar_xlsx_diarias_nao_publicadas(itens, caminho_destino=destino)

    ws = openpyxl.load_workbook(destino)[_NOME_ABA]
    for linha_idx in range(2, 6):
        assert _valores_linha(ws, linha_idx)["Valor da Passagem"] != 0


# --- Rastreabilidade da decisão (item 11 do plano) --------------------------


def test_responsavel_data_observacao_vem_da_decisao(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    decisao = _decisao(
        autor="servidor.investigado@example.com",
        decidido_em=datetime(2026, 5, 25, 10, 15, 0, tzinfo=timezone.utc),
        observacao="Publicar zero — todas fora do escopo desta competência.",
    )
    item = _item(decisao=decisao)
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Responsável pela Decisão"] == "servidor.investigado@example.com"
    # XLSX não representa timezone — o instante UTC é preservado, só sem o
    # marcador de fuso (ver docstring do escritor).
    assert linha["Data/Hora da Decisão"] == datetime(2026, 5, 25, 10, 15, 0)
    assert linha["Observação/Justificativa"] == "Publicar zero — todas fora do escopo desta competência."


def test_categoria_b_preserva_dados_da_decisao_do_grupo(tmp_path):
    """Item 11: Categoria B pode preservar os dados da decisão (quem
    decidiu sobre a CLASSE), mas a categoria em si já deixa claro que não
    foi uma rejeição individual."""
    destino = str(tmp_path / "saida.xlsx")
    decisao = _decisao(autor="servidor.teste", regra="RN04")
    item = _item(categoria=CategoriaNaoPublicada.DUPLICATA_FISICA_IDENTICA_SUPRIMIDA, decisao=decisao)
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    linha = _valores_linha(openpyxl.load_workbook(destino)[_NOME_ABA], 2)
    assert linha["Categoria"] == "Duplicata física idêntica suprimida"
    assert linha["Responsável pela Decisão"] == "servidor.teste"  # rastreável, mas categoria já esclarece


# --- Privacidade -------------------------------------------------------------


def test_cpf_nunca_aparece_em_nenhuma_celula(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    item = _item()
    gerar_xlsx_diarias_nao_publicadas([item], caminho_destino=destino)

    ws = openpyxl.load_workbook(destino)[_NOME_ABA]
    cabecalhos = [c.value for c in ws[1]]
    assert "CPF" not in cabecalhos
    valores_linha_1 = [ws.cell(row=2, column=i + 1).value for i in range(len(cabecalhos))]
    assert "11111111111" not in [str(v) for v in valores_linha_1]


# --- Estrutura/usabilidade XLSX (item 14 do plano) --------------------------


def test_congelamento_do_cabecalho(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino)

    ws = openpyxl.load_workbook(destino)[_NOME_ABA]
    assert ws.freeze_panes == "A2"


def test_autofiltro_cobre_cabecalho_e_dados(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    gerar_xlsx_diarias_nao_publicadas([_item(), _item(ob="2026OB000002", dh="2026DH000002")], caminho_destino=destino)

    ws = openpyxl.load_workbook(destino)[_NOME_ABA]
    assert ws.auto_filter.ref == "A1:T3"


def test_larguras_de_coluna_sao_razoaveis(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino)

    ws = openpyxl.load_workbook(destino)[_NOME_ABA]
    assert ws.column_dimensions["A"].width > 15  # Categoria — texto longo, nunca cortado


# --- Determinismo / reexecução / no-clobber ----------------------------------


def test_gerar_e_deterministico(tmp_path):
    itens = [_item(ob="2026OB000001"), _item(ob="2026OB000002", dh="2026DH000002")]
    destino_1 = str(tmp_path / "a.xlsx")
    destino_2 = str(tmp_path / "b.xlsx")
    gerar_xlsx_diarias_nao_publicadas(itens, caminho_destino=destino_1)
    gerar_xlsx_diarias_nao_publicadas(itens, caminho_destino=destino_2)

    ws1 = openpyxl.load_workbook(destino_1)[_NOME_ABA]
    ws2 = openpyxl.load_workbook(destino_2)[_NOME_ABA]
    for linha_idx in (2, 3):
        assert _valores_linha(ws1, linha_idx) == _valores_linha(ws2, linha_idx)


def test_nao_sobrescreve_destino_existente(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino)

    with pytest.raises(DestinoJaExiste):
        gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino)


def test_reexecucao_apos_falha_no_clobber_funciona_em_novo_destino(tmp_path):
    destino_1 = str(tmp_path / "saida.xlsx")
    gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino_1)

    destino_2 = str(tmp_path / "saida_v2.xlsx")
    resultado = gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino_2)
    assert os.path.exists(destino_2)
    assert resultado.total_ocorrencias == 1


def test_nenhum_tempfile_ou_trava_orfaos_apos_sucesso(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino)

    assert not glob.glob(str(tmp_path / "*.gerando.lock"))
    assert not glob.glob(str(tmp_path / "tmp*"))


def test_nenhuma_trava_orfa_apos_destino_ja_existente(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino)

    try:
        gerar_xlsx_diarias_nao_publicadas([_item()], caminho_destino=destino)
    except DestinoJaExiste:
        pass

    assert not glob.glob(str(tmp_path / "*.gerando.lock"))
