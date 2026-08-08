"""Modo Complemento de Diárias (RN-14) em `escritor_planilha.preencher` —
coluna O (quantidade parcial complementar) e AVISO só escritas neste modo,
duplicidade por SEI deixa de ser bloqueante, e a auditoria obrigatória de
fórmulas (N/P/R/S/V/Z/AA) antes de gravar. O modo Padrão permanece
byte-a-byte o já homologado (ver `tests/unit/test_escritor_planilha.py`)."""

from datetime import date

from openpyxl import Workbook, load_workbook

from app.domain.entities.diaria import Diaria
from app.domain.entities.modo_processamento import ModoProcessamento
from app.domain.entities.processo import Processo, StatusProcesso
from app.planilha.escritor_planilha import (
    AVISO_COMPLEMENTO_DE_DIARIAS,
    CategoriaResultado,
    preencher,
)

_CABECALHO = [
    "DH", "SEI", "DOC ORIGEM", "CPF", "PREVIDÊNCIA", "NOME", "CARGO",
    "PERÍODO INICIAL", "PERÍODO FINAL", "TRECHO", "TRANSPORTE", "MOTIVO",
    "VALOR DA PASSAGEM", "QUANT. DIÁRIA INTEGRAL", "QUANT. DIÁRIA PARCIAL",
    "VALOR TOTAL DAS DIÁRIAS", "LOCAL", "VALOR DAS DIÁRIAS INTEGRAIS",
    "VALOR DAS DIÁRIAS PARCIAIS", "DMTC", "DMOF", "PASSAR DIA",
    "RESTITUIÇÃO", "COMPROVANTES", "CONFIRMAÇÃO DE DIÁRIAS", "CONCATENAÇÃO",
    "AVISO ", "AVISO AUXILIAR",
]  # A..AB, mesma ordem real confirmada na Planilha Oficial


def _criar_planilha(tmp_path, *, linha_livre_com_formulas: bool, nome_arquivo="planilha.xlsx") -> str:
    """Reproduz, só nas colunas relevantes, a estrutura real confirmada por
    auditoria direta da Planilha Oficial: uma linha 2 já existente (com
    fórmulas de N/P/R/S/V/Z/AA) e uma linha 3 livre (coluna B vazia) —
    com ou sem essas mesmas fórmulas pré-copiadas, conforme o parâmetro."""
    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Controle de Diárias"
    planilha.append(_CABECALHO)

    # Linha 2 — diária já existente (mesmo SEI usado nos testes de
    # duplicidade não-bloqueante abaixo).
    planilha.append(
        [
            "2026DH000100", "19.20.0000.0000001/2026-11", "PJTESTE", 11111111111,
            "=VLOOKUP($D2,Dados!$A:$E,4,FALSE)", "=VLOOKUP($D2,Dados!$A:$E,3,FALSE)",
            "=VLOOKUP($D2,Dados!$A:$E,5,FALSE)", date(2026, 3, 17), date(2026, 3, 19),
            "OURICURI-RECIFE", "CARRO", "Motivo original",
            None, "=I2-H2", 1, "=R2+S2", "PERNAMBUCO",
            "=(VLOOKUP(Q2,'Tabela Valores'!A:C,2,FALSE()))*N2",
            "=(VLOOKUP(Q2,'Tabela Valores'!A:C,3,FALSE()))*O2",
            None, None, "=H2-15", None, None, None,
            "=CONCATENATE(D2,H2,I2)", "=VLOOKUP(Z2,Z$1:AB1,3,0)", "JÁ EXISTE DIÁRIA",
        ]
    )

    if linha_livre_com_formulas:
        # Linha 3 — livre (B vazio), mas já com as fórmulas pré-copiadas
        # (mesmo padrão real auditado nas linhas 1670-1699 da planilha
        # oficial): só isso já habilita o modo Complemento a gravar nela.
        planilha.append(
            [
                None, None, None, None, None, None, None, None, None, None, None, None,
                None, "=I3-H3", None, "=R3+S3", None,
                "=(VLOOKUP(Q3,'Tabela Valores'!A:C,2,FALSE()))*N3",
                "=(VLOOKUP(Q3,'Tabela Valores'!A:C,3,FALSE()))*O3",
                None, None, "=H3-15", None, None, None,
                "=CONCATENATE(D3,H3,I3)", "=VLOOKUP(Z3,Z$1:AB2,3,0)", None,
            ]
        )
    # Quando `linha_livre_com_formulas` é False, a linha 3 simplesmente não
    # existe ainda — mesmo caso real de uma linha "além do intervalo com
    # fórmula pré-copiada" citado na investigação.

    workbook.create_sheet("Dados")
    tabela_valores = workbook.create_sheet("Tabela Valores")
    tabela_valores.append(["PERNAMBUCO", 308.19, 154.09])

    caminho = tmp_path / nome_arquivo
    workbook.save(caminho)
    return str(caminho)


def _diaria_complemento(**overrides) -> Diaria:
    base = dict(
        sei="19.20.0000.0000001/2026-11",
        doc_origem="PJTESTE",
        cpf="11111111111",
        periodo_inicial=date(2026, 4, 13),
        periodo_final=date(2026, 4, 13),
        cidade_origem="OURICURI",
        cidade_destino="RECIFE",
        motivo="Pagamento de valor faltante",
        quantidade_diaria_parcial_complementar=0,
        fonte_quantidade_complementar="Padrão institucional do modo Complemento de Diárias.",
    )
    base.update(overrides)
    return Diaria(**base)


def _processo(diaria: Diaria, arquivo: str = "complemento.pdf") -> Processo:
    return Processo(arquivo_origem=arquivo, status=StatusProcesso.SUCESSO, diaria=diaria)


def test_grava_coluna_aviso_e_quantidade_parcial_no_modo_complemento(tmp_path):
    caminho = _criar_planilha(tmp_path, linha_livre_com_formulas=True)
    diaria = _diaria_complemento(quantidade_diaria_parcial_complementar=0)

    resultado = preencher(caminho, [_processo(diaria)], modo=ModoProcessamento.COMPLEMENTO)

    assert resultado.total_gravado == 1
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["AA3"].value == AVISO_COMPLEMENTO_DE_DIARIAS
    assert planilha["O3"].value == 0


def test_grava_quantidade_parcial_maior_que_zero_quando_documentada(tmp_path):
    caminho = _criar_planilha(tmp_path, linha_livre_com_formulas=True)
    diaria = _diaria_complemento(quantidade_diaria_parcial_complementar=1)

    preencher(caminho, [_processo(diaria)], modo=ModoProcessamento.COMPLEMENTO)

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["O3"].value == 1


def test_modo_padrao_nunca_escreve_aviso_nem_quantidade_parcial(tmp_path):
    caminho = _criar_planilha(tmp_path, linha_livre_com_formulas=True)
    diaria = _diaria_complemento(
        sei="19.20.0000.0000002/2026-22", quantidade_diaria_parcial_complementar=1
    )

    preencher(caminho, [_processo(diaria)])  # modo Padrão (default)

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["O3"].value is None
    assert planilha["AA3"].value == "=VLOOKUP(Z3,Z$1:AB2,3,0)"  # fórmula original intacta


def test_modo_complemento_nao_bloqueia_sei_ja_existente_na_planilha(tmp_path):
    """Achado real: SEI nunca é chave obrigatória para um complemento —
    pode repetir o SEI da diária original."""
    caminho = _criar_planilha(tmp_path, linha_livre_com_formulas=True)
    diaria = _diaria_complemento(sei="19.20.0000.0000001/2026-11")  # mesmo SEI da linha 2

    resultado = preencher(caminho, [_processo(diaria)], modo=ModoProcessamento.COMPLEMENTO)

    assert resultado.total_gravado == 1
    assert resultado.itens[0].categoria == CategoriaResultado.INSERIDO
    assert resultado.itens[0].sei_ja_existente_em_linha == 2


def test_modo_padrao_continua_bloqueando_sei_ja_existente_exatamente_como_hoje(tmp_path):
    """Zero regressão: o comportamento já homologado do modo Padrão não
    muda em nada com a introdução do modo Complemento."""
    caminho = _criar_planilha(tmp_path, linha_livre_com_formulas=True)
    diaria = _diaria_complemento(sei="19.20.0000.0000001/2026-11")

    resultado = preencher(caminho, [_processo(diaria)])  # modo Padrão

    assert resultado.total_gravado == 0
    assert resultado.itens[0].categoria == CategoriaResultado.DUPLICADO_NA_PLANILHA
    assert resultado.itens[0].linha == 2


def test_modo_complemento_aborta_gravacao_quando_linha_livre_sem_formulas(tmp_path):
    caminho = _criar_planilha(tmp_path, linha_livre_com_formulas=False)
    diaria = _diaria_complemento()

    resultado = preencher(caminho, [_processo(diaria)], modo=ModoProcessamento.COMPLEMENTO)

    assert resultado.total_gravado == 0
    assert resultado.itens[0].categoria == CategoriaResultado.LINHA_SEM_FORMULAS_NECESSARIAS
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["B3"].value is None  # nada foi gravado na linha


def test_modo_complemento_preserva_linha_original_intacta(tmp_path):
    caminho = _criar_planilha(tmp_path, linha_livre_com_formulas=True)
    diaria = _diaria_complemento()

    preencher(caminho, [_processo(diaria)], modo=ModoProcessamento.COMPLEMENTO)

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["B2"].value == "19.20.0000.0000001/2026-11"
    assert planilha["L2"].value == "Motivo original"
    assert planilha["AA2"].value == "=VLOOKUP(Z2,Z$1:AB1,3,0)"
    assert planilha["AB2"].value == "JÁ EXISTE DIÁRIA"
