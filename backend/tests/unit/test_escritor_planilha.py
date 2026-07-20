import os
from dataclasses import replace
from datetime import date

from openpyxl import load_workbook

from app.domain.entities.diaria import Diaria
from app.domain.entities.fonte_doc_origem import FonteDocOrigem
from app.domain.entities.processo import Processo, StatusProcesso
from app.domain.entities.tipo_processo import TipoProcesso
from app.planilha.escritor_planilha import (
    CategoriaResultado,
    ResultadoPreenchimento,
    _localizar_proxima_linha_livre,
    preencher,
)


def _diaria(sei: str = "19.20.0000.0000001/2025-99") -> Diaria:
    return Diaria(
        sei=sei,
        doc_origem="PJNOVO",
        cpf="99999999999",
        periodo_inicial=date(2025, 3, 1),
        periodo_final=date(2025, 3, 3),
        cidade_origem="RECIFE",
        cidade_destino="OLINDA",
        motivo="Diligência de teste",
    )


def _processo(arquivo: str = "arquivo.pdf", sei: str = "19.20.0000.0000001/2025-99") -> Processo:
    return Processo(arquivo_origem=arquivo, status=StatusProcesso.SUCESSO, diaria=_diaria(sei))


def test_localiza_proxima_linha_livre_ignora_coluna_a(caminho_planilha_teste):
    caminho = caminho_planilha_teste()
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]

    # Linha 3 tem A vazio mas B preenchido — não pode ser considerada livre.
    assert _localizar_proxima_linha_livre(planilha) == 4


def test_preencher_grava_somente_colunas_esperadas(caminho_planilha_teste):
    caminho = caminho_planilha_teste()

    resultado = preencher(caminho, [_processo()])

    assert resultado.total_gravado == 1
    assert resultado.duplicadas_ignoradas == 0
    assert resultado.linha_inicial == 4
    assert resultado.linha_final == 4

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]

    linha = 4
    assert planilha[f"B{linha}"].value == "19.20.0000.0000001/2025-99"
    assert planilha[f"C{linha}"].value == "PJNOVO"
    assert planilha[f"D{linha}"].value == 99999999999
    assert planilha[f"H{linha}"].value.date() == date(2025, 3, 1)
    assert planilha[f"I{linha}"].value.date() == date(2025, 3, 3)
    assert planilha[f"J{linha}"].value == "RECIFE - OLINDA"
    assert planilha[f"L{linha}"].value == "Diligência de teste"

    # E/F/G são fórmulas, nunca escritas. A/K/Q (Sprint 03) ficam vazias
    # porque esta Diaria de teste não tem DH/transporte/estado (campos
    # opcionais, não informados aqui).
    for coluna in ("A", "E", "F", "G", "K", "Q"):
        assert planilha[f"{coluna}{linha}"].value is None


def test_preencher_item_gravado_carrega_a_linha_escrita(caminho_planilha_teste):
    caminho = caminho_planilha_teste()

    resultado = preencher(caminho, [_processo()])

    assert len(resultado.gravados) == 1
    item = resultado.gravados[0]
    assert item.categoria is CategoriaResultado.INSERIDO
    assert item.linha == 4


def test_preencher_nao_altera_linhas_anteriores_nem_estrutura(caminho_planilha_teste):
    caminho = caminho_planilha_teste()
    workbook_antes = load_workbook(caminho)
    planilha_antes = workbook_antes["Controle de Diárias"]
    linha2_antes = [planilha_antes.cell(row=2, column=c).value for c in range(1, 13)]
    linha3_antes = [planilha_antes.cell(row=3, column=c).value for c in range(1, 13)]
    abas_antes = workbook_antes.sheetnames
    cabecalho_antes = [planilha_antes.cell(row=1, column=c).value for c in range(1, 13)]

    preencher(caminho, [_processo()])

    workbook_depois = load_workbook(caminho)
    planilha_depois = workbook_depois["Controle de Diárias"]

    assert workbook_depois.sheetnames == abas_antes
    assert [planilha_depois.cell(row=1, column=c).value for c in range(1, 13)] == cabecalho_antes
    assert [planilha_depois.cell(row=2, column=c).value for c in range(1, 13)] == linha2_antes
    assert [planilha_depois.cell(row=3, column=c).value for c in range(1, 13)] == linha3_antes


def test_preencher_multiplos_processos_ocupa_linhas_sequenciais(caminho_planilha_teste):
    caminho = caminho_planilha_teste()

    resultado = preencher(
        caminho,
        [
            _processo("a.pdf", "19.20.0000.0000001/2025-99"),
            _processo("b.pdf", "19.20.0000.0000002/2025-99"),
            _processo("c.pdf", "19.20.0000.0000003/2025-99"),
        ],
    )

    assert resultado.total_gravado == 3
    assert resultado.linha_inicial == 4
    assert resultado.linha_final == 6
    assert [item.linha for item in resultado.gravados] == [4, 5, 6]

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    for linha in (4, 5, 6):
        assert planilha[f"B{linha}"].value is not None


def test_preencher_lista_vazia_nao_grava_nada(caminho_planilha_teste):
    caminho = caminho_planilha_teste()

    resultado = preencher(caminho, [])

    assert resultado == ResultadoPreenchimento(
        total_gravado=0, duplicadas_ignoradas=0, linha_inicial=0, linha_final=0
    )


def test_preencher_itens_preserva_ordem_de_entrada(caminho_planilha_teste):
    """A tabela de resultados da tela de conclusão depende de `itens` estar
    na mesma ordem/quantidade dos processos enviados, 1:1 — inclusive quando
    o mesmo nome de arquivo aparece mais de uma vez (regressão do Bug 1)."""
    caminho = caminho_planilha_teste()

    resultado = preencher(
        caminho,
        [
            _processo("mesmo_nome.pdf", "19.20.0000.0000030/2025-11"),
            _processo("mesmo_nome.pdf", "19.20.0000.0000030/2025-11"),
            _processo("outro.pdf", "19.20.0000.0000031/2025-11"),
        ],
    )

    assert len(resultado.itens) == 3
    assert resultado.itens[0].categoria is CategoriaResultado.INSERIDO
    assert resultado.itens[1].categoria is CategoriaResultado.DUPLICADO_NO_LOTE
    assert resultado.itens[2].categoria is CategoriaResultado.INSERIDO


# --- BUG 1 — Duplicidade no lote ------------------------------------------------


def test_bug1_mesmo_pdf_anexado_duas_vezes_so_grava_uma_diaria(caminho_planilha_teste):
    """Reprodução exata do BUG 1: o mesmo PDF (mesmo SEI) é anexado duas
    vezes no mesmo lote. Só a primeira ocorrência pode virar linha nova; a
    segunda tem que ser classificada como duplicado no lote, sem gravar."""
    caminho = caminho_planilha_teste()

    resultado = preencher(
        caminho,
        [
            _processo("relatorio.pdf", "19.20.0000.0000009/2025-99"),
            _processo("relatorio.pdf", "19.20.0000.0000009/2025-99"),
        ],
    )

    assert resultado.total_gravado == 1
    assert resultado.duplicadas_ignoradas == 1
    assert len(resultado.duplicados_no_lote) == 1
    assert len(resultado.duplicados_na_planilha) == 0
    assert resultado.duplicados_no_lote[0].arquivo_origem == "relatorio.pdf"
    assert resultado.duplicados_no_lote[0].sei == "19.20.0000.0000009/2025-99"
    assert resultado.duplicados_no_lote[0].linha is None

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["B4"].value == "19.20.0000.0000009/2025-99"
    # Nenhuma segunda linha foi criada para a duplicata.
    assert planilha["B5"].value is None


def test_bug1_pdfs_diferentes_mesmo_sei_so_grava_uma_diaria(caminho_planilha_teste):
    """Variante do BUG 1: dois arquivos com nomes diferentes, mas o mesmo
    número de processo SEI — deve continuar valendo 'um SEI, uma diária'."""
    caminho = caminho_planilha_teste()

    resultado = preencher(
        caminho,
        [
            _processo("processo_via_email.pdf", "19.20.0000.0000010/2025-99"),
            _processo("processo_via_sei.pdf", "19.20.0000.0000010/2025-99"),
            _processo("processo_scanner.pdf", "19.20.0000.0000010/2025-99"),
        ],
    )

    assert resultado.total_gravado == 1
    assert resultado.duplicadas_ignoradas == 2
    assert len(resultado.duplicados_no_lote) == 2
    assert {item.arquivo_origem for item in resultado.duplicados_no_lote} == {
        "processo_via_sei.pdf",
        "processo_scanner.pdf",
    }

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["B4"].value == "19.20.0000.0000010/2025-99"
    assert planilha["B5"].value is None


# --- BUG 2 — Duplicidade na planilha ---------------------------------------------


def test_bug2_sei_ja_existente_na_planilha_nao_e_gravado_e_informa_a_linha(
    caminho_planilha_teste,
):
    """Reprodução exata do BUG 2: a linha 2 da fixture já tem o SEI
    19.20.0000.0000001/2025-00. Uma nova diária com o mesmo SEI não pode
    virar uma segunda linha, e o motivo tem que apontar a linha 2."""
    caminho = caminho_planilha_teste()

    resultado = preencher(caminho, [_processo("duplicado.pdf", "19.20.0000.0000001/2025-00")])

    assert resultado.total_gravado == 0
    assert resultado.duplicadas_ignoradas == 1
    assert len(resultado.duplicados_na_planilha) == 1
    assert len(resultado.duplicados_no_lote) == 0

    duplicado = resultado.duplicados_na_planilha[0]
    assert duplicado.arquivo_origem == "duplicado.pdf"
    assert duplicado.sei == "19.20.0000.0000001/2025-00"
    assert duplicado.linha == 2

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    # Nenhuma linha nova foi criada.
    assert planilha["B4"].value is None


def test_bug2_sei_da_planilha_tem_prioridade_sobre_duplicidade_no_lote(caminho_planilha_teste):
    """Se um SEI já existe na planilha E aparece duas vezes no lote, as duas
    ocorrências são 'duplicado na planilha' (não uma delas 'duplicado no
    lote') — a causa raiz é a mesma para as duas."""
    caminho = caminho_planilha_teste()
    sei_existente = "19.20.0000.0000001/2025-00"

    resultado = preencher(
        caminho,
        [
            _processo("a.pdf", sei_existente),
            _processo("b.pdf", sei_existente),
        ],
    )

    assert resultado.total_gravado == 0
    assert len(resultado.duplicados_na_planilha) == 2
    assert len(resultado.duplicados_no_lote) == 0
    assert all(item.linha == 2 for item in resultado.duplicados_na_planilha)


def test_preencher_nao_salva_arquivo_quando_tudo_e_duplicado(caminho_planilha_teste):
    caminho = caminho_planilha_teste()
    mtime_antes = os.path.getmtime(caminho)

    preencher(caminho, [_processo("duplicado.pdf", "19.20.0000.0000001/2025-00")])

    mtime_depois = os.path.getmtime(caminho)
    assert mtime_antes == mtime_depois


# --- Sprint 03 — DH (A), transporte (K) e local (Q) ------------------------------


def test_preencher_grava_dh_transporte_e_local_quando_presentes(caminho_planilha_teste):
    caminho = caminho_planilha_teste()
    diaria = Diaria(
        sei="19.20.0000.0000050/2025-99",
        doc_origem="UNIDADE TESTE",
        cpf="99999999999",
        periodo_inicial=date(2025, 3, 1),
        periodo_final=date(2025, 3, 1),
        cidade_origem="RECIFE",
        cidade_destino="GARANHUNS",
        motivo="Teste",
        documento_habil="2025DH000102",
        transporte="CARRO",
        estado_destino="PERNAMBUCO",
    )
    processo = Processo(arquivo_origem="a.pdf", status=StatusProcesso.SUCESSO, diaria=diaria)

    preencher(caminho, [processo])

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["A4"].value == "2025DH000102"
    assert planilha["K4"].value == "CARRO"
    assert planilha["Q4"].value == "PERNAMBUCO"
    # E/F/G continuam sendo só fórmula, nunca escritas por nós.
    for coluna in ("E", "F", "G"):
        assert planilha[f"{coluna}4"].value is None


def test_preencher_nao_escreve_dh_transporte_local_quando_ausentes(caminho_planilha_teste):
    caminho = caminho_planilha_teste()

    preencher(caminho, [_processo()])  # DH/transporte/estado = None por padrão

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["A4"].value is None
    assert planilha["K4"].value is None
    assert planilha["Q4"].value is None


# --- Sprint 04 — Diária de Membro: coluna P (valor total das diárias) -----------


def test_preencher_grava_coluna_p_para_diaria_de_membro(caminho_planilha_teste):
    caminho = caminho_planilha_teste()
    diaria = Diaria(
        sei="19.20.0000.0000060/2026-99",
        doc_origem="GABPGJ",
        cpf="99999999999",
        periodo_inicial=date(2026, 2, 11),
        periodo_final=date(2026, 2, 11),
        cidade_origem="RECIFE",
        cidade_destino="BRASÍLIA",
        motivo="Diária de membro de teste",
        tipo_processo=TipoProcesso.DIARIA_MEMBRO,
        valor_total_diarias="652,96",
    )
    processo = Processo(arquivo_origem="membro.pdf", status=StatusProcesso.SUCESSO, diaria=diaria)

    preencher(caminho, [processo])

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["P4"].value == 652.96
    assert "R$" in planilha["P4"].number_format
    # Colunas já homologadas continuam preenchidas normalmente.
    assert planilha["B4"].value == "19.20.0000.0000060/2026-99"


def test_preencher_nao_grava_coluna_p_para_diaria_padrao(caminho_planilha_teste):
    """A coluna P só é preenchida para Diária de Membro — mesmo que
    valor_total_diarias esteja presente por engano, um processo com
    tipo_processo PADRAO nunca escreve na coluna P."""
    caminho = caminho_planilha_teste()
    diaria = Diaria(
        sei="19.20.0000.0000061/2026-99",
        doc_origem="PJTESTE",
        cpf="99999999999",
        periodo_inicial=date(2026, 2, 11),
        periodo_final=date(2026, 2, 11),
        cidade_origem="RECIFE",
        cidade_destino="OLINDA",
        motivo="Diária padrão de teste",
        tipo_processo=TipoProcesso.PADRAO,
        valor_total_diarias="652,96",
    )
    processo = Processo(arquivo_origem="padrao.pdf", status=StatusProcesso.SUCESSO, diaria=diaria)

    preencher(caminho, [processo])

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["P4"].value is None


def test_preencher_diaria_padrao_sem_valor_tambem_nao_grava_coluna_p(caminho_planilha_teste):
    caminho = caminho_planilha_teste()

    preencher(caminho, [_processo()])  # tipo_processo = PADRAO por padrão

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["P4"].value is None


# --- Sprint 06 — DOC DE ORIGEM: histórico da própria planilha do DIMCAC (Prioridade 1) ---


def test_cpf_com_historico_reaproveita_doc_origem_ja_homologado(caminho_planilha_teste):
    """A linha 2 da fixture já tem CPF 11111111111 com DOC ORIGEM 'PJTESTE'.
    Uma nova diária do mesmo CPF, mas com um doc_origem diferente extraído
    do PDF, deve gravar o valor histórico — não o extraído."""
    caminho = caminho_planilha_teste()
    diaria = Diaria(
        sei="19.20.0000.0000099/2025-99",
        doc_origem="TEXTO EXTRAIDO DO PDF, DIFERENTE DO HISTORICO",
        cpf="11111111111",
        periodo_inicial=date(2025, 3, 1),
        periodo_final=date(2025, 3, 3),
        cidade_origem="RECIFE",
        cidade_destino="OLINDA",
        motivo="Diligência de teste",
    )
    processo = Processo(arquivo_origem="a.pdf", status=StatusProcesso.SUCESSO, diaria=diaria)

    resultado = preencher(caminho, [processo])

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["C4"].value == "PJTESTE"
    assert resultado.gravados[0].doc_origem_fonte is FonteDocOrigem.HISTORICO_DIMCAC


def test_cpf_sem_historico_usa_extracao_do_pdf_como_fallback(caminho_planilha_teste):
    """CPF 99999999999 não existe na fixture — nenhuma linha na planilha
    reproduz esse CPF. A lógica de extração do PDF, já homologada, continua
    decidindo o valor normalmente (Prioridade 2, nunca removida)."""
    caminho = caminho_planilha_teste()

    resultado = preencher(caminho, [_processo()])

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["C4"].value == "PJNOVO"
    assert resultado.gravados[0].doc_origem_fonte is FonteDocOrigem.EXTRACAO_PDF


def test_cpf_com_zero_a_esquerda_casa_historico_apesar_da_coluna_ser_numerica(
    caminho_planilha_teste,
):
    """A coluna D é gravada como número (ver `_escrever_celula(..., cpf=True)`),
    então um CPF com zero à esquerda perde esse zero ao ser lido de volta. A
    normalização precisa devolver o zero para casar corretamente com o CPF
    de 11 dígitos extraído do PDF."""
    caminho = caminho_planilha_teste("planilha_cpf_zero.xlsx")

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    planilha["D2"].value = 1234567890  # 10 dígitos: zero à esquerda foi perdido
    planilha["C2"].value = "PJZERO"
    workbook.save(caminho)

    diaria = Diaria(
        sei="19.20.0000.0000098/2025-99",
        doc_origem="EXTRAIDO",
        cpf="01234567890",
        periodo_inicial=date(2025, 3, 1),
        periodo_final=date(2025, 3, 3),
        cidade_origem="RECIFE",
        cidade_destino="OLINDA",
        motivo="Diligência de teste",
    )
    processo = Processo(arquivo_origem="a.pdf", status=StatusProcesso.SUCESSO, diaria=diaria)

    resultado = preencher(caminho, [processo])

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["C4"].value == "PJZERO"
    assert resultado.gravados[0].doc_origem_fonte is FonteDocOrigem.HISTORICO_DIMCAC


# --- Sprint de Correção Crítica — documentação do 2º passo do bug de CPF --


def test_cpf_truncado_se_chegasse_aqui_seria_zero_padded_disfarcando_o_erro(
    caminho_planilha_teste,
):
    """Não é uma regressão a corrigir NESTE arquivo -- o formato de gravação
    do CPF (int + máscara '00000000000') é homologado e permanece
    intencionalmente inalterado (ver `_escrever_celula(..., cpf=True)`).
    Este teste documenta, com o código real e sem alterações, o 2º passo do
    bug do processo 19.20.0001.0000108/2026-08: SE um fragmento truncado de
    CPF (aqui, '333333', o mesmo fragmento de 6 dígitos que a regex do
    parser produz quando o valor é interrompido por uma quebra de
    linha/página) chegasse até `preencher`, o formato numérico já homologado
    o transformaria em '00000333333' -- o valor exato relatado no bug --
    porque `int('333333')` seguido da máscara de 11 dígitos não tem como
    saber que faltam dígitos. É exatamente por isso que a correção real fica
    na validação (`validacao_diaria._cpf_valido`, ver `test_sei_parser_v1.py`
    e `test_validacao_diaria.py`): impedir que esse fragmento chegue até
    aqui, nunca mudar o que acontece aqui."""
    caminho = caminho_planilha_teste()
    diaria = _diaria("19.20.0001.0000108/2026-08")
    diaria_com_fragmento = replace(diaria, cpf="333333")
    processo = Processo(
        arquivo_origem="fragmento.pdf",
        status=StatusProcesso.SUCESSO,
        diaria=diaria_com_fragmento,
    )

    preencher(caminho, [processo])

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    # Sem zero à esquerda a mais no valor bruto (é int, não string) --
    # mas com a máscara aplicada, o Excel EXIBE exatamente "00000333333".
    assert planilha["D4"].value == 333333
    assert planilha["D4"].number_format == "00000000000"


def test_historico_nao_considera_linhas_do_proprio_lote(caminho_planilha_teste):
    """Duas diárias do mesmo CPF no mesmo lote: a segunda não deve
    reaproveitar o doc_origem gravado pela primeira nesta mesma execução —
    só linhas que já existiam na planilha antes do lote contam como
    histórico."""
    caminho = caminho_planilha_teste()

    resultado = preencher(
        caminho,
        [
            _processo("a.pdf", "19.20.0000.0000097/2025-99"),
            _processo("b.pdf", "19.20.0000.0000096/2025-99"),
        ],
    )

    assert [item.doc_origem_fonte for item in resultado.gravados] == [
        FonteDocOrigem.EXTRACAO_PDF,
        FonteDocOrigem.EXTRACAO_PDF,
    ]
