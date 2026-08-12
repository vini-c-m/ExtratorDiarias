"""Micro-sprint A — leitura dinâmica da aba `Tabela Valores` real (nunca
hardcoded). Casos J/K/L do pedido: tabela indisponível/estruturalmente
inválida, UF ausente, valor inválido na tabela."""

from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.consulta.repositorio_tabela_valores import (
    TabelaValoresEstruturalmenteInvalida,
    carregar_tabela_valores,
)


def _planilha_tabela_valores(linhas: list[tuple] | None = None):
    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Tabela Valores"
    planilha.append([None, None, None])
    planilha.append([None, None, None])
    planilha.append([None, "INTEGRAL", "PARCIAL"])
    for linha in linhas or []:
        planilha.append(list(linha))
    return planilha


def test_tabela_valores_real_carregada_corretamente():
    """Estrutura confirmada contra a Planilha Oficial real (2026-08-12):
    cabeçalho na linha 3, dados a partir da linha 4."""
    planilha = _planilha_tabela_valores(
        [
            ("PERNAMBUCO", 308.19, 154.09),
            ("CEARÁ", 539.33, 308.19),
            ("SÃO PAULO", 770.47, 385.24),
        ]
    )
    tabela = carregar_tabela_valores(planilha)

    assert set(tabela.keys()) == {"PERNAMBUCO", "CEARÁ", "SÃO PAULO"}
    assert tabela["PERNAMBUCO"].integral == Decimal("308.19")
    assert tabela["PERNAMBUCO"].parcial == Decimal("154.09")


def test_uf_normalizada_maiuscula_e_sem_espacos_nas_pontas():
    """Achado real da auditoria: 'RORAIMA ' tem espaço final gravado na
    própria célula da Planilha Oficial — a normalização acontece do lado da
    leitura, sem alterar o dado original."""
    planilha = _planilha_tabela_valores([("roraima ", 539.33, 308.19)])
    tabela = carregar_tabela_valores(planilha)
    assert "RORAIMA" in tabela


def test_linha_em_branco_no_meio_e_ignorada_em_silencio():
    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Tabela Valores"
    planilha.append([None, None, None])
    planilha.append([None, None, None])
    planilha.append([None, "INTEGRAL", "PARCIAL"])
    planilha.append(["PERNAMBUCO", 308.19, 154.09])
    planilha.append([None, None, None])  # linha em branco no meio
    planilha.append(["CEARÁ", 539.33, 308.19])

    tabela = carregar_tabela_valores(planilha)
    assert set(tabela.keys()) == {"PERNAMBUCO", "CEARÁ"}


def test_j_tabela_estruturalmente_vazia_levanta_erro_explicito():
    """Nenhuma linha de UF/região reconhecível — mesma aba criada vazia por
    `workbook.create_sheet("Tabela Valores")` na fixture padrão de
    `caminho_planilha_conferencia_teste`."""
    planilha = _planilha_tabela_valores([])
    with pytest.raises(TabelaValoresEstruturalmenteInvalida):
        carregar_tabela_valores(planilha)


def test_l_valor_invalido_na_tabela_levanta_erro_nunca_ignora_a_linha():
    """Um valor unitário ilegível (ex.: texto sem relação com número) NUNCA
    é tratado como "linha vazia" — precisa interromper a leitura com um
    erro explícito, nunca ser descartado em silêncio (P0 da Micro-sprint A:
    um valor financeiro errado não pode ser confundido com um valor
    ausente)."""
    planilha = _planilha_tabela_valores([("PERNAMBUCO", "TEXTO INVALIDO", 154.09)])
    with pytest.raises(TabelaValoresEstruturalmenteInvalida):
        carregar_tabela_valores(planilha)


def test_valor_com_prefixo_r_dinheiro_ainda_e_lido_corretamente():
    """Mesma tolerância já homologada em `repositorio_diarias._valor_como_numero`
    (achado real: 3 linhas da Planilha Oficial gravadas como texto
    'R$ 154,09') — reaproveitada aqui, nunca uma segunda regra de parsing."""
    planilha = _planilha_tabela_valores([("PERNAMBUCO", "R$ 308,19", "R$ 154,09")])
    tabela = carregar_tabela_valores(planilha)
    assert tabela["PERNAMBUCO"].integral == Decimal("308.19")
