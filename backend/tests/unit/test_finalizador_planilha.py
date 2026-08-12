from openpyxl import load_workbook

from app.planilha.finalizador_planilha import (
    EstadoLinhaFinalizacao,
    classificar_estado,
    gravar_finalizacao,
    localizar_candidatos,
)


def _planilha_com_linha_extra(caminho_planilha_teste, sei: str, x=None, y=None, linha: int = 4):
    caminho = caminho_planilha_teste()
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    planilha[f"B{linha}"] = sei
    planilha[f"D{linha}"] = 33333333333
    planilha[f"J{linha}"] = "RECIFE-OLINDA"
    if x is not None:
        planilha[f"X{linha}"] = x
    if y is not None:
        planilha[f"Y{linha}"] = y
    workbook.save(caminho)
    return caminho


# --- classificar_estado -----------------------------------------------------


def test_classificar_estado_vazio_esta_pronta():
    assert classificar_estado(None, None) is EstadoLinhaFinalizacao.PRONTA
    assert classificar_estado("", "") is EstadoLinhaFinalizacao.PRONTA


def test_classificar_estado_sim_sim_ja_finalizada():
    assert classificar_estado("SIM", "SIM") is EstadoLinhaFinalizacao.JA_FINALIZADA
    # Achado real na Planilha Oficial: variações de caixa já digitadas manualmente.
    assert classificar_estado("sim", "SIM") is EstadoLinhaFinalizacao.JA_FINALIZADA


def test_classificar_estado_incompleta_quando_so_uma_coluna_tem_sim():
    """Regra definitiva (Sprint pré-build, 2026-08-11): exatamente uma
    coluna "SIM" e a outra vazia NÃO é mais PARCIAL — é INCOMPLETA, um
    estado gravável (só falta completar a coluna vazia)."""
    assert classificar_estado("SIM", None) is EstadoLinhaFinalizacao.INCOMPLETA
    assert classificar_estado(None, "SIM") is EstadoLinhaFinalizacao.INCOMPLETA
    assert classificar_estado("sim", "") is EstadoLinhaFinalizacao.INCOMPLETA


def test_classificar_estado_parcial_para_valor_fora_do_padrao():
    # Achados reais: "SM", "SI M" — nunca tratados como já finalizado nem
    # sobrescritos silenciosamente.
    assert classificar_estado("SM", "SM") is EstadoLinhaFinalizacao.PARCIAL
    assert classificar_estado("NÃO", "NÃO") is EstadoLinhaFinalizacao.PARCIAL


# --- localizar_candidatos ----------------------------------------------------


def test_localizar_candidatos_sei_inexistente_devolve_lista_vazia(caminho_planilha_teste):
    caminho = caminho_planilha_teste()
    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert localizar_candidatos(planilha, "19.20.9999.9999999/2099-99") == []


def test_localizar_candidatos_uma_linha(caminho_planilha_teste):
    sei = "19.20.0001.0000001/2026-01"
    caminho = _planilha_com_linha_extra(caminho_planilha_teste, sei, x=None, y=None)
    planilha = load_workbook(caminho)["Controle de Diárias"]

    candidatos = localizar_candidatos(planilha, sei)

    assert len(candidatos) == 1
    assert candidatos[0].linha == 4
    assert candidatos[0].estado is EstadoLinhaFinalizacao.PRONTA
    assert candidatos[0].cpf == "33333333333"
    assert candidatos[0].trecho == "RECIFE-OLINDA"


def test_localizar_candidatos_multiplas_linhas_mesmo_sei(caminho_planilha_teste):
    """Reproduz diária + complemento: mesmo SEI em duas linhas."""
    sei = "19.20.0001.0000002/2026-02"
    caminho = caminho_planilha_teste()
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    planilha["B4"] = sei
    planilha["B5"] = sei
    workbook.save(caminho)

    planilha = load_workbook(caminho)["Controle de Diárias"]
    candidatos = localizar_candidatos(planilha, sei)

    assert [c.linha for c in candidatos] == [4, 5]


# --- gravar_finalizacao ------------------------------------------------------


def test_gravar_finalizacao_grava_somente_x_e_y(caminho_planilha_teste):
    sei = "19.20.0001.0000003/2026-03"
    caminho = _planilha_com_linha_extra(caminho_planilha_teste, sei)
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]

    gravar_finalizacao(planilha, 4)
    workbook.save(caminho)

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["X4"].value == "SIM"
    assert planilha["Y4"].value == "SIM"
    # Nenhuma outra coluna desta linha foi tocada.
    assert planilha["B4"].value == sei
    assert planilha["D4"].value == 33333333333
    assert planilha["J4"].value == "RECIFE-OLINDA"


def test_gravar_finalizacao_linha_incompleta_grava_somente_a_coluna_faltante(caminho_planilha_teste):
    """Regra definitiva: X já "SIM" — só Y precisa ser gravado; X nunca é
    reescrito (nem que fosse, o valor seria idêntico — mas a regra pede
    explicitamente não tocar na célula que já está certa)."""
    sei = "19.20.0001.0000006/2026-06"
    caminho = _planilha_com_linha_extra(caminho_planilha_teste, sei, x="SIM", y=None)
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]

    gravar_finalizacao(planilha, 4)
    workbook.save(caminho)

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["X4"].value == "SIM"
    assert planilha["Y4"].value == "SIM"


def test_gravar_finalizacao_linha_incompleta_outro_lado_grava_somente_a_coluna_faltante(
    caminho_planilha_teste,
):
    sei = "19.20.0001.0000007/2026-07"
    caminho = _planilha_com_linha_extra(caminho_planilha_teste, sei, x=None, y="SIM")
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]

    gravar_finalizacao(planilha, 4)
    workbook.save(caminho)

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["X4"].value == "SIM"
    assert planilha["Y4"].value == "SIM"
