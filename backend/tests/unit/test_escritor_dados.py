import os

import pytest
from openpyxl import load_workbook

from app.planilha.escritor_dados import (
    EstadoCadastro,
    _localizar_proxima_linha_livre_dados,
    cadastrar_servidor,
)
from app.planilha import escritor_dados


def _cadastrar(caminho, cpf="111.222.333-44", **overrides):
    parametros = {
        "cpf": cpf,
        "nome": "Novo Servidor de Teste",
        "matricula": "9999",
        "cargo": "Analista Judiciário",
    }
    parametros.update(overrides)
    return cadastrar_servidor(caminho, **parametros)


def test_localiza_proxima_linha_livre_por_coluna_a(caminho_planilha_conferencia_teste):
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    planilha = workbook["Dados"]

    # Fixture já tem 2 linhas de dado (2 e 3) — a próxima livre é a 4.
    assert _localizar_proxima_linha_livre_dados(planilha) == 4


def test_cadastro_bem_sucedido_grava_somente_as_8_colunas_da_linha_nova(
    caminho_planilha_conferencia_teste,
):
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(
        caminho,
        previdencia="INSS",
        orgao_origem="Órgão de Teste",
        regime_previdencia="RGPS",
        situacao="Em exercício",
    )

    assert resultado.estado is EstadoCadastro.SALVO
    assert resultado.linha == 4
    assert resultado.cpf_normalizado == "11122233344"

    workbook = load_workbook(caminho)
    planilha = workbook["Dados"]
    assert planilha["A4"].value == 11122233344
    assert planilha["A4"].number_format == "00000000000"
    # Matrícula gravada como número puro (nunca texto), com o
    # `number_format` real da coluna B (auditoria 2026-08-10 contra a
    # Planilha Oficial: "000000\-0" em 1622/1622 linhas reais).
    assert planilha["B4"].value == 9999
    assert planilha["B4"].number_format == "000000\\-0"
    assert planilha["C4"].value == "Novo Servidor de Teste"
    assert planilha["D4"].value == "INSS"
    assert planilha["E4"].value == "Analista Judiciário"
    assert planilha["F4"].value == "Órgão de Teste"
    assert planilha["G4"].value == "RGPS"
    assert planilha["H4"].value == "Em exercício"


def test_cadastro_sem_campos_opcionais_deixa_colunas_f_g_h_vazias(
    caminho_planilha_conferencia_teste,
):
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(caminho)

    assert resultado.estado is EstadoCadastro.SALVO
    workbook = load_workbook(caminho)
    planilha = workbook["Dados"]
    for coluna in ("D", "F", "G", "H"):
        assert planilha[f"{coluna}4"].value is None


def test_cadastro_preserva_dados_dos_registros_existentes_mesmo_reordenando(
    caminho_planilha_conferencia_teste,
):
    """Desde a reordenação por NOME (item 4, "Ajuste Funcional + UX"), a
    POSIÇÃO (linha) de um registro existente pode mudar — "João da Silva"
    passa a vir antes de "Maria Aparecida Souza" alfabeticamente, por
    exemplo. O que nunca pode acontecer é PERDA/CORRUPÇÃO de dado: as 8
    colunas de cada registro pré-existente continuam intactas, só talvez
    em outra linha — nunca ordenamos só a coluna do nome, sempre a linha
    inteira junto."""
    caminho = caminho_planilha_conferencia_teste()
    workbook_antes = load_workbook(caminho)
    planilha_antes = workbook_antes["Dados"]
    registros_antes = {
        planilha_antes.cell(row=linha, column=1).value: [
            planilha_antes.cell(row=linha, column=c).value for c in range(1, 9)
        ]
        for linha in (2, 3)
    }

    _cadastrar(caminho)

    workbook_depois = load_workbook(caminho)
    planilha_depois = workbook_depois["Dados"]
    # Região cadastral inteira agora tem 3 registros (2 originais + 1
    # novo) — procura cada CPF original em QUALQUER linha da região,
    # nunca assume que ficou onde estava.
    registros_depois_por_cpf = {}
    for linha in range(2, 5):
        cpf_linha = planilha_depois.cell(row=linha, column=1).value
        registros_depois_por_cpf[cpf_linha] = [
            planilha_depois.cell(row=linha, column=c).value for c in range(1, 9)
        ]

    for cpf_original, valores_originais in registros_antes.items():
        assert registros_depois_por_cpf[cpf_original] == valores_originais


def test_cadastro_nao_altera_controle_de_diarias_nem_estrutura_do_workbook(
    caminho_planilha_conferencia_teste,
):
    caminho = caminho_planilha_conferencia_teste()
    workbook_antes = load_workbook(caminho)
    abas_antes = workbook_antes.sheetnames
    controle_antes = workbook_antes["Controle de Diárias"]
    linha2_controle_antes = [controle_antes.cell(row=2, column=c).value for c in range(1, 13)]

    _cadastrar(caminho)

    workbook_depois = load_workbook(caminho)
    assert workbook_depois.sheetnames == abas_antes
    controle_depois = workbook_depois["Controle de Diárias"]
    linha2_controle_depois = [controle_depois.cell(row=2, column=c).value for c in range(1, 13)]
    assert linha2_controle_depois == linha2_controle_antes


def test_cadastro_preserva_cache_de_formula_em_controle_de_diarias(
    caminho_planilha_conferencia_teste, injetar_cache_formula
):
    """Prova a garantia mais importante deste módulo (mesma classe de risco
    do P0 de integridade financeira, agora numa porta lateral): gravar uma
    linha nova em `Dados` reescreve o arquivo inteiro (`workbook.save()`
    sempre reescreve todas as abas) — sem a reinjeção de cache, o valor já
    calculado pelo Excel para a fórmula de `Controle de Diárias` se
    perderia."""
    caminho = caminho_planilha_conferencia_teste()
    injetar_cache_formula(caminho, "E2", 123.45)

    resultado = _cadastrar(caminho)
    assert resultado.estado is EstadoCadastro.SALVO

    workbook_valores = load_workbook(caminho, data_only=True)
    assert workbook_valores["Controle de Diárias"]["E2"].value == 123.45


@pytest.mark.parametrize(
    "cpf_consultado",
    ["123.456.789-01", "12345678901"],
    ids=["mascarado", "sem_mascara"],
)
def test_cpf_duplicado_mascarado_ou_nao_bloqueia_antes_de_escrever(
    caminho_planilha_conferencia_teste, cpf_consultado
):
    caminho = caminho_planilha_conferencia_teste()
    bytes_antes = open(caminho, "rb").read()

    resultado = _cadastrar(caminho, cpf=cpf_consultado)

    assert resultado.estado is EstadoCadastro.DUPLICADO
    assert resultado.linha == 2
    assert resultado.cpf_normalizado == "12345678901"
    # Nada foi escrito: o arquivo permanece byte a byte idêntico.
    assert open(caminho, "rb").read() == bytes_antes


def test_cpf_duplicado_com_cpf_gravado_como_texto_mascarado_na_planilha(
    caminho_planilha_conferencia_teste,
):
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(caminho, cpf="59111111100")

    assert resultado.estado is EstadoCadastro.DUPLICADO
    assert resultado.linha == 3


def test_falha_ao_salvar_arquivo_temporario_mantem_original_intacto(
    caminho_planilha_conferencia_teste, monkeypatch
):
    caminho = caminho_planilha_conferencia_teste()
    bytes_antes = open(caminho, "rb").read()

    # Aponta o caminho temporário para um diretório inexistente — faz
    # `workbook.save(caminho_tmp)` falhar de verdade (sem mockar biblioteca
    # nenhuma), reproduzindo uma falha real antes do `os.replace`.
    diretorio_invalido = os.path.join(os.path.dirname(caminho), "pasta-inexistente")
    monkeypatch.setattr(
        escritor_dados,
        "_caminho_temporario",
        lambda caminho_planilha: os.path.join(diretorio_invalido, "temp.xlsx"),
    )

    resultado = _cadastrar(caminho)

    assert resultado.estado is EstadoCadastro.NAO_SALVO
    assert open(caminho, "rb").read() == bytes_antes
    assert not os.path.isdir(diretorio_invalido)


def test_falha_na_validacao_do_temporario_mantem_original_intacto_e_remove_tmp(
    caminho_planilha_conferencia_teste, monkeypatch
):
    caminho = caminho_planilha_conferencia_teste()
    bytes_antes = open(caminho, "rb").read()

    # `_validar_tmp` devolve a linha encontrada (int) ou `None` em caso de
    # falha — nunca mais um `bool` (a reordenação por NOME pode mover a
    # linha nova para qualquer posição, então a validação passou a buscar
    # o CPF em vez de checar uma linha fixa).
    monkeypatch.setattr(escritor_dados, "_validar_tmp", lambda *args, **kwargs: None)

    resultado = _cadastrar(caminho)

    assert resultado.estado is EstadoCadastro.NAO_SALVO
    assert open(caminho, "rb").read() == bytes_antes
    caminho_tmp = escritor_dados._caminho_temporario(caminho)
    assert not os.path.exists(caminho_tmp)


def test_falha_no_os_replace_mantem_original_intacto_e_remove_tmp(
    caminho_planilha_conferencia_teste, monkeypatch
):
    caminho = caminho_planilha_conferencia_teste()
    bytes_antes = open(caminho, "rb").read()

    def _replace_com_falha(origem, destino):
        raise OSError("falha simulada no os.replace")

    monkeypatch.setattr(escritor_dados.os, "replace", _replace_com_falha)

    resultado = _cadastrar(caminho)

    assert resultado.estado is EstadoCadastro.NAO_SALVO
    assert open(caminho, "rb").read() == bytes_antes
    caminho_tmp = escritor_dados._caminho_temporario(caminho)
    assert not os.path.exists(caminho_tmp)


# ─── MATRÍCULA — normalização para o padrão real da coluna B ───────────────
# Auditoria 2026-08-10 contra a Planilha Oficial real: 1622/1622 linhas
# usam o MESMO `number_format` Excel ("000000\-0") — nunca uma máscara
# embutida no valor/texto em si, e a quantidade de dígitos varia
# legitimamente (6/7/8 observados). Item 1 do pedido "Ajuste Funcional + UX".


def test_matricula_entrada_so_numerica_normaliza_para_numero_com_formato_real(
    caminho_planilha_conferencia_teste,
):
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(caminho, matricula="1891553")

    assert resultado.estado is EstadoCadastro.SALVO
    workbook = load_workbook(caminho)
    celula = workbook["Dados"][f"B{resultado.linha}"]
    assert celula.value == 1891553
    assert isinstance(celula.value, int)
    assert celula.number_format == "000000\\-0"


def test_matricula_colada_com_pontuacao_e_normalizada_para_so_digitos(
    caminho_planilha_conferencia_teste,
):
    """Usuário pode colar um valor já "formatado" (ex.: copiado de outra
    célula/planilha) — só os dígitos são mantidos, o hífen nunca vai para
    o valor armazenado (só para o `number_format`, aplicado à parte)."""
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(caminho, matricula="189.155-3")

    assert resultado.estado is EstadoCadastro.SALVO
    celula = load_workbook(caminho)["Dados"][f"B{resultado.linha}"]
    assert celula.value == 1891553
    assert celula.number_format == "000000\\-0"


def test_matricula_com_menos_digitos_nao_inventa_preenchimento(caminho_planilha_conferencia_teste):
    """A auditoria real mostrou matrículas de 6 e 8 dígitos legítimas, não
    só 7 — nunca força um tamanho fixo (zero à esquerda inventado, corte,
    etc.), só aplica o mesmo `number_format` real."""
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(caminho, matricula="123456")

    assert resultado.estado is EstadoCadastro.SALVO
    celula = load_workbook(caminho)["Dados"][f"B{resultado.linha}"]
    assert celula.value == 123456
    assert celula.number_format == "000000\\-0"


# ─── ORDENAÇÃO por NOME após cadastro (item 4 do pedido) ───────────────────
# Auditoria 2026-08-10: aba `Dados` sem linhas especiais/vazias/subtotal
# depois da região cadastral real — a ordenação cobre só essa região.


def test_ordenacao_novo_registro_entra_na_primeira_posicao_alfabetica(
    caminho_planilha_conferencia_teste,
):
    # Fixture: linha2 "Maria Aparecida Souza", linha3 "João da Silva".
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(caminho, cpf="111.111.111-11", nome="Ana Beatriz Costa")

    assert resultado.linha == 2  # "Ana" < "João" < "Maria"
    planilha = load_workbook(caminho)["Dados"]
    assert planilha["C2"].value == "Ana Beatriz Costa"
    assert planilha["C3"].value == "João da Silva"
    assert planilha["C4"].value == "Maria Aparecida Souza"


def test_ordenacao_novo_registro_entra_no_meio_alfabetico(caminho_planilha_conferencia_teste):
    # "João" < "Lucas" < "Maria" — o novo registro deve ficar entre os dois.
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(caminho, cpf="222.222.222-22", nome="Lucas Pereira")

    assert resultado.linha == 3
    planilha = load_workbook(caminho)["Dados"]
    assert planilha["C2"].value == "João da Silva"
    assert planilha["C3"].value == "Lucas Pereira"
    assert planilha["C4"].value == "Maria Aparecida Souza"


def test_ordenacao_novo_registro_entra_na_ultima_posicao_alfabetica(
    caminho_planilha_conferencia_teste,
):
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(caminho, cpf="333.333.333-33", nome="Zélia Fernandes")

    assert resultado.linha == 4  # "João" < "Maria" < "Zélia"
    planilha = load_workbook(caminho)["Dados"]
    assert planilha["C2"].value == "João da Silva"
    assert planilha["C3"].value == "Maria Aparecida Souza"
    assert planilha["C4"].value == "Zélia Fernandes"


def test_ordenacao_nomes_com_acento_ordenam_junto_das_letras_sem_acento(
    caminho_planilha_conferencia_teste,
):
    """"Álvaro" precisa ordenar perto de "Ana"/"Ana", nunca isolado depois
    de "Z" (o que aconteceria com uma ordenação ingênua por code point
    Unicode, já que 'Á' > 'Z' em ASCII/Unicode puro)."""
    caminho = caminho_planilha_conferencia_teste()
    _cadastrar(caminho, cpf="444.444.444-44", nome="Ana Beatriz Costa")
    _cadastrar(caminho, cpf="555.555.555-55", nome="Zélia Fernandes")

    resultado = _cadastrar(caminho, cpf="666.666.666-66", nome="Álvaro Nunes")

    assert resultado.linha == 2  # "Álvaro" (~ALVARO) < "Ana" < "João" < "Maria" < "Zélia"
    planilha = load_workbook(caminho)["Dados"]
    nomes_em_ordem = [planilha[f"C{linha}"].value for linha in range(2, 7)]
    assert nomes_em_ordem == [
        "Álvaro Nunes",
        "Ana Beatriz Costa",
        "João da Silva",
        "Maria Aparecida Souza",
        "Zélia Fernandes",
    ]


def test_ordenacao_move_a_linha_inteira_nunca_so_a_coluna_nome(caminho_planilha_conferencia_teste):
    """As 8 colunas (CPF/MATRÍCULA/NOME/PREVIDÊNCIA/CARGO/ÓRGÃO/REGIME/
    SITUAÇÃO) precisam viajar JUNTAS — nunca só o nome reordenado enquanto
    o resto da linha fica para trás (o que misturaria CPF de uma pessoa
    com cargo de outra)."""
    caminho = caminho_planilha_conferencia_teste()

    resultado = _cadastrar(
        caminho,
        cpf="777.777.777-77",
        nome="Ana Beatriz Costa",
        matricula="555000",
        cargo="Cargo Exclusivo Da Ana",
        previdencia="FUNAPREV",
        orgao_origem="Órgão Exclusivo Da Ana",
        regime_previdencia="RPPS",
        situacao="Em exercício",
    )

    assert resultado.linha == 2
    planilha = load_workbook(caminho)["Dados"]
    assert planilha["A2"].value == 77777777777
    assert planilha["B2"].value == 555000
    assert planilha["C2"].value == "Ana Beatriz Costa"
    assert planilha["D2"].value == "FUNAPREV"
    assert planilha["E2"].value == "Cargo Exclusivo Da Ana"
    assert planilha["F2"].value == "Órgão Exclusivo Da Ana"
    assert planilha["G2"].value == "RPPS"
    assert planilha["H2"].value == "Em exercício"
    # Os dois registros originais continuam íntegros, só deslocados.
    assert planilha["C3"].value == "João da Silva"
    assert planilha["A3"].value is not None
    assert planilha["C4"].value == "Maria Aparecida Souza"
    assert planilha["A4"].value is not None


def test_ordenacao_preserva_cabecalho_intacto(caminho_planilha_conferencia_teste):
    caminho = caminho_planilha_conferencia_teste()
    cabecalho_antes = [load_workbook(caminho)["Dados"].cell(row=1, column=c).value for c in range(1, 9)]

    _cadastrar(caminho, cpf="888.888.888-88", nome="Ana Beatriz Costa")

    cabecalho_depois = [load_workbook(caminho)["Dados"].cell(row=1, column=c).value for c in range(1, 9)]
    assert cabecalho_depois == cabecalho_antes
