"""Micro-sprint 3F — testes do gerador do XLSX final do Relatório de
Diárias — Portal da Transparência. Fixtures 100% sintéticas (nenhum
CPF/nome/OB/SEI real) — nunca reimplementa RN01-RN07 aqui: `ResultadoRevisado`
é construído diretamente pelos campos do dataclass, exatamente como o
backend homologado (Micro-sprint 3C) já entrega."""

import glob
import os
import threading
from datetime import date, datetime

import openpyxl
import pytest

from app.domain.entities.linha_passagem import EstadoPassagem, LinhaPassagem
from app.domain.entities.ordem_bancaria import OrdemBancaria
from app.domain.entities.registro_controle import ClassificacaoSei, RegistroControle
from app.domain.relatorio_transparencia.decisao import (
    AcaoDecisao,
    CategoriaNaoPublicada,
    DecisaoHumana,
    EstadoRevisao,
    OcorrenciaNaoPublicavel,
    ResultadoRevisado,
    RevisaoComEstado,
)
from app.domain.relatorio_transparencia.linha_publicavel import nome_publicado as _nome_publicado
from app.domain.relatorio_transparencia.motor_decisao import (
    CandidatoNormal,
    OcorrenciaFinanceira,
    RevisaoNecessaria,
)
from app.planilha.escritor_relatorio_transparencia import (
    DestinoJaExiste,
    GeracaoConcorrenteEmAndamento,
    RelatorioBloqueadoPorPendencias,
    TemplateEstruturalmenteInvalido,
    _NOME_ABA_DADOS,
    _NUMBER_FORMAT_VALOR_PASSAGEM,
    _PRIMEIRA_LINHA_DADOS,
    _caminho_trava,
    gerar_xlsx_relatorio_transparencia,
)
from tests.fixtures.template_relatorio_transparencia_sintetico import construir_template_sintetico


# --- Helpers de fixture (nunca reimplementam RN) --------------------------


def _ordem(
    *,
    ob="2026OB000001",
    dh="2026DH000001",
    linha=2,
    nome="FULANO DE TAL",
    credor=None,
    cargo="Cargo de Teste",
    periodo_inicial=date(2026, 1, 10),
    periodo_final=date(2026, 1, 10),
    trecho="RECIFE-GARANHUNS",
    transporte="CARRO",
    motivo="Motivo de teste",
    valor=154.09,
    data_pagamento=datetime(2026, 1, 15, 9, 0, 0),
) -> OrdemBancaria:
    return OrdemBancaria(
        linha=linha,
        dh=dh,
        ob=ob,
        credor=credor if credor is not None else nome,
        situacao="Paga",
        data_pagamento=data_pagamento,
        sei="19.20.0001.0000001/2026-01",
        cpf="11111111111",
        nome=nome,
        cargo=cargo,
        periodo_inicial=periodo_inicial,
        periodo_final=periodo_final,
        cidade_origem=None,
        cidade_destino=None,
        transporte=transporte,
        motivo=motivo,
        valor=valor,
        valor_bruto_sem_formula=valor,
        restituicao=None,
        duplicidade="Único",
        item_de_gasto="33901401",
        trecho=trecho,
    )


def _registro(*, dh="2026DH000001", linha=2, integral=0, parcial=1, cargo=None, transporte="CARRO") -> RegistroControle:
    return RegistroControle(
        linha=linha,
        dh=dh,
        sei="19.20.0001.0000001/2026-01",
        sei_classificacao=ClassificacaoSei.VALIDO,
        doc_origem="PJTESTE",
        cpf="11111111111",
        periodo_inicial=date(2026, 1, 10),
        periodo_final=date(2026, 1, 10),
        cidade_origem="RECIFE",
        cidade_destino="GARANHUNS",
        trecho="RECIFE-GARANHUNS",
        motivo="Motivo de teste",
        transporte=transporte,
        valor_total_diarias=154.09,
        quantidade_diaria_integral=integral,
        quantidade_diaria_parcial=parcial,
        cargo=cargo,
    )


def _candidato(*, registros_controle=None, passagem=None, **kwargs_ordem) -> CandidatoNormal:
    ordem = _ordem(**kwargs_ordem)
    registros = registros_controle if registros_controle is not None else [_registro(dh=ordem.dh, integral=0, parcial=1)]
    ocorrencia = OcorrenciaFinanceira(
        dh=ordem.dh, ordem_bancaria=ordem, registros_controle=registros, passagem=passagem
    )
    return CandidatoNormal(ocorrencia=ocorrencia, registros_controle=registros)


def _linha_passagem(*, ob="2026OB000001", estado, valor=None, transporte="AVIÃO", bruto=None, linha=10):
    """Sprint 3 (HOMO-RT-06) — mesmo formato de `LinhaPassagem` da Sprint 1,
    para testar a coluna J do XLSX oficial com cada estado real."""
    return LinhaPassagem(linha=linha, ob=ob, transporte=transporte, estado=estado, valor=valor, bruto=bruto)


def _resultado(*, candidatos=None, nao_publicaveis=None, revisoes=None) -> ResultadoRevisado:
    return ResultadoRevisado(
        candidatos_publicaveis=candidatos or [],
        ocorrencias_nao_publicaveis=nao_publicaveis or [],
        revisoes=revisoes or [],
        diagnosticos_repeticao_fisica=[],
    )


@pytest.fixture()
def template(tmp_path):
    caminho = str(tmp_path / "template_sintetico.xlsx")
    construir_template_sintetico(caminho)
    return caminho


def _gerar(resultado, template, destino, **overrides):
    parametros = {
        "caminho_template": template,
        "caminho_destino": destino,
        "ano_exercicio": 2026,
        "mes_referencia": 1,
        "hoje": date(2026, 2, 1),
    }
    parametros.update(overrides)
    return gerar_xlsx_relatorio_transparencia(resultado, **parametros)


# --- Bloqueio (item 21, obrigatório) ---------------------------------------


def test_bloqueia_geracao_quando_ha_revisao_pendente_mesmo_chamado_direto(tmp_path, template):
    revisao = RevisaoNecessaria(regra="RN07", motivo="Motivo de teste", ocorrencias=[])
    resultado = _resultado(revisoes=[RevisaoComEstado(revisao=revisao, estado=EstadoRevisao.PENDENTE)])
    destino = str(tmp_path / "saida.xlsx")

    with pytest.raises(RelatorioBloqueadoPorPendencias):
        _gerar(resultado, template, destino)

    assert not os.path.exists(destino)


def test_bloqueia_geracao_quando_decisao_esta_adiada(tmp_path, template):
    revisao = RevisaoNecessaria(regra="RN07", motivo="Motivo de teste", ocorrencias=[])
    resultado = _resultado(revisoes=[RevisaoComEstado(revisao=revisao, estado=EstadoRevisao.ADIADA)])
    destino = str(tmp_path / "saida.xlsx")

    with pytest.raises(RelatorioBloqueadoPorPendencias):
        _gerar(resultado, template, destino)


def test_permite_geracao_quando_todas_as_revisoes_estao_decididas(tmp_path, template):
    revisao = RevisaoNecessaria(regra="RN07", motivo="Motivo de teste", ocorrencias=[])
    decisao = DecisaoHumana(
        identidade_grupo="g1",
        regra="RN07",
        obs_envolvidas=frozenset({"2026OB000001"}),
        fingerprint_evidencias="fp1",
        acao=AcaoDecisao.PUBLICAR,
        autor="teste@teste.com",
        decidido_em=datetime(2026, 1, 20),
    )
    resultado = _resultado(
        candidatos=[_candidato()],
        revisoes=[RevisaoComEstado(revisao=revisao, estado=EstadoRevisao.DECIDIDA, decisao=decisao)],
    )
    destino = str(tmp_path / "saida.xlsx")

    resultado_gerado = _gerar(resultado, template, destino)

    assert os.path.exists(destino)
    assert resultado_gerado.total_ocorrencias == 1


# --- Proteção de destino / template (item 17) ------------------------------


def test_nunca_sobrescreve_destino_existente(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    with open(destino, "wb") as f:
        f.write(b"conteudo previo")

    with pytest.raises(DestinoJaExiste):
        _gerar(_resultado(), template, destino)

    with open(destino, "rb") as f:
        assert f.read() == b"conteudo previo"


def test_template_inexistente_gera_erro_explicito(tmp_path):
    destino = str(tmp_path / "saida.xlsx")
    with pytest.raises(TemplateEstruturalmenteInvalido):
        _gerar(_resultado(), str(tmp_path / "nao_existe.xlsx"), destino)


def test_template_sem_aba_obrigatoria_gera_erro_explicito(tmp_path):
    caminho_template = str(tmp_path / "template_invalido.xlsx")
    wb = openpyxl.Workbook()
    wb.active.title = "Aba Qualquer"
    wb.save(caminho_template)
    destino = str(tmp_path / "saida.xlsx")

    with pytest.raises(TemplateEstruturalmenteInvalido):
        _gerar(_resultado(), caminho_template, destino)


def test_nao_deixa_arquivo_temporario_apos_sucesso(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino)
    temporarios = glob.glob(str(tmp_path / "tmp*.xlsx"))
    assert temporarios == []


def test_nao_deixa_arquivo_temporario_apos_falha_de_bloqueio(tmp_path, template):
    revisao = RevisaoNecessaria(regra="RN07", motivo="Motivo de teste", ocorrencias=[])
    resultado = _resultado(revisoes=[RevisaoComEstado(revisao=revisao, estado=EstadoRevisao.PENDENTE)])
    destino = str(tmp_path / "saida.xlsx")

    with pytest.raises(RelatorioBloqueadoPorPendencias):
        _gerar(resultado, template, destino)

    temporarios = glob.glob(str(tmp_path / "tmp*.xlsx"))
    assert temporarios == []


# --- Concorrência (Confronto QA / homologação operacional) -----------------
#
# A homologação reproduziu 2 requisições HTTP simultâneas para `/gerar` com
# o MESMO `caminho_destino`: as duas retornavam 200, a segunda a terminar
# sobrescrevendo silenciosamente o resultado da primeira — `os.path.exists`
# seguido de escrita não é atômico. A correção usa `os.open(...,
# O_CREAT | O_EXCL)` (ver `_adquirir_trava_destino`) para reservar o destino
# antes de qualquer trabalho pesado. Os testes abaixo reproduzem a corrida
# de verdade (threads reais disputando o mesmo destino), não só simulam a
# sequência de chamadas.


def test_nao_deixa_arquivo_de_trava_apos_sucesso(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    assert not os.path.exists(_caminho_trava(destino))


def test_nao_deixa_arquivo_de_trava_apos_falha_de_bloqueio(tmp_path, template):
    revisao = RevisaoNecessaria(regra="RN07", motivo="Motivo de teste", ocorrencias=[])
    resultado = _resultado(revisoes=[RevisaoComEstado(revisao=revisao, estado=EstadoRevisao.PENDENTE)])
    destino = str(tmp_path / "saida.xlsx")

    with pytest.raises(RelatorioBloqueadoPorPendencias):
        _gerar(resultado, template, destino)

    assert not os.path.exists(_caminho_trava(destino))


def test_nao_deixa_arquivo_de_trava_apos_destino_ja_existente(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    with open(destino, "wb") as f:
        f.write(b"conteudo previo")

    with pytest.raises(DestinoJaExiste):
        _gerar(_resultado(), template, destino)

    assert not os.path.exists(_caminho_trava(destino))


def test_trava_liberada_apos_erro_permite_nova_tentativa_no_mesmo_destino(tmp_path, template):
    """A trava nunca deve ficar órfã bloqueando retentativas legítimas após
    uma falha que não seja uma corrida real (aqui: destino ocupado por um
    arquivo alheio, removido antes da segunda tentativa)."""
    destino = str(tmp_path / "saida.xlsx")
    with open(destino, "wb") as f:
        f.write(b"conteudo previo")
    with pytest.raises(DestinoJaExiste):
        _gerar(_resultado(), template, destino)

    os.remove(destino)
    resultado_gerado = _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    assert os.path.exists(destino)
    assert resultado_gerado.total_ocorrencias == 1


def test_duas_geracoes_simultaneas_no_mesmo_destino_apenas_uma_conclui(tmp_path, template):
    """Reproduz o achado da homologação com concorrência REAL (2 threads),
    não uma sequência de chamadas simulando concorrência. Nunca as duas
    podem retornar sucesso para o mesmo destino."""
    destino = str(tmp_path / "saida.xlsx")
    liberar_barreira = threading.Barrier(2, timeout=10)
    resultados: dict[str, object] = {}

    def _tentar(chave):
        liberar_barreira.wait()  # maximiza a chance real de colisão
        try:
            resultados[chave] = _gerar(_resultado(candidatos=[_candidato()]), template, destino)
        except Exception as erro:  # noqa: BLE001 - captura para asserção, não engolida
            resultados[chave] = erro

    t1 = threading.Thread(target=_tentar, args=("a",))
    t2 = threading.Thread(target=_tentar, args=("b",))
    t1.start()
    t2.start()
    t1.join(timeout=30)
    t2.join(timeout=30)

    valores = list(resultados.values())
    sucessos = [v for v in valores if not isinstance(v, Exception)]
    falhas_concorrencia = [v for v in valores if isinstance(v, GeracaoConcorrenteEmAndamento)]

    assert len(sucessos) == 1, f"esperado exatamente 1 sucesso, obtido {len(sucessos)}: {valores}"
    assert len(falhas_concorrencia) == 1, (
        f"esperado exatamente 1 GeracaoConcorrenteEmAndamento, obtido {len(falhas_concorrencia)}: {valores}"
    )
    assert os.path.exists(destino)
    assert not os.path.exists(_caminho_trava(destino))
    # nunca 200+200: o arquivo final corresponde ao único candidato
    # publicado — nunca um estado indefinido de "quem escreveu por último".
    workbook = openpyxl.load_workbook(destino)
    ws = workbook[_NOME_ABA_DADOS]
    assert ws[f"A{_PRIMEIRA_LINHA_DADOS}"].value == "2026OB000001"
    workbook.close()


def test_duas_geracoes_simultaneas_em_destinos_diferentes_ambas_concluem(tmp_path, template):
    """A trava é por destino, nunca global — duas gerações concorrentes
    para arquivos diferentes não podem se bloquear uma à outra."""
    destino_a = str(tmp_path / "saida_a.xlsx")
    destino_b = str(tmp_path / "saida_b.xlsx")
    liberar_barreira = threading.Barrier(2, timeout=10)
    resultados: dict[str, object] = {}

    def _tentar(chave, destino):
        liberar_barreira.wait()
        try:
            resultados[chave] = _gerar(_resultado(candidatos=[_candidato()]), template, destino)
        except Exception as erro:  # noqa: BLE001
            resultados[chave] = erro

    t1 = threading.Thread(target=_tentar, args=("a", destino_a))
    t2 = threading.Thread(target=_tentar, args=("b", destino_b))
    t1.start()
    t2.start()
    t1.join(timeout=30)
    t2.join(timeout=30)

    assert not isinstance(resultados["a"], Exception), resultados["a"]
    assert not isinstance(resultados["b"], Exception), resultados["b"]
    assert os.path.exists(destino_a)
    assert os.path.exists(destino_b)


def test_quatro_geracoes_simultaneas_no_mesmo_destino_apenas_uma_conclui(tmp_path, template):
    """Confronto QA / revalidação Codex — mesmo teste da corrida com 2
    threads, mas com 4 disputando o mesmo destino: a trava cooperativa
    precisa segurar TODAS as concorrentes, não só a "segunda"."""
    destino = str(tmp_path / "saida.xlsx")
    n = 4
    liberar_barreira = threading.Barrier(n, timeout=10)
    resultados: dict[str, object] = {}

    def _tentar(chave):
        liberar_barreira.wait()
        try:
            resultados[chave] = _gerar(_resultado(candidatos=[_candidato()]), template, destino)
        except Exception as erro:  # noqa: BLE001
            resultados[chave] = erro

    threads = [threading.Thread(target=_tentar, args=(str(i),)) for i in range(n)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=30)

    valores = list(resultados.values())
    sucessos = [v for v in valores if not isinstance(v, Exception)]
    falhas_concorrencia = [v for v in valores if isinstance(v, GeracaoConcorrenteEmAndamento)]

    assert len(sucessos) == 1, f"esperado exatamente 1 sucesso entre {n} concorrentes, obtido {len(sucessos)}: {valores}"
    assert len(falhas_concorrencia) == n - 1, (
        f"esperado {n - 1} GeracaoConcorrenteEmAndamento, obtido {len(falhas_concorrencia)}: {valores}"
    )
    assert os.path.exists(destino)
    assert not os.path.exists(_caminho_trava(destino))


def test_terceiro_cria_destino_entre_verificacao_e_promocao_preserva_arquivo_do_terceiro(tmp_path, template, monkeypatch):
    """Reproduz a janela TOCTOU remanescente: um arquivo aparece no destino
    DEPOIS da checagem inicial (feita sob a trava) e ANTES da promoção
    final — cenário de um terceiro processo/usuário, não de outra chamada
    deste módulo (a trava cooperativa não protege contra isso). A
    promoção precisa falhar de forma controlada, nunca sobrescrever, e o
    arquivo do terceiro precisa sobreviver intacto.

    O monkeypatch injeta a criação do arquivo de terceiro no momento exato
    da promoção e então delega para o `os.rename` REAL — a falha
    observada (`FileExistsError` -> `DestinoJaExiste`) vem do sistema
    operacional de verdade, não é fabricada pelo teste."""
    import app.planilha.escritor_relatorio_transparencia as modulo_escritor

    destino = str(tmp_path / "saida.xlsx")
    conteudo_terceiro = b"CONTEUDO DE UM TERCEIRO PROCESSO - NAO TOCAR"
    os_rename_real = os.rename
    ja_injetou = {"valor": False}

    def _rename_com_terceiro_no_meio(origem, destino_chamado, *args, **kwargs):
        if not ja_injetou["valor"]:
            ja_injetou["valor"] = True
            with open(destino_chamado, "wb") as f:
                f.write(conteudo_terceiro)
        return os_rename_real(origem, destino_chamado, *args, **kwargs)

    monkeypatch.setattr(modulo_escritor.os, "rename", _rename_com_terceiro_no_meio)

    with pytest.raises(DestinoJaExiste):
        _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    with open(destino, "rb") as f:
        assert f.read() == conteudo_terceiro, "arquivo do terceiro foi sobrescrito — TOCTOU não corrigido"
    assert glob.glob(str(tmp_path / "tmp*.xlsx")) == [], "temporário não foi limpo após a falha"
    assert not os.path.exists(_caminho_trava(destino)), "trava não foi liberada após a falha"

    # nova tentativa para um destino DIFERENTE continua funcionando
    # normalmente (a falha do terceiro não deixou o módulo num estado ruim)
    outro_destino = str(tmp_path / "saida_2.xlsx")
    resultado = _gerar(_resultado(candidatos=[_candidato()]), template, outro_destino)
    assert os.path.exists(outro_destino)
    assert resultado.total_ocorrencias == 1


def test_excecao_durante_geracao_com_trava_ja_adquirida_libera_trava_e_permite_retentativa(tmp_path):
    """Distinto do teste de bloqueio por pendências (que falha ANTES da
    trava ser adquirida): aqui o template existe e é aberto, mas não tem a
    aba obrigatória — a exceção acontece DEPOIS de `_adquirir_trava_destino`,
    dentro do trabalho pesado. Confirma que mesmo uma falha "no meio" da
    geração libera a trava (nunca bloqueia retentativas futuras)."""
    caminho_template_invalido = str(tmp_path / "template_sem_aba.xlsx")
    wb = openpyxl.Workbook()
    wb.active.title = "Aba Qualquer"
    wb.save(caminho_template_invalido)
    destino = str(tmp_path / "saida.xlsx")

    with pytest.raises(TemplateEstruturalmenteInvalido):
        _gerar(_resultado(candidatos=[_candidato()]), caminho_template_invalido, destino)

    assert not os.path.exists(_caminho_trava(destino))
    assert glob.glob(str(tmp_path / "tmp*.xlsx")) == []
    assert not os.path.exists(destino)


# --- Conteúdo ---------------------------------------------------------------


def test_zero_ocorrencias_gera_planilha_sem_linhas_de_dados(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    resultado_gerado = _gerar(_resultado(), template, destino)

    assert resultado_gerado.total_ocorrencias == 0
    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"A{_PRIMEIRA_LINHA_DADOS}"].value.startswith("Total")
    assert ws[f"J{_PRIMEIRA_LINHA_DADOS}"].value == 0
    assert ws[f"M{_PRIMEIRA_LINHA_DADOS}"].value == 0


def test_uma_ocorrencia_e_escrita_corretamente(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    candidato = _candidato(ob="2026OB000042", valor=321.55)
    resultado_gerado = _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    linha = _PRIMEIRA_LINHA_DADOS
    assert ws[f"A{linha}"].value == "2026OB000042"
    assert ws[f"M{linha}"].value == 321.55
    # Sprint 3 (HOMO-RT-06) — sem `passagem` anexada (estado E, OB sem
    # correspondência na Planilha de Passagens), a coluna J é célula vazia
    # (`None`), nunca mais o placeholder fixo "-" do histórico pré-Sprint 3.
    assert ws[f"J{linha}"].value is None
    assert resultado_gerado.total_valor_diarias == 321.55
    assert ws[f"A{linha + 1}"].value.startswith("Total")


# --- Sprint 3 (HOMO-RT-06) — coluna J (Valor da Passagem) por estado ------


def test_coluna_j_estado_numerico_publica_o_valor(tmp_path, template):
    passagem = _linha_passagem(ob="2026OB000042", estado=EstadoPassagem.NUMERICO, valor=1657.76)
    candidato = _candidato(ob="2026OB000042", passagem=passagem)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"J{_PRIMEIRA_LINHA_DADOS}"].value == 1657.76


def test_coluna_j_valor_monetario_sem_perda_de_precisao(tmp_path, template):
    """Confirma que o valor chega à célula exatamente como veio da fonte —
    nenhum arredondamento/truncamento no caminho leitor -> escritor."""
    passagem = _linha_passagem(ob="2026OB000042", estado=EstadoPassagem.NUMERICO, valor=4964.32)
    candidato = _candidato(ob="2026OB000042", passagem=passagem)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"J{_PRIMEIRA_LINHA_DADOS}"].value == 4964.32


# --- Correção do bug de formatação (auditoria Setembro/2026: 29 de 60
# valores de passagem exibidos como data em vez de moeda) — ver
# `_NUMBER_FORMAT_VALOR_PASSAGEM` e a correção em `_escrever_linha_
# ocorrencia`. Os 5 testes abaixo reproduzem exatamente os cenários da
# auditoria, sem depender de nenhum dado real.


def test_coluna_j_corrige_formatacao_de_data_herdada_para_a_propria_linha(tmp_path, template):
    """Cenário 1 — a célula que vai receber o valor já está, no template,
    com `number_format` de DATA (reprodução literal do bug real). O valor
    tem que sair como número monetário de verdade, nunca uma data."""
    wb = openpyxl.load_workbook(template)
    ws = wb[_NOME_ABA_DADOS]
    ws[f"J{_PRIMEIRA_LINHA_DADOS}"].number_format = "dd/mm/yyyy"
    wb.save(template)

    passagem = _linha_passagem(ob="2026OB000042", estado=EstadoPassagem.NUMERICO, valor=2200.16)
    candidato = _candidato(ob="2026OB000042", passagem=passagem)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb_g = openpyxl.load_workbook(destino)
    celula = wb_g[_NOME_ABA_DADOS][f"J{_PRIMEIRA_LINHA_DADOS}"]
    assert celula.value == 2200.16
    assert celula.data_type == "n"
    assert celula.number_format == _NUMBER_FORMAT_VALOR_PASSAGEM


def test_coluna_j_linha_nova_nao_depende_do_formato_da_linha_modelo(tmp_path, template):
    """Cenário 2 — a linha-modelo (linha 7, fonte de `_copiar_estilo_
    linha` para toda linha inserida além do tamanho original do template)
    está com `number_format` de DATA na coluna J. Uma linha NOVA (além das
    3 originais) tem que sair com formato monetário mesmo assim — a
    correção nunca pode depender do estilo herdado da linha-modelo."""
    wb = openpyxl.load_workbook(template)
    ws = wb[_NOME_ABA_DADOS]
    ws[f"J{_PRIMEIRA_LINHA_DADOS}"].number_format = "dd/mm/yyyy"
    wb.save(template)

    passagem = _linha_passagem(ob="2026OB000099", estado=EstadoPassagem.NUMERICO, valor=1888.24)
    candidatos = [_candidato(ob=f"2026OB{i:06d}", dh=f"2026DH{i:06d}") for i in range(1, 5)]
    candidatos.append(_candidato(ob="2026OB000099", dh="2026DH000099", passagem=passagem))
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=candidatos), template, destino)

    wb_g = openpyxl.load_workbook(destino)
    ws_g = wb_g[_NOME_ABA_DADOS]
    linha_nova = _PRIMEIRA_LINHA_DADOS + 4  # 5ª linha de dados, além do template original
    assert ws_g[f"A{linha_nova}"].value == "2026OB000099"
    celula = ws_g[f"J{linha_nova}"]
    assert celula.value == 1888.24
    assert celula.number_format == _NUMBER_FORMAT_VALOR_PASSAGEM


def test_coluna_j_linha_preexistente_com_formato_quebrado_e_autocorrigida_na_regeracao(tmp_path, template):
    """Cenário 3 — uma linha DENTRO do tamanho original do template (nunca
    passa por `_copiar_estilo_linha`, que só age em linhas além do
    tamanho original) já chega com valor numérico e `number_format` de
    data — exatamente como uma linha gravada por uma geração de mês
    anterior, antes desta correção. Como `_escrever_linha_ocorrencia`
    reescreve TODAS as linhas a cada chamada (novas e preexistentes), a
    regeração tem que normalizar o formato mesmo sem nenhuma linha ser
    inserida/removida (3 candidatos == 3 linhas do template)."""
    linha_preexistente = _PRIMEIRA_LINHA_DADOS + 1  # linha 8, dentro do template original (7-9)
    wb = openpyxl.load_workbook(template)
    ws = wb[_NOME_ABA_DADOS]
    ws[f"J{linha_preexistente}"] = 1307.84
    ws[f"J{linha_preexistente}"].number_format = "dd/mm/yyyy"
    wb.save(template)

    # `candidatos_ordenados` é ordenado por DATA OB e, em empate, por número
    # da OB (item 14) — todos com a mesma data de pagamento aqui, então
    # "...0002" cai deterministicamente na 2ª posição (linha 8).
    passagem = _linha_passagem(ob="2026OB000002", estado=EstadoPassagem.NUMERICO, valor=1307.84)
    candidatos = [
        _candidato(ob="2026OB000001", dh="2026DH000001"),
        _candidato(ob="2026OB000002", dh="2026DH000002", passagem=passagem),
        _candidato(ob="2026OB000003", dh="2026DH000003"),
    ]
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=candidatos), template, destino)

    wb_g = openpyxl.load_workbook(destino)
    ws_g = wb_g[_NOME_ABA_DADOS]
    assert ws_g[f"A{linha_preexistente}"].value == "2026OB000002"
    celula = ws_g[f"J{linha_preexistente}"]
    assert celula.value == 1307.84
    assert celula.number_format == _NUMBER_FORMAT_VALOR_PASSAGEM


def test_coluna_j_valor_ja_corretamente_formatado_permanece_inalterado(tmp_path, template):
    """Cenário 4 — uma linha cujo valor e formato já estavam corretos no
    template continua com o MESMO valor e o MESMO formato monetário
    depois da geração — a correção nunca reformata por acidente um caso
    que já estava certo."""
    wb = openpyxl.load_workbook(template)
    ws = wb[_NOME_ABA_DADOS]
    ws[f"J{_PRIMEIRA_LINHA_DADOS}"].number_format = _NUMBER_FORMAT_VALOR_PASSAGEM
    wb.save(template)

    passagem = _linha_passagem(ob="2026OB000042", estado=EstadoPassagem.NUMERICO, valor=4964.32)
    candidato = _candidato(ob="2026OB000042", passagem=passagem)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb_g = openpyxl.load_workbook(destino)
    celula = wb_g[_NOME_ABA_DADOS][f"J{_PRIMEIRA_LINHA_DADOS}"]
    assert celula.value == 4964.32
    assert celula.number_format == _NUMBER_FORMAT_VALOR_PASSAGEM


@pytest.mark.parametrize(
    "estado",
    [EstadoPassagem.AUSENTE, EstadoPassagem.NAO_SE_APLICA, EstadoPassagem.INVALIDO],
)
def test_coluna_j_estados_vazios_permanecem_vazios_mesmo_com_linha_modelo_corrompida(tmp_path, template, estado):
    """Cenário 5 (B/C/D) — mesmo com a linha-modelo corrompida com
    `number_format` de data, os estados sem valor numérico nunca ganham
    um valor nem viram monetários por acidente: a correção só age quando
    há um valor real a escrever (`if valor_passagem is not None`)."""
    wb = openpyxl.load_workbook(template)
    ws = wb[_NOME_ABA_DADOS]
    ws[f"J{_PRIMEIRA_LINHA_DADOS}"].number_format = "dd/mm/yyyy"
    wb.save(template)

    passagem = _linha_passagem(ob="2026OB000042", estado=estado, bruto="qualquer coisa")
    candidato = _candidato(ob="2026OB000042", passagem=passagem)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb_g = openpyxl.load_workbook(destino)
    assert wb_g[_NOME_ABA_DADOS][f"J{_PRIMEIRA_LINHA_DADOS}"].value is None


def test_coluna_j_estado_e_permanece_vazio_mesmo_com_linha_modelo_corrompida(tmp_path, template):
    """Cenário 5 (E) — mesma garantia do teste acima, para OB sem
    correspondência na Planilha de Passagens (`passagem=None`)."""
    wb = openpyxl.load_workbook(template)
    ws = wb[_NOME_ABA_DADOS]
    ws[f"J{_PRIMEIRA_LINHA_DADOS}"].number_format = "dd/mm/yyyy"
    wb.save(template)

    candidato = _candidato(ob="2026OB000042", passagem=None)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb_g = openpyxl.load_workbook(destino)
    assert wb_g[_NOME_ABA_DADOS][f"J{_PRIMEIRA_LINHA_DADOS}"].value is None


@pytest.mark.parametrize(
    "estado",
    [EstadoPassagem.AUSENTE, EstadoPassagem.NAO_SE_APLICA],
)
def test_coluna_j_estados_b_e_c_ficam_vazios(tmp_path, template, estado):
    passagem = _linha_passagem(ob="2026OB000042", estado=estado, bruto="qualquer coisa")
    candidato = _candidato(ob="2026OB000042", passagem=passagem)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    # Nunca o texto literal da fonte (ex. "NÃO SE APLICA"), nunca "-", nunca 0.
    assert ws[f"J{_PRIMEIRA_LINHA_DADOS}"].value is None


def test_coluna_j_estado_e_sem_correspondencia_fica_vazio(tmp_path, template):
    """`passagem=None` é literalmente o estado E — OB da Base OB sem
    correspondência na Planilha de Passagens."""
    candidato = _candidato(ob="2026OB000042", passagem=None)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"J{_PRIMEIRA_LINHA_DADOS}"].value is None


# --- Fechamento HOMO-RT-06 (decisão de produto final) — estado D NUNCA
# bloqueia a geração: célula vazia, igual a B/C/E, mas o estado INVALIDO
# permanece preservado no domínio (testado em test_decisao.py/
# test_linha_publicavel.py, não aqui — este arquivo testa só a saída XLSX).


def test_coluna_j_estado_d_invalido_produz_celula_vazia_sem_bloquear(tmp_path, template):
    """Estado D (INVALIDO, ex. "#VALOR!") nunca vira "-", 0 ou o texto bruto
    da fonte — célula vazia, e a geração conclui normalmente (decisão de
    produto final: D não bloqueia)."""
    passagem = _linha_passagem(ob="2026OB000042", estado=EstadoPassagem.INVALIDO, bruto="#VALOR!")
    candidato = _candidato(ob="2026OB000042", passagem=passagem)
    destino = str(tmp_path / "saida.xlsx")

    resultado_gerado = _gerar(_resultado(candidatos=[candidato]), template, destino)

    assert resultado_gerado.total_ocorrencias == 1
    assert os.path.exists(destino)
    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"J{_PRIMEIRA_LINHA_DADOS}"].value is None
    assert ws[f"J{_PRIMEIRA_LINHA_DADOS}"].value != "#VALOR!"
    assert ws[f"J{_PRIMEIRA_LINHA_DADOS}"].value != 0


def test_coluna_j_estado_d_nao_impede_publicacao_de_outras_ocorrencias(tmp_path, template):
    """Uma OB em estado D nunca impede as demais ocorrências (inclusive com
    passagem numérica válida) de serem publicadas normalmente — nunca
    publicação parcial nem bloqueio da geração inteira."""
    passagem_valida = _linha_passagem(ob="2026OB000001", estado=EstadoPassagem.NUMERICO, valor=100.0)
    passagem_invalida = _linha_passagem(ob="2026OB000002", estado=EstadoPassagem.INVALIDO, bruto="#VALOR!")
    candidatos = [
        _candidato(ob="2026OB000001", dh="2026DH000001", passagem=passagem_valida),
        _candidato(ob="2026OB000002", dh="2026DH000002", passagem=passagem_invalida),
    ]
    destino = str(tmp_path / "saida.xlsx")

    resultado_gerado = _gerar(_resultado(candidatos=candidatos), template, destino)

    assert resultado_gerado.total_ocorrencias == 2
    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    valores_j = {ws[f"A{_PRIMEIRA_LINHA_DADOS + i}"].value: ws[f"J{_PRIMEIRA_LINHA_DADOS + i}"].value for i in range(2)}
    assert valores_j["2026OB000001"] == 100.0
    assert valores_j["2026OB000002"] is None


def test_varias_ocorrencias_uma_linha_por_ob_nunca_agregada(tmp_path, template):
    candidatos = [
        _candidato(ob="2026OB000001", dh="2026DH000001", valor=100.0),
        _candidato(ob="2026OB000002", dh="2026DH000002", valor=200.0),
        _candidato(ob="2026OB000003", dh="2026DH000003", valor=300.0),
    ]
    destino = str(tmp_path / "saida.xlsx")
    resultado_gerado = _gerar(_resultado(candidatos=candidatos), template, destino)

    assert resultado_gerado.total_ocorrencias == 3
    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    obs_escritas = [ws[f"A{_PRIMEIRA_LINHA_DADOS + i}"].value for i in range(3)]
    assert obs_escritas == ["2026OB000001", "2026OB000002", "2026OB000003"]


def test_um_dh_com_multiplas_obs_gera_uma_linha_por_ob(tmp_path, template):
    registros = [_registro(dh="2026DH000009", integral=1, parcial=0)]
    candidatos = [
        _candidato(ob="2026OB000010", dh="2026DH000009", registros_controle=registros, valor=50.0),
        _candidato(ob="2026OB000011", dh="2026DH000009", registros_controle=registros, valor=60.0),
    ]
    destino = str(tmp_path / "saida.xlsx")
    resultado_gerado = _gerar(_resultado(candidatos=candidatos), template, destino)

    assert resultado_gerado.total_ocorrencias == 2
    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"A{_PRIMEIRA_LINHA_DADOS}"].value == "2026OB000010"
    assert ws[f"A{_PRIMEIRA_LINHA_DADOS + 1}"].value == "2026OB000011"
    # Mesmo Controle DIMCAC do DH — K/L idênticos nas duas linhas, nunca somados.
    assert ws[f"K{_PRIMEIRA_LINHA_DADOS}"].value == 1
    assert ws[f"K{_PRIMEIRA_LINHA_DADOS + 1}"].value == 1


def test_valor_zero_e_legitimo_e_nunca_vira_branco_ou_none(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    candidato = _candidato(valor=0.0)
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"M{_PRIMEIRA_LINHA_DADOS}"].value == 0.0


def test_campos_ausentes_ficam_em_branco_nunca_inventados(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    candidato = _candidato(cargo=None)
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"D{_PRIMEIRA_LINHA_DADOS}"].value is None


def test_quantidade_diarias_ausente_por_conflito_fica_em_branco(tmp_path, template):
    registros_conflitantes = [
        _registro(dh="2026DH000005", linha=2, integral=1, parcial=1),
        _registro(dh="2026DH000005", linha=3, integral=1, parcial=0),
    ]
    destino = str(tmp_path / "saida.xlsx")
    candidato = _candidato(dh="2026DH000005", registros_controle=registros_conflitantes)
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"K{_PRIMEIRA_LINHA_DADOS}"].value is None
    assert ws[f"L{_PRIMEIRA_LINHA_DADOS}"].value is None


def test_trecho_publicado_e_sempre_o_original_nunca_normalizado(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    candidato = _candidato(trecho="RECIFE - GARANHUNS (viagem oficial)")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"G{_PRIMEIRA_LINHA_DADOS}"].value == "RECIFE - GARANHUNS (viagem oficial)"


def test_datas_sao_preservadas_como_objetos_data(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    candidato = _candidato(
        data_pagamento=datetime(2026, 1, 22, 10, 30, 0),
        periodo_inicial=date(2026, 1, 20),
        periodo_final=date(2026, 1, 21),
    )
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    # Micro-sprint 3F.1 — o modelo oficial real armazena a coluna B como
    # TEXTO literal "DD/MM/YYYY HH:MM:SS" em 100% das linhas (investigação
    # célula a célula), nunca como data/serial Excel — reproduzido aqui.
    assert ws[f"B{_PRIMEIRA_LINHA_DADOS}"].value == "22/01/2026 10:30:00"
    # openpyxl sempre devolve `datetime.datetime` ao reabrir um XLSX salvo em
    # disco (a distinção `date` vs `datetime` não sobrevive ao round-trip,
    # mesmo com number_format só de data) — E/F continuam datas reais,
    # comparamos por data, não por tipo.
    assert ws[f"E{_PRIMEIRA_LINHA_DADOS}"].value == datetime(2026, 1, 20)


def test_data_pagamento_e_texto_nunca_datetime_ou_serial(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    candidato = _candidato(data_pagamento=datetime(2026, 3, 5, 8, 7, 6))
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    celula = ws[f"B{_PRIMEIRA_LINHA_DADOS}"]
    assert celula.data_type == "s"
    assert isinstance(celula.value, str)
    assert celula.value == "05/03/2026 08:07:06"


def test_data_pagamento_ausente_fica_none(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    candidato = _candidato(data_pagamento=None)
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"B{_PRIMEIRA_LINHA_DADOS}"].value is None


def test_decisao_de_excluir_ocorrencia_nao_a_publica(tmp_path, template):
    ocorrencia_excluida = OcorrenciaFinanceira(
        dh="2026DH000099", ordem_bancaria=_ordem(ob="2026OB000099", dh="2026DH000099"), registros_controle=[]
    )
    decisao = DecisaoHumana(
        identidade_grupo="g2",
        regra="RN07",
        obs_envolvidas=frozenset({"2026OB000099"}),
        fingerprint_evidencias="fp2",
        acao=AcaoDecisao.PUBLICAR,
        obs_publicaveis=frozenset(),  # decisão real: publicar nenhuma das OBs envolvidas
        autor="teste@teste.com",
        decidido_em=datetime(2026, 1, 20),
    )
    resultado = _resultado(
        candidatos=[_candidato(ob="2026OB000001")],
        nao_publicaveis=[
            OcorrenciaNaoPublicavel(
                ocorrencia=ocorrencia_excluida,
                decisao=decisao,
                categoria=CategoriaNaoPublicada.NAO_PUBLICADA_POR_DECISAO_HUMANA,
            )
        ],
    )
    destino = str(tmp_path / "saida.xlsx")
    resultado_gerado = _gerar(resultado, template, destino)

    assert resultado_gerado.total_ocorrencias == 1
    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"A{_PRIMEIRA_LINHA_DADOS}"].value == "2026OB000001"
    assert ws[f"A{_PRIMEIRA_LINHA_DADOS + 1}"].value.startswith("Total")


def test_titulo_reflete_competencia_informada_nunca_derivada_dos_dados(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino, ano_exercicio=2026, mes_referencia=7)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws["A3"].value == "PAGAS ATÉ JULHO/2026"


def test_data_de_atualizacao_reflete_data_de_geracao(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    resultado_gerado = _gerar(
        _resultado(candidatos=[_candidato()]), template, destino, hoje=date(2026, 3, 5)
    )

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    celula = ws[f"A{resultado_gerado.linha_total + 2}"]
    assert celula.value == "Data da última atualização: 05.03.2026"


# --- Determinismo (item 22) -------------------------------------------------


def test_mesma_entrada_gera_o_mesmo_conteudo_logico(tmp_path, template):
    candidatos = [_candidato(ob="2026OB000001", valor=154.09), _candidato(ob="2026OB000002", valor=200.0)]
    destino_a = str(tmp_path / "saida_a.xlsx")
    destino_b = str(tmp_path / "saida_b.xlsx")

    _gerar(_resultado(candidatos=list(candidatos)), template, destino_a, hoje=date(2026, 2, 1))
    _gerar(_resultado(candidatos=list(candidatos)), template, destino_b, hoje=date(2026, 2, 1))

    wb_a = openpyxl.load_workbook(destino_a)
    wb_b = openpyxl.load_workbook(destino_b)
    ws_a, ws_b = wb_a[_NOME_ABA_DADOS], wb_b[_NOME_ABA_DADOS]
    for linha in range(_PRIMEIRA_LINHA_DADOS, _PRIMEIRA_LINHA_DADOS + 2):
        for coluna in "ABCDEFGHIJKLM":
            assert ws_a[f"{coluna}{linha}"].value == ws_b[f"{coluna}{linha}"].value


# --- Estrutura / formatação (itens 18-19) -----------------------------------


def test_abas_preservadas_mesmo_nome_e_ordem(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    wb_template = openpyxl.load_workbook(template)
    wb_gerado = openpyxl.load_workbook(destino)
    assert wb_gerado.sheetnames == wb_template.sheetnames


def test_cabecalho_e_subcabecalho_preservados(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    wb_template = openpyxl.load_workbook(template)
    wb_gerado = openpyxl.load_workbook(destino)
    ws_t, ws_g = wb_template[_NOME_ABA_DADOS], wb_gerado[_NOME_ABA_DADOS]
    for coluna in "ABCDEFGHIJKLM":
        assert ws_g[f"{coluna}5"].value == ws_t[f"{coluna}5"].value
        assert ws_g[f"{coluna}6"].value == ws_t[f"{coluna}6"].value


def test_larguras_de_coluna_preservadas(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    wb_template = openpyxl.load_workbook(template)
    wb_gerado = openpyxl.load_workbook(destino)
    ws_t, ws_g = wb_template[_NOME_ABA_DADOS], wb_gerado[_NOME_ABA_DADOS]
    for coluna in "ABCDEFGHIJKLM":
        assert ws_g.column_dimensions[coluna].width == ws_t.column_dimensions[coluna].width


def test_freeze_panes_e_page_setup_preservados(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    wb_template = openpyxl.load_workbook(template)
    wb_gerado = openpyxl.load_workbook(destino)
    ws_t, ws_g = wb_template[_NOME_ABA_DADOS], wb_gerado[_NOME_ABA_DADOS]
    assert ws_g.freeze_panes == ws_t.freeze_panes
    assert ws_g.page_setup.orientation == ws_t.page_setup.orientation
    assert ws_g.page_setup.paperSize == ws_t.page_setup.paperSize
    assert ws_g.page_margins.left == ws_t.page_margins.left
    assert ws_g.page_margins.right == ws_t.page_margins.right


def test_merges_estruturais_de_titulo_preservados(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    wb_gerado = openpyxl.load_workbook(destino)
    ws_g = wb_gerado[_NOME_ABA_DADOS]
    intervalos = {str(r) for r in ws_g.merged_cells.ranges}
    assert "A2:C2" in intervalos
    assert "E5:F5" in intervalos
    assert "K5:L5" in intervalos


def test_notas_de_rodape_preservadas_sem_alteracao(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    resultado_gerado = _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    wb_template = openpyxl.load_workbook(template)
    wb_gerado = openpyxl.load_workbook(destino)
    ws_t, ws_g = wb_template[_NOME_ABA_DADOS], wb_gerado[_NOME_ABA_DADOS]
    # As notas no template sintético começam 4 linhas abaixo do total original
    # (linha 10) — no gerado, a linha de total pode ter mudado; procuramos o
    # texto da 1ª nota em ambos e comparamos o conteúdo, não a posição.
    textos_template = [c.value for row in ws_t.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("(a)")]
    textos_gerado = [c.value for row in ws_g.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("(a)")]
    assert textos_template == textos_gerado


def test_estilo_da_primeira_linha_de_dados_bate_com_o_template(tmp_path, template):
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[_candidato()]), template, destino)

    wb_template = openpyxl.load_workbook(template)
    wb_gerado = openpyxl.load_workbook(destino)
    ws_t, ws_g = wb_template[_NOME_ABA_DADOS], wb_gerado[_NOME_ABA_DADOS]
    linha = _PRIMEIRA_LINHA_DADOS
    for coluna in "ABCDEFGHIJKLM":
        cel_t, cel_g = ws_t[f"{coluna}{linha}"], ws_g[f"{coluna}{linha}"]
        assert cel_g.font.name == cel_t.font.name
        assert cel_g.font.size == cel_t.font.size
        assert cel_g.border.left.style == cel_t.border.left.style
        assert cel_g.alignment.horizontal == cel_t.alignment.horizontal
    assert ws_g[f"M{linha}"].number_format == ws_t[f"M{linha}"].number_format


def test_estilo_de_linha_nova_alem_do_tamanho_original_do_template_e_copiado(tmp_path, template):
    # Template sintético tem só 3 linhas de dados (7-9) — 5 candidatos força
    # a inserção de 2 linhas novas (10 e 11) além do tamanho original.
    candidatos = [_candidato(ob=f"2026OB{i:06d}", dh=f"2026DH{i:06d}") for i in range(1, 6)]
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=candidatos), template, destino)

    wb_template = openpyxl.load_workbook(template)
    wb_gerado = openpyxl.load_workbook(destino)
    ws_t, ws_g = wb_template[_NOME_ABA_DADOS], wb_gerado[_NOME_ABA_DADOS]
    linha_modelo = _PRIMEIRA_LINHA_DADOS
    linha_nova = _PRIMEIRA_LINHA_DADOS + 4  # 5ª linha de dados, além do template original
    for coluna in "ABCDEFGHIJKLM":
        cel_modelo, cel_nova = ws_t[f"{coluna}{linha_modelo}"], ws_g[f"{coluna}{linha_nova}"]
        assert cel_nova.font.name == cel_modelo.font.name
        assert cel_nova.font.size == cel_modelo.font.size
        assert cel_nova.border.left.style == cel_modelo.border.left.style
        assert cel_nova.alignment.horizontal == cel_modelo.alignment.horizontal


def test_menos_ocorrencias_que_o_template_original_remove_linhas_extras(tmp_path, template):
    # Template sintético tem 3 linhas de dados; só 1 candidato deve deixar a
    # planilha final com 1 linha de dados, nunca 2 linhas históricas residuais.
    destino = str(tmp_path / "saida.xlsx")
    resultado_gerado = _gerar(_resultado(candidatos=[_candidato(ob="2026OB000001")]), template, destino)

    wb_gerado = openpyxl.load_workbook(destino)
    ws_g = wb_gerado[_NOME_ABA_DADOS]
    assert resultado_gerado.linha_total == _PRIMEIRA_LINHA_DADOS + 1
    assert ws_g[f"A{_PRIMEIRA_LINHA_DADOS + 1}"].value.startswith("Total")


def test_formula_de_total_e_recalculada_para_a_faixa_real_nunca_herda_faixa_velha(tmp_path, template):
    candidatos = [_candidato(ob="2026OB000001", valor=100.0), _candidato(ob="2026OB000002", valor=50.0)]
    destino = str(tmp_path / "saida.xlsx")
    resultado_gerado = _gerar(_resultado(candidatos=candidatos), template, destino)

    wb_gerado = openpyxl.load_workbook(destino)
    ws_g = wb_gerado[_NOME_ABA_DADOS]
    formula = ws_g[f"M{resultado_gerado.linha_total}"].value
    assert formula == f"=SUM(M{_PRIMEIRA_LINHA_DADOS}:M{_PRIMEIRA_LINHA_DADOS + 1})"


# --- Micro-sprint 3G.1 — precedência Cargo/Transporte: Controle DIMCAC primeiro, Base OB como fallback


def test_cargo_do_controle_prevalece_sobre_base_ob(tmp_path, template):
    registros = [_registro(cargo="Promotor de Justiça de 1a. Entrância")]
    candidato = _candidato(registros_controle=registros, cargo="Cargo da Base OB")
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"D{_PRIMEIRA_LINHA_DADOS}"].value == "Promotor de Justiça de 1a. Entrância"


def test_cargo_vazio_no_controle_usa_base_ob(tmp_path, template):
    registros = [_registro(cargo=None)]
    candidato = _candidato(registros_controle=registros, cargo="Cargo da Base OB")
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"D{_PRIMEIRA_LINHA_DADOS}"].value == "Cargo da Base OB"


def test_cargo_ausente_em_ambas_as_fontes_fica_vazio(tmp_path, template):
    registros = [_registro(cargo=None)]
    candidato = _candidato(registros_controle=registros, cargo=None)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"D{_PRIMEIRA_LINHA_DADOS}"].value is None


def test_transporte_do_controle_prevalece_sobre_base_ob(tmp_path, template):
    registros = [_registro(transporte="AVIÃO")]
    candidato = _candidato(registros_controle=registros, transporte="CARRO")
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"H{_PRIMEIRA_LINHA_DADOS}"].value == "AVIÃO"


def test_transporte_vazio_no_controle_usa_base_ob(tmp_path, template):
    registros = [_registro(transporte=None)]
    candidato = _candidato(registros_controle=registros, transporte="CARRO")
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"H{_PRIMEIRA_LINHA_DADOS}"].value == "CARRO"


def test_transporte_ausente_em_ambas_as_fontes_fica_vazio(tmp_path, template):
    registros = [_registro(transporte=None)]
    candidato = _candidato(registros_controle=registros, transporte=None)
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"H{_PRIMEIRA_LINHA_DADOS}"].value is None


def test_cargo_conflitante_entre_registros_de_controle_cai_no_fallback(tmp_path, template):
    registros = [
        _registro(linha=2, cargo="Cargo A do Controle"),
        _registro(linha=3, cargo="Cargo B do Controle"),
    ]
    candidato = _candidato(registros_controle=registros, cargo="Cargo da Base OB")
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"D{_PRIMEIRA_LINHA_DADOS}"].value == "Cargo da Base OB"


def test_demais_campos_continuam_com_as_fontes_ja_homologadas(tmp_path, template):
    """Cargo/Transporte mudaram de fonte (3G.1) — período, trecho, motivo e
    valor continuam vindo da Base OB mesmo quando o Controle DIMCAC tem um
    valor DIFERENTE nesses mesmos campos (Controle também carrega
    trecho/motivo próprios, usados só para a reconciliação, nunca para a
    escrita do relatório)."""
    registros = [_registro()]
    candidato = _candidato(
        registros_controle=registros,
        periodo_inicial=date(2026, 3, 1),
        periodo_final=date(2026, 3, 2),
        trecho="TRECHO DA BASE OB",
        motivo="Motivo da Base OB",
        valor=999.99,
    )
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"E{_PRIMEIRA_LINHA_DADOS}"].value == datetime(2026, 3, 1)
    assert ws[f"F{_PRIMEIRA_LINHA_DADOS}"].value == datetime(2026, 3, 2)
    assert ws[f"G{_PRIMEIRA_LINHA_DADOS}"].value == "TRECHO DA BASE OB"
    assert ws[f"I{_PRIMEIRA_LINHA_DADOS}"].value == "Motivo da Base OB"
    assert ws[f"M{_PRIMEIRA_LINHA_DADOS}"].value == 999.99


# --- Nome x Credor (Micro-sprint corretiva, decisão de negócio aprovada) ---
#
# Base OB.NOME continua a fonte primária do campo Nome; só cai para
# Base OB.CREDOR quando houver relação de truncamento/prefixo COMPROVADA
# (nunca por NOME ter 30 caracteres, nunca só por serem diferentes — ver
# `_nome_publicado` e o relatório da investigação read-only dedicada).


def test_nome_igual_credor_publica_nome():
    ordem = _ordem(nome="FULANO DE TAL", credor="FULANO DE TAL")
    assert _nome_publicado(ordem) == "FULANO DE TAL"


def test_nome_prefixo_de_credor_publica_credor():
    ordem = _ordem(
        nome="JOÃO CARLOS ALBUQUERQUE DE OLI",
        credor="JOÃO CARLOS ALBUQUERQUE DE OLIVEIRA",
    )
    assert _nome_publicado(ordem) == "JOÃO CARLOS ALBUQUERQUE DE OLIVEIRA"


def test_nome_de_30_caracteres_mas_nao_prefixo_publica_nome():
    nome_30 = "A" * 30
    assert len(nome_30) == 30
    ordem = _ordem(nome=nome_30, credor="B" * 40)
    assert _nome_publicado(ordem) == nome_30


def test_nome_menor_que_30_mas_prefixo_comprovado_publica_credor():
    ordem = _ordem(nome="ANA LIMA", credor="ANA LIMA DE ALBUQUERQUE SOUZA")
    assert len(ordem.nome) < 30
    assert _nome_publicado(ordem) == "ANA LIMA DE ALBUQUERQUE SOUZA"


def test_nome_maior_que_30_mas_prefixo_comprovado_publica_credor():
    """Regra é sobre RELAÇÃO, não tamanho — mesmo NOME já tendo mais de 30
    caracteres, se ainda for prefixo comprovado de um CREDOR mais longo,
    publica CREDOR."""
    nome_35 = "MARIA DAS GRAÇAS PEREIRA DE ALBUQ"
    assert len(nome_35) > 30
    credor_completo = "MARIA DAS GRAÇAS PEREIRA DE ALBUQUERQUE"
    ordem = _ordem(nome=nome_35, credor=credor_completo)
    assert _nome_publicado(ordem) == credor_completo


def test_nome_maior_que_30_e_nao_prefixo_publica_nome():
    """Mesmo tamanho grande, sem relação de prefixo comprovada -> mantém
    NOME (não decide por tamanho)."""
    nome_35 = "MARIA DAS GRAÇAS PEREIRA DE ALBUQ"
    assert len(nome_35) > 30
    ordem = _ordem(nome=nome_35, credor="OUTRO NOME COMPLETAMENTE DIFERENTE")
    assert _nome_publicado(ordem) == nome_35


def test_nome_e_credor_com_trim_externo_ainda_detecta_continuacao():
    ordem = _ordem(nome="  JOÃO CARLOS ALBUQUERQUE DE OLI  ", credor="  JOÃO CARLOS ALBUQUERQUE DE OLIVEIRA  ")
    assert _nome_publicado(ordem) == "JOÃO CARLOS ALBUQUERQUE DE OLIVEIRA"


def test_nome_e_credor_grafias_diferentes_publica_nome():
    """Divergência de grafia (não truncamento) -> mantém NOME, nunca tenta
    decidir qual grafia é correta."""
    ordem = _ordem(nome="LUIS ALBUQUERQUE PEREIRA", credor="LUIZ ALBUQUERQUE PEREIRA DE SOUZA")
    assert _nome_publicado(ordem) == "LUIS ALBUQUERQUE PEREIRA"


def test_nome_com_acento_diferente_publica_nome():
    ordem = _ordem(nome="JOSE ALBUQUERQUE", credor="JOSÉ ALBUQUERQUE PEREIRA")
    assert _nome_publicado(ordem) == "JOSE ALBUQUERQUE"


def test_credor_menor_que_nome_publica_nome():
    ordem = _ordem(nome="JOÃO CARLOS ALBUQUERQUE DE OLIVEIRA", credor="JOÃO CARLOS")
    assert _nome_publicado(ordem) == "JOÃO CARLOS ALBUQUERQUE DE OLIVEIRA"


def test_ausencia_de_credor_publica_nome():
    ordem = _ordem(nome="FULANO DE TAL", credor=None)
    assert _nome_publicado(ordem) == "FULANO DE TAL"


def test_ausencia_de_nome_preserva_comportamento_atual():
    """NOME ausente não aciona nenhuma regra nova — comportamento
    preservado, compatível com a arquitetura atual (a mesma coisa que
    aconteceria hoje ao escrever `ordem.nome` direto: célula fica vazia,
    nunca inventa um valor a partir de CREDOR)."""
    ordem = _ordem(nome=None, credor="ALGUEM DE TAL")
    assert _nome_publicado(ordem) is None


def test_nome_publicado_e_escrito_de_fato_no_xlsx_gerado(tmp_path, template):
    """Confirma a fiação ponta-a-ponta (não só a função pura isolada)."""
    candidato = _candidato(
        nome="JOÃO CARLOS ALBUQUERQUE DE OLI",
        credor="JOÃO CARLOS ALBUQUERQUE DE OLIVEIRA",
    )
    destino = str(tmp_path / "saida.xlsx")
    _gerar(_resultado(candidatos=[candidato]), template, destino)

    wb = openpyxl.load_workbook(destino)
    ws = wb[_NOME_ABA_DADOS]
    assert ws[f"C{_PRIMEIRA_LINHA_DADOS}"].value == "JOÃO CARLOS ALBUQUERQUE DE OLIVEIRA"
