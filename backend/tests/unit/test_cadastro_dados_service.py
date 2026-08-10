import pytest
from openpyxl import load_workbook

from app.consulta import base_local
from app.infra.config import configuracao_app
from app.planilha import escritor_dados
from app.services import cadastro_dados_service
from app.services.cadastro_dados_service import (
    EstadoCadastroServico,
    cadastrar_servidor_membro,
    obter_opcoes_cadastro,
)
from app.services.conferencia_service import BaseLocalIndisponivel


def _isolar_base_local(monkeypatch, tmp_path, nome_pasta="dados_usuario"):
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / nome_pasta))


def _apontar_planilha_oficial(monkeypatch, caminho):
    monkeypatch.setattr(configuracao_app, "resolver_caminho_para_preenchimento", lambda: caminho)


def test_cadastro_novo_cpf_sincroniza_e_confirma_concluido(
    monkeypatch, tmp_path, caminho_planilha_conferencia_teste
):
    caminho = caminho_planilha_conferencia_teste()
    _isolar_base_local(monkeypatch, tmp_path)
    _apontar_planilha_oficial(monkeypatch, caminho)

    resultado = cadastrar_servidor_membro(
        cpf="111.222.333-44",
        nome="Novo Servidor de Teste",
        matricula="9999",
        cargo="Analista Judiciário",
    )

    assert resultado.estado is EstadoCadastroServico.CONCLUIDO
    assert resultado.servidor is not None
    assert resultado.servidor.nome == "Novo Servidor de Teste"
    assert resultado.servidor.cargo == "Analista Judiciário"

    workbook = load_workbook(caminho)
    assert workbook["Dados"]["A4"].value == 11122233344


def test_cadastro_cpf_duplicado_carrega_servidor_ja_cadastrado(
    monkeypatch, tmp_path, caminho_planilha_conferencia_teste
):
    caminho = caminho_planilha_conferencia_teste()
    _isolar_base_local(monkeypatch, tmp_path)
    _apontar_planilha_oficial(monkeypatch, caminho)
    # Sincroniza a Base Local ANTES do cadastro — a fixture já traz o CPF
    # 12345678901 (Maria Aparecida Souza) na linha 2 da aba Dados.
    base_local.sincronizar(caminho)

    resultado = cadastrar_servidor_membro(
        cpf="123.456.789-01",  # mascarado, mesmo CPF já cadastrado
        nome="Outro Nome Qualquer",
        matricula="1",
        cargo="Outro Cargo",
    )

    assert resultado.estado is EstadoCadastroServico.DUPLICADO
    assert resultado.servidor is not None
    assert resultado.servidor.nome == "Maria Aparecida Souza"
    # Nada novo foi escrito — só a linha 2 (já existente) e a 3 continuam.
    workbook = load_workbook(caminho)
    assert workbook["Dados"]["A4"].value is None


def test_cadastro_planilha_oficial_indisponivel_vira_nao_salvo(monkeypatch):
    mensagem = "Nenhuma planilha oficial configurada. Configure em 'Alterar Planilha Oficial'."

    def _levantar():
        raise configuracao_app.PlanilhaOficialIndisponivel(mensagem)

    monkeypatch.setattr(configuracao_app, "resolver_caminho_para_preenchimento", _levantar)

    resultado = cadastrar_servidor_membro(cpf="11122233344", nome="X", matricula="1", cargo="Y")

    assert resultado.estado is EstadoCadastroServico.NAO_SALVO
    assert resultado.mensagem == mensagem


def test_cadastro_com_sincronizacao_falhando_retorna_salvo_sync_pendente(
    monkeypatch, tmp_path, caminho_planilha_conferencia_teste
):
    """Rule 10 do pedido aprovado: cadastro salvo + sincronização falhada
    NUNCA é reportado como falha — mas também nunca é `concluido` sem o
    CPF ter sido reconfirmado pela mesma fonte que a Conferência usa."""
    caminho = caminho_planilha_conferencia_teste()
    _isolar_base_local(monkeypatch, tmp_path, nome_pasta="dados_usuario_sem_sync_previo")
    _apontar_planilha_oficial(monkeypatch, caminho)

    def _sincronizar_com_falha(*args, **kwargs):
        raise RuntimeError("falha simulada de sincronização")

    monkeypatch.setattr(cadastro_dados_service.base_local, "sincronizar", _sincronizar_com_falha)

    resultado = cadastrar_servidor_membro(
        cpf="555.666.777-88", nome="Servidor Sem Sync", matricula="42", cargo="Técnico Judiciário"
    )

    assert resultado.estado is EstadoCadastroServico.SALVO_SYNC_PENDENTE
    assert resultado.mensagem == cadastro_dados_service._MENSAGEM_SYNC_PENDENTE
    assert resultado.servidor is None

    # A escrita em si teve sucesso — só a sincronização falhou.
    workbook = load_workbook(caminho)
    assert workbook["Dados"]["A4"].value == 55566677788


def test_reenviar_apos_salvo_sync_pendente_sincroniza_de_verdade_e_confirma(
    monkeypatch, tmp_path, caminho_planilha_conferencia_teste
):
    """§H — "Tentar sincronizar novamente" reenvia a MESMA requisição (nunca
    um endpoint separado): na 2ª tentativa a checagem de duplicidade já
    encontra a linha gravada na 1ª (== `DUPLICADO`), mas isso não pode virar
    uma sincronização "de mentira" — o retry só faz sentido se realmente
    tentar `sincronizar()` de novo."""
    caminho = caminho_planilha_conferencia_teste()
    _isolar_base_local(monkeypatch, tmp_path, nome_pasta="dados_usuario_retry")
    _apontar_planilha_oficial(monkeypatch, caminho)
    sincronizar_real = base_local.sincronizar  # capturado ANTES do monkeypatch abaixo

    def _sincronizar_com_falha(*args, **kwargs):
        raise RuntimeError("falha simulada de sincronização")

    monkeypatch.setattr(cadastro_dados_service.base_local, "sincronizar", _sincronizar_com_falha)
    parametros = {
        "cpf": "555.666.777-88",
        "nome": "Servidor Sem Sync",
        "matricula": "42",
        "cargo": "Técnico Judiciário",
    }
    primeira_tentativa = cadastrar_servidor_membro(**parametros)
    assert primeira_tentativa.estado is EstadoCadastroServico.SALVO_SYNC_PENDENTE

    # Restaura a sincronização real (só a 1ª tentativa falhava) e reenvia
    # exatamente a mesma requisição.
    monkeypatch.setattr(cadastro_dados_service.base_local, "sincronizar", sincronizar_real)

    segunda_tentativa = cadastrar_servidor_membro(**parametros)

    assert segunda_tentativa.estado is EstadoCadastroServico.DUPLICADO
    assert segunda_tentativa.servidor is not None
    assert segunda_tentativa.servidor.nome == "Servidor Sem Sync"
    # Nenhuma segunda linha foi gravada — só a 4 (da 1ª tentativa).
    workbook = load_workbook(caminho)
    assert workbook["Dados"]["A5"].value is None


def test_cadastro_escrita_falha_nunca_tenta_sincronizar(
    monkeypatch, tmp_path, caminho_planilha_conferencia_teste
):
    caminho = caminho_planilha_conferencia_teste()
    _isolar_base_local(monkeypatch, tmp_path)
    _apontar_planilha_oficial(monkeypatch, caminho)

    monkeypatch.setattr(
        cadastro_dados_service.escritor_dados,
        "cadastrar_servidor",
        lambda *args, **kwargs: escritor_dados.ResultadoCadastroDados(
            estado=escritor_dados.EstadoCadastro.NAO_SALVO, mensagem="falha simulada de escrita"
        ),
    )
    chamou_sincronizar = False

    def _sincronizar_espiao(*args, **kwargs):
        nonlocal chamou_sincronizar
        chamou_sincronizar = True

    monkeypatch.setattr(cadastro_dados_service.base_local, "sincronizar", _sincronizar_espiao)

    resultado = cadastrar_servidor_membro(cpf="111.222.333-44", nome="X", matricula="1", cargo="Y")

    assert resultado.estado is EstadoCadastroServico.NAO_SALVO
    assert resultado.mensagem == "falha simulada de escrita"
    assert chamou_sincronizar is False


def test_obter_opcoes_cadastro_le_valores_reais_da_base_local_sincronizada(
    monkeypatch, tmp_path, caminho_planilha_conferencia_teste
):
    """Mesma fonte (Base Local, RN-01) já usada por `_consultar_servidor_atual`
    — nunca uma lista hardcoded no backend nem no frontend (item 8 do pedido)."""
    caminho = caminho_planilha_conferencia_teste()
    _isolar_base_local(monkeypatch, tmp_path)
    base_local.sincronizar(caminho)

    opcoes = obter_opcoes_cadastro()

    assert opcoes.cargos == ["Analista Judiciário", "Técnico Judiciário"]
    assert opcoes.previdencias == ["FUNAFIN", "INSS"]
    assert opcoes.regimes_previdencia == ["RGPS", "RPPS"]
    assert opcoes.situacoes == ["Em exercício"]


def test_obter_opcoes_cadastro_levanta_base_local_indisponivel_sem_sync_previo(monkeypatch, tmp_path):
    _isolar_base_local(monkeypatch, tmp_path, nome_pasta="dados_usuario_sem_opcoes")

    with pytest.raises(BaseLocalIndisponivel):
        obter_opcoes_cadastro()
