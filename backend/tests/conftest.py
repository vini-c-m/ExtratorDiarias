import os
import re
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook

from app.consulta.base_local import VAR_AMBIENTE_USER_DATA


@pytest.fixture(autouse=True)
def _isolar_base_local_da_appdata_real(tmp_path, monkeypatch):
    """Sprint 13 — o hook de sincronização da Base Local
    (`app/services/planilha_service.py`) é engolido em silêncio quando
    falha, então sem isto qualquer teste que chame `preencher_planilha`
    escreveria de verdade em `%APPDATA%/ExtratorDiarias` na máquina que
    roda os testes, sem nenhum teste perceber. `autouse=True` garante que
    isso nunca aconteça, mesmo em testes que não sabem que a Central
    Inteligente existe. Testes que precisam controlar o caminho da Base
    Local explicitamente (`tests/integration/test_base_local.py`) sobrescrevem
    esta variável com o seu próprio `monkeypatch`."""
    monkeypatch.setenv(VAR_AMBIENTE_USER_DATA, str(tmp_path / "_base_local_isolada_de_testes"))

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures", "pdfs_exemplo")
FIXTURES_DIR_DIARIA_MEMBRO = os.path.join(
    os.path.dirname(__file__), "fixtures", "pdfs_diaria_membro"
)
FIXTURES_DIR_LOCAL_DESTINO = os.path.join(
    os.path.dirname(__file__), "fixtures", "pdfs_local_destino"
)
FIXTURES_DIR_PADROES_DOCUMENTAIS = os.path.join(
    os.path.dirname(__file__), "fixtures", "pdfs_padroes_documentais"
)


@pytest.fixture
def pdfs_exemplo() -> list[str]:
    return [
        os.path.join(FIXTURES_DIR, arquivo)
        for arquivo in sorted(os.listdir(FIXTURES_DIR))
        if arquivo.lower().endswith(".pdf")
    ]


@pytest.fixture
def pdfs_diaria_membro() -> list[str]:
    """3 processos fictícios de Diária de Membro (Sprint 10 — dados
    sintéticos, ver `gerar_fixtures_sinteticas.py`), reproduzindo o padrão
    documental estudado na implementação original da Sprint 04 — mantidos
    num diretório próprio, separado de `pdfs_exemplo`, para não alterar a
    contagem/composição usada pelos testes já existentes do fluxo Padrão."""
    return [
        os.path.join(FIXTURES_DIR_DIARIA_MEMBRO, arquivo)
        for arquivo in sorted(os.listdir(FIXTURES_DIR_DIARIA_MEMBRO))
        if arquivo.lower().endswith(".pdf")
    ]


@pytest.fixture
def caminho_pdf():
    def _caminho(nome: str) -> str:
        return os.path.join(FIXTURES_DIR, nome)

    return _caminho


@pytest.fixture
def caminho_pdf_processo_sta_cruz_capibaribe() -> str:
    """Processo fictício (Sprint 10 — dados sintéticos) reproduzindo o
    padrão do caso original da Sprint 04.1 (refinamento de LOCAL):
    'CIDADE(S) DESTINO:' vem como 'Sta Cuz do Capibaribe' (abreviação + erro
    de digitação combinados), e deve resolver para PERNAMBUCO. Mantido num
    diretório próprio, separado de `pdfs_exemplo`, para não alterar a
    contagem/composição usada pelos testes já existentes."""
    return os.path.join(
        FIXTURES_DIR_LOCAL_DESTINO, "SEI_19.20.0001.0000910_2025_10.pdf"
    )


# Sprint 04.2 — processos fictícios (Sprint 10 — dados sintéticos)
# reproduzindo os 4 padrões documentais identificados na homologação
# original. Num diretório próprio, separado de `pdfs_exemplo`, para não
# alterar a contagem/composição usada pelos testes já existentes.
@pytest.fixture
def caminho_pdf_multiplos_requerimentos_membro() -> str:
    """Processo fictício (Caso 1): dois "CMFC - Requerimento de Diária -
    Membro" anexados — o primeiro com a data de ida registrada com o ano
    errado (2025), corrigido por um segundo requerimento (2026) depois que a
    DIMCAC apontou o erro num despacho. Só o último deve valer."""
    return os.path.join(
        FIXTURES_DIR_PADROES_DOCUMENTAIS, "SEI_19.20.0001.0000103_2026_03.pdf"
    )


@pytest.fixture
def caminho_pdf_data_retorno_confirmacao_servicos() -> str:
    """Processo fictício (Caso 2): a "Solicitação de Diária" original tem
    "DATA DE RETORNO: 22/05/2026" corretamente, mas o valor cai bem numa
    quebra de página (rodapé do SEI entre o rótulo e o valor) — sem a
    tolerância a essa quebra, a extração pulava para a "Confirmação dos
    Serviços" (documento diferente, preenchido depois), que tem um erro de
    digitação registrando o retorno em 22/06/2026."""
    return os.path.join(
        FIXTURES_DIR_PADROES_DOCUMENTAIS, "SEI_19.20.0001.0000104_2026_04.pdf"
    )


@pytest.fixture
def caminho_pdf_duas_validacoes_diaria() -> str:
    """Processo fictício (Caso 3): duas "CMFC - Validação de Diária -
    Servidores" anexadas — a segunda traz literalmente "DESCONSIDERAR
    VALIDAÇÃO ANTERIOR" nas Observações e um Documento Hábil diferente do da
    primeira. Só o DH da última deve valer."""
    return os.path.join(
        FIXTURES_DIR_PADROES_DOCUMENTAIS, "SEI_19.20.0001.0000105_2026_05.pdf"
    )


@pytest.fixture
def caminho_pdf_le_de_outro_processo_anexada_por_engano() -> str:
    """Processo fictício (Sprint — Refatoração DH): uma única "CMFC -
    Validação de Diária - Servidores" válida, mas DUAS Notas de Liquidação
    anexadas — a primeira pertence, ficticiamente, a OUTRO processo (credor
    diferente), anexada por engano; a segunda é a correta. Como a fonte
    oficial do Documento Hábil é sempre o Despacho de Validação — nunca a
    Nota de Liquidação —, a extração correta ignora as duas Notas de
    Liquidação por completo e usa o DH citado nas Observações da própria
    Validação."""
    return os.path.join(
        FIXTURES_DIR_PADROES_DOCUMENTAIS, "SEI_19.20.0001.0000106_2023_06.pdf"
    )


@pytest.fixture
def caminho_pdf_destino_pj_palmares() -> str:
    """Processo fictício (Caso 4): "CIDADE(S) DESTINO: PJ PALMARES" — "PJ" é
    o prefixo institucional (Procuradoria/Promotoria de Justiça), não faz
    parte do nome do município."""
    return os.path.join(
        FIXTURES_DIR_PADROES_DOCUMENTAIS, "SEI_19.20.0001.0000107_2026_07.pdf"
    )


@pytest.fixture
def caminho_pdf_cpf_com_espaco_interno() -> str:
    """Processo fictício (Sprint de Correção Crítica — CPF): o texto
    extraído do PDF traz 'CPF:\\n333.333 333-01' -- um espaço comum no lugar
    do ponto entre o 2º e o 3º grupo do CPF. O CPF correto é
    333.333.333-01."""
    return os.path.join(
        FIXTURES_DIR_PADROES_DOCUMENTAIS, "SEI_19.20.0001.0000108_2026_08.pdf"
    )


@pytest.fixture
def caminho_pdf_destino_pj_bom_jardim() -> str:
    """Processo fictício (Caso 4): "CIDADE(S) DESTINO: PJ Bom Jardim" com
    "CIDADE ORIGEM: São José do Egito" — "Bom Jardim" sozinho é ambíguo
    (existe em PE/MA/RJ), resolvido pelo contexto da cidade de origem depois
    que o prefixo "PJ" é removido."""
    return os.path.join(
        FIXTURES_DIR_PADROES_DOCUMENTAIS, "SEI_19.20.0001.0000109_2026_09.pdf"
    )


_CABECALHO_PLANILHA = [
    "DH",
    "SEI",
    "DOC ORIGEM",
    "CPF",
    "PREVIDÊNCIA",
    "NOME",
    "CARGO",
    "PERÍODO INICIAL",
    "PERÍODO FINAL",
    "TRECHO",
    "TRANSPORTE",
    "MOTIVO",
]


@pytest.fixture
def caminho_planilha_teste(tmp_path):
    """Planilha sintética (sem dados reais) reproduzindo a estrutura real:
    mesma aba/cabeçalho, uma linha totalmente preenchida, e uma linha com a
    coluna A (DH) vazia mas B em diante preenchida — o mesmo padrão de gap
    encontrado na planilha real, que invalida usar a coluna A como indicador
    de linha livre."""

    def _criar(nome_arquivo: str = "planilha_teste.xlsx") -> str:
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Controle de Diárias"
        planilha.append(_CABECALHO_PLANILHA)

        # Linha 2: registro totalmente preenchido (formulas simuladas como texto).
        planilha.append(
            [
                "2025DH000001",
                "19.20.0000.0000001/2025-00",
                "PJTESTE",
                11111111111,
                "=VLOOKUP($D2,Dados!$A:$E,4,FALSE)",
                "=VLOOKUP($D2,Dados!$A:$E,3,FALSE)",
                "=VLOOKUP($D2,Dados!$A:$E,5,FALSE)",
                "08/01/2025",
                "10/01/2025",
                "CIDADE A-CIDADE B",
                "CARRO",
                "Motivo de teste",
            ]
        )

        # Linha 3: coluna A vazia, mas B em diante preenchida (gap real).
        planilha.append(
            [
                None,
                "19.20.0000.0000002/2025-11",
                "PJTESTE2",
                22222222222,
                "=VLOOKUP($D3,Dados!$A:$E,4,FALSE)",
                "=VLOOKUP($D3,Dados!$A:$E,3,FALSE)",
                "=VLOOKUP($D3,Dados!$A:$E,5,FALSE)",
                "09/01/2025",
                "11/01/2025",
                "CIDADE C-CIDADE D",
                "AVIÃO",
                "Outro motivo de teste",
            ]
        )

        # P0 FINAL (2026-08-09) — preparação estrutural automática (ver
        # `escritor_planilha._preparar_estrutura_futura`) agora audita
        # N/P/R/S/V/Z/AA em TODO fluxo, não só no Complemento: sem uma
        # linha-modelo íntegra com essas 7 colunas em fórmula, `preencher()`
        # levanta `EstruturaFuturaIndisponivel`. Linhas 2 e 3 desta fixture
        # ganham essas fórmulas (mesmo padrão real confirmado em
        # `test_escritor_planilha_complemento.py`) só para servir de modelo —
        # nenhuma coluna A-L existente muda, então nenhuma asserção já
        # homologada sobre B/C/D/H/I/J/L é afetada.
        for linha_modelo in (2, 3):
            planilha[f"N{linha_modelo}"] = f"=I{linha_modelo}-H{linha_modelo}"
            planilha[f"P{linha_modelo}"] = f"=R{linha_modelo}+S{linha_modelo}"
            planilha[f"R{linha_modelo}"] = (
                f"=(VLOOKUP(Q{linha_modelo},'Tabela Valores'!A:C,2,FALSE()))*N{linha_modelo}"
            )
            planilha[f"S{linha_modelo}"] = (
                f"=(VLOOKUP(Q{linha_modelo},'Tabela Valores'!A:C,3,FALSE()))*O{linha_modelo}"
            )
            planilha[f"V{linha_modelo}"] = f"=CONCATENATE(D{linha_modelo},H{linha_modelo},I{linha_modelo})"
            planilha[f"Z{linha_modelo}"] = f"=H{linha_modelo}-15"
            planilha[f"AA{linha_modelo}"] = (
                f"=VLOOKUP(Z{linha_modelo},Z$1:AB{linha_modelo - 1},3,0)"
            )

        workbook.create_sheet("Dados")
        workbook.create_sheet("Tabela Valores")

        caminho = tmp_path / nome_arquivo
        workbook.save(caminho)
        return str(caminho)

    return _criar


# Sprint 15 — Configuração Inicial da Planilha Oficial. Cabeçalho completo
# (inclui P/Q, que `caminho_planilha_teste` acima não tem) porque a
# validação de configuração (§5 do doc arquitetural) confere exatamente as
# colunas que `escritor_planilha.preencher` escreve, e nada além delas.
_CABECALHO_PLANILHA_COMPLETO = {
    "A": "DH",
    "B": "SEI",
    "C": "DOC ORIGEM",
    "D": "CPF",
    "E": "PREVIDÊNCIA",
    "F": "NOME",
    "G": "CARGO",
    "H": "PERÍODO INICIAL",
    "I": "PERÍODO FINAL",
    "J": "TRECHO",
    "K": "TRANSPORTE",
    "L": "MOTIVO",
    "M": "VALOR DA PASSAGEM",
    "N": "QUANT. DIÁRIA INTEGRAL",
    "O": "QUANT. DIÁRIA PARCIAL",
    "P": "VALOR TOTAL DAS DIÁRIAS",
    "Q": "LOCAL",
}


@pytest.fixture
def caminho_planilha_configuracao_teste(tmp_path):
    """Planilha sintética própria da Sprint 15 (Configuração Inicial) —
    cabeçalho completo, incluindo as colunas P/Q que a fixture de
    preenchimento (`caminho_planilha_teste`) não tem. Aceita um cabeçalho
    alternativo para os testes negativos (rótulo errado/faltando)."""

    def _criar(
        nome_arquivo: str = "planilha_configuracao_teste.xlsx",
        *,
        cabecalho: dict[str, str] | None = None,
        abas_extra: tuple[str, ...] = ("Dados", "Tabela Valores"),
    ) -> str:
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Controle de Diárias"
        for coluna, rotulo in (cabecalho if cabecalho is not None else _CABECALHO_PLANILHA_COMPLETO).items():
            planilha[f"{coluna}1"] = rotulo

        for aba in abas_extra:
            workbook.create_sheet(aba)

        caminho = tmp_path / nome_arquivo
        workbook.save(caminho)
        return str(caminho)

    return _criar


_CABECALHO_DADOS = [
    "SEI",  # rotulada errado na planilha oficial — conteúdo real é CPF (RN-01)
    "MATRÍCULA",
    "NOME",
    "PREVIDÊNCIA",
    "CARGO",
    "ÓRGÃO DE ORIGEM",
    "REGIME DE PREVIDÊNCIA",
    "SITUAÇÃO DO SERVIDOR",
]


@pytest.fixture
def remover_tag_dimension_da_planilha():
    """Micro-sprint 18.7.1 — reproduz o caso real que quebrava
    `base_local_sqlite.capturar_sinal_frescor`: planilhas que passaram
    pelo Google Sheets (confirmado numa planilha real desta instalação,
    via a relação `.../customschemas.google.com/...` em
    `xl/_rels/workbook.xml.rels`) podem não gravar a tag `<dimension>` no
    XML da aba — sem ela, `openpyxl` com `read_only=True` não calcula
    `max_row` sem escanear célula por célula, e o atributo volta `None`.
    Remove a tag diretamente do `.xlsx` (arquivo ZIP) para reproduzir essa
    condição de forma determinística, sem depender de gerar o arquivo via
    outra ferramenta."""

    def _remover(caminho_xlsx: str, nome_arquivo_interno: str = "xl/worksheets/sheet1.xml") -> None:
        with zipfile.ZipFile(caminho_xlsx, "r") as zf:
            entradas = {info.filename: zf.read(info.filename) for info in zf.infolist()}

        conteudo_original = entradas[nome_arquivo_interno].decode("utf-8")
        conteudo_sem_dimension = re.sub(r"<dimension[^>]*/>", "", conteudo_original)
        assert conteudo_sem_dimension != conteudo_original, "a tag <dimension> não foi encontrada para remoção"
        entradas[nome_arquivo_interno] = conteudo_sem_dimension.encode("utf-8")

        with zipfile.ZipFile(caminho_xlsx, "w", zipfile.ZIP_DEFLATED) as zf:
            for nome, dado in entradas.items():
                zf.writestr(nome, dado)

    return _remover


@pytest.fixture
def injetar_cache_formula():
    """URGENTE — Central Analítica (2026-08-06): reproduz o que o Excel
    real grava ao salvar uma célula de fórmula — o elemento `<v>` com o
    último valor calculado, ao lado de `<f>` — algo que `openpyxl` nunca
    calcula sozinho (ao escrever `ws["P2"] = "=R2+S2"` e salvar, o XML sai
    com `<v />` vazio, nunca com o número). Sem essa injeção, um teste que
    monta a planilha só com `openpyxl` nunca reproduziria o caso real que
    causou a perda de valores (coluna `P` = `VALOR TOTAL DAS DIÁRIAS`,
    fórmula `=R+S`, cache já calculado pelo Excel — 100% das 2.217 linhas
    de fórmula da Planilha Oficial real tinham cache; a perda era só de
    leitura, nunca de cálculo). Omitir a chamada a esta fixture para uma
    célula de fórmula reproduz o outro caso real, genuinamente sem cache
    (`data_only=True` devolve `None`, nunca inventa um número)."""

    def _injetar(
        caminho_xlsx: str,
        referencia_celula: str,
        valor_cache: float,
        nome_arquivo_interno: str = "xl/worksheets/sheet1.xml",
    ) -> None:
        with zipfile.ZipFile(caminho_xlsx, "r") as zf:
            entradas = {info.filename: zf.read(info.filename) for info in zf.infolist()}

        conteudo_original = entradas[nome_arquivo_interno].decode("utf-8")
        padrao = re.compile(rf'(<c r="{referencia_celula}"[^>]*><f>[^<]*</f>)<v\s*/>')
        conteudo_novo, total = padrao.subn(rf"\1<v>{valor_cache}</v>", conteudo_original)
        assert total == 1, (
            f"célula {referencia_celula} com fórmula (e cache vazio `<v />`) não encontrada "
            "para injeção — confira se a célula foi escrita como fórmula antes de salvar"
        )
        entradas[nome_arquivo_interno] = conteudo_novo.encode("utf-8")

        with zipfile.ZipFile(caminho_xlsx, "w", zipfile.ZIP_DEFLATED) as zf:
            for nome, dado in entradas.items():
                zf.writestr(nome, dado)

    return _injetar


@pytest.fixture
def caminho_planilha_conferencia_teste(tmp_path):
    """Planilha sintética própria da Central de Conferência (Sprint 13) —
    separada de `caminho_planilha_teste` para não alterar o comportamento
    dos testes já homologados de `escritor_planilha`. Tem a aba `Dados`
    populada (cabeçalho real confirmado no doc arquitetural, § 2.2) e a aba
    `Controle de Diárias` com uma linha cujo CPF+período serve de "duplicata"
    nos testes que precisarem — os dois servidores cobrem a inconsistência
    real de armazenamento de CPF (um como `int`, outro como texto mascarado,
    § 5 do doc arquitetural)."""

    def _criar(nome_arquivo: str = "planilha_conferencia_teste.xlsx") -> str:
        workbook = Workbook()
        controle = workbook.active
        controle.title = "Controle de Diárias"
        controle.append(_CABECALHO_PLANILHA)
        controle.append(
            [
                "2026DH000100",
                "19.20.0902.0004521/2026-11",
                "PJTESTE",
                12345678901,
                "=VLOOKUP($D2,Dados!$A:$E,4,FALSE)",
                "=VLOOKUP($D2,Dados!$A:$E,3,FALSE)",
                "=VLOOKUP($D2,Dados!$A:$E,5,FALSE)",
                date(2026, 7, 21),
                date(2026, 7, 23),
                "RECIFE - GARANHUNS",
                "CARRO",
                "Motivo de teste",
            ]
        )

        dados = workbook.create_sheet("Dados")
        dados.append(_CABECALHO_DADOS)
        dados.append(
            [
                12345678901,
                1001,
                "Maria Aparecida Souza",
                "FUNAFIN",
                "Analista Judiciário",
                None,
                "RPPS",
                "Em exercício",
            ]
        )
        dados.append(
            [
                "591.111.111-00",
                1002,
                "João da Silva",
                "INSS",
                "Técnico Judiciário",
                None,
                "RGPS",
                "Em exercício",
            ]
        )

        workbook.create_sheet("Tabela Valores")

        caminho = tmp_path / nome_arquivo
        workbook.save(caminho)
        return str(caminho)

    return _criar
