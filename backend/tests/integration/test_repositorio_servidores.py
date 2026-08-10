from openpyxl import load_workbook

from app.consulta.repositorio_servidores import NOME_ABA, opcoes_cadastro, servidores_por_cpf


def test_servidores_por_cpf_le_dados_cpf_int_e_texto_mascarado(caminho_planilha_conferencia_teste):
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    servidores = servidores_por_cpf(workbook[NOME_ABA])

    assert "12345678901" in servidores
    servidor_int = servidores["12345678901"]
    assert servidor_int.nome == "Maria Aparecida Souza"
    assert servidor_int.cargo == "Analista Judiciário"
    assert servidor_int.previdencia == "FUNAFIN"

    assert "59111111100" in servidores
    servidor_texto = servidores["59111111100"]
    assert servidor_texto.nome == "João da Silva"


def test_servidores_por_cpf_cpf_inexistente_nao_aparece(caminho_planilha_conferencia_teste):
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    servidores = servidores_por_cpf(workbook[NOME_ABA])

    assert "00000000000" not in servidores


def test_servidores_por_cpf_le_matricula_e_situacao(caminho_planilha_conferencia_teste):
    """Sprint 17.3 — colunas B/H, usadas pela Central Analítica (busca por
    matrícula, filtro por situação); a Central de Conferência não exibe
    esses campos, mas eles já são lidos aqui (`docs/central-analitica/
    README.md § 2`)."""
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    servidores = servidores_por_cpf(workbook[NOME_ABA])

    servidor_int = servidores["12345678901"]
    assert servidor_int.matricula == "1001"
    assert servidor_int.situacao == "Em exercício"

    servidor_texto = servidores["59111111100"]
    assert servidor_texto.matricula == "1002"
    assert servidor_texto.situacao == "Em exercício"


def test_opcoes_cadastro_devolve_valores_distintos_e_ordenados(caminho_planilha_conferencia_teste):
    """Fonte real do modal de Cadastro de Servidor/Membro (item 8 do
    pedido): nunca lista hardcoded — os valores vêm exatamente dos dois
    servidores já cadastrados na fixture."""
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    opcoes = opcoes_cadastro(workbook[NOME_ABA])

    assert opcoes.cargos == ["Analista Judiciário", "Técnico Judiciário"]
    assert opcoes.previdencias == ["FUNAFIN", "INSS"]
    assert opcoes.regimes_previdencia == ["RGPS", "RPPS"]
    assert opcoes.situacoes == ["Em exercício"]


def test_opcoes_cadastro_ignora_celulas_vazias(caminho_planilha_conferencia_teste):
    """Os dois servidores da fixture têm `orgao_origem` em branco (coluna F,
    `None`) — nunca deve virar uma opção `""` ou `"None"` na lista."""
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    opcoes = opcoes_cadastro(workbook[NOME_ABA])

    assert opcoes.orgaos_origem == []
