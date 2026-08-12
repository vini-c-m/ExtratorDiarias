"""Micro-sprint D ("Escrita Exclusiva do DH na Central de Regularização de
DH") — a única rota de escrita desta Central. Testes direcionados desta
Sprint: gravação, SEI não encontrado, linha desatualizada, DH já
preenchido (idêntico/divergente), falha de sincronização, confirmação de
UI (frontend, arquivo separado) e integridade de linha (só a coluna DH
muda). NUNCA testa contra a Planilha Oficial real — sempre fixture/cópia
em `tmp_path`, nunca sincroniza a AppData real do usuário desta máquina."""

from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.consulta import base_local, base_local_sqlite
from app.infra.config import configuracao_app
from app.main import app

client = TestClient(app)

_CABECALHO_CONTROLE = [
    "DH", "SEI", "DOC ORIGEM", "CPF", "PREVIDÊNCIA", "NOME", "CARGO",
    "PERÍODO INICIAL", "PERÍODO FINAL", "TRECHO", "TRANSPORTE", "MOTIVO",
    "VALOR DA PASSAGEM", "QUANT. DIÁRIA INTEGRAL", "QUANT. DIÁRIA PARCIAL",
    "VALOR TOTAL DAS DIÁRIAS", "LOCAL",
]
_CABECALHO_DADOS = [
    "SEI", "MATRÍCULA", "NOME", "PREVIDÊNCIA", "CARGO",
    "ÓRGÃO DE ORIGEM", "REGIME DE PREVIDÊNCIA", "SITUAÇÃO DO SERVIDOR",
]

_SEI_TESTE = "19.20.0099.0000001/2026-01"
_CPF_TESTE = 11111111111


def _linha_controle(*, dh=None, sei=_SEI_TESTE):
    return [
        dh, sei, "PJTESTE", _CPF_TESTE, None, None, None,
        date(2026, 1, 1), date(2026, 1, 3), "RECIFE - GARANHUNS", "CARRO", "Reunião de trabalho",
        None, 2, 1, 308.19, "PERNAMBUCO",
    ]


def _montar_planilha_oficial(tmp_path, *, linhas_controle, nome_arquivo="planilha_oficial.xlsx"):
    workbook = Workbook()
    controle = workbook.active
    controle.title = "Controle de Diárias"
    controle.append(_CABECALHO_CONTROLE)
    for linha in linhas_controle:
        controle.append(linha)

    dados = workbook.create_sheet("Dados")
    dados.append(_CABECALHO_DADOS)
    dados.append([_CPF_TESTE, 1001, "Maria Souza", "RGPS", "Analista Judiciário", None, "RPPS", "Em exercício"])

    workbook.create_sheet("Tabela Valores")

    caminho = tmp_path / nome_arquivo
    workbook.save(caminho)
    return str(caminho)


def _configurar_e_sincronizar(monkeypatch, tmp_path, caminho_planilha):
    monkeypatch.setenv(configuracao_app.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    fingerprint = configuracao_app.validar_planilha(caminho_planilha)
    configuracao_app.persistir(caminho_planilha, fingerprint)
    base_local.sincronizar(caminho_planilha, origem="teste")


def _snapshot_linha(caminho_planilha: str, linha: int) -> dict:
    workbook = load_workbook(caminho_planilha, data_only=True)
    planilha = workbook["Controle de Diárias"]
    snapshot = {
        coluna: planilha[f"{coluna}{linha}"].value
        for coluna in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"]
    }
    workbook.close()
    return snapshot


def _registrar_dh(sei: str, dh: str):
    return client.post(f"/api/v1/regularizacao-dh/{sei}/dh", json={"dh": dh})


def test_gravacao_bem_sucedida_conclui_e_processo_some_da_lista(tmp_path, monkeypatch):
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh=None)])
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    resposta_lista_antes = client.get("/api/v1/regularizacao-dh")
    assert resposta_lista_antes.json()["total"] == 1

    resposta = _registrar_dh(_SEI_TESTE, "2026DH000123")

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "concluido"
    assert corpo["sei"] == _SEI_TESTE
    assert corpo["dh"] == "2026DH000123"

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["A2"].value == "2026DH000123"

    # Atualização da lista após sucesso — mesma fonte homologada, reconsultada.
    resposta_lista_depois = client.get("/api/v1/regularizacao-dh")
    assert resposta_lista_depois.json()["total"] == 0


def test_integridade_de_linha_so_a_coluna_dh_muda(tmp_path, monkeypatch):
    """Teste de integridade controlado (pedido explícito) — snapshot antes/
    depois da linha inteira, prova que a operação altera SOMENTE a coluna
    DH (A), nunca qualquer outra célula da linha."""
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh=None)])
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    snapshot_antes = _snapshot_linha(caminho, 2)

    resposta = _registrar_dh(_SEI_TESTE, "2026DH000999")
    assert resposta.status_code == 200
    assert resposta.json()["situacao"] == "concluido"

    snapshot_depois = _snapshot_linha(caminho, 2)

    diferencas = {
        coluna: (snapshot_antes[coluna], snapshot_depois[coluna])
        for coluna in snapshot_antes
        if snapshot_antes[coluna] != snapshot_depois[coluna]
    }
    assert diferencas == {"A": (None, "2026DH000999")}


def test_sei_nao_encontrado_404(tmp_path, monkeypatch):
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh=None)])
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    resposta = _registrar_dh("19.20.9999.9999999/2099-99", "2026DH000001")

    assert resposta.status_code == 404
    # Nada foi escrito na única linha real existente.
    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["A2"].value is None


def test_linha_desatualizada_recusa_gravacao_sem_escrever(tmp_path, monkeypatch):
    """A Base Local (SQLite) foi sincronizada apontando `linha_planilha=2`
    para este SEI; se a Planilha Oficial mudar depois (ex.: alguém editou
    manualmente e o SEI daquela linha agora é outro) sem nova
    sincronização, a operação precisa se recusar a gravar — nunca escrever
    na linha errada."""
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh=None)])
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    # Simula a Planilha real mudando SEM nova sincronização da Base Local.
    workbook = load_workbook(caminho)
    workbook["Controle de Diárias"]["B2"] = "19.20.0000.0000000/2020-00"
    workbook.save(caminho)

    resposta = _registrar_dh(_SEI_TESTE, "2026DH000001")

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "nao_salvo"
    assert "desatualizada" in corpo["mensagem"]

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["A2"].value is None  # nada foi gravado


def test_dh_ja_preenchido_com_mesmo_valor_e_idempotente(tmp_path, monkeypatch):
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh="2026DH000555")])
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    resposta_1 = _registrar_dh(_SEI_TESTE, "2026DH000555")
    assert resposta_1.status_code == 200
    assert resposta_1.json()["situacao"] == "concluido"

    resposta_2 = _registrar_dh(_SEI_TESTE, "2026DH000555")
    assert resposta_2.status_code == 200
    assert resposta_2.json()["situacao"] == "concluido"

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["A2"].value == "2026DH000555"


def test_dh_ja_preenchido_com_valor_diferente_nunca_sobrescreve(tmp_path, monkeypatch):
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh="2026DH000111")])
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    resposta = _registrar_dh(_SEI_TESTE, "2026DH000222")

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "nao_salvo"
    assert "2026DH000111" in corpo["mensagem"]

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["A2"].value == "2026DH000111"  # valor original preservado


def test_falha_de_sincronizacao_retorna_salvo_sync_pendente(tmp_path, monkeypatch):
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh=None)])
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    def _sincronizar_falho(caminho_planilha, *, diretorio_base=None, origem="pos_preenchimento", ao_falhar_sqlite=None):
        if ao_falhar_sqlite is not None:
            ao_falhar_sqlite(RuntimeError("falha simulada de sincronização"))
        return caminho_planilha

    monkeypatch.setattr(base_local, "sincronizar", _sincronizar_falho)

    resposta = _registrar_dh(_SEI_TESTE, "2026DH000777")

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "salvo_sync_pendente"

    # O DH FOI gravado na Planilha Oficial, mesmo com a sincronização falha.
    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["A2"].value == "2026DH000777"


def test_dh_vazio_retorna_422():
    resposta = _registrar_dh(_SEI_TESTE, "   ")
    assert resposta.status_code == 422


def test_planilha_oficial_nao_configurada_retorna_422(tmp_path, monkeypatch):
    monkeypatch.setenv(configuracao_app.VAR_AMBIENTE_USER_DATA, str(tmp_path / "vazio_nunca_configurado"))
    # Precisa existir na Base Local para passar da primeira checagem — sem
    # isso o teste provaria só "SEI não encontrado", não "Planilha não
    # configurada". Sincroniza uma planilha qualquer só para popular a
    # Base Local, mas SEM chamar `configuracao_app.persistir` (nunca
    # configurada oficialmente).
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh=None)])
    base_local.sincronizar(caminho, diretorio_base=tmp_path / "vazio_nunca_configurado", origem="teste")

    resposta = _registrar_dh(_SEI_TESTE, "2026DH000001")

    assert resposta.status_code == 422


def test_historico_e_valor_continuam_disponiveis_apos_registrar_dh(tmp_path, monkeypatch):
    """Preservação do Histórico e Valor da própria Central (item explícito
    do pedido) — a escrita do DH não pode quebrar as capacidades de
    leitura já homologadas nas Micro-sprints B/C."""
    caminho = _montar_planilha_oficial(tmp_path, linhas_controle=[_linha_controle(dh=None)])
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    resposta_registro = _registrar_dh(_SEI_TESTE, "2026DH000321")
    assert resposta_registro.json()["situacao"] == "concluido"

    resposta_detalhe = client.get(f"/api/v1/regularizacao-dh/{_SEI_TESTE}")
    assert resposta_detalhe.status_code == 200
    corpo = resposta_detalhe.json()
    assert corpo["ainda_aguardando_dh"] is False  # já foi resolvido
    assert corpo["historico"]["situacao"] == "gerado"
    assert corpo["valor"]["situacao"] == "disponivel"
    assert corpo["valor"]["valor"] == "308.19"


# --- GATE de integridade da v2.1 (release) --------------------------------
# Réplica da estrutura real da Planilha Oficial (colunas de fórmula N/P/R/S
# com cache já calculado pelo Excel, mesma técnica de
# `test_integridade_valor_financeiro.py`) — prova que registrar um DH nunca
# toca fórmula/cache, nem da própria linha nem de uma linha ADJACENTE.
_SEI_ADJACENTE = "19.20.0099.0000002/2026-02"


def _montar_planilha_com_formulas(tmp_path, injetar_cache_formula):
    workbook = Workbook()
    controle = workbook.active
    controle.title = "Controle de Diárias"
    controle.append(_CABECALHO_CONTROLE)
    # Linha 2 — linha alvo do registro de DH, com fórmulas reais N/P/R/S.
    controle.append([
        None, _SEI_TESTE, "PJTESTE", _CPF_TESTE, None, None, None,
        date(2026, 1, 1), date(2026, 1, 3), "RECIFE - GARANHUNS", "CARRO", "Reunião de trabalho",
        None, "=I2-H2", 1, "=R2+S2", "PERNAMBUCO",
    ])
    # Linha 3 — linha ADJACENTE, SEI diferente, DH já preenchido, também com
    # fórmulas reais e cache — nunca deve ser tocada pelo registro da linha 2.
    controle.append([
        "2026DH000001", _SEI_ADJACENTE, "PJTESTE", _CPF_TESTE, None, None, None,
        date(2026, 2, 1), date(2026, 2, 3), "RECIFE - CARUARU", "CARRO", "Diligência",
        None, "=I3-H3", 1, "=R3+S3", "PERNAMBUCO",
    ])

    dados = workbook.create_sheet("Dados")
    dados.append(_CABECALHO_DADOS)
    dados.append([_CPF_TESTE, 1001, "Maria Souza", "RGPS", "Analista Judiciário", None, "RPPS", "Em exercício"])

    tabela_valores = workbook.create_sheet("Tabela Valores")
    tabela_valores.append([None, None, None])
    tabela_valores.append([None, None, None])
    tabela_valores.append([None, "INTEGRAL", "PARCIAL"])
    tabela_valores.append(["PERNAMBUCO", 308.19, 154.09])

    caminho = tmp_path / "planilha_gate_integridade.xlsx"
    workbook.save(caminho)
    injetar_cache_formula(str(caminho), "N2", 2)
    injetar_cache_formula(str(caminho), "P2", 462.28)
    injetar_cache_formula(str(caminho), "N3", 2)
    injetar_cache_formula(str(caminho), "P3", 462.28)
    return str(caminho)


def _snapshot_linha_formulas(caminho_planilha: str, linha: int) -> dict:
    """Snapshot duplo — texto da fórmula (`data_only=False`) e valor
    cacheado (`data_only=True`) — para provar que NEM a fórmula NEM o cache
    mudam, não só o valor final lido."""
    wb_formulas = load_workbook(caminho_planilha, data_only=False)
    wb_valores = load_workbook(caminho_planilha, data_only=True)
    colunas = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"]
    snapshot = {
        coluna: (
            wb_formulas["Controle de Diárias"][f"{coluna}{linha}"].value,
            wb_valores["Controle de Diárias"][f"{coluna}{linha}"].value,
        )
        for coluna in colunas
    }
    wb_formulas.close()
    wb_valores.close()
    return snapshot


def test_gate_v21_integridade_registrar_dh_linha_alvo_e_linha_adjacente(
    tmp_path, monkeypatch, injetar_cache_formula
):
    """GATE de release v2.1 (§4 do pedido) — usa fixture com fórmulas E
    cache reais (nunca a Planilha Oficial real). Resultado obrigatório:
    célula de conteúdo alterada = 1 (DH da linha alvo); nenhuma outra
    célula da própria linha, nenhuma fórmula, nenhum cache de fórmula,
    nenhuma célula da linha adjacente."""
    caminho = _montar_planilha_com_formulas(tmp_path, injetar_cache_formula)
    _configurar_e_sincronizar(monkeypatch, tmp_path, caminho)

    snapshot_alvo_antes = _snapshot_linha_formulas(caminho, 2)
    snapshot_adjacente_antes = _snapshot_linha_formulas(caminho, 3)

    resposta = _registrar_dh(_SEI_TESTE, "2026DH000999")
    assert resposta.status_code == 200
    assert resposta.json()["situacao"] == "concluido"

    snapshot_alvo_depois = _snapshot_linha_formulas(caminho, 2)
    snapshot_adjacente_depois = _snapshot_linha_formulas(caminho, 3)

    diferencas_alvo = {
        c: (snapshot_alvo_antes[c], snapshot_alvo_depois[c])
        for c in snapshot_alvo_antes
        if snapshot_alvo_antes[c] != snapshot_alvo_depois[c]
    }
    diferencas_adjacente = {
        c: (snapshot_adjacente_antes[c], snapshot_adjacente_depois[c])
        for c in snapshot_adjacente_antes
        if snapshot_adjacente_antes[c] != snapshot_adjacente_depois[c]
    }

    # CÉLULAS DE CONTEÚDO ALTERADAS: 1 — exatamente a coluna DH (A) da linha alvo.
    assert set(diferencas_alvo.keys()) == {"A"}
    assert diferencas_alvo["A"] == ((None, None), ("2026DH000999", "2026DH000999"))
    # FÓRMULAS/CACHE da própria linha alvo (N/P) preservados intactos.
    assert snapshot_alvo_depois["N"] == ("=I2-H2", 2)
    assert snapshot_alvo_depois["P"] == ("=R2+S2", 462.28)
    # LINHA ADJACENTE: zero diferenças — nem conteúdo, nem fórmula, nem cache.
    assert diferencas_adjacente == {}
