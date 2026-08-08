"""Modo Complemento de Diárias (RN-14) ponta a ponta pela API real:
`POST /api/v1/processamento` (extração + resolução de quantidade parcial
complementar) e `POST /api/v1/planilha/preencher` (gravação de AVISO/O,
duplicidade por SEI não-bloqueante, linha original preservada). Usa os
mesmos PDFs reais de referência já homologados (`pdfs_exemplo`) — o que
está sob teste é o `modo`, não um documento específico de complemento."""

import os

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app

client = TestClient(app)


def test_processamento_modo_complemento_expoe_quantidade_parcial_resolvida(pdfs_exemplo):
    resposta = client.post(
        "/api/v1/processamento", json={"arquivos": pdfs_exemplo, "modo": "COMPLEMENTO"}
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    for processo in corpo["processos"]:
        assert processo["diaria"]["quantidade_diaria_parcial_complementar"] is not None
        assert processo["diaria"]["fonte_quantidade_complementar"] is not None


def test_processamento_modo_padrao_nao_resolve_quantidade_parcial(pdfs_exemplo):
    """Zero regressão: sem `modo` (default PADRAO), os dois campos novos
    continuam `None`, exatamente como antes desta feature."""
    resposta = client.post("/api/v1/processamento", json={"arquivos": pdfs_exemplo})

    assert resposta.status_code == 200
    for processo in resposta.json()["processos"]:
        assert processo["diaria"]["quantidade_diaria_parcial_complementar"] is None
        assert processo["diaria"]["fonte_quantidade_complementar"] is None


def _planilha_com_formulas(caminho, *, sei_existente: str) -> None:
    """Mesma estrutura mínima de `test_escritor_planilha_complemento.py`:
    uma linha existente e uma linha livre já com as fórmulas de
    N/P/R/S/V/Z/AA pré-copiadas (auditoria obrigatória do modo Complemento)."""
    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Controle de Diárias"
    planilha.append(
        [
            "DH", "SEI", "DOC ORIGEM", "CPF", "PREVIDÊNCIA", "NOME", "CARGO",
            "PERÍODO INICIAL", "PERÍODO FINAL", "TRECHO", "TRANSPORTE", "MOTIVO",
            "VALOR DA PASSAGEM", "QUANT. DIÁRIA INTEGRAL", "QUANT. DIÁRIA PARCIAL",
            "VALOR TOTAL DAS DIÁRIAS", "LOCAL", "VALOR DAS DIÁRIAS INTEGRAIS",
            "VALOR DAS DIÁRIAS PARCIAIS", "DMTC", "DMOF", "PASSAR DIA",
            "RESTITUIÇÃO", "COMPROVANTES", "CONFIRMAÇÃO DE DIÁRIAS", "CONCATENAÇÃO",
            "AVISO ", "AVISO AUXILIAR",
        ]
    )
    planilha.append(
        [
            "2026DH000100", sei_existente, "PJTESTE", 11111111111,
            None, None, None, "17/03/2026", "19/03/2026", "OURICURI-RECIFE", "CARRO",
            "Motivo original", None, "=I2-H2", 1, "=R2+S2", "PERNAMBUCO",
            "=(VLOOKUP(Q2,'Tabela Valores'!A:C,2,FALSE()))*N2",
            "=(VLOOKUP(Q2,'Tabela Valores'!A:C,3,FALSE()))*O2",
            None, None, "=H2-15", None, None, None,
            "=CONCATENATE(D2,H2,I2)", "=VLOOKUP(Z2,Z$1:AB1,3,0)", "JÁ EXISTE DIÁRIA",
        ]
    )
    planilha.append(
        [
            None, None, None, None, None, None, None, None, None, None, None, None,
            None, "=I3-H3", None, "=R3+S3", None,
            "=(VLOOKUP(Q3,'Tabela Valores'!A:C,2,FALSE()))*N3",
            "=(VLOOKUP(Q3,'Tabela Valores'!A:C,3,FALSE()))*O3",
            None, None, "=H3-15", None, None, None,
            "=CONCATENATE(D3,H3,I3)", "=VLOOKUP(Z3,Z$1:AB2,3,0)", None,
        ]
    )
    workbook.create_sheet("Dados")
    tabela_valores = workbook.create_sheet("Tabela Valores")
    tabela_valores.append(["PERNAMBUCO", 308.19, 154.09])
    workbook.save(caminho)


def _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_planilha):
    """Mesmo padrão de `tests/integration/test_api_planilha.py` — desde a
    Sprint 15 o backend só preenche uma planilha já configurada/validada."""
    from app.infra.config import configuracao_app

    monkeypatch.setenv(configuracao_app.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    fingerprint = configuracao_app.validar_planilha(caminho_planilha)
    configuracao_app.persistir(caminho_planilha, fingerprint)


def test_preencher_modo_complemento_grava_aviso_e_o_nao_bloqueia_sei(
    monkeypatch, tmp_path, pdfs_exemplo
):
    lote = client.post(
        "/api/v1/processamento", json={"arquivos": pdfs_exemplo[:1], "modo": "COMPLEMENTO"}
    ).json()
    diaria = lote["processos"][0]["diaria"]

    caminho_planilha = str(tmp_path / "planilha_oficial.xlsx")
    _planilha_com_formulas(caminho_planilha, sei_existente=diaria["sei"])
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_planilha)

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={
            "itens": [{"arquivo_origem": pdfs_exemplo[0], "diaria": diaria}],
            "modo": "COMPLEMENTO",
        },
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_gravado"] == 1
    assert corpo["itens"][0]["categoria"] == "inserido"
    assert corpo["itens"][0]["sei_ja_existente_em_linha"] == 2

    workbook = load_workbook(caminho_planilha)
    planilha = workbook["Controle de Diárias"]
    assert planilha["AA3"].value == "COMPLEMENTO DE DIÁRIAS"
    assert planilha["O3"].value is not None
    # Linha original (2) permanece intocada.
    assert planilha["B2"].value == diaria["sei"]
    assert planilha["L2"].value == "Motivo original"
    assert planilha["AA2"].value == "=VLOOKUP(Z2,Z$1:AB1,3,0)"


def test_preencher_modo_padrao_continua_bloqueando_sei_ja_existente(
    monkeypatch, tmp_path, pdfs_exemplo
):
    """Zero regressão: sem `modo` (default PADRAO), SEI repetido continua
    bloqueando a gravação exatamente como antes desta feature."""
    lote = client.post("/api/v1/processamento", json={"arquivos": pdfs_exemplo[:1]}).json()
    diaria = lote["processos"][0]["diaria"]

    caminho_planilha = str(tmp_path / "planilha_oficial.xlsx")
    _planilha_com_formulas(caminho_planilha, sei_existente=diaria["sei"])
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_planilha)

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={"itens": [{"arquivo_origem": pdfs_exemplo[0], "diaria": diaria}]},
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_gravado"] == 0
    assert corpo["itens"][0]["categoria"] == "duplicado_na_planilha"
