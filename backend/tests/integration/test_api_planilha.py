import os

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.infra.config import configuracao_app
from app.main import app

client = TestClient(app)

_DIARIA = {
    "sei": "19.20.0000.0000001/2025-99",
    "doc_origem": "PJNOVO",
    "cpf": "99999999999",
    "periodo_inicial": "2025-03-01",
    "periodo_final": "2025-03-03",
    "cidade_origem": "RECIFE",
    "cidade_destino": "OLINDA",
    "motivo": "Diligência de teste",
}


def _com_cabecalho_completo(caminho_planilha: str) -> str:
    """`caminho_planilha_teste` (conftest) não tem as colunas P/Q — exigidas
    pela validação de configuração da Sprint 15 (§5 do doc arquitetural),
    que agora é pré-requisito para qualquer preenchimento. Completa o
    cabeçalho aqui, sem alterar a fixture compartilhada — os demais testes
    que já a usam (escritor, diária de membro etc.) continuam intocados."""
    workbook = load_workbook(caminho_planilha)
    planilha = workbook["Controle de Diárias"]
    planilha["P1"] = "VALOR TOTAL DAS DIÁRIAS"
    planilha["Q1"] = "LOCAL"
    workbook.save(caminho_planilha)
    return caminho_planilha


def _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_planilha: str) -> None:
    """Sprint 15 — desde que o backend passou a ser a única autoridade
    sobre o caminho da planilha (§7.1), qualquer teste de preenchimento
    precisa primeiro configurar uma planilha oficial válida, exatamente
    como a aplicação real exige."""
    monkeypatch.setenv(configuracao_app.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    caminho = _com_cabecalho_completo(caminho_planilha)
    fingerprint = configuracao_app.validar_planilha(caminho)
    configuracao_app.persistir(caminho, fingerprint)


def test_post_planilha_preencher_grava_na_copia_de_teste(monkeypatch, tmp_path, caminho_planilha_teste):
    caminho = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho)

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={"itens": [{"arquivo_origem": "arquivo.pdf", "diaria": _DIARIA}]},
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_gravado"] == 1
    assert corpo["duplicadas_ignoradas"] == 0
    assert corpo["linha_inicial"] == 4
    assert corpo["linha_final"] == 4
    assert corpo["itens"] == [
        {
            "arquivo_origem": "arquivo.pdf",
            "sei": _DIARIA["sei"],
            "categoria": "inserido",
            "linha": 4,
            "sei_ja_existente_em_linha": None,
            # Organização Automática de PDFs — item INSERIDO sem
            # `caminho_absoluto` informado nesta requisição: organização
            # tentada e reportada como não encontrada, nunca silenciosa.
            "organizacao": {
                "sucesso": False,
                "status": "arquivo_nao_encontrado",
                "detalhe": "Caminho de origem não informado.",
            },
        }
    ]
    assert os.path.isdir(corpo["pasta_resultados"])
    assert corpo["caminho_relatorio"].endswith(".html")
    assert os.path.isfile(corpo["caminho_log"])

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["B4"].value == _DIARIA["sei"]


def test_post_planilha_preencher_bug1_duplicidade_no_lote(monkeypatch, tmp_path, caminho_planilha_teste):
    """Mesmo PDF (mesmo SEI) enviado duas vezes na mesma requisição — só uma
    diária pode ser gravada, a outra vira 'duplicado no lote'."""
    caminho = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho)
    sei_novo = "19.20.0000.0000020/2025-00"

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={
            "itens": [
                {"arquivo_origem": "a.pdf", "diaria": {**_DIARIA, "sei": sei_novo}},
                {"arquivo_origem": "b.pdf", "diaria": {**_DIARIA, "sei": sei_novo}},
            ],
        },
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_gravado"] == 1
    assert corpo["duplicadas_ignoradas"] == 1
    assert corpo["itens"][0]["categoria"] == "inserido"
    assert corpo["itens"][0]["linha"] == 4
    assert corpo["itens"][1]["categoria"] == "duplicado_no_lote"
    assert corpo["itens"][1]["linha"] is None

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["B4"].value == sei_novo
    assert planilha["B5"].value is None


def test_post_planilha_preencher_bug2_duplicidade_na_planilha_informa_linha(
    monkeypatch, tmp_path, caminho_planilha_teste,
):
    """SEI que já existe na planilha (linha 2 da fixture) não pode virar
    linha nova, e a resposta precisa apontar em qual linha ele já existe."""
    caminho = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho)
    sei_ja_existente = "19.20.0000.0000001/2025-00"

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={
            "itens": [
                {"arquivo_origem": "a.pdf", "diaria": {**_DIARIA, "sei": sei_ja_existente}},
            ],
            "erros": [{"arquivo_origem": "c.pdf", "motivos": ["Falha ao ler o PDF"]}],
        },
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_gravado"] == 0
    assert corpo["duplicadas_ignoradas"] == 1
    assert corpo["itens"] == [
        {
            "arquivo_origem": "a.pdf",
            "sei": sei_ja_existente,
            "categoria": "duplicado_na_planilha",
            "linha": 2,
            "sei_ja_existente_em_linha": None,
            # Organização Automática de PDFs — nunca tentada para
            # duplicidade (duplicidade não prova que a linha existente
            # esteja correta/completa).
            "organizacao": None,
        }
    ]

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["B4"].value is None

    conteudo = open(corpo["caminho_relatorio"], encoding="utf-8").read()
    assert "PDFs analisados" in conteudo
    assert "a.pdf" in conteudo

    conteudo_log = open(corpo["caminho_log"], encoding="utf-8").read()
    assert "já existe na linha 2" in conteudo_log


def test_post_planilha_preencher_grava_dh_transporte_local_e_lista_avisos_no_relatorio(
    monkeypatch, tmp_path, caminho_planilha_teste,
):
    """Sprint 03: quando a diária vem com documento_habil/transporte/
    estado_destino preenchidos, eles são gravados nas colunas A/K/Q. Quando
    um arquivo tem avisos (campo não identificado com segurança), eles
    aparecem no relatório e no log — sem afetar a deduplicação por SEI."""
    caminho = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho)
    sei_completo = "19.20.0000.0000030/2025-00"
    sei_incompleto = "19.20.0000.0000031/2025-00"

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={
            "itens": [
                {
                    "arquivo_origem": "completo.pdf",
                    "diaria": {
                        **_DIARIA,
                        "sei": sei_completo,
                        "documento_habil": "2025DH009999",
                        "transporte": "AVIÃO",
                        "estado_destino": "SÃO PAULO",
                    },
                },
                {
                    "arquivo_origem": "incompleto.pdf",
                    "diaria": {**_DIARIA, "sei": sei_incompleto},
                },
            ],
            "avisos": [
                {
                    "arquivo_origem": "incompleto.pdf",
                    "mensagens": [
                        "Transporte não identificado no processo — necessita revisão manual."
                    ],
                }
            ],
        },
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_gravado"] == 2

    planilha = load_workbook(caminho)["Controle de Diárias"]
    assert planilha["A4"].value == "2025DH009999"
    assert planilha["K4"].value == "AVIÃO"
    assert planilha["Q4"].value == "SÃO PAULO"
    # A segunda diária não trouxe DH/transporte/estado — colunas ficam vazias.
    assert planilha["A5"].value is None
    assert planilha["K5"].value is None
    assert planilha["Q5"].value is None

    conteudo_relatorio = open(corpo["caminho_relatorio"], encoding="utf-8").read()
    # Sprint 05.1: a hierarquia de revisão substituiu a antiga seção plana
    # "Revisão manual necessária" pelo selo "Revisão recomendada".
    assert "Revisão recomendada" in conteudo_relatorio
    assert "incompleto.pdf" in conteudo_relatorio

    conteudo_log = open(corpo["caminho_log"], encoding="utf-8").read()
    assert "REVISAO_MANUAL arquivo: incompleto.pdf" in conteudo_log


# --- Sprint 15 — backend como autoridade sobre o caminho (§7.1) -------------


def test_post_planilha_preencher_sem_configuracao_retorna_422(monkeypatch, tmp_path):
    monkeypatch.setenv(configuracao_app.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={"itens": [{"arquivo_origem": "arquivo.pdf", "diaria": _DIARIA}]},
    )

    assert resposta.status_code == 422
    assert "Nenhuma planilha oficial configurada" in resposta.json()["detail"]


def test_post_planilha_preencher_planilha_configurada_mas_inacessivel_retorna_422(
    monkeypatch, tmp_path, caminho_planilha_teste,
):
    caminho = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho)
    os.remove(caminho)

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={"itens": [{"arquivo_origem": "arquivo.pdf", "diaria": _DIARIA}]},
    )

    assert resposta.status_code == 422
    assert "inacessível" in resposta.json()["detail"]


def test_post_planilha_preencher_requer_revalidacao_retorna_422(
    monkeypatch, tmp_path, caminho_planilha_teste,
):
    caminho = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho)

    # Estrutura divergente da configurada — mesmo caminho, arquivo diferente.
    workbook = Workbook()
    workbook.active.title = "Controle de Diárias"
    workbook.save(caminho)

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={"itens": [{"arquivo_origem": "arquivo.pdf", "diaria": _DIARIA}]},
    )

    assert resposta.status_code == 422
    assert "estrutura da planilha oficial mudou" in resposta.json()["detail"]


def test_post_planilha_preencher_ignora_qualquer_caminho_enviado_pelo_cliente(
    monkeypatch, tmp_path, caminho_planilha_teste,
):
    """O cliente não tem mais como influenciar qual arquivo é gravado — um
    `caminho_planilha` extra no corpo é ignorado pelo schema (campo
    removido de `RequisicaoPreenchimento`), o caminho usado é sempre o da
    configuração."""
    caminho_oficial = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_oficial)

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={
            "caminho_planilha": "C:/tentativa/de/enviar/outro/caminho.xlsx",
            "itens": [{"arquivo_origem": "arquivo.pdf", "diaria": _DIARIA}],
        },
    )

    assert resposta.status_code == 200
    planilha = load_workbook(caminho_oficial)["Controle de Diárias"]
    assert planilha["B4"].value == _DIARIA["sei"]


# --- Organização Automática de PDFs -----------------------------------------


def test_post_planilha_preencher_move_pdf_inserido_para_processados(
    monkeypatch, tmp_path, caminho_planilha_teste
):
    caminho_oficial = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_oficial)

    pasta_a_processar = tmp_path / "a_processar"
    pasta_processados = tmp_path / "processados"
    pasta_a_processar.mkdir()
    pasta_processados.mkdir()
    client.post(
        "/api/v1/configuracao/pastas",
        json={
            "caminho_pdfs_a_processar": str(pasta_a_processar),
            "caminho_pdfs_processados": str(pasta_processados),
        },
    )

    pdf_origem = pasta_a_processar / "arquivo.pdf"
    pdf_origem.write_bytes(b"conteudo do pdf de teste")

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={
            "itens": [
                {
                    "arquivo_origem": "arquivo.pdf",
                    "diaria": _DIARIA,
                    "caminho_absoluto": str(pdf_origem),
                }
            ]
        },
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["itens"][0]["organizacao"] == {
        "sucesso": True,
        "status": "movido",
        "detalhe": None,
    }
    assert not pdf_origem.exists()
    assert (pasta_processados / "arquivo.pdf").read_bytes() == b"conteudo do pdf de teste"


def test_post_planilha_preencher_nao_move_pdf_duplicado_no_lote(
    monkeypatch, tmp_path, caminho_planilha_teste
):
    caminho_oficial = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_oficial)

    pasta_a_processar = tmp_path / "a_processar"
    pasta_processados = tmp_path / "processados"
    pasta_a_processar.mkdir()
    pasta_processados.mkdir()
    client.post(
        "/api/v1/configuracao/pastas",
        json={
            "caminho_pdfs_a_processar": str(pasta_a_processar),
            "caminho_pdfs_processados": str(pasta_processados),
        },
    )

    sei_repetido = "19.20.0000.0000099/2025-00"
    pdf_a = pasta_a_processar / "a.pdf"
    pdf_b = pasta_a_processar / "b.pdf"
    pdf_a.write_bytes(b"a")
    pdf_b.write_bytes(b"b")

    resposta = client.post(
        "/api/v1/planilha/preencher",
        json={
            "itens": [
                {
                    "arquivo_origem": "a.pdf",
                    "diaria": {**_DIARIA, "sei": sei_repetido},
                    "caminho_absoluto": str(pdf_a),
                },
                {
                    "arquivo_origem": "b.pdf",
                    "diaria": {**_DIARIA, "sei": sei_repetido},
                    "caminho_absoluto": str(pdf_b),
                },
            ]
        },
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["itens"][0]["categoria"] == "inserido"
    assert corpo["itens"][0]["organizacao"]["status"] == "movido"
    assert corpo["itens"][1]["categoria"] == "duplicado_no_lote"
    # Duplicidade nunca prova que a linha existente esteja correta/completa
    # — o PDF duplicado nunca é organizado (movido/excluído).
    assert corpo["itens"][1]["organizacao"] is None
    assert not pdf_a.exists()
    assert pdf_b.exists()
