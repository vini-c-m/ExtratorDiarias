"""Bloco de Refinamento da Central Analítica (2026-08-05), item 1 —
qualidade de apresentação dos relatórios PDF/Excel. Reproduz, com dados
sintéticos, os problemas reais relatados pelo usuário (PDF com conteúdo
cortado/tabelas maiores que a página; Excel sem quebra de texto, larguras,
congelamento ou autofiltro) e prova que a correção os elimina — sem alterar
nenhum dado/cálculo (mesmos `Indicadores`/`Distribuicao`/`RegistroAnalitico`
de sempre, só a apresentação)."""

from datetime import date

import pdfplumber
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.consulta import base_local_sqlite
from app.consulta.repositorio_diarias import ResultadoDiariasCompletas
from app.domain.entities.registro_diaria_indexada import RegistroDiariaIndexada
from app.domain.entities.servidor import Servidor
from app.main import app

client = TestClient(app)

_MOTIVO_LONGO = (
    "Participação no Encontro Nacional de Coordenadores de Inteligência do Ministério Público, "
    "com apresentação de relatório de atividades do núcleo regional e discussão de protocolos "
    "de cooperação interinstitucional entre os órgãos de segurança pública do estado"
)
_NOME_LONGO = "Maria Aparecida da Silva Souza Nascimento de Albuquerque Filho"


def _registro_longo(indice: int) -> RegistroDiariaIndexada:
    cpf = f"{indice:011d}"
    inicio = date(2026, 1, 1)
    return RegistroDiariaIndexada(
        linha=indice + 2,
        sei=f"19.20.0001.{indice:07d}/2026-{indice % 100:02d}",
        doc_origem="DIMOT",
        cpf=cpf,
        periodo_inicial=inicio,
        periodo_final=inicio,
        cidade_origem="RECIFE",
        cidade_destino="GARANHUNS",
        motivo=_MOTIVO_LONGO,
        transporte="CARRO",
        valor_total_diarias=1234.56,
        dh=f"2026DH{indice:06d}",
    )


def _sincronizar_muitas_linhas(quantidade: int) -> None:
    registros = [_registro_longo(indice) for indice in range(quantidade)]
    servidores = {
        registro.cpf: Servidor(
            cpf=registro.cpf, nome=_NOME_LONGO, cargo="Analista Judiciário",
            previdencia="FUNAFIN", matricula=f"{indice:07d}", situacao="Em exercício",
        )
        for indice, registro in enumerate(registros)
    }
    resultado = ResultadoDiariasCompletas(
        diarias={r.sei: r for r in registros}, avisos=[], linhas_ignoradas=0
    )
    sinal = base_local_sqlite.SinalFrescor(tamanho_bytes=1, mtime=1.0, linhas_controle_diarias=len(registros) + 1)
    destino = base_local_sqlite.caminho_base_local_sqlite()
    base_local_sqlite.reconstruir(resultado, servidores, sinal, origem="teste", caminho_destino=destino)


def _gerar_relatorio(*, tipo="detalhado", formato="pdf", incluir_graficos=None, filtros=None):
    corpo = {
        "filtros": filtros or {},
        "tipo": tipo,
        "formato": formato,
        "modo": "final",
        "incluir_graficos": incluir_graficos if incluir_graficos is not None else (formato == "pdf"),
    }
    resposta = client.post("/api/v1/analitica/relatorios", json=corpo)
    assert resposta.status_code == 200, resposta.text
    return resposta.json()["caminho_arquivo_temporario"]


class _PaginaExtraida:
    """`pdfplumber.Page` é preguiçoso (só decodifica o layout no primeiro
    acesso a `.chars`/`.extract_text()`) e fecha o arquivo assim que o
    `with pdfplumber.open(...)` termina — teria que manter o arquivo
    aberto pelo teste inteiro, ou (mais simples) materializar tudo que os
    testes precisam enquanto o arquivo ainda está aberto."""

    def __init__(self, largura: float, altura: float, chars: list, texto: str) -> None:
        self.width = largura
        self.height = altura
        self.chars = chars
        self._texto = texto

    def extract_text(self) -> str:
        return self._texto


def _paginas_pdf(caminho: str) -> list[_PaginaExtraida]:
    with pdfplumber.open(caminho) as pdf:
        return [
            _PaginaExtraida(pagina.width, pagina.height, pagina.chars, pagina.extract_text() or "")
            for pagina in pdf.pages
        ]


# --- PDF -------------------------------------------------------------------


def test_pdf_relatorio_detalhado_com_muitas_linhas_gera_multiplas_paginas():
    """500 linhas — volume real de produção é da mesma ordem de grandeza
    (a base real tem 1.806 diárias)."""
    _sincronizar_muitas_linhas(500)
    caminho = _gerar_relatorio(tipo="detalhado", formato="pdf")

    paginas = _paginas_pdf(caminho)
    assert len(paginas) > 1


def test_pdf_nenhuma_pagina_tem_conteudo_fora_da_area_util():
    """Critério de aceite central desta correção: nenhum caractere pode
    cair fora dos limites físicos da página em nenhuma página — a causa
    raiz confirmada do "conteúdo cortado" era a tabela de 7 colunas sem
    `colWidths`, desenhada no tamanho natural do texto (~990pt, quase o
    dobro dos ~482pt úteis de uma A4), sempre a partir da 3ª/4ª página
    (onde a tabela de registros aparece)."""
    _sincronizar_muitas_linhas(80)
    caminho = _gerar_relatorio(tipo="detalhado", formato="pdf")

    paginas = _paginas_pdf(caminho)
    assert len(paginas) > 3  # garante que o teste realmente cobre a tabela de registros
    for indice, pagina in enumerate(paginas):
        if not pagina.chars:
            continue
        x0_min = min(c["x0"] for c in pagina.chars)
        x1_max = max(c["x1"] for c in pagina.chars)
        assert x0_min >= -0.5, f"página {indice + 1}: conteúdo começa antes da margem esquerda (x0={x0_min:.1f})"
        assert x1_max <= pagina.width + 0.5, (
            f"página {indice + 1}: conteúdo ultrapassa a largura da página "
            f"(x1_max={x1_max:.1f}, largura={pagina.width:.1f})"
        )


def test_pdf_textos_longos_aparecem_por_completo_sem_truncamento():
    """Um Motivo de ~240 caracteres precisa aparecer integralmente no PDF
    (quebrado em várias linhas dentro da célula, nunca cortado) — prova
    indireta de que as células usam `Paragraph` com quebra de linha, não
    texto cru de largura fixa."""
    _sincronizar_muitas_linhas(3)
    caminho = _gerar_relatorio(tipo="detalhado", formato="pdf")

    texto_completo = " ".join(pagina.extract_text() or "" for pagina in _paginas_pdf(caminho))
    # O texto quebra dentro da célula em várias linhas (é exatamente o que
    # provamos aqui) — pdfplumber devolve cada linha visual separada por
    # `\n`, então normalizamos para espaço antes de comparar um trecho
    # inequívoco do início e do fim do Motivo inteiro.
    texto_normalizado = " ".join(texto_completo.split())
    assert "Participação no Encontro Nacional" in texto_normalizado
    assert "segurança pública do estado" in texto_normalizado


def test_pdf_cabecalho_da_tabela_se_repete_em_paginas_novas():
    _sincronizar_muitas_linhas(80)
    caminho = _gerar_relatorio(tipo="detalhado", formato="pdf")

    paginas = _paginas_pdf(caminho)
    paginas_com_cabecalho = sum(
        1 for pagina in paginas if "SEI" in (pagina.extract_text() or "") and "Servidor" in (pagina.extract_text() or "")
    )
    assert paginas_com_cabecalho >= 2  # apareceu na página da tabela E numa página de continuação


def test_pdf_rodape_com_numero_de_pagina_em_todas_as_paginas():
    _sincronizar_muitas_linhas(80)
    caminho = _gerar_relatorio(tipo="detalhado", formato="pdf")

    paginas = _paginas_pdf(caminho)
    total = len(paginas)
    assert total > 1
    for indice, pagina in enumerate(paginas, start=1):
        texto = pagina.extract_text() or ""
        assert f"Página {indice} de {total}" in texto


def test_pdf_indicadores_em_destaque_e_resumo_do_recorte():
    _sincronizar_muitas_linhas(10)
    caminho = _gerar_relatorio(tipo="detalhado", formato="pdf")

    texto = _paginas_pdf(caminho)[0].extract_text() or ""
    assert "Relatório Detalhado" in texto
    assert "Recorte:" in texto
    assert "10 diária" in texto
    assert "Total de diárias" in texto
    assert "Total de servidores" in texto


def test_pdf_recorte_vazio_nao_corta_nem_quebra():
    _sincronizar_muitas_linhas(5)
    caminho = _gerar_relatorio(tipo="detalhado", formato="pdf", filtros={"motivo": "Inexistente De Verdade"})

    texto = (_paginas_pdf(caminho)[0].extract_text()) or ""
    assert "Sem dados para este recorte." in texto
    assert "0 diária" in texto


def test_pdf_abre_como_documento_valido():
    _sincronizar_muitas_linhas(5)
    caminho = _gerar_relatorio(tipo="executivo", formato="pdf")
    with open(caminho, "rb") as arquivo:
        assert arquivo.read(5) == b"%PDF-"
    # abrir com pdfplumber sem exceção já prova que a estrutura do PDF é válida
    with pdfplumber.open(caminho) as pdf:
        assert len(pdf.pages) >= 1


# --- Excel -------------------------------------------------------------------


def test_excel_wrap_text_em_coluna_de_texto_longo():
    _sincronizar_muitas_linhas(5)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    aba = load_workbook(caminho)["Tabela do recorte"]
    coluna_motivo = 7  # SEI, Período Inicial, Período Final, Servidor, CPF, Destino, Motivo, Valor
    celula = aba.cell(row=2, column=coluna_motivo)
    assert celula.value == _MOTIVO_LONGO
    assert celula.alignment.wrap_text is True


def test_excel_larguras_de_coluna_tem_limite_maximo():
    _sincronizar_muitas_linhas(5)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    aba = load_workbook(caminho)["Tabela do recorte"]
    for dimensao in aba.column_dimensions.values():
        if dimensao.width:
            assert dimensao.width <= 60 + 0.01
    # a coluna de Motivo (texto de ~240 chars) tem que ter sido limitada,
    # não alargada até caber o texto inteiro numa linha só
    from openpyxl.utils import get_column_letter

    assert aba.column_dimensions[get_column_letter(7)].width <= 60


def test_excel_nenhuma_coluna_usa_a_largura_padrao_do_excel():
    """Causa raiz confirmada do "conteúdo truncado visualmente": sem
    nenhuma largura definida, o Excel usa ~8.43 caracteres — muito menos
    que qualquer campo real (SEI já tem 26 caracteres sozinho)."""
    _sincronizar_muitas_linhas(5)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    aba = load_workbook(caminho)["Tabela do recorte"]
    for letra in ("A", "B", "C", "D", "E", "F", "G", "H"):
        assert aba.column_dimensions[letra].width is not None
        assert aba.column_dimensions[letra].width > 8.43


def test_excel_formato_de_data_nativo_dd_mm_aaaa():
    _sincronizar_muitas_linhas(2)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    aba = load_workbook(caminho)["Tabela do recorte"]
    celula_data = aba.cell(row=2, column=2)  # Período Inicial
    # openpyxl sempre devolve `datetime.datetime` ao ler uma célula de data
    # de volta (Excel não distingue data de data+hora internamente) — o
    # valor ESCRITO (`registro.periodo_inicial`) é um `date` nativo, o que
    # importa aqui é que o dia/mês/ano batem e o formato é o brasileiro.
    assert celula_data.value.date() == date(2026, 1, 1)
    assert celula_data.number_format == "DD/MM/YYYY"


def test_excel_formato_de_moeda_nativo_padrao_brasileiro():
    _sincronizar_muitas_linhas(2)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    aba = load_workbook(caminho)["Tabela do recorte"]
    celula_valor = aba.cell(row=2, column=8)  # Valor
    assert celula_valor.value == 1234.56
    assert "R$" in celula_valor.number_format
    assert celula_valor.number_format != "General"


def test_excel_filtro_e_congelamento_na_tabela_do_recorte():
    _sincronizar_muitas_linhas(5)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    aba = load_workbook(caminho)["Tabela do recorte"]
    assert aba.freeze_panes == "A2"
    assert len(aba.tables) == 1
    tabela = next(iter(aba.tables.values()))
    assert tabela.autoFilter is not None
    assert tabela.tableStyleInfo.showRowStripes is True


def test_excel_cabecalho_estilizado():
    _sincronizar_muitas_linhas(5)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    aba = load_workbook(caminho)["Tabela do recorte"]
    celula_cabecalho = aba.cell(row=1, column=1)
    assert celula_cabecalho.font.bold is True
    assert celula_cabecalho.font.color.rgb.endswith("FFFFFF")
    assert celula_cabecalho.fill.fgColor.rgb.endswith("0F172A")


def test_excel_uma_aba_por_secao_nunca_tudo_empilhado():
    """Via HTTP, `incluir_graficos=True` é rejeitado (422) para qualquer
    formato que não seja PDF (`schemas.analitica._graficos_apenas_em_pdf`,
    regra já homologada, não alterada por esta correção) — então hoje o
    máximo alcançável por essa via é "Resumo" + "Tabela do recorte".
    Chamamos o serviço diretamente (sem o contrato HTTP no meio) só para
    confirmar que `_escrever_excel` já separa CADA distribuição em sua
    própria aba — comportamento correto e testado, mesmo que a regra de
    contrato atual não deixe o usuário chegar até ele por enquanto."""
    _sincronizar_muitas_linhas(5)
    caminho_http = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)
    workbook_http = load_workbook(caminho_http)
    assert workbook_http.sheetnames == ["Resumo", "Tabela do recorte"]

    from app.services import analitica_service

    parametros = analitica_service.ParametrosRelatorio(
        tipo="detalhado", formato="excel", modo="final",
        incluir_indicadores=True, incluir_graficos=True, incluir_tabela=True,
        incluir_inconsistencias=False,
    )
    resultado = analitica_service.gerar_relatorio(analitica_service.FiltrosDiaria(), parametros)
    workbook_direto = load_workbook(resultado.caminho_arquivo_temporario)
    assert "Resumo" in workbook_direto.sheetnames
    assert "Tabela do recorte" in workbook_direto.sheetnames
    assert len(workbook_direto.sheetnames) > 2  # Resumo + 6 distribuições + Tabela do recorte


def test_excel_resumo_do_recorte_na_aba_resumo():
    _sincronizar_muitas_linhas(7)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    valores = [
        str(celula.value)
        for linha in load_workbook(caminho)["Resumo"].iter_rows()
        for celula in linha
        if celula.value is not None
    ]
    assert any("Recorte:" in valor and "7 diária" in valor for valor in valores)


def test_excel_com_muitas_linhas_abre_corretamente():
    _sincronizar_muitas_linhas(500)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    aba = load_workbook(caminho)["Tabela do recorte"]
    assert aba.max_row == 501  # 1 cabeçalho + 500 linhas, nenhuma perdida


def test_excel_abre_como_documento_valido():
    _sincronizar_muitas_linhas(5)
    caminho = _gerar_relatorio(tipo="detalhado", formato="excel", incluir_graficos=False)

    with open(caminho, "rb") as arquivo:
        assinatura = arquivo.read(4)
    assert assinatura == b"PK\x03\x04"  # xlsx é um zip válido
    # `load_workbook` já usado nos testes acima levanta exceção se o
    # arquivo estiver corrompido — nenhuma chamada adicional necessária
    # aqui além da assinatura de arquivo.
