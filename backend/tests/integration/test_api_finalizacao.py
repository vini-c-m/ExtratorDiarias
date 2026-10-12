from fastapi.testclient import TestClient
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.infra.config import configuracao_app
from app.main import app

client = TestClient(app)


def _com_cabecalho_completo(caminho_planilha: str) -> str:
    workbook = load_workbook(caminho_planilha)
    planilha = workbook["Controle de Diárias"]
    planilha["P1"] = "VALOR TOTAL DAS DIÁRIAS"
    planilha["Q1"] = "LOCAL"
    workbook.save(caminho_planilha)
    return caminho_planilha


def _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_planilha: str) -> None:
    monkeypatch.setenv(configuracao_app.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    caminho = _com_cabecalho_completo(caminho_planilha)
    fingerprint = configuracao_app.validar_planilha(caminho)
    configuracao_app.persistir(caminho, fingerprint)


def _gerar_pdf(caminho: str, paginas: list[list[str]]) -> None:
    c = canvas.Canvas(caminho, pagesize=A4)
    _largura, altura = A4
    for pagina in paginas:
        y = altura - 50
        for linha in pagina:
            c.drawString(50, y, linha)
            y -= 14
        c.showPage()
    c.save()


def _pdf_confirmacao_e_recibo(tmp_path, sei: str) -> str:
    caminho = str(tmp_path / "processo.pdf")
    _gerar_pdf(
        caminho,
        [
            [
                "CMFC - Confirmação dos Serviços - Diária -Servidor - 1000001",
                "Confirmação dos serviços:",
                "Sim",
                f"CMFC - Confirmação dos Serviços - Diária -Servidor 1000001 SEI {sei} / pg. 1",
            ],
            [f"Recibo (1000002) SEI {sei} / pg. 2"],
        ],
    )
    return caminho


def test_post_finalizacao_analisar_e_confirmar_fluxo_completo(monkeypatch, tmp_path, caminho_planilha_teste):
    caminho_planilha = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_planilha)

    sei = "19.20.0003.0000001/2026-01"
    workbook = load_workbook(caminho_planilha)
    workbook["Controle de Diárias"]["B4"] = sei
    workbook.save(caminho_planilha)

    pdf = _pdf_confirmacao_e_recibo(tmp_path, sei)

    resposta_analise = client.post("/api/v1/finalizacao/analisar", json={"arquivos": [pdf]})
    assert resposta_analise.status_code == 200
    item = resposta_analise.json()["itens"][0]
    assert item["status"] == "pronto"
    assert item["avaliacao"]["apto"] is True
    assert len(item["candidatos"]) == 1
    linha = item["candidatos"][0]["linha"]

    resposta_confirmar = client.post(
        "/api/v1/finalizacao/confirmar",
        json={"itens": [{"arquivo_origem": "processo.pdf", "sei": sei}]},
    )
    assert resposta_confirmar.status_code == 200
    corpo = resposta_confirmar.json()
    assert corpo["total_finalizado"] == 1
    assert corpo["itens"][0]["categoria"] == "finalizado"
    assert corpo["itens"][0]["linhas_atualizadas"] == [linha]

    planilha = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha[f"X{linha}"].value == "SIM"
    assert planilha[f"Y{linha}"].value == "SIM"


def test_post_finalizacao_analisar_sem_configuracao_retorna_422(monkeypatch, tmp_path):
    monkeypatch.setenv(configuracao_app.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))

    resposta = client.post("/api/v1/finalizacao/analisar", json={"arquivos": []})

    assert resposta.status_code == 422
    assert "Nenhuma planilha oficial configurada" in resposta.json()["detail"]


# --- Organização Automática de PDFs -----------------------------------------


def test_post_finalizacao_confirmar_exclui_pdf_finalizado_da_pasta_de_origem(
    monkeypatch, tmp_path, caminho_planilha_teste
):
    caminho_planilha = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_planilha)

    sei = "19.20.0003.0000001/2026-01"
    workbook = load_workbook(caminho_planilha)
    workbook["Controle de Diárias"]["B4"] = sei
    workbook.save(caminho_planilha)

    pasta_a_confirmar = tmp_path / "a_confirmar"
    pasta_a_confirmar.mkdir()
    client.post(
        "/api/v1/configuracao/pastas",
        json={"caminho_pdfs_diarias_a_confirmar": str(pasta_a_confirmar)},
    )

    pdf = pasta_a_confirmar / "processo.pdf"
    _gerar_pdf(
        str(pdf),
        [
            [
                "CMFC - Confirmação dos Serviços - Diária -Servidor - 1000001",
                "Confirmação dos serviços:",
                "Sim",
                f"CMFC - Confirmação dos Serviços - Diária -Servidor 1000001 SEI {sei} / pg. 1",
            ],
            [f"Recibo (1000002) SEI {sei} / pg. 2"],
        ],
    )

    resposta_confirmar = client.post(
        "/api/v1/finalizacao/confirmar",
        json={"itens": [{"arquivo_origem": "processo.pdf", "sei": sei, "caminho_absoluto": str(pdf)}]},
    )

    assert resposta_confirmar.status_code == 200
    corpo = resposta_confirmar.json()
    assert corpo["itens"][0]["categoria"] == "finalizado"
    assert corpo["itens"][0]["organizacao"] == {
        "sucesso": True,
        "status": "excluido",
        "detalhe": None,
    }
    assert not pdf.exists()


def test_post_finalizacao_confirmar_nao_exclui_pdf_ignorado(
    monkeypatch, tmp_path, caminho_planilha_teste
):
    """SEI não localizado na Planilha Oficial — categoria IGNORADO nunca
    exclui o PDF da pasta de origem."""
    caminho_planilha = caminho_planilha_teste()
    _configurar_planilha_oficial(monkeypatch, tmp_path, caminho_planilha)

    pasta_a_confirmar = tmp_path / "a_confirmar"
    pasta_a_confirmar.mkdir()
    client.post(
        "/api/v1/configuracao/pastas",
        json={"caminho_pdfs_diarias_a_confirmar": str(pasta_a_confirmar)},
    )

    pdf = pasta_a_confirmar / "processo.pdf"
    pdf.write_bytes(b"conteudo")

    resposta_confirmar = client.post(
        "/api/v1/finalizacao/confirmar",
        json={
            "itens": [
                {
                    "arquivo_origem": "processo.pdf",
                    "sei": "19.20.9999.9999999/2026-99",
                    "caminho_absoluto": str(pdf),
                }
            ]
        },
    )

    assert resposta_confirmar.status_code == 200
    corpo = resposta_confirmar.json()
    assert corpo["itens"][0]["categoria"] == "ignorado"
    assert corpo["itens"][0]["organizacao"] is None
    assert pdf.exists()
