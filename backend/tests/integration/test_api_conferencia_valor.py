"""Micro-sprint A ("Valor Automático da Diária na Central de Conferência") —
testes de integração service+API do endpoint `POST /conferencia/valor`.
Casos G/H/I/J/K/N do pedido (destino resolvido/não resolvido/ambíguo, Tabela
Valores indisponível, UF ausente, Membro não usa a fórmula de Servidor) —
os casos A-F/M (regra financeira pura) já estão em
`tests/unit/test_valor_diaria_servidor.py`."""

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.consulta import base_local
from app.main import app

client = TestClient(app)

_CABECALHO_CONTROLE = [
    "DH", "SEI", "DOC ORIGEM", "CPF", "PREVIDÊNCIA", "NOME", "CARGO",
    "PERÍODO INICIAL", "PERÍODO FINAL", "TRECHO", "TRANSPORTE", "MOTIVO",
]
_CABECALHO_DADOS = [
    "SEI", "MATRÍCULA", "NOME", "PREVIDÊNCIA", "CARGO",
    "ÓRGÃO DE ORIGEM", "REGIME DE PREVIDÊNCIA", "SITUAÇÃO DO SERVIDOR",
]

_CARGO_SERVIDOR = "Analista Judiciário"
_CARGO_MEMBRO = "Promotor de Justiça de 1a. Entrancia"


def _montar_planilha(tmp_path, *, servidores: list[tuple], tabela_valores: list[tuple] | None):
    workbook = Workbook()
    controle = workbook.active
    controle.title = "Controle de Diárias"
    controle.append(_CABECALHO_CONTROLE)

    dados = workbook.create_sheet("Dados")
    dados.append(_CABECALHO_DADOS)
    for cpf, nome, previdencia, cargo in servidores:
        dados.append([cpf, 1001, nome, previdencia, cargo, None, "RPPS", "Em exercício"])

    tabela = workbook.create_sheet("Tabela Valores")
    tabela.append([None, None, None])
    tabela.append([None, None, None])
    tabela.append([None, "INTEGRAL", "PARCIAL"])
    for uf, integral, parcial in tabela_valores or []:
        tabela.append([uf, integral, parcial])

    caminho = tmp_path / "planilha_valor_diaria.xlsx"
    workbook.save(caminho)
    return str(caminho)


def _sincronizar_base_local(caminho_planilha: str, tmp_path, monkeypatch):
    destino_base = tmp_path / "dados_usuario"
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(destino_base))
    base_local.sincronizar(caminho_planilha)


def _requisitar_valor(**overrides):
    corpo = {
        "cpf": "12345678901",
        "periodo_inicial": "2026-07-21",
        "periodo_final": "2026-07-23",
        "destino": "Garanhuns",
    }
    corpo.update(overrides)
    return client.post("/api/v1/conferencia/valor", json=corpo)


def test_g_servidor_destino_resolvido_automaticamente_calcula_valor(tmp_path, monkeypatch):
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=[("PERNAMBUCO", 308.19, 154.09)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor()

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "calculado"
    assert corpo["categoria_funcional"] == "SERVIDOR"
    assert corpo["uf_destino"] == "PERNAMBUCO"
    # periodo 21/07 a 23/07 = 2 diárias integrais (RN-03) + 1 parcial (RN-04, modo Padrão)
    # `Decimal` é serializado pelo Pydantic como string (preserva precisão
    # exata ponta a ponta, nunca reintroduz o risco de ponto flutuante que
    # `Decimal` existe para evitar — ver `valor_diaria_servidor.py`).
    assert corpo["valor_integral"] == "616.38"
    assert corpo["valor_parcial"] == "154.09"
    assert corpo["valor_total"] == "770.47"


def test_h_destino_nao_resolvido_nunca_escolhe_silenciosamente(tmp_path, monkeypatch):
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=[("PERNAMBUCO", 308.19, 154.09)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor(destino="Xyzalandia Inexistente 123")

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "uf_nao_resolvida"
    assert corpo["valor_total"] is None
    assert corpo["ufs_disponiveis"] == ["PERNAMBUCO"]


def test_i_destino_ambiguo_nunca_escolhe_silenciosamente(tmp_path, monkeypatch):
    """'Rio Branco' existe no Acre e em Mato Grosso — achado real da
    auditoria contra a Planilha Oficial (linha 62)."""
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=[("PERNAMBUCO", 308.19, 154.09), ("ACRE", 539.33, 308.19)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor(destino="Rio Branco")

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "uf_nao_resolvida"
    assert corpo["valor_total"] is None


def test_uf_confirmada_manualmente_apos_ambiguidade_calcula_normalmente(tmp_path, monkeypatch):
    """Fluxo completo da RN de segurança: 1) ambíguo -> `uf_nao_resolvida`;
    2) usuário confirma; 3) mesma consulta com `uf_confirmada` calcula."""
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=[("ACRE", 539.33, 308.19)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor(destino="Rio Branco", uf_confirmada="ACRE")

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "calculado"
    assert corpo["uf_destino"] == "ACRE"


def test_k_uf_resolvida_mas_ausente_da_tabela(tmp_path, monkeypatch):
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=[("CEARÁ", 539.33, 308.19)],  # sem Pernambuco de propósito
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor(destino="Garanhuns")  # resolve para PERNAMBUCO

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "uf_fora_da_tabela"
    assert corpo["uf_destino"] == "PERNAMBUCO"
    assert corpo["valor_total"] is None
    assert corpo["ufs_disponiveis"] == ["CEARÁ"]


def test_j_tabela_valores_indisponivel_retorna_503_mas_nao_impede_historico(
    tmp_path, monkeypatch
):
    """Tabela Valores estruturalmente vazia (mesma aba criada vazia por
    `workbook.create_sheet`, achado real desta Micro-sprint): a falha
    financeira é explícita (503) — e, separadamente, o Histórico e-Fisco
    (endpoint independente) continua funcionando normalmente."""
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=None,  # aba vazia
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta_valor = _requisitar_valor()
    assert resposta_valor.status_code == 503

    resposta_historico = client.post(
        "/api/v1/conferencia/historico",
        json={
            "cpf": "12345678901",
            "periodo_inicial": "2026-07-21",
            "periodo_final": "2026-07-23",
            "sei": "19.20.0903.0030720/2024-36",
            "destino": "Garanhuns",
        },
    )
    assert resposta_historico.status_code == 200


def test_n_membro_nunca_usa_a_formula_de_servidor(tmp_path, monkeypatch):
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Herbet Gonçalves Santos", "FUNAFIN", _CARGO_MEMBRO)],
        tabela_valores=[("PERNAMBUCO", 308.19, 154.09)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor()

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "indisponivel_para_categoria"
    assert corpo["categoria_funcional"] == "MEMBRO"
    assert corpo["valor_total"] is None
    assert corpo["valor_integral"] is None
    assert "Membro" in corpo["motivo"]


def test_nao_identificado_tambem_nao_calcula(tmp_path, monkeypatch):
    """CPF cadastrado mas sem cargo preenchido — RN-01, situação real já
    documentada, nunca um 'chute' para Servidor nem Membro."""
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Fulano Sem Cargo", "FUNAFIN", "")],
        tabela_valores=[("PERNAMBUCO", 308.19, 154.09)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor()

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["situacao"] == "indisponivel_para_categoria"
    assert corpo["categoria_funcional"] == "NAO_IDENTIFICADO"


def test_servidor_nao_encontrado_404(tmp_path, monkeypatch):
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=[("PERNAMBUCO", 308.19, 154.09)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor(cpf="00000000000")

    assert resposta.status_code == 404


def test_base_local_indisponivel_503(tmp_path, monkeypatch):
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / "vazio"))

    resposta = _requisitar_valor()

    assert resposta.status_code == 503


def test_422_periodo_final_antes_do_inicial():
    resposta = _requisitar_valor(periodo_inicial="2026-07-23", periodo_final="2026-07-21")
    assert resposta.status_code == 422


def test_422_destino_vazio():
    resposta = _requisitar_valor(destino="   ")
    assert resposta.status_code == 422


def test_modo_complemento_usa_quantidade_diaria_parcial_informada(tmp_path, monkeypatch):
    """RN-16 — mesmo default institucional já homologado para o Histórico
    de Complemento: sem `quantidade_diaria_parcial`, vira 0; com o valor
    explícito, usa exatamente esse valor (nunca 1, que é exclusivo do modo
    Padrão / RN-04)."""
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=[("PERNAMBUCO", 308.19, 154.09)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta_default = _requisitar_valor(modo="COMPLEMENTO")
    assert resposta_default.status_code == 200
    corpo_default = resposta_default.json()
    # 2 integrais (21/07-23/07) + 0 parcial (default institucional do Complemento)
    assert corpo_default["valor_parcial"] == "0.00"
    assert corpo_default["valor_total"] == "616.38"

    resposta_explicita = _requisitar_valor(modo="COMPLEMENTO", quantidade_diaria_parcial=1)
    assert resposta_explicita.status_code == 200
    corpo_explicito = resposta_explicita.json()
    assert corpo_explicito["valor_parcial"] == "154.09"
    assert corpo_explicito["valor_total"] == "770.47"


def test_modo_padrao_ignora_quantidade_diaria_parcial_enviada_rn04(tmp_path, monkeypatch):
    """RN-04 é regra fixa, não negociável — mesmo que o cliente envie
    `quantidade_diaria_parcial=0` no modo Padrão (implícito, nem precisa ser
    enviado no schema real), o cálculo continua somando exatamente 1
    parcial, nunca 0."""
    caminho = _montar_planilha(
        tmp_path,
        servidores=[(12345678901, "Maria Aparecida Souza", "FUNAFIN", _CARGO_SERVIDOR)],
        tabela_valores=[("PERNAMBUCO", 308.19, 154.09)],
    )
    _sincronizar_base_local(caminho, tmp_path, monkeypatch)

    resposta = _requisitar_valor(modo="PADRAO", quantidade_diaria_parcial=0)

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["valor_parcial"] == "154.09"  # RN-04: sempre 1 parcial, nunca 0, no modo Padrão
    assert corpo["valor_total"] == "770.47"
