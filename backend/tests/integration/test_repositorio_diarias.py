from datetime import date

from openpyxl import load_workbook

from app.consulta.repositorio_diarias import NOME_ABA, diarias_por_cpf


def test_diarias_por_cpf_indexa_registro_existente(caminho_planilha_conferencia_teste):
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    resultado = diarias_por_cpf(workbook[NOME_ABA])

    assert "12345678901" in resultado.diarias
    registros = resultado.diarias["12345678901"]
    assert len(registros) == 1
    registro = registros[0]
    assert registro.linha == 2
    assert registro.sei == "19.20.0902.0004521/2026-11"
    assert registro.periodo_inicial == date(2026, 7, 21)
    assert registro.periodo_final == date(2026, 7, 23)
    assert registro.trecho == "RECIFE - GARANHUNS"
    assert resultado.avisos == []


def test_diarias_por_cpf_sem_registro_para_cpf_desconhecido(caminho_planilha_conferencia_teste):
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    resultado = diarias_por_cpf(workbook[NOME_ABA])

    assert "00000000000" not in resultado.diarias


def test_diarias_por_cpf_ignora_linha_com_periodo_invalido_e_gera_aviso(
    caminho_planilha_conferencia_teste,
):
    """Reproduz um caso real encontrado na planilha de produção: uma linha
    histórica com o período gravado como texto malformado (ano com dígito
    faltando), em vez de uma data — antes desta correção, isso derrubava
    com `ValueError` a indexação da aba inteira, quebrando a conferência
    para qualquer CPF, não só o da linha problemática."""
    caminho = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho)
    planilha = workbook[NOME_ABA]
    planilha.append(
        [
            "2026DH000101",
            "19.20.0902.0009999/2026-22",
            "PJTESTE",
            22222222222,
            "=VLOOKUP($D3,Dados!$A:$E,4,FALSE)",
            "=VLOOKUP($D3,Dados!$A:$E,3,FALSE)",
            "=VLOOKUP($D3,Dados!$A:$E,5,FALSE)",
            date(2026, 8, 1),
            "22/05/0206",
            "RECIFE - CARUARU",
            "CARRO",
            "Motivo de teste",
        ]
    )

    resultado = diarias_por_cpf(planilha)

    assert "22222222222" not in resultado.diarias
    assert "12345678901" in resultado.diarias  # linha válida da fixture continua indexada
    assert len(resultado.avisos) == 1
    assert resultado.avisos[0].aba == NOME_ABA
    assert resultado.avisos[0].linha == 3
    assert "22/05/0206" in resultado.avisos[0].motivo
