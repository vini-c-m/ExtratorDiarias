"""ETAPA 3 (Auditoria P0, 2026-08-09) — Finalizar/Confirmar dispara a
sincronização da Base Local, reutilizando integralmente o mecanismo já
consolidado na ETAPA 2 (`base_local.sincronizar`, `LOCK_SINCRONIZACAO`,
retry/validação/Last Known Good) e a preservação de cache de fórmula da
ETAPA 1. Nenhum dado real — tudo em `tmp_path`, mesmos padrões de
`test_finalizacao_service.py`/`test_sincronizacao_concorrencia_p0.py`."""

import sqlite3
import threading
from datetime import date

from openpyxl import load_workbook

from app.consulta import base_local, base_local_sqlite
from app.services.finalizacao_service import ItemConfirmacao, confirmar


def _planilha_com_linha(caminho_planilha_teste, sei: str) -> str:
    """`_planilha_com_linhas` de `test_finalizacao_service.py` só grava o
    SEI (suficiente para `confirmar()`, que não olha CPF/período) — aqui
    a linha precisa ficar completa o bastante para `diarias_completas()`
    (chamada por `base_local.sincronizar()`) também aceitá-la, já que
    estes testes encadeiam até a sincronização de verdade."""
    caminho = caminho_planilha_teste()
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    planilha["B4"] = sei
    planilha["D4"] = 11111111111
    planilha["H4"] = date(2026, 1, 8)
    planilha["I4"] = date(2026, 1, 10)
    planilha["L4"] = "Motivo de teste"
    workbook.save(caminho)
    workbook.close()
    return caminho


def _total_diarias(caminho_sqlite) -> int:
    conexao = sqlite3.connect(caminho_sqlite)
    try:
        return conexao.execute("SELECT COUNT(*) FROM diarias").fetchone()[0]
    finally:
        conexao.close()


def test_confirmar_com_gravacao_real_dispara_sincronizacao_e_atualiza_base_local(
    tmp_path, caminho_planilha_teste, monkeypatch
):
    """(A) Alteração real -> Planilha alterada -> sincronização chamada ->
    Base Local atualizada."""
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    sei = "19.20.0002.0000007/2026-07"
    caminho_planilha = _planilha_com_linha(caminho_planilha_teste, sei)

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    assert resultado.total_finalizado == 1
    assert resultado.sincronizacao_analitica_falhou is False

    caminho_sqlite = base_local_sqlite.caminho_base_local_sqlite(diretorio_base=tmp_path / "dados_usuario")
    assert caminho_sqlite.is_file()
    assert _total_diarias(caminho_sqlite) == 1


def test_falha_de_sincronizacao_nao_desfaz_confirmacao_e_preserva_base_antiga(
    tmp_path, caminho_planilha_teste, monkeypatch
):
    """(B) Falha da sincronização após finalização -> finalização
    permanece concluída -> base antiga íntegra permanece disponível ->
    falha analítica sinalizada separadamente (nunca reportada como
    'confirmação falhou')."""
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    sei = "19.20.0002.0000008/2026-08"
    caminho_planilha = _planilha_com_linha(caminho_planilha_teste, sei)

    # Base "antiga" já sincronizada com sucesso, antes da falha simulada.
    base_local.sincronizar(caminho_planilha, diretorio_base=tmp_path / "dados_usuario", origem="teste")
    caminho_sqlite = base_local_sqlite.caminho_base_local_sqlite(diretorio_base=tmp_path / "dados_usuario")
    total_antes = _total_diarias(caminho_sqlite)

    def reconstruir_falha(*args, **kwargs):
        raise RuntimeError("falha simulada na reconstrução SQLite")

    monkeypatch.setattr("app.consulta.base_local_sqlite.reconstruir", reconstruir_falha)

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    # Finalização concluída de verdade — nunca desfeita pela falha de sync.
    assert resultado.total_finalizado == 1
    assert resultado.sincronizacao_analitica_falhou is True

    planilha_depois = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha_depois["X4"].value == "SIM"
    assert planilha_depois["Y4"].value == "SIM"

    # Base Local antiga continua íntegra e disponível (mesma contagem).
    assert _total_diarias(caminho_sqlite) == total_antes


def test_confirmar_sem_alteracao_efetiva_nao_dispara_sincronizacao(
    tmp_path, caminho_planilha_teste, monkeypatch
):
    """(C) Nenhuma alteração efetiva (SEI não encontrado) -> não disparar
    sincronização inútil."""
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    caminho_planilha = caminho_planilha_teste()

    chamou_sincronizar = {"valor": False}
    sincronizar_original = base_local.sincronizar

    def sincronizar_espiao(*args, **kwargs):
        chamou_sincronizar["valor"] = True
        return sincronizar_original(*args, **kwargs)

    monkeypatch.setattr(base_local, "sincronizar", sincronizar_espiao)

    resultado = confirmar(
        [ItemConfirmacao(arquivo_origem="p.pdf", sei="19.20.9999.9999999/2099-99")], caminho_planilha
    )

    assert resultado.total_finalizado == 0
    assert chamou_sincronizar["valor"] is False


def test_confirmar_e_sincronizacao_manual_concorrentes_usam_o_mesmo_lock(
    tmp_path, caminho_planilha_teste, monkeypatch
):
    """(D) Confirma reuso do mesmo lock/mecanismo da ETAPA 2: `confirmar()`
    (que chama `base_local.sincronizar` internamente) e uma sincronização
    manual concorrente, sobre o mesmo destino, nunca colidem (mesma prova
    de `test_sincronizacao_concorrencia_p0.py`, agora envolvendo
    Finalizar/Confirmar)."""
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    sei = "19.20.0002.0000009/2026-09"
    caminho_planilha = _planilha_com_linha(caminho_planilha_teste, sei)

    barreira = threading.Barrier(2)
    erros: dict[str, Exception | None] = {}

    def tarefa_confirmar() -> None:
        try:
            barreira.wait(timeout=5)
            confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)
            erros["confirmar"] = None
        except Exception as erro:  # noqa: BLE001
            erros["confirmar"] = erro

    def tarefa_sincronizar_manual() -> None:
        try:
            barreira.wait(timeout=5)
            base_local.sincronizar(caminho_planilha, diretorio_base=tmp_path / "dados_usuario", origem="manual")
            erros["manual"] = None
        except Exception as erro:  # noqa: BLE001
            erros["manual"] = erro

    t1 = threading.Thread(target=tarefa_confirmar)
    t2 = threading.Thread(target=tarefa_sincronizar_manual)
    t1.start()
    t2.start()
    t1.join(timeout=15)
    t2.join(timeout=15)

    assert erros == {"confirmar": None, "manual": None}
    caminho_sqlite = base_local_sqlite.caminho_base_local_sqlite(diretorio_base=tmp_path / "dados_usuario")
    assert caminho_sqlite.is_file()
    assert not caminho_sqlite.with_suffix(caminho_sqlite.suffix + ".tmp").exists()
