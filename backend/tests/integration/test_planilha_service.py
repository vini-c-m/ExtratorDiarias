import os
from datetime import date

import pytest
from openpyxl import load_workbook

from app.domain.entities.diaria import Diaria
from app.domain.entities.processo import Processo, StatusProcesso
from app.planilha.relatorio import ItemErro
from app.services.planilha_service import ErroPreenchimentoPlanilha, preencher_planilha


def _processo(arquivo: str = "arquivo.pdf", sei: str = "19.20.0000.0000001/2025-99") -> Processo:
    diaria = Diaria(
        sei=sei,
        doc_origem="PJNOVO",
        cpf="99999999999",
        periodo_inicial=date(2025, 3, 1),
        periodo_final=date(2025, 3, 3),
        cidade_origem="RECIFE",
        cidade_destino="OLINDA",
        motivo="Diligência de teste",
    )
    return Processo(arquivo_origem=arquivo, status=StatusProcesso.SUCESSO, diaria=diaria)


def test_preencher_planilha_com_sucesso_monta_pasta_de_resultados(caminho_planilha_teste):
    caminho = caminho_planilha_teste()

    resultado, arquivos_gerados, _ = preencher_planilha(caminho, [_processo()], [])

    assert resultado.total_gravado == 1
    assert resultado.linha_inicial == 4

    assert os.path.isdir(arquivos_gerados.pasta_resultados)
    assert os.path.dirname(arquivos_gerados.pasta_resultados) == os.path.dirname(caminho)

    assert os.path.isfile(arquivos_gerados.caminho_planilha_gerada)
    assert os.path.isfile(arquivos_gerados.caminho_relatorio)
    assert arquivos_gerados.caminho_relatorio.endswith(".html")
    assert os.path.isfile(arquivos_gerados.caminho_log)

    # A cópia na pasta de resultados reflete a planilha já preenchida.
    copia = load_workbook(arquivos_gerados.caminho_planilha_gerada)["Controle de Diárias"]
    assert copia["B4"].value == _processo().diaria.sei

    conteudo_relatorio = open(arquivos_gerados.caminho_relatorio, encoding="utf-8").read()
    assert "PDFs analisados" in conteudo_relatorio
    assert "Diárias inseridas" in conteudo_relatorio
    assert "<html" in conteudo_relatorio.lower()

    conteudo_log = open(arquivos_gerados.caminho_log, encoding="utf-8").read()
    assert "INSERIDO" in conteudo_log
    assert "gravado na linha 4" in conteudo_log


def test_preencher_planilha_relatorio_distingue_duplicidade_no_lote_e_na_planilha(
    caminho_planilha_teste,
):
    caminho = caminho_planilha_teste()
    erros = [ItemErro(arquivo_origem="quebrado.pdf", motivos=["Falha ao ler o PDF"])]

    resultado, arquivos_gerados, _ = preencher_planilha(
        caminho,
        [
            _processo("novo.pdf", "19.20.0000.0000005/2025-99"),
            _processo("novo_copia.pdf", "19.20.0000.0000005/2025-99"),  # duplicado no lote
            _processo("ja_na_planilha.pdf", "19.20.0000.0000001/2025-00"),  # duplicado na planilha
        ],
        erros,
    )

    assert resultado.total_gravado == 1
    assert resultado.duplicadas_ignoradas == 2
    assert len(resultado.duplicados_no_lote) == 1
    assert len(resultado.duplicados_na_planilha) == 1
    assert resultado.duplicados_na_planilha[0].linha == 2

    conteudo = open(arquivos_gerados.caminho_relatorio, encoding="utf-8").read()
    assert "quebrado.pdf" in conteudo
    assert "novo_copia.pdf" in conteudo
    assert "ja_na_planilha.pdf" in conteudo

    conteudo_log = open(arquivos_gerados.caminho_log, encoding="utf-8").read()
    assert "DUPLICADO_NO_LOTE" in conteudo_log
    assert "DUPLICADO_NA_PLANILHA" in conteudo_log
    assert "já existe na linha 2" in conteudo_log
    assert "ERRO arquivo: quebrado.pdf" in conteudo_log


def test_preencher_planilha_inexistente_gera_erro_claro(tmp_path):
    caminho = str(tmp_path / "nao_existe.xlsx")

    with pytest.raises(ErroPreenchimentoPlanilha, match="não encontrada"):
        preencher_planilha(caminho, [_processo()], [])


# --- Sprint 06 — DOC DE ORIGEM: aviso de fallback chega ao relatório e ao log ---


def test_doc_origem_sem_historico_gera_aviso_no_relatorio_e_no_log(caminho_planilha_teste):
    """O CPF de `_processo()` (99999999999) não existe na fixture — sem
    histórico na planilha oficial do DIMCAC, a coluna C é preenchida pela
    extração do PDF (Prioridade 2), e o relatório/log precisam registrar o
    aviso automaticamente, mesmo sem nenhum aviso vindo da extração."""
    caminho = caminho_planilha_teste()

    resultado, arquivos_gerados, _ = preencher_planilha(caminho, [_processo()], [])

    assert resultado.total_gravado == 1

    conteudo_relatorio = open(arquivos_gerados.caminho_relatorio, encoding="utf-8").read()
    assert "⚠ Revisão recomendada" in conteudo_relatorio
    assert "Não foi localizado histórico deste servidor na planilha oficial do DIMCAC" in (
        conteudo_relatorio
    )

    conteudo_log = open(arquivos_gerados.caminho_log, encoding="utf-8").read()
    assert "REVISAO_MANUAL arquivo: arquivo.pdf" in conteudo_log
    assert "planilha oficial do DIMCAC" in conteudo_log


def test_doc_origem_com_historico_dimcac_nao_gera_aviso(caminho_planilha_teste):
    """CPF 11111111111 já tem histórico na linha 2 da fixture (DOC ORIGEM
    'PJTESTE') — a coluna C reaproveita esse valor e nenhum aviso é gerado,
    exatamente como uma diária já totalmente resolvida."""
    caminho = caminho_planilha_teste()
    diaria = Diaria(
        sei="19.20.0000.0000008/2025-99",
        doc_origem="TEXTO EXTRAIDO, DIFERENTE DO HISTORICO",
        cpf="11111111111",
        periodo_inicial=date(2025, 3, 1),
        periodo_final=date(2025, 3, 3),
        cidade_origem="RECIFE",
        cidade_destino="OLINDA",
        motivo="Diligência de teste",
    )
    processo = Processo(arquivo_origem="com_historico.pdf", status=StatusProcesso.SUCESSO, diaria=diaria)

    resultado, arquivos_gerados, _ = preencher_planilha(caminho, [processo], [])

    assert resultado.total_gravado == 1

    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    assert planilha["C4"].value == "PJTESTE"

    conteudo_relatorio = open(arquivos_gerados.caminho_relatorio, encoding="utf-8").read()
    assert "com_historico.pdf" in conteudo_relatorio
    assert "planilha oficial do DIMCAC" not in conteudo_relatorio

    conteudo_log = open(arquivos_gerados.caminho_log, encoding="utf-8").read()
    assert "REVISAO_MANUAL arquivo: com_historico.pdf" not in conteudo_log
