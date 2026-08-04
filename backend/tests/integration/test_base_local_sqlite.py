import os
import sqlite3
from datetime import date

import pytest
from openpyxl import load_workbook

from app.consulta import base_local_sqlite
from app.consulta.repositorio_diarias import NOME_ABA as NOME_ABA_DIARIAS
from app.consulta.repositorio_diarias import diarias_completas
from app.consulta.repositorio_servidores import NOME_ABA as NOME_ABA_SERVIDORES
from app.consulta.repositorio_servidores import servidores_por_cpf


def _reconstruir_a_partir_de(caminho_planilha, caminho_destino):
    workbook = load_workbook(caminho_planilha)
    try:
        resultado_diarias = diarias_completas(workbook[NOME_ABA_DIARIAS])
        servidores = servidores_por_cpf(workbook[NOME_ABA_SERVIDORES])
    finally:
        workbook.close()
    sinal = base_local_sqlite.capturar_sinal_frescor(caminho_planilha)
    return base_local_sqlite.reconstruir(
        resultado_diarias, servidores, sinal, origem="teste", caminho_destino=caminho_destino
    )


def test_reconstruir_grava_diaria_denormalizada_com_servidor(
    tmp_path, caminho_planilha_conferencia_teste
):
    caminho_planilha = caminho_planilha_conferencia_teste()
    destino = tmp_path / "base_local.sqlite"

    resumo = _reconstruir_a_partir_de(caminho_planilha, destino)

    assert destino.is_file()
    assert resumo.linhas_diarias == 1
    assert resumo.linhas_ignoradas == 0
    assert resumo.status == "sucesso"

    conexao = sqlite3.connect(destino)
    try:
        linha = conexao.execute(
            "SELECT sei, cpf, servidor_nome, servidor_matricula, servidor_situacao FROM diarias"
        ).fetchone()
    finally:
        conexao.close()

    assert linha == (
        "19.20.0902.0004521/2026-11",
        "12345678901",
        "Maria Aparecida Souza",
        "1001",
        "Em exercício",
    )


def test_reconstruir_diaria_com_cpf_sem_cadastro_fica_com_campos_servidor_nulos(
    tmp_path, caminho_planilha_conferencia_teste
):
    caminho_planilha = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho_planilha)
    planilha = workbook[NOME_ABA_DIARIAS]
    planilha.append(
        [
            "2026DH000102",
            "19.20.0902.0009999/2026-33",
            "PJTESTE",
            99999999999,
            None,
            None,
            None,
            date(2026, 8, 1),
            date(2026, 8, 3),
            "RECIFE - OLINDA",
            "CARRO",
            "Motivo sem cadastro",
        ]
    )
    workbook.save(caminho_planilha)
    workbook.close()

    destino = tmp_path / "base_local.sqlite"
    resumo = _reconstruir_a_partir_de(caminho_planilha, destino)

    assert resumo.linhas_diarias == 2
    conexao = sqlite3.connect(destino)
    try:
        linha = conexao.execute(
            "SELECT servidor_nome, servidor_matricula FROM diarias WHERE sei = ?",
            ("19.20.0902.0009999/2026-33",),
        ).fetchone()
    finally:
        conexao.close()

    assert linha == (None, None)


def test_reconstruir_grava_dh_da_diaria(tmp_path, caminho_planilha_conferencia_teste):
    caminho_planilha = caminho_planilha_conferencia_teste()
    destino = tmp_path / "base_local.sqlite"

    _reconstruir_a_partir_de(caminho_planilha, destino)

    conexao = sqlite3.connect(destino)
    try:
        dh = conexao.execute(
            "SELECT dh FROM diarias WHERE sei = ?", ("19.20.0902.0004521/2026-11",)
        ).fetchone()[0]
    finally:
        conexao.close()

    assert dh == "2026DH000100"


def test_reconstruir_diaria_com_dh_vazio_grava_null(tmp_path, caminho_planilha_conferencia_teste):
    """DH vazio é esperado (processo ainda não chegou à fase de
    liquidação) — a reconstrução deve gravar `NULL`, nunca inventar ou
    descartar a linha por causa disso."""
    caminho_planilha = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho_planilha)
    planilha = workbook[NOME_ABA_DIARIAS]
    planilha.append(
        [
            None,
            "19.20.0902.0008888/2026-55",
            "PJTESTE",
            55555555555,
            None,
            None,
            None,
            date(2026, 8, 10),
            date(2026, 8, 12),
            "RECIFE - CARUARU",
            "CARRO",
            "Motivo com DH vazio",
        ]
    )
    workbook.save(caminho_planilha)
    workbook.close()

    destino = tmp_path / "base_local.sqlite"
    resumo = _reconstruir_a_partir_de(caminho_planilha, destino)

    assert resumo.linhas_diarias == 2
    conexao = sqlite3.connect(destino)
    try:
        dh = conexao.execute(
            "SELECT dh FROM diarias WHERE sei = ?", ("19.20.0902.0008888/2026-55",)
        ).fetchone()[0]
    finally:
        conexao.close()

    assert dh is None


def test_reconstruir_grava_avisos_sincronizacao(tmp_path, caminho_planilha_conferencia_teste):
    caminho_planilha = caminho_planilha_conferencia_teste()
    workbook = load_workbook(caminho_planilha)
    planilha = workbook[NOME_ABA_DIARIAS]
    planilha.append(
        [
            "2026DH000103",
            "19.20.0902.0007777/2026-44",
            "PJTESTE",
            11122233344,
            None,
            None,
            None,
            date(2026, 8, 5),
            "22/05/0206",  # período malformado — mesmo caso já coberto em diarias_por_cpf
            "RECIFE - CARUARU",
            "CARRO",
            "Motivo com período inválido",
        ]
    )
    workbook.save(caminho_planilha)
    workbook.close()

    destino = tmp_path / "base_local.sqlite"
    resumo = _reconstruir_a_partir_de(caminho_planilha, destino)

    assert resumo.linhas_ignoradas == 1
    assert len(resumo.avisos) == 1

    conexao = sqlite3.connect(destino)
    try:
        total_avisos = conexao.execute("SELECT COUNT(*) FROM avisos_sincronizacao").fetchone()[0]
    finally:
        conexao.close()
    assert total_avisos == 1


def test_reconstruir_atomico_nao_toca_destino_se_falhar_no_meio(
    tmp_path, caminho_planilha_conferencia_teste, monkeypatch
):
    caminho_planilha = caminho_planilha_conferencia_teste()
    destino = tmp_path / "base_local.sqlite"
    _reconstruir_a_partir_de(caminho_planilha, destino)
    conteudo_original = destino.read_bytes()

    def _falhar(_conexao):
        raise RuntimeError("falha simulada durante a reconstrução")

    monkeypatch.setattr(base_local_sqlite, "criar_esquema", _falhar)

    with pytest.raises(RuntimeError):
        _reconstruir_a_partir_de(caminho_planilha, destino)

    assert destino.read_bytes() == conteudo_original


def test_reconstruir_nao_deixa_sidecar_de_journal(tmp_path, caminho_planilha_conferencia_teste):
    caminho_planilha = caminho_planilha_conferencia_teste()
    destino = tmp_path / "base_local.sqlite"

    _reconstruir_a_partir_de(caminho_planilha, destino)

    arquivos_no_destino = os.listdir(destino.parent)
    assert not any(nome.endswith((".tmp", "-wal", "-shm")) for nome in arquivos_no_destino)


def test_capturar_sinal_frescor_reflete_linhas_da_planilha(caminho_planilha_conferencia_teste):
    caminho = caminho_planilha_conferencia_teste()
    sinal = base_local_sqlite.capturar_sinal_frescor(caminho)
    assert sinal.linhas_controle_diarias == 2  # cabeçalho + 1 linha de dado


def test_capturar_sinal_frescor_planilha_normal_nao_aciona_fallback(
    monkeypatch, caminho_planilha_conferencia_teste
):
    """Micro-sprint 18.7.1 — requisito 1: o caminho rápido (`read_only=True`)
    continua sendo o único usado quando `max_row` já vem como inteiro
    válido; o fallback (`read_only=False`) nunca é acionado à toa."""
    caminho = caminho_planilha_conferencia_teste()
    chamadas_read_only = []
    load_workbook_original = base_local_sqlite.load_workbook

    def _load_workbook_rastreado(*args, **kwargs):
        chamadas_read_only.append(kwargs.get("read_only"))
        return load_workbook_original(*args, **kwargs)

    monkeypatch.setattr(base_local_sqlite, "load_workbook", _load_workbook_rastreado)

    sinal = base_local_sqlite.capturar_sinal_frescor(caminho)

    assert chamadas_read_only == [True]
    assert sinal.linhas_controle_diarias == 2


def test_capturar_sinal_frescor_usa_fallback_quando_dimension_ausente(
    caminho_planilha_conferencia_teste, remover_tag_dimension_da_planilha
):
    """Micro-sprint 18.7.1 — requisito 2: planilha sem a tag `<dimension>`
    (caso real, Google Sheets) faz `read_only=True` devolver `max_row =
    None`; `capturar_sinal_frescor` deve cair no fallback completo e
    devolver a quantidade real de linhas, nunca `None`."""
    caminho = caminho_planilha_conferencia_teste()
    remover_tag_dimension_da_planilha(caminho)

    # Confirma a premissa do bug antes de testar a correção.
    workbook_read_only = load_workbook(caminho, read_only=True, data_only=False)
    try:
        assert workbook_read_only[NOME_ABA_DIARIAS].max_row is None
    finally:
        workbook_read_only.close()

    sinal = base_local_sqlite.capturar_sinal_frescor(caminho)

    assert sinal.linhas_controle_diarias == 2  # cabeçalho + 1 linha de dado — fallback escaneou de verdade
    assert isinstance(sinal.linhas_controle_diarias, int)


def test_capturar_sinal_frescor_levanta_erro_claro_quando_ambos_os_modos_falham(
    monkeypatch, caminho_planilha_conferencia_teste
):
    """Micro-sprint 18.7.1 — requisito 4: se nem o fallback conseguir um
    inteiro válido, `capturar_sinal_frescor` nunca devolve `None`/`0`
    silenciosamente — levanta `SinalFrescorIndeterminado`, uma exceção
    própria e diagnosticável, antes de qualquer tentativa de gravar no
    SQLite."""
    caminho = caminho_planilha_conferencia_teste()

    class _PlanilhaSemMaxRow:
        max_row = None

    class _WorkbookFalso:
        def __getitem__(self, _nome):
            return _PlanilhaSemMaxRow()

        def close(self):
            pass

    monkeypatch.setattr(base_local_sqlite, "load_workbook", lambda *args, **kwargs: _WorkbookFalso())

    with pytest.raises(base_local_sqlite.SinalFrescorIndeterminado):
        base_local_sqlite.capturar_sinal_frescor(caminho)


def test_reconstruir_com_planilha_sem_dimension_publica_base_local_sqlite(
    tmp_path, caminho_planilha_conferencia_teste, remover_tag_dimension_da_planilha
):
    """Micro-sprint 18.7.1 — requisito 3 (reconstrução completa): a
    planilha real que expôs o bug original agora deve concluir a
    reconstrução, publicar `base_local.sqlite`, não deixar `.tmp` órfão, e
    gravar um inteiro válido em `linhas_controle_diarias_fonte` — as
    diárias históricas continuam consultáveis."""
    caminho_planilha = caminho_planilha_conferencia_teste()
    remover_tag_dimension_da_planilha(caminho_planilha)

    destino = tmp_path / "base_local.sqlite"
    resumo = _reconstruir_a_partir_de(caminho_planilha, destino)

    assert resumo.status == "sucesso"
    assert destino.is_file()
    assert not any(nome.endswith(".tmp") for nome in os.listdir(destino.parent))

    conexao = sqlite3.connect(destino)
    try:
        linhas_fonte = conexao.execute(
            "SELECT linhas_controle_diarias_fonte FROM sincronizacoes"
        ).fetchone()[0]
        diaria = conexao.execute("SELECT sei FROM diarias").fetchone()
    finally:
        conexao.close()

    assert isinstance(linhas_fonte, int)
    assert linhas_fonte == 2
    assert diaria == ("19.20.0902.0004521/2026-11",)


def test_sinal_mudou_registrado_none_sempre_true():
    sinal = base_local_sqlite.SinalFrescor(tamanho_bytes=1, mtime=1.0, linhas_controle_diarias=1)
    assert base_local_sqlite.sinal_mudou(sinal, None) is True


def test_sinal_mudou_false_quando_identico():
    sinal = base_local_sqlite.SinalFrescor(tamanho_bytes=1, mtime=1.0, linhas_controle_diarias=1)
    assert base_local_sqlite.sinal_mudou(sinal, sinal) is False


def test_sinal_mudou_true_quando_diferente():
    atual = base_local_sqlite.SinalFrescor(tamanho_bytes=2, mtime=1.0, linhas_controle_diarias=1)
    registrado = base_local_sqlite.SinalFrescor(tamanho_bytes=1, mtime=1.0, linhas_controle_diarias=1)
    assert base_local_sqlite.sinal_mudou(atual, registrado) is True


def test_ultimo_sinal_frescor_none_quando_nunca_sincronizado(tmp_path):
    assert base_local_sqlite.ultimo_sinal_frescor(tmp_path / "nao_existe.sqlite") is None


def test_ultimo_sinal_frescor_le_o_registrado(tmp_path, caminho_planilha_conferencia_teste):
    caminho_planilha = caminho_planilha_conferencia_teste()
    destino = tmp_path / "base_local.sqlite"
    _reconstruir_a_partir_de(caminho_planilha, destino)

    sinal = base_local_sqlite.ultimo_sinal_frescor(destino)
    esperado = base_local_sqlite.capturar_sinal_frescor(caminho_planilha)
    assert sinal == esperado


# --- Sprint 19.9 — `esquema_desatualizado` (achado da auditoria: Electron
# real + Planilha Oficial real revelou um `.sqlite` de instalação anterior
# a uma mudança de schema, nunca detectado como desatualizado enquanto a
# Planilha Oficial em si não mudava). ------------------------------------


def test_esquema_desatualizado_false_quando_nunca_sincronizado(tmp_path):
    assert base_local_sqlite.esquema_desatualizado(tmp_path / "nao_existe.sqlite") is False


def test_esquema_desatualizado_false_apos_reconstrucao_normal(tmp_path, caminho_planilha_conferencia_teste):
    caminho_planilha = caminho_planilha_conferencia_teste()
    destino = tmp_path / "base_local.sqlite"
    _reconstruir_a_partir_de(caminho_planilha, destino)

    assert base_local_sqlite.esquema_desatualizado(destino) is False


def test_esquema_desatualizado_true_quando_user_version_diverge(tmp_path, caminho_planilha_conferencia_teste):
    caminho_planilha = caminho_planilha_conferencia_teste()
    destino = tmp_path / "base_local.sqlite"
    _reconstruir_a_partir_de(caminho_planilha, destino)

    # Simula um `.sqlite` de uma versão de código anterior, sem qualquer
    # mudança na Planilha Oficial — reproduz exatamente o caso encontrado
    # na auditoria (schema desatualizado, conteúdo "fresco").
    conexao = sqlite3.connect(destino)
    try:
        conexao.execute("PRAGMA user_version = 1")
        conexao.commit()
    finally:
        conexao.close()

    assert base_local_sqlite.esquema_desatualizado(destino) is True


def test_caminho_base_local_sqlite_mesma_pasta_do_xlsx(tmp_path, monkeypatch):
    from app.consulta import base_local

    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path))
    caminho_xlsx = base_local.caminho_base_local()
    caminho_sqlite = base_local_sqlite.caminho_base_local_sqlite()

    assert caminho_sqlite.parent == caminho_xlsx.parent
    assert caminho_sqlite.name == "base_local.sqlite"
