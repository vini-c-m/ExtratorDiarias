"""ETAPA 4 (Auditoria P0, 2026-08-09) — invariantes permanentes da Base
Local, formalizadas como testes. Não introduz mecanismo novo: só prova, de
forma explícita e nomeada, garantias que as Etapas 1-3 já implementaram.

Mapa das 8 invariantes e onde cada uma é provada (algumas só aqui, outras
já cobertas em detalhe nas suítes das Etapas anteriores — não duplicadas):

1. LAST KNOWN GOOD — falha de sincronização nunca invalida a última Base
   Local íntegra. Provado aqui (`test_invariante_1_...`) e em
   `test_sincronizacao_concorrencia_p0.py::
   test_permissionerror_persistente_preserva_base_antiga_e_nao_falseia_metadata`.

2. READ ONLY — consulta da Central Analítica nunca altera Planilha
   Oficial nem Base Local. Provado aqui
   (`test_invariante_2_leitura_da_central_nao_altera_planilha_nem_base_local`)
   por assinatura de arquivo (mtime+hash) antes/depois de várias consultas
   reais via `repositorio_analitico`.

3. ATOMICIDADE — nova Base Local só se torna oficial depois de construída,
   fechada e validada por completo. Provado aqui
   (`test_invariante_3_falha_durante_construcao_nao_promove_base_parcial`,
   falha ANTES do commit) e em `test_sincronizacao_concorrencia_p0.py`
   (falha no `os.replace`, DEPOIS do commit).

4. FALHA EXPLÍCITA — falha de sincronização nunca vira sucesso silencioso.
   Já provado em `test_sincronizacao_concorrencia_p0.py` (metadata) e
   `test_finalizacao_sincronizacao_p0.py`/`test_integridade_valor_
   financeiro.py` (sinalização ao chamador) — não duplicado aqui.

5. CONSISTÊNCIA ANALÍTICA — escrita legítima altera só o dado
   correspondente, nunca derruba valores históricos não relacionados. Já
   provado em `test_integridade_valor_financeiro.py` (processamento e
   complemento) e `test_finalizacao_sincronizacao_p0.py` (finalização) —
   não duplicado aqui, exceto o caso de fórmula específico da invariante 8.

6. RECUPERAÇÃO — sincronização válida após uma falha recupera sozinha.
   Já provado em `test_sincronizacao_concorrencia_p0.py::
   test_sincronizacao_valida_apos_falha_anterior_recupera_normalmente`.

7. SINCRONIZAÇÃO DE TODA ESCRITA RELEVANTE — processamento, complemento e
   Finalizar/Confirmar usam o MESMO mecanismo (`base_local.sincronizar`),
   nunca uma implementação paralela. Provado aqui
   (`test_invariante_7_...`).

8. PRESERVAÇÃO DE FÓRMULAS/CACHE — todo fluxo homologado que salva a
   Planilha Oficial preserva fórmula e cache necessário à leitura
   analítica. Já provado para processamento/complemento em
   `test_integridade_valor_financeiro.py`; provado aqui para Finalizar/
   Confirmar (`test_invariante_8_...`), o único caso ainda sem teste
   dedicado."""

import hashlib
import sqlite3
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.consulta import base_local, base_local_sqlite, repositorio_analitico
from app.consulta.repositorio_analitico import FiltrosDiaria, ParametrosOrdenacao, ParametrosPaginacao
from app.consulta.repositorio_diarias import ResultadoDiariasCompletas
from app.domain.entities.diaria import Diaria
from app.domain.entities.processo import Processo, StatusProcesso
from app.domain.entities.registro_diaria_indexada import RegistroDiariaIndexada
from app.services import planilha_service
from app.services.finalizacao_service import ItemConfirmacao, confirmar


def _resultado_sintetico(sei: str) -> ResultadoDiariasCompletas:
    registro = RegistroDiariaIndexada(
        linha=2, sei=sei, doc_origem="PJTESTE", cpf="11111111111",
        periodo_inicial=date(2026, 1, 8), periodo_final=date(2026, 1, 10),
        cidade_origem="RECIFE", cidade_destino="GARANHUNS",
        motivo="teste", transporte="CARRO", valor_total_diarias=100.0, dh=None,
    )
    return ResultadoDiariasCompletas(diarias={sei: registro}, avisos=[], linhas_ignoradas=0)


def _sinal_frescor_dummy(linhas: int = 2) -> base_local_sqlite.SinalFrescor:
    return base_local_sqlite.SinalFrescor(tamanho_bytes=1, mtime=1.0, linhas_controle_diarias=linhas)


def _sei_atual(caminho_sqlite) -> str:
    conexao = sqlite3.connect(caminho_sqlite)
    try:
        return conexao.execute("SELECT sei FROM diarias").fetchone()[0]
    finally:
        conexao.close()


def _assinatura(caminho) -> tuple:
    caminho = Path(caminho)
    estatisticas = caminho.stat()
    return (estatisticas.st_mtime_ns, estatisticas.st_size, hashlib.sha256(caminho.read_bytes()).hexdigest())


def _planilha_com_linha(caminho_planilha_teste, sei: str) -> str:
    caminho = caminho_planilha_teste()
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    planilha["B4"] = sei
    planilha["D4"] = 11111111111
    planilha["H4"] = date(2026, 1, 8)
    planilha["I4"] = date(2026, 1, 10)
    planilha["L4"] = "Motivo de teste"
    workbook.save(caminho)
    workbook.close()
    return caminho


def test_invariante_1_falha_persistente_preserva_ultima_base_integra(tmp_path, monkeypatch):
    destino = tmp_path / "base_local.sqlite"
    base_local_sqlite.reconstruir(
        _resultado_sintetico("ORIGINAL"), {}, _sinal_frescor_dummy(), origem="teste", caminho_destino=destino
    )

    def sempre_falha(origem, dest):
        raise PermissionError(13, "Acesso negado (simulado)")

    monkeypatch.setattr(base_local_sqlite.os, "replace", sempre_falha)

    with pytest.raises(PermissionError):
        base_local_sqlite.reconstruir(
            _resultado_sintetico("NUNCA_PROMOVIDO"), {}, _sinal_frescor_dummy(),
            origem="teste", caminho_destino=destino,
        )

    assert _sei_atual(destino) == "ORIGINAL"


def test_invariante_2_leitura_da_central_nao_altera_planilha_nem_base_local(
    tmp_path, caminho_planilha_conferencia_teste
):
    origem_planilha = caminho_planilha_conferencia_teste()
    destino_base = tmp_path / "dados_usuario"
    base_local.sincronizar(origem_planilha, diretorio_base=destino_base, origem="teste")
    caminho_sqlite = base_local_sqlite.caminho_base_local_sqlite(diretorio_base=destino_base)

    assinatura_planilha_antes = _assinatura(origem_planilha)
    assinatura_sqlite_antes = _assinatura(caminho_sqlite)

    filtros = FiltrosDiaria()
    paginacao = ParametrosPaginacao(pagina=1, tamanho_pagina=10)
    ordenacao = ParametrosOrdenacao(campo="sei", ordem="asc")
    repositorio_analitico.listar_diarias(caminho_sqlite, filtros, paginacao, ordenacao)
    repositorio_analitico.agregar_indicadores(caminho_sqlite, filtros)
    repositorio_analitico.opcoes_filtro(caminho_sqlite)
    repositorio_analitico.ultima_sincronizacao(caminho_sqlite)
    repositorio_analitico.estado_sincronizacao_registrado(caminho_sqlite)

    assert _assinatura(origem_planilha) == assinatura_planilha_antes
    assert _assinatura(caminho_sqlite) == assinatura_sqlite_antes


def test_invariante_3_falha_durante_construcao_nao_promove_base_parcial(tmp_path):
    """Falha ANTES do `commit()` (diferente das falhas em `os.replace`, já
    cobertas noutras suítes): um valor de tipo não suportado pelo driver
    `sqlite3` levanta `ProgrammingError` no meio do `INSERT` — a base
    antiga precisa permanecer intocada e o `.tmp` nunca deve ser
    promovido."""
    destino = tmp_path / "base_local.sqlite"
    base_local_sqlite.reconstruir(
        _resultado_sintetico("ORIGINAL"), {}, _sinal_frescor_dummy(), origem="teste", caminho_destino=destino
    )

    registro_invalido = RegistroDiariaIndexada(
        linha=2, sei="QUEBRA", doc_origem="PJTESTE", cpf="22222222222",
        periodo_inicial=date(2026, 1, 1), periodo_final=date(2026, 1, 2),
        cidade_origem="X", cidade_destino="Y", motivo="m", transporte="CARRO",
        valor_total_diarias=object(),  # type: ignore[arg-type] — tipo não vinculável pelo sqlite3, de propósito
        dh=None,
    )
    resultado_quebrado = ResultadoDiariasCompletas(
        diarias={"QUEBRA": registro_invalido}, avisos=[], linhas_ignoradas=0
    )

    with pytest.raises(sqlite3.ProgrammingError):
        base_local_sqlite.reconstruir(
            resultado_quebrado, {}, _sinal_frescor_dummy(), origem="teste", caminho_destino=destino
        )

    assert _sei_atual(destino) == "ORIGINAL"
    assert destino.with_suffix(destino.suffix + ".tmp").exists()


def test_invariante_7_processamento_e_finalizacao_usam_o_mesmo_mecanismo_de_sincronizacao(
    tmp_path, caminho_planilha_teste, monkeypatch
):
    """Complemento não precisa de um caso à parte aqui: `modo` só muda o
    que `escritor_planilha.preencher()` escreve, nunca qual função de
    sincronização `planilha_service.preencher_planilha` chama — mesmo
    `preencher_planilha` para os dois modos."""
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    chamadas = []
    sincronizar_original = base_local.sincronizar

    def sincronizar_espiao(*args, **kwargs):
        chamadas.append(sincronizar_original)
        return sincronizar_original(*args, **kwargs)

    monkeypatch.setattr(base_local, "sincronizar", sincronizar_espiao)

    caminho1 = caminho_planilha_teste()
    diaria1 = Diaria(
        sei="19.20.0003.0000001/2026-01", doc_origem="PJ", cpf="10101010101",
        periodo_inicial=date(2026, 1, 1), periodo_final=date(2026, 1, 2),
        cidade_origem="A", cidade_destino="B", motivo="m",
    )
    planilha_service.preencher_planilha(
        caminho1, [Processo(arquivo_origem="a.pdf", status=StatusProcesso.SUCESSO, diaria=diaria1)], erros=[]
    )

    sei2 = "19.20.0003.0000002/2026-02"
    caminho2 = _planilha_com_linha(caminho_planilha_teste, sei2)
    confirmar([ItemConfirmacao(arquivo_origem="b.pdf", sei=sei2)], caminho2)

    assert len(chamadas) == 2
    assert all(chamada is sincronizar_original for chamada in chamadas)


def test_invariante_8_finalizar_confirmar_preserva_valor_historico_com_cache_de_formula(
    tmp_path, caminho_planilha_teste, injetar_cache_formula, monkeypatch
):
    """Único fluxo (dos 3 que escrevem na Planilha Oficial) ainda sem
    teste dedicado de preservação de fórmula — processamento/complemento
    já cobertos em `test_integridade_valor_financeiro.py`."""
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    caminho = caminho_planilha_teste()
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    # Linha histórica de SERVIDOR com fórmula, NÃO relacionada ao SEI confirmado.
    planilha["B4"] = "19.20.0004.0000001/2025-01"
    planilha["D4"] = 99999999999
    planilha["H4"] = date(2025, 1, 8)
    planilha["I4"] = date(2025, 1, 10)
    planilha["L4"] = "Motivo histórico"
    planilha["P4"] = "=R4+S4"
    planilha["R4"] = 300.00
    planilha["S4"] = 162.28
    sei_confirmar = "19.20.0004.0000002/2026-02"
    planilha["B5"] = sei_confirmar
    workbook.save(caminho)
    workbook.close()
    injetar_cache_formula(caminho, "P4", 462.28)

    valor_antes = load_workbook(caminho, data_only=True)["Controle de Diárias"]["P4"].value
    assert valor_antes == 462.28

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei_confirmar)], caminho)
    assert resultado.total_finalizado == 1

    valor_depois = load_workbook(caminho, data_only=True)["Controle de Diárias"]["P4"].value
    assert valor_depois == 462.28
