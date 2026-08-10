from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.consulta import base_local
from app.infra.config import configuracao_app as config
from app.main import app

client = TestClient(app)

_CABECALHO_DADOS = [
    "CPF",
    "MATRÍCULA",
    "NOME",
    "PREVIDÊNCIA",
    "CARGO",
    "ÓRGÃO DE ORIGEM",
    "REGIME DE PREVIDÊNCIA",
    "SITUAÇÃO DO SERVIDOR",
]


def _configurar(monkeypatch, tmp_path, caminho_planilha_configuracao_teste, *, linha_dados_existente=None):
    monkeypatch.setenv(config.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario"))
    caminho = caminho_planilha_configuracao_teste()

    if linha_dados_existente is not None:
        workbook = load_workbook(caminho)
        dados = workbook["Dados"]
        dados.append(_CABECALHO_DADOS)
        dados.append(linha_dados_existente)
        workbook.save(caminho)

    resposta = client.post("/api/v1/configuracao", json={"caminho_planilha_oficial": caminho})
    assert resposta.status_code == 200
    return caminho


def test_cadastro_novo_servidor_por_api_conclui_e_grava_na_planilha(
    monkeypatch, tmp_path, caminho_planilha_configuracao_teste
):
    caminho = _configurar(monkeypatch, tmp_path, caminho_planilha_configuracao_teste)

    resposta = client.post(
        "/api/v1/dados/servidores",
        json={
            "cpf": "111.222.333-44",
            "nome": "Novo Servidor de Teste",
            "matricula": "9999",
            "cargo": "Analista Judiciário",
            "previdencia": "INSS",
        },
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["estado"] == "concluido"
    assert corpo["servidor"]["nome"] == "Novo Servidor de Teste"
    assert corpo["servidor"]["cargo"] == "Analista Judiciário"

    workbook = load_workbook(caminho)
    dados = workbook["Dados"]
    assert dados["A2"].value == 11122233344
    assert dados["C2"].value == "Novo Servidor de Teste"


def test_cadastro_cpf_ja_existente_retorna_duplicado(
    monkeypatch, tmp_path, caminho_planilha_configuracao_teste
):
    caminho = _configurar(
        monkeypatch,
        tmp_path,
        caminho_planilha_configuracao_teste,
        linha_dados_existente=[
            12345678901,
            1001,
            "Maria Aparecida Souza",
            "FUNAFIN",
            "Analista Judiciário",
            None,
            "RPPS",
            "Em exercício",
        ],
    )

    resposta = client.post(
        "/api/v1/dados/servidores",
        json={
            "cpf": "123.456.789-01",
            "nome": "Outro Nome",
            "matricula": "1",
            "cargo": "Outro Cargo",
        },
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["estado"] == "duplicado"
    assert corpo["servidor"]["nome"] == "Maria Aparecida Souza"

    # Nenhuma linha nova foi gravada (a original tem só a linha 2).
    workbook = load_workbook(caminho)
    assert workbook["Dados"]["A3"].value is None


def test_cadastro_campo_obrigatorio_ausente_e_rejeitado_pelo_schema(
    monkeypatch, tmp_path, caminho_planilha_configuracao_teste
):
    _configurar(monkeypatch, tmp_path, caminho_planilha_configuracao_teste)

    resposta = client.post(
        "/api/v1/dados/servidores",
        json={"cpf": "111.222.333-44", "nome": "  ", "matricula": "1", "cargo": "Analista"},
    )

    assert resposta.status_code == 422


def test_cadastro_cpf_com_menos_de_11_digitos_e_rejeitado_pelo_schema(
    monkeypatch, tmp_path, caminho_planilha_configuracao_teste
):
    """Ajuste Funcional + UX (item 2) — CPF deixou de ser bloqueado no
    modal, então pode chegar editado/incompleto: mesma regra de 11
    dígitos já usada em todo o projeto (`cpfValido` no frontend),
    defendida agora também no backend."""
    _configurar(monkeypatch, tmp_path, caminho_planilha_configuracao_teste)

    resposta = client.post(
        "/api/v1/dados/servidores",
        json={"cpf": "123.456", "nome": "Fulano", "matricula": "1", "cargo": "Analista"},
    )

    assert resposta.status_code == 422


def test_cadastro_cpf_editado_para_um_ja_existente_retorna_duplicado(
    monkeypatch, tmp_path, caminho_planilha_configuracao_teste
):
    """Item 2 do pedido: se o usuário corrigir o CPF no modal para um que
    já existe, o backend precisa bloquear como duplicidade — a checagem já
    usa o CPF FINAL recebido na requisição, nunca um CPF "original"
    guardado à parte, então isso já funciona sem lógica nova; este teste
    prova o contrato ponta a ponta."""
    caminho = _configurar(
        monkeypatch,
        tmp_path,
        caminho_planilha_configuracao_teste,
        linha_dados_existente=[
            12345678901,
            1001,
            "Maria Aparecida Souza",
            "FUNAFIN",
            "Analista Judiciário",
            None,
            "RPPS",
            "Em exercício",
        ],
    )

    # Simula: usuário pesquisou um CPF (111...), o modal veio com esse
    # valor pré-preenchido, mas o usuário EDITOU para o CPF que na
    # verdade já está cadastrado (123.456.789-01) antes de salvar.
    resposta = client.post(
        "/api/v1/dados/servidores",
        json={"cpf": "123.456.789-01", "nome": "Outro Nome", "matricula": "1", "cargo": "Outro Cargo"},
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["estado"] == "duplicado"
    assert corpo["servidor"]["nome"] == "Maria Aparecida Souza"
    workbook = load_workbook(caminho)
    assert workbook["Dados"]["A3"].value is None


def test_cadastro_sem_planilha_oficial_configurada_retorna_nao_salvo(monkeypatch, tmp_path):
    monkeypatch.setenv(config.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario_vazio"))

    resposta = client.post(
        "/api/v1/dados/servidores",
        json={"cpf": "111.222.333-44", "nome": "X", "matricula": "1", "cargo": "Y"},
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["estado"] == "nao_salvo"
    assert "Nenhuma planilha oficial configurada" in corpo["mensagem"]


def test_opcoes_cadastro_devolve_valores_reais_ja_levantados(
    monkeypatch, tmp_path, caminho_planilha_configuracao_teste
):
    """GET /api/v1/dados/opcoes-cadastro — fonte do combobox/seletor do
    modal de Cadastro de Servidor/Membro (item 8 do pedido): nunca uma
    lista arbitrária, sempre os valores reais da aba `Dados`."""
    _configurar(
        monkeypatch,
        tmp_path,
        caminho_planilha_configuracao_teste,
        linha_dados_existente=[
            12345678901,
            1001,
            "Maria Aparecida Souza",
            "FUNAFIN",
            "Analista Judiciário",
            "Secretaria de Educação",
            "RPPS",
            "Em exercício",
        ],
    )

    resposta = client.get("/api/v1/dados/opcoes-cadastro")

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["cargos"] == ["Analista Judiciário"]
    assert corpo["orgaos_origem"] == ["Secretaria de Educação"]
    assert corpo["previdencias"] == ["FUNAFIN"]
    assert corpo["regimes_previdencia"] == ["RPPS"]
    assert corpo["situacoes"] == ["Em exercício"]


def test_opcoes_cadastro_503_quando_base_local_indisponivel(monkeypatch, tmp_path):
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path / "dados_usuario_sem_opcoes"))

    resposta = client.get("/api/v1/dados/opcoes-cadastro")

    assert resposta.status_code == 503
