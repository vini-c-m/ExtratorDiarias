"""Validação com dois processos reais de referência da funcionalidade
"Finalizar Diárias" (um Recibo, um Anexo Comprovante). Roda só nesta
máquina, contra uma CÓPIA SEGURA da Planilha Oficial real (nunca o arquivo
original) — pulado automaticamente onde a configuração local não estiver
disponível (ex.: CI).

Confronto QA Codex / QA-RT-06 (fechamento) — os dois SEIs de referência
também eram hardcoded (junto dos caminhos, já externalizados numa correção
anterior desta mesma sprint). Passam a vir de
`EXTRATOR_DIARIAS_TESTE_SEIS_REFERENCIA` (variável de ambiente, nunca
commitada, cada desenvolvedor define a sua localmente com os SEIs reais que
tem em `EXTRATOR_DIARIAS_TESTE_LEGADO_DIR`) — mesmo comportamento de
"disponível só nesta máquina" de antes, agora sem nenhum identificador real
versionado. O nome do PDF de cada caso é derivado do SEI (nunca informado à
parte): mesma convenção real de nomenclatura que os PDFs exportados do SEI
já seguem em todo este projeto (ver `_nome_arquivo_pdf_do_sei` abaixo e
`tests/fixtures/gerar_fixtures_sinteticas.py`), então não há como o
identificador e o nome do arquivo divergirem."""

import os
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.finalizacao_service import ItemConfirmacao, StatusFinalizacao, analisar_lote, confirmar

def _caminho_opcional(variavel_ambiente: str) -> Path | None:
    """`None` quando a variável não está definida — nunca `Path("")`, que o
    `pathlib` normaliza silenciosamente para `Path(".")` (o diretório atual
    sempre existe, o que quebraria a condição de skip abaixo — "não
    configurado" precisa continuar distinto de "diretório atual")."""
    valor = os.environ.get(variavel_ambiente)
    return Path(valor) if valor else None


def _nome_arquivo_pdf_do_sei(sei: str) -> str:
    """Mesma convenção real de nomenclatura que os PDFs exportados do SEI já
    seguem em todo este projeto (`SEI_<numero>.pdf`, com "/" e "-"
    trocados por "_" — ver `tests/fixtures/gerar_fixtures_sinteticas.py`,
    que gera fixtures sintéticas com o mesmo padrão): nunca um segundo dado
    configurado à parte que pudesse divergir do SEI."""
    return f"SEI_{sei.replace('/', '_').replace('-', '_')}.pdf"


_LEGADO_DIR = _caminho_opcional("EXTRATOR_DIARIAS_TESTE_LEGADO_DIR")
_PLANILHA_OFICIAL_REAL = _caminho_opcional("EXTRATOR_DIARIAS_TESTE_PLANILHA_OFICIAL_REAL")

_SEIS_REFERENCIA = [
    sei.strip()
    for sei in os.environ.get("EXTRATOR_DIARIAS_TESTE_SEIS_REFERENCIA", "").split(",")
    if sei.strip()
]
_CASOS = [(sei, _nome_arquivo_pdf_do_sei(sei)) for sei in _SEIS_REFERENCIA]

pytestmark = pytest.mark.skipif(
    _LEGADO_DIR is None
    or _PLANILHA_OFICIAL_REAL is None
    or len(_CASOS) != 2
    or not _LEGADO_DIR.is_dir()
    or not _PLANILHA_OFICIAL_REAL.is_file(),
    reason="PDFs/planilha/SEIs reais de referência não disponíveis nesta máquina "
    "(defina EXTRATOR_DIARIAS_TESTE_LEGADO_DIR, "
    "EXTRATOR_DIARIAS_TESTE_PLANILHA_OFICIAL_REAL e "
    "EXTRATOR_DIARIAS_TESTE_SEIS_REFERENCIA (2 SEIs separados por vírgula) "
    "localmente para rodar este teste)",
)


@pytest.fixture
def copia_segura_planilha_oficial(tmp_path):
    """Copia a Planilha Oficial real para uma pasta temporária e zera X/Y só
    das duas linhas de teste na CÓPIA — o arquivo original nunca é aberto em
    modo de escrita por este teste."""
    destino = tmp_path / "planilha_copia.xlsx"
    shutil.copy(_PLANILHA_OFICIAL_REAL, destino)

    workbook = load_workbook(destino)
    planilha = workbook["Controle de Diárias"]
    seis_do_teste = {sei for sei, _arquivo in _CASOS}
    linha_por_sei = {}
    for linha in range(2, planilha.max_row + 1):
        valor = str(planilha[f"B{linha}"].value or "").strip()
        if valor in seis_do_teste:
            linha_por_sei[valor] = linha
            planilha[f"X{linha}"] = None
            planilha[f"Y{linha}"] = None
    workbook.save(destino)

    assert set(linha_por_sei) == seis_do_teste, "SEIs de referência não encontrados na planilha real"
    return str(destino), linha_por_sei


def test_dois_processos_reais_sao_identificados_como_aptos(copia_segura_planilha_oficial):
    caminho_planilha, _linha_por_sei = copia_segura_planilha_oficial
    caminhos_pdf = [str(_LEGADO_DIR / arquivo) for _sei, arquivo in _CASOS]

    itens = analisar_lote(caminhos_pdf, caminho_planilha)

    assert len(itens) == 2
    for item in itens:
        assert item.avaliacao is not None
        assert item.avaliacao.confirmacao_valida, item.sei
        assert item.avaliacao.comprovante_encontrado, item.sei
        assert item.status is StatusFinalizacao.PRONTO, (item.sei, item.status, item.motivo)
        assert len(item.candidatos) == 1


def test_dois_processos_reais_confirmacao_grava_x_e_y_sem_alterar_o_resto(
    copia_segura_planilha_oficial,
):
    caminho_planilha, linha_por_sei = copia_segura_planilha_oficial
    caminhos_pdf = [str(_LEGADO_DIR / arquivo) for _sei, arquivo in _CASOS]

    antes = load_workbook(caminho_planilha)["Controle de Diárias"]
    instantaneos_antes = {
        sei: [antes.cell(row=linha, column=c).value for c in range(1, 24)]
        for sei, linha in linha_por_sei.items()
    }

    itens = analisar_lote(caminhos_pdf, caminho_planilha)
    confirmacoes = [
        ItemConfirmacao(arquivo_origem=item.arquivo_origem, sei=item.sei) for item in itens
    ]
    resultado = confirmar(confirmacoes, caminho_planilha)

    assert resultado.total_finalizado == 2
    assert resultado.total_ignorado == 0

    depois = load_workbook(caminho_planilha)["Controle de Diárias"]
    for sei, linha in linha_por_sei.items():
        assert depois[f"X{linha}"].value == "SIM", sei
        assert depois[f"Y{linha}"].value == "SIM", sei
        # Nenhuma outra coluna (A..W, colunas 1-23) foi alterada.
        instantaneo_depois = [depois.cell(row=linha, column=c).value for c in range(1, 24)]
        assert instantaneo_depois == instantaneos_antes[sei], sei
