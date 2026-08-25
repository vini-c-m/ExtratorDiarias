"""Micro-sprint 3D — testes de API do Relatório de Diárias — Portal da
Transparência. Dados 100% sintéticos (identificadores fictícios, sem
nenhuma base real)."""

import os
import threading
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app

client = TestClient(app)

_BASE_URL = "/api/v1/relatorio-transparencia"

_CABECALHO_BASE_OB = [
    "DH",
    "OB",
    "SITUAÇÃO",
    "DATA OB",
    "SEI",
    "CPF",
    "NOME",
    "CARGO",
    "PERÍODO INICIAL",
    "PERÍODO FINAL",
    "TRECHO",
    "TRANSPORTE",
    "MOTIVO",
    "VALOR DA DIÁRIA",
    "RESTITUIÇÃO",
    "DUPLICIDADE",
    "ITEM DE GASTO",
]


def _linha_ob(
    *,
    dh,
    ob,
    cpf=11111111111,
    periodo_inicial=date(2026, 5, 10),
    periodo_final=date(2026, 5, 10),
    trecho="RECIFE - GARANHUNS",
    valor=154.09,
    data_ob="15/05/2026 09:00:00",
    motivo="Motivo de teste",
):
    return [
        dh, ob, "Paga", data_ob, "19.20.0001.0000001/2026-01", cpf, "FULANO DE TAL",
        "Cargo de Teste", periodo_inicial, periodo_final, trecho, "CARRO", motivo,
        valor, None, "Único", "33901401",
    ]


def _criar_base_ob(caminho, linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(_CABECALHO_BASE_OB)
    for linha in linhas:
        ws.append(linha)
    wb.save(caminho)


def _criar_controle_vazio(caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = "Controle de Diárias"
    wb.save(caminho)


def _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes, mes_referencia=5):
    return {
        "caminho_base_ordens_bancarias": str(caminho_base_ob),
        "caminho_controle_dimcac": str(caminho_controle),
        "ano_exercicio": 2026,
        "mes_referencia": mes_referencia,
        "caminho_decisoes": str(caminho_decisoes),
    }


def _preparar_fontes(tmp_path, linhas):
    caminho_base_ob = tmp_path / "base_ob.xlsx"
    caminho_controle = tmp_path / "controle.xlsx"
    _criar_base_ob(caminho_base_ob, linhas)
    _criar_controle_vazio(caminho_controle)
    caminho_decisoes = tmp_path / "decisoes.json"
    return caminho_base_ob, caminho_controle, caminho_decisoes


# --- Contrato básico: POST /preparar ----------------------------------------


def test_preparar_via_api_sem_revisao(tmp_path):
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(
        tmp_path, [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    )

    resposta = client.post(f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes))

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_candidatos_normais"] == 1
    assert corpo["total_revisoes"] == 0
    assert corpo["pode_gerar_relatorio"] is True


def test_preparar_via_api_com_revisao_nao_expoe_fingerprint(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19, trecho="RECIFE - ARCOVERDE", periodo_inicial=date(2026, 5, 11), periodo_final=date(2026, 5, 11)),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19, trecho="RECIFE - ARCOVERDE", periodo_inicial=date(2026, 5, 11), periodo_final=date(2026, 5, 11)),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)

    resposta = client.post(f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes))

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_revisoes"] == 1
    assert corpo["pode_gerar_relatorio"] is False
    texto_bruto = resposta.text.lower()
    assert "fingerprint" not in texto_bruto


# --- GET /revisoes -----------------------------------------------------------


def test_get_revisoes_via_api(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)

    resposta = client.get(f"{_BASE_URL}/revisoes", params=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes))

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total"] == 1
    assert corpo["revisoes"][0]["regra"] == "RN07"
    assert corpo["revisoes"][0]["estado"] == "pendente"


# --- Erros seguem o padrão da aplicação (item 28) ---------------------------


def test_base_ob_inexistente_devolve_404(tmp_path):
    caminho_controle = tmp_path / "controle.xlsx"
    _criar_controle_vazio(caminho_controle)

    resposta = client.post(
        f"{_BASE_URL}/preparar",
        json=_payload_base(tmp_path / "nao_existe.xlsx", caminho_controle, tmp_path / "decisoes.json"),
    )
    assert resposta.status_code == 404


def test_controle_inexistente_devolve_404(tmp_path):
    caminho_base_ob = tmp_path / "base_ob.xlsx"
    _criar_base_ob(caminho_base_ob, [_linha_ob(dh="2026DH000001", ob="2026OB000001")])

    resposta = client.post(
        f"{_BASE_URL}/preparar",
        json=_payload_base(caminho_base_ob, tmp_path / "nao_existe.xlsx", tmp_path / "decisoes.json"),
    )
    assert resposta.status_code == 404


def test_mes_referencia_invalido_devolve_422(tmp_path):
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(
        tmp_path, [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    )
    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload["mes_referencia"] = 13

    resposta = client.post(f"{_BASE_URL}/preparar", json=payload)
    assert resposta.status_code == 422


def test_acao_desconhecida_devolve_422(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update({"acao": "aprovar_tudo", "autor": "servidor.teste", "obs_selecionadas": []})

    resposta = client.post(f"{_BASE_URL}/revisoes/qualquer-identidade/decisao", json=payload)
    assert resposta.status_code == 422


def test_grupo_inexistente_devolve_404(tmp_path):
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(
        tmp_path, [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    )
    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update({"acao": "publicar", "autor": "servidor.teste", "obs_selecionadas": ["2026OB000001"]})

    resposta = client.post(f"{_BASE_URL}/revisoes/grupo-inexistente/decisao", json=payload)
    assert resposta.status_code == 404


def test_ob_fora_do_grupo_devolve_422(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000001", ob="2026OB000001"),
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    corpo_preparar = client.post(
        f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    ).json()
    identidade = corpo_preparar["revisoes"][0]["identidade_grupo"]

    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update({"acao": "publicar", "autor": "servidor.teste", "obs_selecionadas": ["2026OB000001"]})

    resposta = client.post(f"{_BASE_URL}/revisoes/{identidade}/decisao", json=payload)
    assert resposta.status_code == 422


def test_autor_vazio_devolve_422(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update({"acao": "publicar", "autor": "  ", "obs_selecionadas": ["2026OB000002"]})

    resposta = client.post(f"{_BASE_URL}/revisoes/qualquer/decisao", json=payload)
    assert resposta.status_code == 422


# --- Micro-sprint 3E.2 — PUBLICAR com zero OBs exige justificativa --------
# Prova que a garantia vale também batendo direto na API (sem passar pelo
# frontend/React) — nenhuma UI pode contornar.


def test_publicar_zero_obs_sem_observacao_devolve_422(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    corpo_preparar = client.post(
        f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    ).json()
    identidade = corpo_preparar["revisoes"][0]["identidade_grupo"]

    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update({"acao": "publicar", "autor": "servidor.teste", "obs_selecionadas": []})

    resposta = client.post(f"{_BASE_URL}/revisoes/{identidade}/decisao", json=payload)
    assert resposta.status_code == 422


def test_publicar_zero_obs_com_observacao_vazia_devolve_422(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    corpo_preparar = client.post(
        f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    ).json()
    identidade = corpo_preparar["revisoes"][0]["identidade_grupo"]

    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update(
        {"acao": "publicar", "autor": "servidor.teste", "obs_selecionadas": [], "observacao": "   "}
    )

    resposta = client.post(f"{_BASE_URL}/revisoes/{identidade}/decisao", json=payload)
    assert resposta.status_code == 422


def test_publicar_zero_obs_com_justificativa_valida_devolve_200(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    corpo_preparar = client.post(
        f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    ).json()
    identidade = corpo_preparar["revisoes"][0]["identidade_grupo"]

    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update(
        {
            "acao": "publicar",
            "autor": "servidor.teste",
            "obs_selecionadas": [],
            "observacao": "Ambas as OBs excluídas até apuração interna.",
        }
    )

    resposta = client.post(f"{_BASE_URL}/revisoes/{identidade}/decisao", json=payload)
    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["revisoes"][0]["estado"] == "decidida"
    assert corpo["pode_gerar_relatorio"] is True


def test_publicar_com_obs_selecionadas_continua_sem_exigir_observacao(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    corpo_preparar = client.post(
        f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    ).json()
    identidade = corpo_preparar["revisoes"][0]["identidade_grupo"]

    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update(
        {"acao": "publicar", "autor": "servidor.teste", "obs_selecionadas": ["2026OB000002"]}
    )

    resposta = client.post(f"{_BASE_URL}/revisoes/{identidade}/decisao", json=payload)
    assert resposta.status_code == 200


def test_adiar_continua_sem_exigir_observacao(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    corpo_preparar = client.post(
        f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    ).json()
    identidade = corpo_preparar["revisoes"][0]["identidade_grupo"]

    payload = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    payload.update({"acao": "adiar", "autor": "servidor.teste", "obs_selecionadas": []})

    resposta = client.post(f"{_BASE_URL}/revisoes/{identidade}/decisao", json=payload)
    assert resposta.status_code == 200
    assert resposta.json()["revisoes"][0]["estado"] == "adiada"


def test_persistencia_corrompida_devolve_503(tmp_path):
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(
        tmp_path, [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    )
    caminho_decisoes.write_text("{isso nao e json valido", encoding="utf-8")

    resposta = client.post(f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes))
    assert resposta.status_code == 503


# --- Teste ponta a ponta do backend (§13 da especificação) -----------------


def test_ponta_a_ponta_preparar_decidir_reaproveitar_invalidar(tmp_path):
    linhas = [
        _linha_ob(
            dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19,
            trecho="RECIFE - ARCOVERDE", periodo_inicial=date(2026, 5, 11), periodo_final=date(2026, 5, 11),
        ),
        _linha_ob(
            dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19,
            trecho="RECIFE - ARCOVERDE", periodo_inicial=date(2026, 5, 11), periodo_final=date(2026, 5, 11),
        ),
        _linha_ob(
            dh="2026DH000004", ob="2026OB000004", cpf=33333333333, valor=200.00,
            trecho="RECIFE - CARUARU", periodo_inicial=date(2026, 5, 12), periodo_final=date(2026, 5, 12),
        ),
        _linha_ob(
            dh="2026DH000005", ob="2026OB000005", cpf=33333333333, valor=200.00,
            trecho="RECIFE - CARUARU", periodo_inicial=date(2026, 5, 12), periodo_final=date(2026, 5, 12),
        ),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    payload_fontes = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)

    # PREPARAR -> existem 2 revisões, pode_gerar=False
    #
    # NÃO presumir que `revisoes[0]` é sempre o grupo OB2/OB3: a ordem de
    # `ResultadoReconciliacao.documentos` depende de `set(...) | set(...)`
    # (`motor_reconciliacao.reconciliar`), cuja iteração varia entre
    # processos por causa da randomização de hash de string do Python —
    # identificar cada grupo pelas OBs que ele de fato contém, nunca por
    # posição na lista.
    corpo = client.post(f"{_BASE_URL}/preparar", json=payload_fontes).json()
    assert corpo["total_revisoes"] == 2
    assert corpo["pode_gerar_relatorio"] is False

    def _revisao_do_grupo(corpo, obs_esperadas):
        return next(r for r in corpo["revisoes"] if {oc["ob"] for oc in r["ocorrencias"]} == obs_esperadas)

    obs_grupo_1 = {"2026OB000002", "2026OB000003"}
    obs_grupo_2 = {"2026OB000004", "2026OB000005"}
    identidade_1 = _revisao_do_grupo(corpo, obs_grupo_1)["identidade_grupo"]
    identidade_2 = _revisao_do_grupo(corpo, obs_grupo_2)["identidade_grupo"]

    # DECIDIR revisão 1 -> resta 1, pode_gerar=False
    payload_decisao_1 = dict(payload_fontes, acao="publicar", autor="servidor.teste", obs_selecionadas=sorted(obs_grupo_1))
    corpo = client.post(f"{_BASE_URL}/revisoes/{identidade_1}/decisao", json=payload_decisao_1).json()
    assert corpo["pode_gerar_relatorio"] is False
    estados = {r["identidade_grupo"]: r["estado"] for r in corpo["revisoes"]}
    assert estados[identidade_1] == "decidida"
    assert estados[identidade_2] == "pendente"

    # DECIDIR revisão 2 -> nenhuma pendência, pode_gerar=True
    payload_decisao_2 = dict(payload_fontes, acao="publicar", autor="servidor.teste", obs_selecionadas=sorted(obs_grupo_2))
    corpo = client.post(f"{_BASE_URL}/revisoes/{identidade_2}/decisao", json=payload_decisao_2).json()
    assert corpo["pode_gerar_relatorio"] is True

    # PREPARAR NOVAMENTE ("mês seguinte") -> decisões anteriores reaproveitadas
    payload_fontes_mes_seguinte = dict(payload_fontes, mes_referencia=8)
    corpo = client.post(f"{_BASE_URL}/preparar", json=payload_fontes_mes_seguinte).json()
    assert corpo["pode_gerar_relatorio"] is True
    assert all(r["estado"] == "decidida" for r in corpo["revisoes"])

    # Alterar materialmente uma evidência (motivo do grupo 1) -> decisão
    # correspondente vira DESATUALIZADA -> pode_gerar=False
    linhas_alteradas = list(linhas)
    linhas_alteradas[0] = _linha_ob(
        dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19,
        trecho="RECIFE - ARCOVERDE", periodo_inicial=date(2026, 5, 11), periodo_final=date(2026, 5, 11),
        motivo="Motivo revisado após o fechamento",
    )
    _criar_base_ob(caminho_base_ob, linhas_alteradas)

    corpo = client.post(f"{_BASE_URL}/preparar", json=payload_fontes_mes_seguinte).json()
    estados = {r["identidade_grupo"]: r["estado"] for r in corpo["revisoes"]}
    assert estados[identidade_1] == "desatualizada"
    assert estados[identidade_2] == "decidida"


# --- Micro-sprint 3F: POST /gerar — endpoint aditivo, não conectado ao frontend


def _configurar_template(tmp_path, nome_arquivo="template.xlsx"):
    """Micro-sprint 3H — `POST /template` configura o Template Oficial uma
    vez (via a própria API, não um atalho direto ao módulo de config) e
    devolve `caminho_configuracao_template`, pronto para reuso em qualquer
    chamada de `/gerar` do teste."""
    from tests.fixtures.template_relatorio_transparencia_sintetico import construir_template_sintetico

    caminho_template = tmp_path / nome_arquivo
    construir_template_sintetico(str(caminho_template))
    caminho_configuracao_template = str(tmp_path / "template_configuracao.json")

    resposta = client.post(
        f"{_BASE_URL}/template",
        json={
            "caminho_configuracao_template": caminho_configuracao_template,
            "caminho_template": str(caminho_template),
        },
    )
    assert resposta.status_code == 200

    return caminho_configuracao_template


def test_post_gerar_produz_arquivo(tmp_path):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = tmp_path / "saida.xlsx"

    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=str(caminho_destino),
    )
    resposta = client.post(f"{_BASE_URL}/gerar", json=payload)

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_ocorrencias"] == 1
    assert caminho_destino.exists()

    # Item 16 (teste ponta a ponta) — confirma que o arquivo tem a
    # estrutura mínima esperada: mesma aba do template, cabeçalho
    # preservado, e a linha de dados escrita de fato.
    from openpyxl import load_workbook

    workbook = load_workbook(str(caminho_destino))
    assert "Controle de Diárias" in workbook.sheetnames
    ws = workbook["Controle de Diárias"]
    assert ws["A7"].value == "2026OB000001"


def test_post_gerar_bloqueia_com_revisao_pendente_409(tmp_path):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = tmp_path / "saida.xlsx"

    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=str(caminho_destino),
    )
    resposta = client.post(f"{_BASE_URL}/gerar", json=payload)

    assert resposta.status_code == 409
    assert not caminho_destino.exists()


def test_post_gerar_nao_sobrescreve_destino_existente_409(tmp_path):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = tmp_path / "saida.xlsx"
    caminho_destino.write_bytes(b"conteudo previo")

    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=str(caminho_destino),
    )
    resposta = client.post(f"{_BASE_URL}/gerar", json=payload)

    assert resposta.status_code == 409
    assert caminho_destino.read_bytes() == b"conteudo previo"


def test_post_gerar_duas_requisicoes_http_simultaneas_mesmo_destino_apenas_uma_200(tmp_path):
    """Confronto QA / homologação operacional — reproduz exatamente o
    cenário original (2 requisições `/gerar` simultâneas para o mesmo
    `caminho_destino`) na própria camada HTTP, não só na função interna do
    escritor. Antes da correção: 200 + 200, a segunda sobrescrevendo a
    primeira silenciosamente. Esperado agora: exatamente um 200, o outro
    409 — nunca as duas concluindo."""
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = tmp_path / "saida_concorrente.xlsx"

    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=str(caminho_destino),
    )

    liberar_barreira = threading.Barrier(2, timeout=10)
    respostas: dict[str, object] = {}

    def _chamar(chave):
        liberar_barreira.wait()
        respostas[chave] = client.post(f"{_BASE_URL}/gerar", json=payload)

    t1 = threading.Thread(target=_chamar, args=("a",))
    t2 = threading.Thread(target=_chamar, args=("b",))
    t1.start()
    t2.start()
    t1.join(timeout=30)
    t2.join(timeout=30)

    status = sorted(r.status_code for r in respostas.values())
    assert status == [200, 409], f"esperado [200, 409], obtido {status}"
    assert caminho_destino.exists()


def test_post_gerar_sem_template_configurado_404(tmp_path):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    caminho_configuracao_template = str(tmp_path / "template_configuracao_nunca_criado.json")
    caminho_destino = tmp_path / "saida.xlsx"

    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=str(caminho_destino),
    )
    resposta = client.post(f"{_BASE_URL}/gerar", json=payload)

    assert resposta.status_code == 404
    assert not caminho_destino.exists()
    # Nunca expõe caminho interno na mensagem (item 12/privacidade).
    assert str(tmp_path) not in resposta.json()["detail"]


def test_post_gerar_template_configurado_mas_arquivo_sumiu_404(tmp_path):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    caminho_configuracao_template = _configurar_template(tmp_path)

    caminho_template_removido = tmp_path / "template.xlsx"
    assert caminho_template_removido.exists()
    caminho_template_removido.unlink()

    caminho_destino = tmp_path / "saida.xlsx"
    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=str(caminho_destino),
    )
    resposta = client.post(f"{_BASE_URL}/gerar", json=payload)

    assert resposta.status_code == 404
    assert not caminho_destino.exists()


# --- Micro-sprint 3H: GET/POST /template — configuração do Template Oficial


def test_get_template_nao_configurado(tmp_path):
    caminho_configuracao_template = str(tmp_path / "nunca_configurado.json")

    resposta = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": caminho_configuracao_template}
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo == {"configurado": False, "disponivel": False, "caminho_template": None}


def test_post_template_configura_e_get_reflete(tmp_path):
    caminho_configuracao_template = _configurar_template(tmp_path)

    resposta = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": caminho_configuracao_template}
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["configurado"] is True
    assert corpo["disponivel"] is True
    assert corpo["caminho_template"] == str(tmp_path / "template.xlsx")


def test_post_template_arquivo_invalido_422(tmp_path):
    caminho_configuracao_template = str(tmp_path / "template_configuracao.json")
    caminho_invalido = tmp_path / "nao_existe.xlsx"

    resposta = client.post(
        f"{_BASE_URL}/template",
        json={
            "caminho_configuracao_template": caminho_configuracao_template,
            "caminho_template": str(caminho_invalido),
        },
    )

    assert resposta.status_code == 422
    # Nada persistido quando a validação falha.
    assert not os.path.exists(caminho_configuracao_template)


def test_get_template_configurado_mas_arquivo_sumiu(tmp_path):
    caminho_configuracao_template = _configurar_template(tmp_path)
    (tmp_path / "template.xlsx").unlink()

    resposta = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": caminho_configuracao_template}
    )

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo == {"configurado": True, "disponivel": False, "caminho_template": None}


# --- Confronto QA Codex / QA-RT-03: configuração de Template Oficial
# corrompida vira erro de domínio controlado (503, mensagem amigável),
# nunca HTTP 500 genérico. Achado real (reproduzido antes da correção):
# `JSONDecodeError`/`KeyError` escapavam de `carregar_configuracao_template`
# até o middleware padrão do FastAPI, que devolvia 500 sem nenhuma
# orientação para o usuário.


def test_get_template_json_corrompido_503_amigavel(tmp_path):
    caminho_configuracao_template = tmp_path / "corrompido.json"
    caminho_configuracao_template.write_text("{ isso nao e json valido ][", encoding="utf-8")

    resposta = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": str(caminho_configuracao_template)}
    )

    assert resposta.status_code == 503
    corpo = resposta.json()
    assert "detail" in corpo
    # Nunca o nome cru da exceção Python nem um traceback.
    assert "Traceback" not in corpo["detail"]
    assert "JSONDecodeError" not in corpo["detail"]


def test_get_template_estrutura_incompleta_503_amigavel(tmp_path):
    caminho_configuracao_template = tmp_path / "incompleto.json"
    caminho_configuracao_template.write_text('{"outra_coisa": 1}', encoding="utf-8")

    resposta = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": str(caminho_configuracao_template)}
    )

    assert resposta.status_code == 503
    corpo = resposta.json()
    assert "KeyError" not in corpo["detail"]


def test_post_gerar_com_configuracao_template_corrompida_503_nunca_500(tmp_path):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    caminho_configuracao_template = tmp_path / "corrompido.json"
    caminho_configuracao_template.write_text("{ isso nao e json valido ][", encoding="utf-8")
    caminho_destino = tmp_path / "saida.xlsx"

    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=str(caminho_configuracao_template),
        caminho_destino=str(caminho_destino),
    )
    resposta = client.post(f"{_BASE_URL}/gerar", json=payload)

    assert resposta.status_code == 503
    assert not caminho_destino.exists()


def test_post_template_reconfigurar_corrige_configuracao_corrompida(tmp_path):
    """Recuperação natural: selecionar o template de novo sobrescreve o
    JSON corrompido (o `POST /template` nunca lê o arquivo anterior)."""
    caminho_configuracao_template = tmp_path / "corrompido.json"
    caminho_configuracao_template.write_text("{ isso nao e json valido ][", encoding="utf-8")

    resposta_corrompida = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": str(caminho_configuracao_template)}
    )
    assert resposta_corrompida.status_code == 503

    from tests.fixtures.template_relatorio_transparencia_sintetico import construir_template_sintetico

    caminho_template = tmp_path / "novo_template.xlsx"
    construir_template_sintetico(str(caminho_template))

    resposta_config = client.post(
        f"{_BASE_URL}/template",
        json={
            "caminho_configuracao_template": str(caminho_configuracao_template),
            "caminho_template": str(caminho_template),
        },
    )
    assert resposta_config.status_code == 200

    resposta_get = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": str(caminho_configuracao_template)}
    )
    assert resposta_get.status_code == 200
    assert resposta_get.json()["configurado"] is True


# --- Confronto QA Codex / revalidação final — QA-RT-03 incompleto: bytes
# inválidos em UTF-8 (achado real do Codex) também precisam virar 503
# amigável, nunca 500 genérico. As mesmas checagens de status/mensagem já
# feitas acima para JSON malformado, agora para o novo caso.


def test_get_template_bytes_utf8_invalidos_503_amigavel(tmp_path):
    caminho_configuracao_template = tmp_path / "bytes_invalidos.json"
    caminho_configuracao_template.write_bytes(b"\x80\x81\x82\x83 nao decodifica como utf-8")

    resposta = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": str(caminho_configuracao_template)}
    )

    assert resposta.status_code == 503
    corpo = resposta.json()
    assert "detail" in corpo
    assert "Traceback" not in corpo["detail"]
    assert "UnicodeDecodeError" not in corpo["detail"]


def test_post_gerar_com_configuracao_bytes_utf8_invalidos_503_nunca_500(tmp_path):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    caminho_configuracao_template = tmp_path / "bytes_invalidos.json"
    caminho_configuracao_template.write_bytes(b"\x80\x81\x82\x83 nao decodifica como utf-8")
    caminho_destino = tmp_path / "saida.xlsx"

    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=str(caminho_configuracao_template),
        caminho_destino=str(caminho_destino),
    )
    resposta = client.post(f"{_BASE_URL}/gerar", json=payload)

    assert resposta.status_code == 503
    assert not caminho_destino.exists()


def test_post_template_reconfigurar_corrige_configuracao_com_bytes_invalidos(tmp_path):
    """Mesma recuperação natural de `test_post_template_reconfigurar_corrige_configuracao_corrompida`,
    agora para o caso de bytes inválidos em UTF-8."""
    caminho_configuracao_template = tmp_path / "bytes_invalidos.json"
    caminho_configuracao_template.write_bytes(b"\x80\x81\x82\x83 nao decodifica como utf-8")

    resposta_corrompida = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": str(caminho_configuracao_template)}
    )
    assert resposta_corrompida.status_code == 503

    from tests.fixtures.template_relatorio_transparencia_sintetico import construir_template_sintetico

    caminho_template = tmp_path / "novo_template.xlsx"
    construir_template_sintetico(str(caminho_template))

    resposta_config = client.post(
        f"{_BASE_URL}/template",
        json={
            "caminho_configuracao_template": str(caminho_configuracao_template),
            "caminho_template": str(caminho_template),
        },
    )
    assert resposta_config.status_code == 200

    resposta_get = client.get(
        f"{_BASE_URL}/template", params={"caminho_configuracao_template": str(caminho_configuracao_template)}
    )
    assert resposta_get.status_code == 200
    assert resposta_get.json()["configurado"] is True


# --- Confronto QA Codex — arquivo de DECISÕES corrompido/ilegível também
# precisa virar 503 amigável em TODOS os endpoints que o leem (nunca 500
# genérico). Achado real: `UnicodeDecodeError` (e, por extensão,
# `PermissionError`/`OSError` — já cobertos no nível unitário de
# `test_persistencia_decisoes.py`) escapavam de `carregar_decisoes`.


def _decisoes_json_invalido(tmp_path) -> str:
    caminho = tmp_path / "decisoes_corrompidas.json"
    caminho.write_text("{ isso nao e json valido ][", encoding="utf-8")
    return str(caminho)


def _decisoes_bytes_invalidos(tmp_path) -> str:
    caminho = tmp_path / "decisoes_bytes_invalidos.json"
    caminho.write_bytes(b"\x80\x81\x82\x83 nao decodifica como utf-8")
    return str(caminho)


def _assert_503_amigavel(resposta):
    assert resposta.status_code == 503
    corpo = resposta.json()
    assert "detail" in corpo
    assert "Traceback" not in corpo["detail"]
    assert "UnicodeDecodeError" not in corpo["detail"]
    assert "JSONDecodeError" not in corpo["detail"]


@pytest.mark.parametrize("fabrica_decisoes", [_decisoes_json_invalido, _decisoes_bytes_invalidos])
def test_post_preparar_com_decisoes_corrompidas_503_amigavel(tmp_path, fabrica_decisoes):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, _ = _preparar_fontes(tmp_path, linhas)
    caminho_decisoes = fabrica_decisoes(tmp_path)

    resposta = client.post(
        f"{_BASE_URL}/preparar", json=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    )
    _assert_503_amigavel(resposta)


@pytest.mark.parametrize("fabrica_decisoes", [_decisoes_json_invalido, _decisoes_bytes_invalidos])
def test_get_revisoes_com_decisoes_corrompidas_503_amigavel(tmp_path, fabrica_decisoes):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, _ = _preparar_fontes(tmp_path, linhas)
    caminho_decisoes = fabrica_decisoes(tmp_path)

    resposta = client.get(
        f"{_BASE_URL}/revisoes", params=_payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)
    )
    _assert_503_amigavel(resposta)


@pytest.mark.parametrize("fabrica_decisoes", [_decisoes_json_invalido, _decisoes_bytes_invalidos])
def test_registrar_decisao_com_decisoes_corrompidas_503_amigavel(tmp_path, fabrica_decisoes):
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, _ = _preparar_fontes(tmp_path, linhas)
    caminho_decisoes = fabrica_decisoes(tmp_path)
    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        acao="publicar", autor="servidor.teste", obs_selecionadas=["2026OB000002"],
    )

    resposta = client.post(f"{_BASE_URL}/revisoes/qualquer-identidade/decisao", json=payload)
    _assert_503_amigavel(resposta)


@pytest.mark.parametrize("fabrica_decisoes", [_decisoes_json_invalido, _decisoes_bytes_invalidos])
def test_post_gerar_com_decisoes_corrompidas_503_nunca_500(tmp_path, fabrica_decisoes):
    linhas = [_linha_ob(dh="2026DH000001", ob="2026OB000001")]
    caminho_base_ob, caminho_controle, _ = _preparar_fontes(tmp_path, linhas)
    caminho_decisoes = fabrica_decisoes(tmp_path)
    caminho_configuracao_template = _configurar_template(tmp_path)
    caminho_destino = tmp_path / "saida.xlsx"

    payload = dict(
        _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes),
        caminho_configuracao_template=caminho_configuracao_template,
        caminho_destino=str(caminho_destino),
    )
    resposta = client.post(f"{_BASE_URL}/gerar", json=payload)

    _assert_503_amigavel(resposta)
    assert not caminho_destino.exists()


def test_decisoes_validas_continuam_funcionando_apos_a_correcao(tmp_path):
    """Controle — a correção não pode quebrar o caminho feliz já
    homologado: preparar, decidir e listar continuam funcionando com um
    arquivo de decisões válido."""
    linhas = [
        _linha_ob(dh="2026DH000002", ob="2026OB000002", cpf=22222222222, valor=308.19),
        _linha_ob(dh="2026DH000003", ob="2026OB000003", cpf=22222222222, valor=308.19),
    ]
    caminho_base_ob, caminho_controle, caminho_decisoes = _preparar_fontes(tmp_path, linhas)
    payload_fontes = _payload_base(caminho_base_ob, caminho_controle, caminho_decisoes)

    corpo = client.post(f"{_BASE_URL}/preparar", json=payload_fontes).json()
    identidade = corpo["revisoes"][0]["identidade_grupo"]

    payload_decisao = dict(
        payload_fontes, acao="publicar", autor="servidor.teste", obs_selecionadas=["2026OB000002", "2026OB000003"]
    )
    resposta = client.post(f"{_BASE_URL}/revisoes/{identidade}/decisao", json=payload_decisao)
    assert resposta.status_code == 200
    assert resposta.json()["pode_gerar_relatorio"] is True

    resposta_revisoes = client.get(f"{_BASE_URL}/revisoes", params=payload_fontes)
    assert resposta_revisoes.status_code == 200
    assert resposta_revisoes.json()["revisoes"][0]["estado"] == "decidida"
