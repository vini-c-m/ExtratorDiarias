"""Sprint 04 — Diária de Membro.

Cobre o critério de aceite: os 3 processos reais de referência são
identificados como Diária de Membro, o despacho GABPGJ é localizado, o valor
é extraído e a coluna P é preenchida — enquanto os processos Padrão já
homologados continuam se comportando exatamente como antes (sem tipo
Membro, sem valor, sem tocar a coluna P)."""

import os

from openpyxl import Workbook, load_workbook

from app.domain.entities.tipo_processo import TipoProcesso
from app.services.planilha_service import preencher_planilha
from app.services.processamento_service import processar_lote

_ESPERADOS = {
    "SEI_19.20.0001.0000114_2026_14.pdf": {
        "sei": "19.20.0001.0000114/2026-14",
        "cpf": "22222222201",
        "valor": "652,96",
    },
    "SEI_19.20.0001.0000115_2026_15.pdf": {
        "sei": "19.20.0001.0000115/2026-15",
        "cpf": "22222222202",
        "valor": "1.069,14",
    },
    "SEI_19.20.0001.0000916_2025_16.pdf": {
        "sei": "19.20.0001.0000916/2025-16",
        "cpf": "22222222203",
        "valor": "1.069,14",
    },
}


def _por_arquivo(lote, nome_arquivo):
    return next(p for p in lote.processos if p.arquivo_origem == nome_arquivo)


def test_os_3_processos_de_referencia_sao_diaria_de_membro_com_valor(pdfs_diaria_membro):
    lote = processar_lote(pdfs_diaria_membro)

    assert lote.sucesso == len(pdfs_diaria_membro)
    assert lote.erro == 0

    for processo in lote.processos:
        esperado = _ESPERADOS[processo.arquivo_origem]
        diaria = processo.diaria

        assert diaria.tipo_processo is TipoProcesso.DIARIA_MEMBRO
        assert diaria.valor_total_diarias == esperado["valor"]

        # Demais campos continuam usando exatamente a lógica já homologada —
        # SEI, CPF e datas seguem extraídos e validados normalmente.
        assert diaria.sei == esperado["sei"]
        assert diaria.cpf == esperado["cpf"]
        assert diaria.periodo_inicial is not None
        assert diaria.periodo_final is not None


def test_processos_padrao_nao_ganham_tipo_membro_nem_valor(pdfs_exemplo):
    """Nenhuma regressão: todos os PDFs já homologados continuam com tipo
    Padrão e sem valor de diárias — exatamente o comportamento anterior a
    esta Sprint."""
    lote = processar_lote(pdfs_exemplo)

    assert lote.sucesso == len(pdfs_exemplo)
    for processo in lote.processos:
        assert processo.diaria.tipo_processo is TipoProcesso.PADRAO
        assert processo.diaria.valor_total_diarias is None


def _planilha_vazia(caminho) -> None:
    """P0 FINAL (2026-08-09) — "vazia" quer dizer sem nenhuma diária ainda,
    não sem nenhuma linha: a preparação estrutural automática (ver
    `escritor_planilha._preparar_estrutura_futura`) exige uma linha-modelo
    íntegra (N/P/R/S/V/Z/AA em fórmula) acima da primeira linha livre para
    QUALQUER gravação, em qualquer modo — sem isso, `preencher()` aborta
    antes de gravar (comportamento novo e intencional, não um bug deste
    teste). Uma linha 2 puramente estrutural (nenhum dado de diária real)
    supre esse modelo sem afetar nada do que este teste verifica."""
    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Controle de Diárias"
    planilha.append(
        [
            "DH",
            "SEI",
            "DOC ORIGEM",
            "CPF",
            "PREVIDÊNCIA",
            "NOME",
            "CARGO",
            "PERÍODO INICIAL",
            "PERÍODO FINAL",
            "TRECHO",
            "TRANSPORTE",
            "MOTIVO",
        ]
    )
    planilha.append(
        [
            None, "SEI-MODELO-ESTRUTURAL", "PJMODELO", 0, None, None, None,
            "01/01/2020", "03/01/2020", "MODELO-MODELO", "CARRO", "Linha-modelo estrutural",
        ]
    )
    for coluna, formula in (
        ("N", "=I2-H2"),
        ("P", "=R2+S2"),
        ("R", "=(VLOOKUP(Q2,'Tabela Valores'!A:C,2,FALSE()))*N2"),
        ("S", "=(VLOOKUP(Q2,'Tabela Valores'!A:C,3,FALSE()))*O2"),
        ("V", "=CONCATENATE(D2,H2,I2)"),
        ("Z", "=H2-15"),
        ("AA", "=VLOOKUP(Z2,Z$1:AB1,3,0)"),
    ):
        planilha[f"{coluna}2"] = formula
    workbook.save(caminho)


def test_coluna_p_e_preenchida_somente_para_diaria_de_membro(
    tmp_path, pdfs_diaria_membro, pdfs_exemplo
):
    caminho_planilha = str(tmp_path / "planilha.xlsx")
    _planilha_vazia(caminho_planilha)

    lote_membro = processar_lote(pdfs_diaria_membro[:1])
    lote_padrao = processar_lote(pdfs_exemplo[:1])
    processos = [*lote_membro.processos, *lote_padrao.processos]

    resultado, arquivos_gerados = preencher_planilha(caminho_planilha, processos, [])
    assert resultado.total_gravado == 2

    workbook = load_workbook(arquivos_gerados.caminho_planilha_gerada)
    planilha = workbook["Controle de Diárias"]

    linha_membro = next(
        item.linha
        for item in resultado.gravados
        if item.arquivo_origem == os.path.basename(pdfs_diaria_membro[0])
    )
    linha_padrao = next(
        item.linha
        for item in resultado.gravados
        if item.arquivo_origem == os.path.basename(pdfs_exemplo[0])
    )

    assert planilha[f"P{linha_membro}"].value is not None
    # P0 FINAL (2026-08-09): P do processo Padrão nunca recebe um VALOR
    # literal (regra intacta) — mas deixa de ficar vazio, porque a
    # preparação estrutural automática já copiou a fórmula da linha-modelo
    # antes desta gravação (mesmo achado de `test_escritor_planilha.py::
    # test_preencher_nao_grava_coluna_p_para_diaria_padrao`).
    assert planilha[f"P{linha_padrao}"].value == f"=R{linha_padrao}+S{linha_padrao}"

    # Demais colunas já homologadas continuam preenchidas normalmente para
    # os dois tipos de processo.
    for linha in (linha_membro, linha_padrao):
        assert planilha[f"B{linha}"].value  # SEI
        assert planilha[f"D{linha}"].value  # CPF
