import multiprocessing
import os
import sqlite3
import time

from openpyxl import Workbook, load_workbook

from app.consulta import base_local
from tests._multiprocesso_util import executar_e_reportar, segurar_trava_ate_ser_morto


def test_caminho_base_local_usa_env_var_quando_definida(tmp_path, monkeypatch):
    monkeypatch.setenv(base_local.VAR_AMBIENTE_USER_DATA, str(tmp_path))
    caminho = base_local.caminho_base_local()
    assert caminho == tmp_path / "central-conferencia" / base_local.NOME_ARQUIVO_BASE_LOCAL


def test_caminho_base_local_fallback_appdata_quando_env_var_ausente(tmp_path, monkeypatch):
    monkeypatch.delenv(base_local.VAR_AMBIENTE_USER_DATA, raising=False)
    monkeypatch.setenv("APPDATA", str(tmp_path))
    caminho = base_local.caminho_base_local()
    assert caminho == tmp_path / base_local.NOME_PASTA_DADOS_FALLBACK / "central-conferencia" / base_local.NOME_ARQUIVO_BASE_LOCAL


def test_sincronizar_copia_planilha_para_destino_fixo(tmp_path):
    origem = tmp_path / "planilha_preenchida.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "conteudo A"
    workbook.save(origem)

    destino_base = tmp_path / "dados_usuario"
    caminho_destino = base_local.sincronizar(str(origem), diretorio_base=destino_base)

    assert caminho_destino.is_file()
    assert load_workbook(caminho_destino).active["A1"].value == "conteudo A"
    # nenhum arquivo temporário deve sobrar depois da troca atômica
    assert not any(nome.endswith(".tmp") for nome in os.listdir(caminho_destino.parent))


def test_sincronizar_sobrescreve_versao_anterior(tmp_path):
    origem = tmp_path / "planilha_preenchida.xlsx"
    destino_base = tmp_path / "dados_usuario"

    workbook_v1 = Workbook()
    workbook_v1.active["A1"] = "versão 1"
    workbook_v1.save(origem)
    base_local.sincronizar(str(origem), diretorio_base=destino_base)

    workbook_v2 = Workbook()
    workbook_v2.active["A1"] = "versão 2"
    workbook_v2.save(origem)
    caminho_destino = base_local.sincronizar(str(origem), diretorio_base=destino_base)

    assert load_workbook(caminho_destino).active["A1"].value == "versão 2"


def test_base_local_disponivel(tmp_path):
    destino_base = tmp_path / "dados_usuario"
    assert base_local.base_local_disponivel(diretorio_base=destino_base) is False

    origem = tmp_path / "planilha_preenchida.xlsx"
    Workbook().save(origem)
    base_local.sincronizar(str(origem), diretorio_base=destino_base)

    assert base_local.base_local_disponivel(diretorio_base=destino_base) is True


def test_sincronizar_tambem_reconstroi_base_local_em_sqlite(
    tmp_path, caminho_planilha_conferencia_teste
):
    """Sprint 17.3 — o mesmo gancho que já copia o `.xlsx` passa também a
    reconstruir a segunda materialização da Base Local, em SQLite."""
    origem = caminho_planilha_conferencia_teste()
    destino_base = tmp_path / "dados_usuario"

    caminho_xlsx = base_local.sincronizar(origem, diretorio_base=destino_base)

    caminho_sqlite = caminho_xlsx.parent / "base_local.sqlite"
    assert caminho_sqlite.is_file()


def test_sincronizar_grava_versao_app_da_variavel_de_ambiente(
    tmp_path, caminho_planilha_conferencia_teste, monkeypatch
):
    """URGENTE — Central Analítica (2026-08-06, achado real do 2º
    rastreamento): `sincronizar()` repassa `EXTRATOR_DIARIAS_APP_VERSION`
    (a mesma variável que `electron/main.ts` passa via `app.getVersion()`)
    para `base_local_sqlite.reconstruir`, gravada em
    `sincronizacoes.versao_app` — é o que permite `GET /sincronizacao`
    detectar uma atualização do app mesmo quando a Planilha Oficial e o
    schema de tabelas não mudam."""
    import sqlite3

    monkeypatch.setenv(base_local.VAR_AMBIENTE_VERSAO_APP, "2.1.0")
    origem = caminho_planilha_conferencia_teste()
    destino_base = tmp_path / "dados_usuario"

    caminho_xlsx = base_local.sincronizar(origem, diretorio_base=destino_base)

    caminho_sqlite = caminho_xlsx.parent / "base_local.sqlite"
    conexao = sqlite3.connect(caminho_sqlite)
    try:
        (versao_gravada,) = conexao.execute(
            "SELECT versao_app FROM sincronizacoes ORDER BY id DESC LIMIT 1"
        ).fetchone()
    finally:
        conexao.close()
    assert versao_gravada == "2.1.0"


def test_sincronizar_sem_variavel_de_ambiente_grava_versao_app_vazia(
    tmp_path, caminho_planilha_conferencia_teste, monkeypatch
):
    """Backend rodando sem Electron (standalone/dev/teste) — nunca inventa
    uma versão que não foi informada."""
    import sqlite3

    monkeypatch.delenv(base_local.VAR_AMBIENTE_VERSAO_APP, raising=False)
    origem = caminho_planilha_conferencia_teste()
    destino_base = tmp_path / "dados_usuario"

    caminho_xlsx = base_local.sincronizar(origem, diretorio_base=destino_base)

    caminho_sqlite = caminho_xlsx.parent / "base_local.sqlite"
    conexao = sqlite3.connect(caminho_sqlite)
    try:
        (versao_gravada,) = conexao.execute(
            "SELECT versao_app FROM sincronizacoes ORDER BY id DESC LIMIT 1"
        ).fetchone()
    finally:
        conexao.close()
    assert versao_gravada == ""


def test_sincronizar_publica_sqlite_mesmo_com_planilha_sem_dimension(
    tmp_path, caminho_planilha_conferencia_teste, remover_tag_dimension_da_planilha
):
    """Micro-sprint 18.7.1 — regressão do bug real (auditoria do Gatilho
    A): uma planilha sem a tag `<dimension>` (ex. exportada pelo Google
    Sheets) fazia `capturar_sinal_frescor` devolver `None`, e
    `sincronizar()` engolia o `sqlite3.IntegrityError` resultante em
    silêncio — `base_local.sqlite` nunca era publicado, não importa qual
    dos três chamadores (pós-preenchimento, botão manual, Configuração
    Inicial) disparasse. Testado neste nível (não só em
    `test_base_local_sqlite.py`) porque é exatamente esta função,
    `base_local.sincronizar()`, que os três gatilhos chamam."""
    origem = caminho_planilha_conferencia_teste()
    remover_tag_dimension_da_planilha(origem)
    destino_base = tmp_path / "dados_usuario"

    caminho_xlsx = base_local.sincronizar(origem, diretorio_base=destino_base)

    caminho_sqlite = caminho_xlsx.parent / "base_local.sqlite"
    assert caminho_sqlite.is_file()
    assert not any(nome.endswith(".tmp") for nome in os.listdir(caminho_xlsx.parent))


def test_sincronizar_planilha_sem_abas_esperadas_nao_quebra_copia_xlsx(tmp_path):
    """A cópia `.xlsx` (já homologada desde a Sprint 13) nunca pode ser
    afetada por uma falha na etapa nova de reconstrução SQLite (risco 4,
    `docs/central-analitica/README.md § 10`) — aqui simulada com uma
    planilha sem nenhuma das abas esperadas."""
    origem = tmp_path / "planilha_sem_abas.xlsx"
    Workbook().save(origem)  # só a aba padrão "Sheet"

    destino_base = tmp_path / "dados_usuario"
    caminho_destino = base_local.sincronizar(str(origem), diretorio_base=destino_base)

    assert caminho_destino.is_file()
    caminho_sqlite = caminho_destino.parent / "base_local.sqlite"
    assert not caminho_sqlite.exists()


def _origem_gravada(caminho_sqlite):
    conexao = sqlite3.connect(caminho_sqlite)
    try:
        return conexao.execute(
            "SELECT origem FROM sincronizacoes ORDER BY id DESC LIMIT 1"
        ).fetchone()[0]
    finally:
        conexao.close()


def test_sincronizar_grava_origem_padrao_pos_preenchimento_quando_nao_informada(
    tmp_path, caminho_planilha_conferencia_teste
):
    """Sprint 17.4.6.1 — preservação integral do fluxo consolidado: uma
    chamada sem o parâmetro `origem` (exatamente como
    `app/services/planilha_service.py` sempre chamou) continua gravando
    `"pos_preenchimento"`."""
    origem_planilha = caminho_planilha_conferencia_teste()
    destino_base = tmp_path / "dados_usuario"

    caminho_xlsx = base_local.sincronizar(origem_planilha, diretorio_base=destino_base)

    caminho_sqlite = caminho_xlsx.parent / "base_local.sqlite"
    assert _origem_gravada(caminho_sqlite) == "pos_preenchimento"


def test_sincronizar_grava_origem_informada_explicitamente(tmp_path, caminho_planilha_conferencia_teste):
    """Sprint 17.4.6.1 — `origem` persiste de verdade na Base Local em
    SQLite (não é simulada/sobrescrita só na resposta HTTP): o valor
    passado por `executar_sincronizacao_manual` (`"manual"`) é o mesmo
    valor lido de volta do banco."""
    origem_planilha = caminho_planilha_conferencia_teste()
    destino_base = tmp_path / "dados_usuario"

    caminho_xlsx = base_local.sincronizar(origem_planilha, diretorio_base=destino_base, origem="manual")

    caminho_sqlite = caminho_xlsx.parent / "base_local.sqlite"
    assert _origem_gravada(caminho_sqlite) == "manual"


# ============================================================================
# QA Codex (revalidação cross-process) — `LOCK_SINCRONIZACAO` (`RLock`,
# correção ETAPA 2 original) só protegia threads do mesmo processo; sob
# `WEB_CONCURRENCY`/múltiplos workers, dois processos distintos podiam
# colidir na mesma cópia `.xlsx`/reconstrução SQLite sem nenhuma
# coordenação. Corrigido com `_trava_sincronizacao_arquivo`
# (`filelock.FileLock`, cross-process) — o `RLock` continua existindo só
# como sinal rápido em memória para `executar_sincronizacao_manual` (ver
# `app/services/analitica_service.py`), não mais como a proteção real do
# arquivo. Testes com `multiprocessing.Process` de verdade.
# ============================================================================


def test_sincronizar_multiprocesso_nao_corrompe_nem_falha(tmp_path):
    origem_a = tmp_path / "planilha_a.xlsx"
    workbook_a = Workbook()
    workbook_a.active["A1"] = "conteudo A"
    workbook_a.save(origem_a)

    origem_b = tmp_path / "planilha_b.xlsx"
    workbook_b = Workbook()
    workbook_b.active["A1"] = "conteudo B"
    workbook_b.save(origem_b)

    destino_base = tmp_path / "dados_usuario"

    fila = multiprocessing.Queue()
    p1 = multiprocessing.Process(
        target=executar_e_reportar,
        args=(fila, base_local.sincronizar, (str(origem_a),), {"diretorio_base": destino_base}),
    )
    p2 = multiprocessing.Process(
        target=executar_e_reportar,
        args=(fila, base_local.sincronizar, (str(origem_b),), {"diretorio_base": destino_base}),
    )
    p1.start()
    p2.start()
    resultados = [fila.get(timeout=30) for _ in range(2)]
    p1.join(timeout=30)
    p2.join(timeout=30)

    assert all(r[0] == "ok" for r in resultados), resultados  # nenhuma exceção em nenhum processo

    caminho_destino = base_local.caminho_base_local(diretorio_base=destino_base)
    assert caminho_destino.is_file()
    # arquivo final sempre válido e coerente — nunca um estado misto/
    # corrompido entre as duas cópias concorrentes (last-write-wins
    # determinístico, mesma semântica de "sincronizar de novo sobrescreve").
    conteudo_final = load_workbook(caminho_destino).active["A1"].value
    assert conteudo_final in {"conteudo A", "conteudo B"}
    assert not any(nome.endswith(".tmp") for nome in os.listdir(caminho_destino.parent))


def test_sincronizar_processo_morto_segurando_trava_recupera_depois(tmp_path):
    origem = tmp_path / "planilha_preenchida.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "conteudo recuperacao"
    workbook.save(origem)

    destino_base = tmp_path / "dados_usuario"
    destino = base_local.caminho_base_local(diretorio_base=destino_base)
    destino.parent.mkdir(parents=True, exist_ok=True)

    fila = multiprocessing.Queue()
    processo = multiprocessing.Process(
        target=segurar_trava_ate_ser_morto,
        args=(fila, base_local._trava_sincronizacao_arquivo, (destino,)),
    )
    processo.start()
    sinal = fila.get(timeout=15)
    assert sinal == "trava_obtida"

    processo.kill()
    processo.join(timeout=15)
    assert not processo.is_alive()

    inicio = time.monotonic()
    caminho_destino = base_local.sincronizar(str(origem), diretorio_base=destino_base)
    duracao = time.monotonic() - inicio
    assert duracao < 8  # recupera bem antes do timeout de 10s — nenhum lock órfão
    assert load_workbook(caminho_destino).active["A1"].value == "conteudo recuperacao"
