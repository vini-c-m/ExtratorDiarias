from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.services.finalizacao_service import (
    CategoriaConfirmacao,
    ItemConfirmacao,
    StatusFinalizacao,
    analisar_lote,
    confirmar,
)

# Fixtures 100% sintéticas (sem dado real) — mesma política de
# `tests/fixtures/gerar_fixtures_sinteticas.py`, geradas aqui mesmo por não
# precisarem de campos do formulário de Solicitação (esta funcionalidade só
# lê SEI/Confirmação/Comprovante, nunca CPF/período/motivo).


def _gerar_pdf(caminho: str, paginas: list[list[str]]) -> None:
    c = canvas.Canvas(caminho, pagesize=A4)
    _largura, altura = A4
    for pagina in paginas:
        y = altura - 50
        for linha in pagina:
            c.drawString(50, y, linha)
            y -= 14
        c.showPage()
    c.save()


def _pdf_confirmacao_e_recibo(tmp_path, sei: str, nome_arquivo: str = "processo.pdf") -> str:
    caminho = str(tmp_path / nome_arquivo)
    _gerar_pdf(
        caminho,
        [
            [
                "CMFC - Confirmação dos Serviços - Diária -Servidor - 1000001",
                "Confirmação dos serviços:",
                "Sim",
                f"CMFC - Confirmação dos Serviços - Diária -Servidor 1000001 SEI {sei} / pg. 1",
            ],
            [f"Recibo (1000002) SEI {sei} / pg. 2"],
        ],
    )
    return caminho


def _pdf_sem_confirmacao(tmp_path, sei: str, nome_arquivo: str = "sem_confirmacao.pdf") -> str:
    caminho = str(tmp_path / nome_arquivo)
    _gerar_pdf(caminho, [[f"CMFC - Solicitação de Diária - Servidor {sei} SEI {sei} / pg. 1"]])
    return caminho


def _pdf_confirmacao_e_anexo_externo_generico(
    tmp_path, sei: str, nome_arquivo: str = "alerta.pdf"
) -> str:
    """Comprovante com título só "EXTERNO" — gera alerta (Sprint pré-build,
    2026-08-11), nunca mais bloqueio."""
    caminho = str(tmp_path / nome_arquivo)
    _gerar_pdf(
        caminho,
        [
            [
                "CMFC - Confirmação dos Serviços - Diária -Servidor - 1000001",
                "Confirmação dos serviços:",
                "Sim",
                f"CMFC - Confirmação dos Serviços - Diária -Servidor 1000001 SEI {sei} / pg. 1",
            ],
            [f"EXTERNO (1000002) SEI {sei} / pg. 2"],
        ],
    )
    return caminho


def _planilha_com_linhas(caminho_planilha_teste, linhas: list[tuple[str, object, object]]) -> str:
    """`linhas`: lista de (sei, x, y), uma por linha a partir da linha 4."""
    caminho = caminho_planilha_teste()
    workbook = load_workbook(caminho)
    planilha = workbook["Controle de Diárias"]
    for indice, (sei, x, y) in enumerate(linhas):
        linha = 4 + indice
        planilha[f"B{linha}"] = sei
        if x is not None:
            planilha[f"X{linha}"] = x
        if y is not None:
            planilha[f"Y{linha}"] = y
    workbook.save(caminho)
    return caminho


# --- analisar_lote -----------------------------------------------------------
#
# Regra de negócio definitiva (Sprint pré-build, 2026-08-11): o único
# critério que decide `status` é a localização do SEI na planilha — a
# avaliação documental (`avaliar`) nunca mais bloqueia, só popula
# `avaliacao.alertas`, sempre informativo.


def test_analisar_lote_sei_nao_localizado_bloqueia_mesmo_com_documentacao_completa(
    tmp_path, caminho_planilha_teste
):
    caminho_planilha = caminho_planilha_teste()
    sei = "19.20.0002.0000002/2026-02"
    pdf = _pdf_confirmacao_e_recibo(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.SEI_NAO_ENCONTRADO
    assert itens[0].avaliacao.apto  # documentação em si está ok — só a localização bloqueia


def test_analisar_lote_apto_uma_linha_pronta(tmp_path, caminho_planilha_teste):
    sei = "19.20.0002.0000003/2026-03"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, None, None)])
    pdf = _pdf_confirmacao_e_recibo(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.PRONTO
    assert itens[0].avaliacao.alertas == ()
    assert len(itens[0].candidatos) == 1
    assert itens[0].candidatos[0].linha == 4


def test_analisar_lote_sem_confirmacao_documental_mas_sei_localizado_fica_pronto_com_alerta(
    tmp_path, caminho_planilha_teste
):
    """Regra definitiva: documentação incompleta NUNCA bloqueia mais — SEI
    localizado é sempre PRONTO, com o alerta anexado para revisão do
    servidor."""
    sei = "19.20.0002.0000001/2026-01"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, None, None)])
    pdf = _pdf_sem_confirmacao(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.PRONTO
    assert len(itens[0].avaliacao.alertas) > 0
    assert not itens[0].avaliacao.apto  # documental continua registrando o problema, só não bloqueia


def test_analisar_lote_anexo_externo_generico_fica_pronto_com_alerta(tmp_path, caminho_planilha_teste):
    sei = "19.20.0003.0000001/2026-01"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, None, None)])
    pdf = _pdf_confirmacao_e_anexo_externo_generico(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.PRONTO
    assert any("revisar manualmente" in a for a in itens[0].avaliacao.alertas)
    assert len(itens[0].candidatos) == 1


def test_analisar_lote_ja_finalizado(tmp_path, caminho_planilha_teste):
    sei = "19.20.0002.0000004/2026-04"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, "SIM", "SIM")])
    pdf = _pdf_confirmacao_e_recibo(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.JA_FINALIZADO
    assert itens[0].motivo == "Já confirmado na Planilha Oficial."


def test_analisar_lote_estado_parcial_valor_fora_do_padrao_ainda_bloqueia(
    tmp_path, caminho_planilha_teste
):
    """Único caso de linha que ainda bloqueia: valor fora do padrão (ex.
    "SM"), nunca "SIM"/vazio (isso agora é INCOMPLETA, gravável)."""
    sei = "19.20.0002.0000005/2026-05"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, "SM", "SM")])
    pdf = _pdf_confirmacao_e_recibo(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.ESTADO_PARCIAL


def test_analisar_lote_multiplas_linhas_mesmo_sei_todas_vazias_fica_pronto(
    tmp_path, caminho_planilha_teste
):
    """Ajuste de 2026-08-07 — RN do processo, não da linha: mesmo SEI em
    duas linhas (diária + Complemento de Diárias) e ambas vazias já conta
    como pronto para o SEI inteiro, sem exigir escolha manual."""
    sei = "19.20.0002.0000006/2026-06"
    caminho_planilha = _planilha_com_linhas(
        caminho_planilha_teste, [(sei, None, None), (sei, None, None)]
    )
    pdf = _pdf_confirmacao_e_recibo(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.PRONTO
    assert [c.linha for c in itens[0].candidatos] == [4, 5]


def test_analisar_lote_multiplas_linhas_mistura_pronta_e_ja_finalizada_fica_pronto(
    tmp_path, caminho_planilha_teste
):
    """Uma linha já SIM/SIM e outra vazia — o SEI ainda está pronto (a já
    finalizada só é mantida, nunca bloqueia as demais)."""
    sei = "19.20.0002.0000012/2026-12"
    caminho_planilha = _planilha_com_linhas(
        caminho_planilha_teste, [(sei, "SIM", "SIM"), (sei, None, None)]
    )
    pdf = _pdf_confirmacao_e_recibo(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.PRONTO


def test_analisar_lote_multiplas_linhas_uma_incompleta_fica_pronto_nao_parcial(
    tmp_path, caminho_planilha_teste
):
    """SIM/vazio numa das linhas NÃO bloqueia mais o SEI (era ESTADO_PARCIAL
    antes da regra definitiva) — o SEI fica PRONTO, pronto para completar."""
    sei = "19.20.0002.0000013/2026-13"
    caminho_planilha = _planilha_com_linhas(
        caminho_planilha_teste, [(sei, None, None), (sei, "SIM", None)]
    )
    pdf = _pdf_confirmacao_e_recibo(tmp_path, sei)

    itens = analisar_lote([pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.PRONTO


def test_analisar_lote_arquivo_ilegivel_vira_erro_leitura(caminho_planilha_teste):
    caminho_planilha = caminho_planilha_teste()
    itens = analisar_lote(["/caminho/inexistente.pdf"], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.ERRO_LEITURA
    assert itens[0].sei is None


# --- confirmar ----------------------------------------------------------------


def test_confirmar_uma_linha_vazia_grava_sim_sim(caminho_planilha_teste):
    sei = "19.20.0002.0000007/2026-07"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, None, None)])
    workbook = load_workbook(caminho_planilha)
    planilha_antes = workbook["Controle de Diárias"]
    valores_linha_antes = [planilha_antes.cell(row=4, column=c).value for c in range(1, 13)]

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    assert resultado.total_finalizado == 1
    assert resultado.total_ja_confirmado == 0
    assert resultado.total_ignorado == 0
    assert resultado.itens[0].categoria is CategoriaConfirmacao.FINALIZADO
    assert resultado.itens[0].linhas_atualizadas == [4]
    assert resultado.itens[0].linhas_mantidas == []

    planilha_depois = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha_depois["X4"].value == "SIM"
    assert planilha_depois["Y4"].value == "SIM"
    valores_linha_depois = [planilha_depois.cell(row=4, column=c).value for c in range(1, 13)]
    assert valores_linha_depois == valores_linha_antes


# --- Testes A-D do pedido: SIM/SIM, SIM/vazio, vazio/SIM, vazio/vazio ------


def test_a_sim_sim_nao_escreve_nada_e_vira_ja_confirmado(caminho_planilha_teste):
    sei = "19.20.0004.0000001/2026-01"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, "SIM", "SIM")])

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    assert resultado.total_finalizado == 0
    assert resultado.total_ja_confirmado == 1
    assert resultado.itens[0].categoria is CategoriaConfirmacao.JA_CONFIRMADO
    assert resultado.itens[0].motivo == "Já confirmado na Planilha Oficial."
    assert resultado.itens[0].linhas_atualizadas == []
    assert resultado.itens[0].linhas_mantidas == [4]


def test_b_sim_vazio_grava_somente_coluna_faltante(caminho_planilha_teste):
    sei = "19.20.0004.0000002/2026-02"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, "SIM", None)])

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    assert resultado.total_finalizado == 1
    assert resultado.itens[0].categoria is CategoriaConfirmacao.FINALIZADO
    assert resultado.itens[0].linhas_atualizadas == [4]

    planilha = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha["X4"].value == "SIM"
    assert planilha["Y4"].value == "SIM"


def test_c_vazio_sim_grava_somente_coluna_faltante(caminho_planilha_teste):
    sei = "19.20.0004.0000003/2026-03"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, None, "SIM")])

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    assert resultado.total_finalizado == 1
    planilha = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha["X4"].value == "SIM"
    assert planilha["Y4"].value == "SIM"


def test_d_vazio_vazio_grava_sim_sim(caminho_planilha_teste):
    sei = "19.20.0004.0000004/2026-04"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, None, None)])

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    assert resultado.total_finalizado == 1
    planilha = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha["X4"].value == "SIM"
    assert planilha["Y4"].value == "SIM"


# --- Testes E-H do pedido: alertas nunca bloqueiam a gravação --------------


def test_e_alerta_documental_com_sei_localizado_finaliza_normalmente(tmp_path, caminho_planilha_teste):
    sei = "19.20.0004.0000005/2026-05"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, None, None)])
    pdf = _pdf_sem_confirmacao(tmp_path, sei)

    analise = analisar_lote([pdf], caminho_planilha)
    assert analise[0].status is StatusFinalizacao.PRONTO
    assert len(analise[0].avaliacao.alertas) > 0

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)
    assert resultado.total_finalizado == 1
    planilha = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha["X4"].value == "SIM" and planilha["Y4"].value == "SIM"


def test_f_sei_nao_localizado_nao_grava(caminho_planilha_teste):
    caminho_planilha = caminho_planilha_teste()

    resultado = confirmar(
        [ItemConfirmacao(arquivo_origem="p.pdf", sei="19.20.9999.9999999/2099-99")],
        caminho_planilha,
    )

    assert resultado.total_finalizado == 0
    assert resultado.itens[0].categoria is CategoriaConfirmacao.IGNORADO
    assert resultado.itens[0].motivo == "SEI não localizado na Planilha Oficial."


def test_g_sei_nao_extraido_nao_grava(tmp_path, caminho_planilha_teste):
    caminho_planilha = caminho_planilha_teste()
    caminho_pdf = str(tmp_path / "sem_sei.pdf")
    _gerar_pdf(caminho_pdf, [["nenhum numero de processo aqui"]])

    itens = analisar_lote([caminho_pdf], caminho_planilha)

    assert itens[0].status is StatusFinalizacao.SEI_NAO_EXTRAIDO
    assert itens[0].sei is None


def test_h_erro_real_de_escrita_nao_e_mascarado_como_alerta(caminho_planilha_teste):
    """Falha técnica de escrita continua propagando como exceção real —
    nunca convertida em item "com alerta"."""
    import pytest

    from app.services.finalizacao_service import ErroFinalizacaoDiarias

    with pytest.raises(ErroFinalizacaoDiarias):
        confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei="x")], "/caminho/inexistente.xlsx")


def test_confirmar_uma_linha_em_estado_parcial_aborta_todas(caminho_planilha_teste):
    """Atomicidade por SEI: uma linha com valor fora do padrão bloqueia
    TODO o SEI — mesmo as demais linhas, prontas, ficam sem alteração."""
    sei = "19.20.0002.0000005/2026-05"
    caminho_planilha = _planilha_com_linhas(
        caminho_planilha_teste, [(sei, None, None), (sei, "SM", "SM")]
    )

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    assert resultado.total_finalizado == 0
    assert resultado.itens[0].categoria is CategoriaConfirmacao.IGNORADO
    assert resultado.itens[0].motivo == "Estado parcial — revisão manual necessária."
    assert resultado.itens[0].linhas_atualizadas == []

    planilha = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha["X4"].value is None and planilha["Y4"].value is None
    assert planilha["X5"].value == "SM" and planilha["Y5"].value == "SM"


def test_confirmar_nao_altera_nenhuma_outra_coluna(caminho_planilha_teste):
    sei = "19.20.0002.0000017/2026-17"
    caminho_planilha = _planilha_com_linhas(
        caminho_planilha_teste, [(sei, None, None), (sei, None, None)]
    )
    antes = load_workbook(caminho_planilha)["Controle de Diárias"]
    instantaneo_antes = {
        linha: [antes.cell(row=linha, column=c).value for c in range(1, 24)] for linha in (4, 5)
    }

    confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    depois = load_workbook(caminho_planilha)["Controle de Diárias"]
    for linha in (4, 5):
        instantaneo_depois = [depois.cell(row=linha, column=c).value for c in range(1, 24)]
        assert instantaneo_depois == instantaneo_antes[linha]


def test_confirmar_revalida_estado_no_momento_da_gravacao(caminho_planilha_teste):
    """SEI já finalizado por outra operação entre a análise e a confirmação
    — nunca regravado."""
    sei = "19.20.0002.0000009/2026-09"
    caminho_planilha = _planilha_com_linhas(caminho_planilha_teste, [(sei, "SIM", "SIM")])

    resultado = confirmar([ItemConfirmacao(arquivo_origem="p.pdf", sei=sei)], caminho_planilha)

    assert resultado.total_finalizado == 0
    assert resultado.total_ja_confirmado == 1
    assert resultado.itens[0].motivo == "Já confirmado na Planilha Oficial."
    assert resultado.itens[0].linhas_mantidas == [4]


# --- I) lote misto: localizados finalizam, não localizados ficam de fora ---


def test_i_lote_misto_localizados_finalizam_nao_localizados_ficam_de_fora_nenhum_desaparece(
    caminho_planilha_teste,
):
    sei_bloqueado = "19.20.0002.0000010/2026-10"
    sei_ok = "19.20.0002.0000011/2026-11"
    caminho_planilha = _planilha_com_linhas(
        caminho_planilha_teste, [(sei_bloqueado, "SM", "SM"), (sei_ok, None, None)]
    )

    resultado = confirmar(
        [
            ItemConfirmacao(arquivo_origem="a.pdf", sei=sei_bloqueado),
            ItemConfirmacao(arquivo_origem="b.pdf", sei="19.20.9999.9999999/2099-99"),
            ItemConfirmacao(arquivo_origem="c.pdf", sei=sei_ok),
        ],
        caminho_planilha,
    )

    assert len(resultado.itens) == 3  # J) entrada = resultado, nenhum desaparece
    assert resultado.total_finalizado == 1
    assert resultado.total_ignorado == 2
    planilha = load_workbook(caminho_planilha)["Controle de Diárias"]
    assert planilha["X5"].value == "SIM" and planilha["Y5"].value == "SIM"
    # SEI bloqueado permanece intocado.
    assert planilha["X4"].value == "SM" and planilha["Y4"].value == "SM"
