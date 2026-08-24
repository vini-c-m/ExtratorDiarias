"""Micro-sprint 3F — template SINTÉTICO do Relatório de Diárias — Portal da
Transparência, para testes automatizados. Reproduz a ESTRUTURA do modelo
oficial real (investigação read-only, nunca versionada) — abas, células de
título/cabeçalho, mesclagens, estilos, totais com fórmula deliberadamente
desatualizada (para provar que o gerador recalcula), rodapé/notas, page
setup, freeze panes, autofiltro — com dados 100% fictícios (3 linhas de
exemplo, nenhum CPF/nome/OB/SEI real).

Nunca importar isto fora de `tests/` — é só fixture, não parte do
gerador."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

NOME_ABA = "Controle de Diárias"

_BORDA_FINA = Border(*(Side(style="thin") for _ in range(4)))
_ALINHAMENTO_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FONTE_TITULO = Font(name="Libre Franklin", size=18, bold=True)
_FONTE_SUBTITULO = Font(name="Libre Franklin", size=13, bold=True)
_FONTE_CABECALHO = Font(name="Libre Franklin", size=11, bold=True, color="FF7F7F7F")
_PREENCHIMENTO_CABECALHO = PatternFill(fill_type="solid", fgColor="FFFFC000")
_FONTE_SUBCABECALHO = Font(name="Libre Franklin", size=11, bold=True, color="FFFFFFFF")
_PREENCHIMENTO_SUBCABECALHO = PatternFill(fill_type="solid", fgColor="FFD8D8D8")
_FONTE_DADOS = Font(name="Tahoma", size=10)

_CABECALHOS = {
    "A": "Ordem Bancária",
    "B": "Data de Pagamento",
    "C": "Nome",
    "D": "Cargo",
    "E": "Período",
    "G": "Trecho",
    "H": "Transporte",
    "I": "Motivo",
    "J": "Valor da Passagem",
    "K": "Nº Diárias                                            Integral              Parcial",
    "M": "Valor total diárias",
}
_SUBCABECALHOS = {"C": "(a)", "D": "(b)", "E": "(c)", "G": "(d)", "H": "(e)", "I": "(f)", "J": "(g)", "K": "(h)", "M": "(i)"}

_LARGURAS = {
    "A": 14.85, "B": 16.0, "C": 38.14, "D": 32.71, "E": 16.42, "F": 14.71,
    "G": 29.85, "H": 17.14, "I": 66.14, "J": 17.14, "K": 16.42, "L": 14.57, "M": 17.28,
}

_DADOS_EXEMPLO = [
    # (ob, data_pagamento, nome, cargo, periodo_ini, periodo_fim, trecho, transporte,
    #  motivo, valor_passagem, integral, parcial, valor_total)
    ("2026OB000001", "01/01/2026 08:00:00", "FULANO DE TAL", "Cargo de Teste",
     "07/01/2026", "07/01/2026", "RECIFE-GARANHUNS", "CARRO", "Motivo de teste", "-", 0, 1, 154.09),
    ("2026OB000002", "02/01/2026 08:00:00", "BELTRANO DE TAL", "Cargo de Teste 2",
     "10/01/2026", "12/01/2026", "RECIFE-FORTALEZA", "AVIÃO", "Outro motivo de teste", "-", 2, 1, 900.00),
    ("2026OB000003", "03/01/2026 08:00:00", "CICRANO DE TAL", "Cargo de Teste 3",
     "05/01/2026", "05/01/2026", "RECIFE-ARCOVERDE", "CARRO", "Terceiro motivo de teste", "-", 0, 1, 200.50),
]


def _escrever_linha_dados(ws: Worksheet, linha: int, dados: tuple) -> None:
    from datetime import datetime as _dt

    valores = [
        dados[0], _dt.strptime(dados[1], "%d/%m/%Y %H:%M:%S"), dados[2], dados[3],
        _dt.strptime(dados[4], "%d/%m/%Y").date(), _dt.strptime(dados[5], "%d/%m/%Y").date(),
        dados[6], dados[7], dados[8], dados[9], dados[10], dados[11], dados[12],
    ]
    for coluna, valor in zip("ABCDEFGHIJKLM", valores):
        celula = ws[f"{coluna}{linha}"]
        celula.value = valor
        celula.font = _FONTE_DADOS
        celula.border = _BORDA_FINA
        celula.alignment = _ALINHAMENTO_CENTRO
        if coluna in ("E", "F"):
            celula.number_format = "dd/mm/yyyy"
        elif coluna == "M":
            celula.number_format = '"R$ "#,##0.00'
    ws.row_dimensions[linha].height = 51.75


def construir_template_sintetico(caminho: str) -> None:
    """Salva um template sintético completo em `caminho`. Chamado só por
    testes — nunca em código de produção."""
    wb = Workbook()
    ws = wb.active
    ws.title = NOME_ABA
    wb.create_sheet("Planilha1")

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_margins.left = 0.7
    ws.page_margins.right = 0.7
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.75

    for coluna, largura in _LARGURAS.items():
        ws.column_dimensions[coluna].width = largura

    ws.row_dimensions[1].height = 16.5
    ws.row_dimensions[2].height = 26.25
    ws.row_dimensions[3].height = 26.25
    ws.row_dimensions[4].height = 16.5

    ws.merge_cells("A2:C2")
    ws["A2"] = "6.1.8 Diárias e Passagens"
    ws["A2"].font = _FONTE_TITULO
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells("A3:C3")
    ws["A3"] = "PAGAS ATÉ JANEIRO/2026"
    ws["A3"].font = _FONTE_SUBTITULO
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A3"].number_format = "@"

    ws.merge_cells("E5:F5")
    ws.merge_cells("K5:L5")
    ws.row_dimensions[5].height = 50.25
    for coluna, texto in _CABECALHOS.items():
        celula = ws[f"{coluna}5"]
        celula.value = texto
        celula.font = _FONTE_CABECALHO
        celula.fill = _PREENCHIMENTO_CABECALHO
        celula.border = _BORDA_FINA
        celula.alignment = _ALINHAMENTO_CENTRO
        if coluna in ("J", "M"):
            celula.number_format = '"R$ "#,##0.00'
        elif coluna == "E":
            celula.number_format = "mm/dd/yyyy"

    ws.merge_cells("E6:F6")
    ws.merge_cells("K6:L6")
    ws.row_dimensions[6].height = 17.25
    for coluna, texto in _SUBCABECALHOS.items():
        celula = ws[f"{coluna}6"]
        celula.value = texto
        celula.font = _FONTE_SUBCABECALHO
        celula.fill = _PREENCHIMENTO_SUBCABECALHO
        celula.border = _BORDA_FINA
        celula.alignment = _ALINHAMENTO_CENTRO

    linha = 7
    for dados in _DADOS_EXEMPLO:
        _escrever_linha_dados(ws, linha, dados)
        linha += 1
    linha_total = linha  # 10, com as 3 linhas de exemplo acima

    ws.merge_cells(f"A{linha_total}:I{linha_total}")
    ws[f"A{linha_total}"] = "Total (j)"
    # Fórmula DELIBERADAMENTE desatualizada (só até a 2ª linha, não a 3ª) —
    # mesmo padrão real encontrado no modelo oficial (`=SUM(J7:J1524)`
    # quando os dados já iam até a linha 1818) — prova que o gerador
    # recalcula, nunca herda a fórmula velha do template.
    ws[f"J{linha_total}"] = f"=SUM(J7:J{linha_total - 2})"
    ws[f"M{linha_total}"] = f"=SUM(M7:M{linha_total - 2})"

    ws.merge_cells(f"A{linha_total + 1}:M{linha_total + 1}")
    ws[f"A{linha_total + 1}"] = "Fonte de Informação: ÓRGÃO FICTÍCIO DE TESTE"

    ws.merge_cells(f"A{linha_total + 2}:M{linha_total + 2}")
    ws[f"A{linha_total + 2}"] = "Data da última atualização: 01.01.2026"

    notas = [
        "(a) Nome – nota de teste.",
        "(b) Cargo – nota de teste.",
        "(c) Período – nota de teste.",
        "(d) Trecho – nota de teste.",
        "(e) Transporte – nota de teste.",
        "(f) Motivo – nota de teste.",
        "(g) Valor da passagem – nota de teste.",
        "(h) Nº de diárias – nota de teste.",
        "(i) Valor total das diárias – nota de teste.",
        "(j) Total – nota de teste.",
    ]
    linha_nota = linha_total + 4
    for texto in notas:
        ws.merge_cells(f"A{linha_nota}:L{linha_nota}")
        ws[f"A{linha_nota}"] = texto
        linha_nota += 1

    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A5:O{linha_total - 1}"

    wb.save(caminho)
